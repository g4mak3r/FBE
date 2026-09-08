from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .repository import (
    list_all_products,
    list_name_formats,
    list_names,
    list_samples,
    record_import,
    save_name_formats_batch,
    save_names_batch,
    save_products_batch,
    save_samples_batch,
)


DEFAULT_COLUMNS = {
    "seller_article": "Артикул продавца",
    "wb_article": "Артикул WB",
    "subcategory": "Подкатегория",
    "product": "Продукт",
    "volume": "Объем (парфюмерия)",
    "sample": "Пробник",
    "sample_volume": "Объем пробника",
    "category": "Категория",
    "marking_profile": "Профиль маркировки",
    "gtin": "GTIN",
    "product_group": "Товарная группа ЧЗ",
    "template_id": "templateId",
    "cis_type": "Тип КИЗ",
    "tnved_code": "ТН ВЭД",
    "certificate_type": "Тип документа",
    "certificate_number": "Номер документа",
    "certificate_date": "Дата начала действия документа",
    "certificate_valid_until": "Дата окончания действия документа",
    "shelf_life_months": "Срок годности, мес.",
    "alcohol_volume_pct": "Этиловый спирт, %",
    "notes": "Комментарий",
    "active": "Активен",
    "marketing_title": "Маркетинговое название",
    "fragrance_family": "Семейство аромата",
    "target_audience": "Аудитория",
    "marketing_description": "Маркетинговое описание",
    "marketing_keywords": "Ключевые слова",
    "package_length_cm": "Длина упаковки, см",
    "package_width_cm": "Ширина упаковки, см",
    "package_height_cm": "Высота упаковки, см",
    "package_weight_g": "Вес с упаковкой, г",
    "units_per_box": "Штук в коробе",
    "cost_price_rub": "Себестоимость, руб",
}

# Legacy input is still accepted, but it is deliberately not exported anymore.
# Since 0.77.0 the profile is the source of truth and controls kiz_required.
IMPORT_ONLY_COLUMNS = {
    "excel_kiz": "КИЗ",
}

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "cost_price_rub": ("Себестоимость, руб", "Себестоимость итого, руб"),
    "marking_profile": ("Профиль маркировки", "Профиль ЧЗ"),
    "excel_kiz": ("КИЗ", "Нужен КИЗ", "Нужен КИЗ?"),
    "tnved_code": ("ТН ВЭД", "ТН ВЭД 10", "Код ТН ВЭД"),
    "certificate_type": (
        "Тип документа", "Тип документа соответствия", "Вид документа соответствия",
    ),
    "certificate_number": (
        "Номер документа", "Номер документа соответствия", "Номер декларации / сертификата",
    ),
    "certificate_date": (
        "Дата начала действия документа", "Дата регистрации документа соответствия",
        "Дата регистрации", "Дата выдачи документа",
    ),
    "certificate_valid_until": (
        "Дата окончания действия документа", "Документ действует до", "Действует до",
    ),
    "shelf_life_months": ("Срок годности, мес.", "Срок годности, месяцев", "Срок годности (мес.)"),
    "alcohol_volume_pct": ("Этиловый спирт, %", "Содержание этилового спирта, %", "Спирт, %"),
}



CALCULATED_COLUMNS: tuple[str, ...] = ()

MARKING_PROFILE_EXPORT_LABELS = {
    "AUTO": "Автоматически",
    "PERFUMERY": "Парфюмерия",
    "CHEMISTRY": "Дезодоранты и косметика",
    "NONE": "Не требует маркировки",
    "CUSTOM": "Пользовательский профиль",
}

CERTIFICATE_TYPE_EXPORT_LABELS = {
    "CONFORMITY_DECLARATION": "Декларация о соответствии",
    "CONFORMITY_CERTIFICATE": "Сертификат соответствия",
}

MARKING_PROFILE_IMPORT_ALIASES = {
    "AUTO": "AUTO",
    "АВТО": "AUTO",
    "АВТОМАТИЧЕСКИ": "AUTO",
    "PERFUMERY": "PERFUMERY",
    "ПАРФЮМЕРИЯ": "PERFUMERY",
    "ДУХИ": "PERFUMERY",
    "CHEMISTRY": "CHEMISTRY",
    "ДЕЗОДОРАНТЫ": "CHEMISTRY",
    "ДЕЗОДОРАНТЫ И КОСМЕТИКА": "CHEMISTRY",
    "КОСМЕТИКА": "CHEMISTRY",
    "NONE": "NONE",
    "БЕЗ МАРКИРОВКИ": "NONE",
    "НЕ ТРЕБУЕТ МАРКИРОВКИ": "NONE",
    "НЕ ТРЕБУЕТСЯ МАРКИРОВКА": "NONE",
    "НЕ ТРЕБУЕТСЯ": "NONE",
    "CUSTOM": "CUSTOM",
    "ПОЛЬЗОВАТЕЛЬСКИЙ": "CUSTOM",
    "ПОЛЬЗОВАТЕЛЬСКИЙ ПРОФИЛЬ": "CUSTOM",
}

CERTIFICATE_TYPE_IMPORT_ALIASES = {
    "": "",
    "CONFORMITY_DECLARATION": "CONFORMITY_DECLARATION",
    "ДЕКЛАРАЦИЯ": "CONFORMITY_DECLARATION",
    "ДЕКЛАРАЦИЯ О СООТВЕТСТВИИ": "CONFORMITY_DECLARATION",
    "CONFORMITY_CERTIFICATE": "CONFORMITY_CERTIFICATE",
    "СЕРТИФИКАТ": "CONFORMITY_CERTIFICATE",
    "СЕРТИФИКАТ СООТВЕТСТВИЯ": "CONFORMITY_CERTIFICATE",
}



def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _yes(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    text = _clean(value).lower()
    if not text:
        return default
    return text in {"1", "true", "yes", "да", "+", "on"}


def _normalize_profile_from_excel(value: Any) -> str:
    text = _clean(value).upper().replace("Ё", "Е")
    if not text:
        return ""
    return MARKING_PROFILE_IMPORT_ALIASES.get(text, "INVALID")


def _normalize_certificate_type_from_excel(value: Any) -> str:
    text = _clean(value).upper().replace("Ё", "Е")
    return CERTIFICATE_TYPE_IMPORT_ALIASES.get(text, "INVALID")


def _legacy_profile_from_kiz(item: dict[str, Any], raw_kiz: Any) -> str:
    """Translate old КИЗ spreadsheets into an explicit marking profile.

    Empty/No means NONE. Yes is mapped to a known profile where possible and
    remains AUTO only when the old spreadsheet does not contain enough context.
    """
    if not _yes(raw_kiz, False):
        return "NONE"
    text = " ".join(
        _clean(item.get(key)).lower().replace("ё", "е")
        for key in ("category", "subcategory", "product")
    )
    if "дезодорант" in text or "косметик" in text:
        return "CHEMISTRY"
    if any(token in text for token in ("парфюмер", "духи", "туалетная вода")):
        return "PERFUMERY"
    return "AUTO"


RAW_META_KEY = "__fbe_cell_meta__"
RAW_HEADERS_KEY = "__fbe_header_order__"


def _serialize_raw_cell(cell: Any) -> tuple[Any, dict[str, Any] | None]:
    """Store the user's original Excel value separately from normalized DB fields."""
    value = cell.value
    number_format = str(getattr(cell, "number_format", "General") or "General")
    if value is None:
        meta = {"type": "blank", "number_format": number_format} if number_format != "General" else None
        return "", meta
    if getattr(cell, "data_type", "") == "f" or (isinstance(value, str) and value.startswith("=")):
        return str(value), {"type": "formula", "number_format": number_format}
    if isinstance(value, datetime):
        return value.isoformat(), {"type": "datetime", "number_format": number_format}
    if isinstance(value, date):
        return value.isoformat(), {"type": "date", "number_format": number_format}
    if isinstance(value, bool):
        return value, {"type": "bool", "number_format": number_format}
    if isinstance(value, (int, float)):
        return value, {"type": "number", "number_format": number_format}
    meta = {"type": "text", "number_format": number_format} if number_format != "General" else None
    return str(value), meta


def read_catalog_xlsx(path: str | Path, *, sheet_name: str = "mainSheet", columns: dict[str, str] | None = None) -> dict[str, Any]:
    source = Path(path)
    # One workbook provides cached formula results for internal normalization,
    # the second preserves the exact user-entered formulas/types/formats for export.
    wb = load_workbook(source, data_only=True, read_only=False)
    raw_wb = load_workbook(source, data_only=False, read_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Лист {sheet_name} не найден. Есть: {', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]
        raw_ws = raw_wb[sheet_name]
        import_columns = {**DEFAULT_COLUMNS, **IMPORT_ONLY_COLUMNS}
        mapping = dict(import_columns)
        if columns:
            mapping.update(columns)
        headers = [_clean(ws.cell(row=1, column=column).value) for column in range(1, ws.max_column + 1)]
        indexes = {header: i for i, header in enumerate(headers) if header}

        resolved_headers: dict[str, str] = {}
        for internal, default_header in import_columns.items():
            candidates = [mapping.get(internal, ""), default_header, *COLUMN_ALIASES.get(internal, ())]
            for candidate in candidates:
                if candidate and candidate in indexes:
                    resolved_headers[internal] = candidate
                    break

        article_header = resolved_headers.get("seller_article", mapping["seller_article"])
        if article_header not in indexes:
            raise ValueError(f"Не найден обязательный столбец «{mapping['seller_article']}»")

        products: list[dict[str, Any]] = []
        warnings: list[str] = []
        seen: set[str] = set()
        for row_number in range(2, ws.max_row + 1):
            values = [ws.cell(row=row_number, column=column).value for column in range(1, ws.max_column + 1)]
            raw: dict[str, Any] = {RAW_HEADERS_KEY: list(headers)}
            raw_meta: dict[str, Any] = {}
            for header, index in indexes.items():
                cell = raw_ws.cell(row=row_number, column=index + 1)
                serialized, meta = _serialize_raw_cell(cell)
                raw[header] = serialized
                if meta:
                    raw_meta[header] = meta
            if raw_meta:
                raw[RAW_META_KEY] = raw_meta

            article = _clean(values[indexes[article_header]]) if indexes[article_header] < len(values) else ""
            if not article:
                continue
            if article in seen:
                warnings.append(f"Строка {row_number}: повтор артикула {article}; используется последняя строка")
            seen.add(article)

            item: dict[str, Any] = {"seller_article": article, "raw_excel": raw, "active": True}
            legacy_kiz_value: Any = ""
            for internal in import_columns:
                header = resolved_headers.get(internal, "")
                value = values[indexes[header]] if header and indexes[header] < len(values) else ""
                if internal == "excel_kiz":
                    legacy_kiz_value = value
                elif internal == "active":
                    item["active"] = _yes(value, True)
                elif internal == "template_id":
                    try:
                        item[internal] = int(value or 0)
                    except (TypeError, ValueError):
                        item[internal] = 0
                        warnings.append(f"Строка {row_number}: некорректный templateId для {article}")
                elif internal == "marking_profile":
                    normalized_profile = _normalize_profile_from_excel(value)
                    if normalized_profile == "INVALID":
                        raise ValueError(
                            f"Строка {row_number}, {article}: неизвестный профиль маркировки «{_clean(value)}»"
                        )
                    item[internal] = normalized_profile
                elif internal == "certificate_type":
                    normalized_type = _normalize_certificate_type_from_excel(value)
                    if normalized_type == "INVALID":
                        raise ValueError(
                            f"Строка {row_number}, {article}: неизвестный тип документа «{_clean(value)}»"
                        )
                    item[internal] = normalized_type
                else:
                    item[internal] = _clean(value)

            # Typed fields may be derived for FBE's internal work, but the exact
            # user-entered cells above remain untouched in raw_excel and are used
            # for the next XLSX export.
            if not item.get("marking_profile"):
                item["marking_profile"] = _legacy_profile_from_kiz(item, legacy_kiz_value)
                warnings.append(
                    f"Строка {row_number}, {article}: профиль маркировки не заполнен; "
                    f"применен профиль {MARKING_PROFILE_EXPORT_LABELS.get(item['marking_profile'], item['marking_profile'])}"
                )
            item["kiz_required"] = item["marking_profile"] not in {"NONE"}
            if resolved_headers.get("excel_kiz") and resolved_headers.get("marking_profile"):
                legacy_required = _yes(legacy_kiz_value, False)
                profile_required = item["marking_profile"] != "NONE"
                if _clean(legacy_kiz_value) and legacy_required != profile_required:
                    warnings.append(
                        f"Строка {row_number}, {article}: значение старого столбца КИЗ проигнорировано; "
                        "использован Профиль маркировки"
                    )
            if not item.get("product_group"):
                item["product_group"] = "perfumery" if item.get("category") == "Парфюмерия" else ""
            if not item.get("cis_type"):
                item["cis_type"] = "UNIT"
            products.append(item)

        samples: list[dict[str, Any]] = []
        if "samples" in wb.sheetnames:
            ws_samples = wb["samples"]
            sample_headers = [_clean(cell.value) for cell in next(ws_samples.iter_rows(min_row=1, max_row=1))]
            sample_idx = {header: i for i, header in enumerate(sample_headers) if header}
            for row in ws_samples.iter_rows(min_row=2, values_only=True):
                values = list(row)
                record = {header: _clean(values[i]) if i < len(values) else "" for header, i in sample_idx.items()}
                if record.get("Name"):
                    record["active"] = _yes(record.get("Активен"), True)
                    samples.append(record)

        names: list[dict[str, Any]] = []
        if "names" in wb.sheetnames:
            ws_names = wb["names"]
            name_headers = [_clean(cell.value) for cell in next(ws_names.iter_rows(min_row=1, max_row=1))]
            idx = {header: i for i, header in enumerate(name_headers) if header}
            for row in ws_names.iter_rows(min_row=2, values_only=True):
                values = list(row)
                record = {header: _clean(values[i]) if i < len(values) else "" for header, i in idx.items()}
                if record.get("Название"):
                    record["active"] = _yes(record.get("Активен"), True)
                    names.append(record)

        formats: list[dict[str, Any]] = []
        if "formats" in wb.sheetnames:
            ws_formats = wb["formats"]
            format_headers = [_clean(cell.value) for cell in next(ws_formats.iter_rows(min_row=1, max_row=1))]
            idx = {header: i for i, header in enumerate(format_headers) if header}
            for row in ws_formats.iter_rows(min_row=2, values_only=True):
                values = list(row)
                record = {header: _clean(values[i]) if i < len(values) else "" for header, i in idx.items()}
                if record.get("Формат"):
                    record["active"] = _yes(record.get("Активен"), True)
                    formats.append(record)
        return {
            "products": products,
            "samples": samples,
            "names": names,
            "formats": formats,
            "samples_sheet_present": "samples" in wb.sheetnames,
            "names_sheet_present": "names" in wb.sheetnames,
            "formats_sheet_present": "formats" in wb.sheetnames,
            "warnings": warnings,
        }
    finally:
        wb.close()
        raw_wb.close()

def import_catalog_xlsx(
    db_path: str | Path,
    xlsx_path: str | Path,
    *,
    mode: str = "merge",
    sheet_name: str = "mainSheet",
    columns: dict[str, str] | None = None,
    source_filename: str | None = None,
) -> dict[str, Any]:
    parsed = read_catalog_xlsx(xlsx_path, sheet_name=sheet_name, columns=columns)
    display_filename = source_filename or Path(xlsx_path).name
    product_result = save_products_batch(db_path, parsed["products"], mode=mode, filename=display_filename)
    sample_result = save_samples_batch(
        db_path, parsed["samples"], mode=mode if parsed["samples_sheet_present"] else "merge"
    )
    name_result = save_names_batch(
        db_path, parsed["names"], mode=mode if parsed["names_sheet_present"] else "merge"
    )
    format_result = save_name_formats_batch(
        db_path, parsed["formats"], mode=mode if parsed["formats_sheet_present"] else "merge"
    )
    result = {
        "filename": display_filename,
        "mode": mode,
        "products_read": len(parsed["products"]),
        "products_written": product_result["written"],
        "samples_read": len(parsed["samples"]),
        "samples_written": sample_result["written"],
        "names_read": len(parsed["names"]),
        "names_written": name_result["written"],
        "formats_read": len(parsed["formats"]),
        "formats_written": format_result["written"],
        "warnings": parsed["warnings"],
    }
    record_import(db_path, **result)
    return result


def _style_workbook(wb: Workbook) -> None:
    dark = "1F2937"
    green = "15803D"
    border = Border(bottom=Side(style="thin", color="D1D5DB"))
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=green if ws.title == "mainSheet" else dark)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 34
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        for column_cells in ws.columns:
            letter = column_cells[0].column_letter
            max_len = max(len(_clean(cell.value)) for cell in column_cells[:300]) if column_cells else 10
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 38)


def _parse_raw_excel(item: dict[str, Any]) -> dict[str, Any]:
    raw = item.get("raw_json") or {}
    if isinstance(raw, dict):
        return dict(raw)
    try:
        parsed = json.loads(str(raw))
    except Exception:
        return {}
    return dict(parsed) if isinstance(parsed, dict) else {}


def _typed_export_value(internal: str, item: dict[str, Any]) -> Any:
    if internal == "active":
        return "Да" if int(item.get("active") or 0) else "Нет"
    if internal == "marking_profile":
        code = str(item.get(internal) or "AUTO").upper()
        return MARKING_PROFILE_EXPORT_LABELS.get(code, code)
    if internal == "certificate_type":
        code = str(item.get(internal) or "").upper()
        return CERTIFICATE_TYPE_EXPORT_LABELS.get(code, code)
    return item.get(internal, "")


def _restore_raw_cell(raw: dict[str, Any], header: str, fallback: Any) -> tuple[Any, str | None]:
    # If a row came from Excel, absence of a header/value means blank. We never
    # synthesize technical fields into the user's exported catalog.
    imported_row = RAW_HEADERS_KEY in raw or any(candidate in raw for candidate in DEFAULT_COLUMNS.values())
    if imported_row:
        if header not in raw:
            return "", None
        value = raw.get(header, "")
        meta = raw.get(RAW_META_KEY, {})
        spec = meta.get(header, {}) if isinstance(meta, dict) else {}
        cell_type = str(spec.get("type") or "")
        number_format = str(spec.get("number_format") or "") or None
        if value in (None, ""):
            return "", number_format
        try:
            if cell_type == "datetime":
                return datetime.fromisoformat(str(value)), number_format
            if cell_type == "date":
                return date.fromisoformat(str(value)), number_format
            if cell_type == "number":
                return value, number_format
            if cell_type == "bool":
                return bool(value), number_format
            if cell_type == "formula":
                formula = str(value)
                return formula if formula.startswith("=") else f"={formula}", number_format
        except (TypeError, ValueError):
            pass
        return value, number_format
    return fallback, None


def export_catalog_xlsx(db_path: str | Path, output_path: str | Path) -> Path:
    products = list_all_products(db_path, active_only=False)
    samples = list_samples(db_path, active_only=False)
    names = list_names(db_path, active_only=False)
    formats = list_name_formats(db_path, active_only=False)
    wb = Workbook()
    ws = wb.active
    ws.title = "mainSheet"
    headers = [*DEFAULT_COLUMNS.values(), *CALCULATED_COLUMNS]
    ws.append(headers)
    header_index = {header: idx + 1 for idx, header in enumerate(headers)}

    for row_number, item in enumerate(products, start=2):
        raw = _parse_raw_excel(item)
        for column_number, (internal, header) in enumerate(DEFAULT_COLUMNS.items(), start=1):
            fallback = _typed_export_value(internal, item)
            value, number_format = _restore_raw_cell(raw, header, fallback)
            cell = ws.cell(row=row_number, column=column_number, value=value)
            if number_format:
                cell.number_format = number_format

    sample_ws = wb.create_sheet("samples")
    sample_ws.append(["№", "Name", "WB code", "QR link", "Активен"])
    for item in samples:
        sample_ws.append([
            item.get("sample_no", ""), item.get("name", ""), item.get("wb_code", ""),
            item.get("qr_link", ""), "Да" if int(item.get("active") or 0) else "Нет",
        ])

    names_ws = wb.create_sheet("names")
    names_ws.append(["№", "Название", "Название мотива", "Активен"])
    for item in names:
        names_ws.append([
            item.get("name_no", ""), item.get("name", ""), item.get("inspiration_name", ""),
            "Да" if int(item.get("active") or 0) else "Нет",
        ])

    formats_ws = wb.create_sheet("formats")
    formats_ws.append(["№", "Формат", "Активен"])
    for item in formats:
        formats_ws.append([
            item.get("format_no", ""), item.get("name", ""),
            "Да" if int(item.get("active") or 0) else "Нет",
        ])

    guide_ws = wb.create_sheet("Справочники")
    guide_ws.append(["Профили маркировки", "Активность", "Тип документа"])
    profiles = ["Парфюмерия", "Дезодоранты и косметика", "Не требует маркировки", "Пользовательский профиль"]
    document_types = ["Декларация о соответствии", "Сертификат соответствия"]
    for i in range(max(len(profiles), len(document_types), 2)):
        guide_ws.append([
            profiles[i] if i < len(profiles) else "",
            ["Да", "Нет"][i] if i < 2 else "",
            document_types[i] if i < len(document_types) else "",
        ])
    guide_ws.sheet_state = "hidden"

    if ws.max_row >= 2:
        profile_col = get_column_letter(header_index[DEFAULT_COLUMNS["marking_profile"]])
        active_col = get_column_letter(header_index[DEFAULT_COLUMNS["active"]])
        certificate_type_col = get_column_letter(header_index[DEFAULT_COLUMNS["certificate_type"]])
        profile_validation = DataValidation(type="list", formula1="'Справочники'!$A$2:$A$5", allow_blank=False)
        yes_no_validation = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
        certificate_type_validation = DataValidation(type="list", formula1="'Справочники'!$C$2:$C$3", allow_blank=True)
        ws.add_data_validation(profile_validation)
        ws.add_data_validation(yes_no_validation)
        ws.add_data_validation(certificate_type_validation)
        profile_validation.add(f"{profile_col}2:{profile_col}{ws.max_row + 500}")
        yes_no_validation.add(f"{active_col}2:{active_col}{ws.max_row + 500}")
        certificate_type_validation.add(f"{certificate_type_col}2:{certificate_type_col}{ws.max_row + 500}")

    _style_workbook(wb)

    # 0.77.2 deliberately exports ordinary filtered ranges, not structured
    # Excel Tables. This avoids the tableN.xml repair errors seen in desktop Excel.
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    wb.close()
    return target
