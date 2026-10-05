"""Assortment acceptance checks use isolated sellers and XLSX files."""

import base64
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from fbe_flow.core.catalog_models import CatalogProduct, normalize_gtin
from fbe_flow.core.errors import Conflict, InvalidInput, NotFound
from fbe_flow.modules import catalog_xlsx
from fbe_flow.modules.catalog import check_source_binding

from .conftest import normalized_batch


def gtin(number=1):
    prefix = f"{number:013d}"
    check = (-sum(int(v) * (3 if i % 2 == 0 else 1) for i, v in enumerate(prefix))) % 10
    return prefix + str(check)


def data(number=1, **changes):
    return {
        "title": f"Товар {number}",
        "sku": f"{number:03d}",
        "gtins": [gtin(number)],
        "tnved": "3303001000",
        **changes,
    }


@pytest.fixture
def assortment(client):
    state = client.app.state
    seller = state.sellers.create("Ассортимент")["id"]
    other = state.sellers.create("Другой")["id"]
    return state, client, seller, other


def create(workspace, number=1, **changes):
    state, _, seller, _ = workspace
    return state.catalog.save_product(seller, data(number, **changes))


def book_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def edit(content, sheet, rows):
    book = load_workbook(io.BytesIO(content))
    keys = {c.value: c.column for c in book[sheet][2]}
    for row, fields in rows.items():
        for key, value in fields.items():
            if isinstance(value, list):
                value = "; ".join(value) if key in catalog_xlsx.LIST_FIELDS else json.dumps(value)
            book[sheet].cell(row, keys[key], value)
    return book_bytes(book)


def preview(workspace, content, **options):
    state, _, seller, _ = workspace
    return catalog_xlsx.preview(state.catalog, seller, content, **options)


def apply(workspace, plan):
    return workspace[0].catalog.apply_import(workspace[2], plan["id"], plan["digest"])


def test_identity_revision_archive_and_seller_scope(assortment):
    state, client, seller, other = assortment
    product = create(assortment)
    same = state.catalog.save_product(seller, data(), product["id"], product["revision"])
    assert same["id"] == product["id"] and same["revision"] == 1
    changed = state.catalog.save_product(seller, data(title="После"), product["id"], 1)
    assert changed["id"] == product["id"] and changed["revision"] == 2
    with pytest.raises(Conflict):
        state.catalog.save_product(seller, data(), product["id"], 1)
    with pytest.raises(NotFound):
        state.catalog.detail(other, product["id"])
    assert client.get(f"/api/sellers/{other}/catalog/products/{product['id']}").status_code == 404
    state.catalog.save_product(seller, data(title="После", archived=True), product["id"], 2)
    assert state.catalog.list(seller)["total"] == 0
    assert state.catalog.list(seller, archived=True)["total"] == 1


@pytest.mark.parametrize(
    "changes",
    [
        {"gtins": ["00000000000001"]},
        {"gtins": [123]},
        {"gross_weight_g": "10", "net_weight_g": "20"},
        {"length_mm": "-1"},
        {"length_mm": "Infinity"},
        {"title": "Bad\x01value"},
        {"tnved": "3303"},
        {"package_quantity": 1.2},
        {"attributes": {"key": float("nan")}},
    ],
)
def test_invalid_trade_items_are_rejected(changes):
    with pytest.raises(ValueError):
        CatalogProduct.model_validate(data(**changes))


def test_gtin_preserves_zeroes_and_uniqueness_per_seller(assortment):
    state, _, seller, other = assortment
    p = create(assortment)
    assert normalize_gtin(gtin()) == gtin()
    with pytest.raises(Conflict):
        state.catalog.save_product(seller, data(2, gtins=p["gtins"]))
    assert state.catalog.list(seller)["total"] == 1
    assert state.catalog.save_product(other, data())["gtins"] == p["gtins"]


def test_full_export_roundtrip_is_formula_safe_and_lossless(assortment):
    state, _, seller, _ = assortment
    p = create(
        assortment,
        title="=1+2",
        description="#CLEAR",
        length_mm="12.3456",
        attributes={"text": "#CLEAR", "null": None, "number": 1.25, "__proto__": {"x": 1}},
    )
    state.catalog.save_document(
        seller, {"kind": "declaration", "number": "ЕАЭС N 001", "product_ids": [p["id"]]}
    )
    state.catalog.save_batch(seller, {"product_id": p["id"], "name": "Партия 001", "quantity": 5})
    content = catalog_xlsx.export(state.catalog, seller)
    book = load_workbook(io.BytesIO(content), data_only=False)
    sheet = book["Товары"]
    keys = {c.value: c.column for c in sheet[2]}
    assert sheet.cell(3, keys["title"]).data_type == "s"
    assert sheet.cell(3, keys["gtins"]).value == gtin()
    assert sheet.cell(3, keys["gtins"]).number_format == "@"
    plan = preview(assortment, content)
    assert not plan["errors"] and plan["operations"] == []
    assert apply(assortment, plan)["state"] == "applied"


def test_native_template_creates_product_documents_attributes_and_batch(assortment):
    state, _, seller, _ = assortment
    content = catalog_xlsx.export(state.catalog, seller, empty=True)
    content = edit(content, "Товары", {3: data()})
    content = edit(
        content, "Характеристики", {3: {"product_id": "001", "field": "Цвет", "value": '"Синий"'}}
    )
    content = edit(content, "Документы", {3: {"kind": "declaration", "number": "ДоС-001"}})
    content = edit(
        content,
        "Применимость",
        {3: {"product_id": "001", "document_id": "ДоС-001", "linked": True}},
    )
    content = edit(content, "Партии", {3: {"product_id": "001", "name": "Партия", "quantity": 2}})
    plan = preview(assortment, content)
    assert not plan["errors"]
    apply(assortment, plan)
    p = state.catalog.list(seller)["items"][0]
    detail = state.catalog.detail(seller, p["id"])
    assert p["attributes"] == {"Цвет": "Синий"}
    assert detail["documents"][0]["number"] == "ДоС-001"
    assert detail["batches"][0]["quantity"] == 2
    assert preview(assortment, catalog_xlsx.export(state.catalog, seller))["operations"] == []


def test_blank_cells_preserve_explicit_clear_removes_and_import_is_single_use(assortment):
    state, _, seller, _ = assortment
    p = create(assortment, brand="Бренд", description="Старое")
    content = edit(
        catalog_xlsx.export(state.catalog, seller),
        "Товары",
        {3: {"brand": None, "description": "#CLEAR", "title": "Новое"}},
    )
    plan = preview(assortment, content)
    assert not plan["errors"]
    assert state.catalog.product(seller, p["id"])["title"] == "Товар 1"
    apply(assortment, plan)
    new = state.catalog.product(seller, p["id"])
    assert new["brand"] == "Бренд" and new["description"] is None and new["title"] == "Новое"
    with pytest.raises(Conflict):
        apply(assortment, plan)


def test_import_rolls_back_every_change_if_a_later_row_changed(assortment):
    state, _, seller, _ = assortment
    p, q = create(assortment), create(assortment, 2)
    content = edit(
        catalog_xlsx.export(state.catalog, seller),
        "Товары",
        {3: {"title": "Первый"}, 4: {"title": "Второй"}},
    )
    plan = preview(assortment, content)
    assert not plan["errors"]
    state.catalog.save_product(seller, data(2, title="Параллельно"), q["id"], q["revision"])
    with pytest.raises(Conflict):
        apply(assortment, plan)
    assert state.catalog.product(seller, p["id"])["title"] == p["title"]


@pytest.mark.parametrize("field,value", [("gtins", 123), ("title", "=1+2"), ("length_mm", -1)])
def test_cell_errors_block_whole_import_with_sheet_and_row(assortment, field, value):
    state, _, seller, _ = assortment
    p = create(assortment)
    content = edit(catalog_xlsx.export(state.catalog, seller), "Товары", {3: {field: value}})
    plan = preview(assortment, content)
    assert plan["errors"] and plan["errors"][0]["sheet"] == "Товары"
    assert plan["errors"][0]["row"] == 3
    with pytest.raises(Conflict):
        apply(assortment, plan)
    assert state.catalog.product(seller, p["id"])["revision"] == p["revision"]


def test_import_scope_cancel_digest_and_local_origin_boundary(assortment):
    state, client, seller, other = assortment
    create(assortment)
    content = catalog_xlsx.export(state.catalog, seller)
    with pytest.raises(InvalidInput):
        catalog_xlsx.preview(state.catalog, other, content)
    plan = preview(assortment, content)
    with pytest.raises(NotFound):
        state.catalog.import_plan(other, plan["id"])
    with pytest.raises(Conflict):
        state.catalog.apply_import(seller, plan["id"], "0" * 64)
    state.catalog.cancel_import(seller, plan["id"])
    with pytest.raises(Conflict):
        apply(assortment, plan)
    body = {"filename": "base.xlsx", "content": base64.b64encode(content).decode()}
    path = f"/api/sellers/{seller}/catalog/xlsx/preview"
    assert (
        client.post(path, json=body, headers={"Origin": "https://foreign.test"}).status_code == 403
    )
    assert client.post(path, json=body).status_code == 201


def external():
    book = Workbook()
    sheet = book.active
    sheet.title = "Товары WB"
    sheet.append(
        [
            "Артикул продавца",
            "Название товара",
            "Длина упаковки, см",
            "Вес с упаковкой, кг",
            "Баркод",
            "Номер ДоС",
            "Цвет",
        ]
    )
    sheet.append(["001", "Товар", "12.5", "0.25", gtin(), "ДоС-001", "Синий"])
    return book


def test_external_template_mapping_units_documents_and_unknown_attributes(assortment):
    state, _, seller, _ = assortment
    content = book_bytes(external())
    found = catalog_xlsx.inspect(content)["sheets"][0]
    assert found["mapping"][2]["scale"] == "10"
    plan = preview(
        assortment,
        content,
        profile="wb",
        sheet_name=found["name"],
        header_row=found["header_row"],
        mapping=found["mapping"],
    )
    assert not plan["errors"]
    apply(assortment, plan)
    p = state.catalog.list(seller)["items"][0]
    assert p["length_mm"] == "125.0" and p["gross_weight_g"] == "250.00"
    assert p["attributes"]["Цвет"] == "Синий"
    assert state.catalog.documents(seller, p["id"])[0]["number"] == "ДоС-001"


def test_fill_keeps_other_sheets_styles_and_validation_but_replaces_examples(assortment):
    state, _, seller, _ = assortment
    p = create(
        assortment,
        length_mm="125",
        gross_weight_g="250",
        barcodes=[gtin()],
        attributes={"Цвет": "Синий"},
    )
    book = external()
    sheet = book.active
    sheet["A2"].fill = PatternFill("solid", fgColor="FF0000")
    validation = DataValidation(type="list", formula1='"Синий,Белый"')
    sheet.add_data_validation(validation)
    validation.add("G2:G100")
    sheet.append(["obsolete", "Old example"])
    book.create_sheet("Справочник").append(["Сохранить"])
    content = book_bytes(book)
    found = catalog_xlsx.inspect(content)["sheets"][0]
    mapping = [v for v in found["mapping"] if v["field"] != "document_number"]
    output = catalog_xlsx.fill_template(
        state.catalog, seller, content, [p["id"]], found["name"], found["header_row"], mapping
    )
    result = load_workbook(io.BytesIO(output))
    assert result["Справочник"]["A1"].value == "Сохранить"
    assert result[found["name"]]["A2"].fill.fgColor.rgb == "00FF0000"
    assert result[found["name"]]["C2"].value == 12.5
    assert result[found["name"]]["D2"].value == 0.25
    assert result[found["name"]]["A3"].value is None
    assert len(result[found["name"]].data_validations.dataValidation) == 1


@pytest.mark.parametrize("part", ["xl/externalLinks/externalLink1.xml", "xl/vbaProject.bin"])
def test_external_links_and_macros_are_rejected(part):
    content = io.BytesIO(book_bytes(external()))
    with zipfile.ZipFile(content, "a") as archive:
        archive.writestr(part, "invalid")
    with pytest.raises(InvalidInput):
        catalog_xlsx.inspect(content.getvalue())


def test_entity_declarations_are_rejected_before_xml_parsing():
    original = zipfile.ZipFile(io.BytesIO(book_bytes(external())))
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        for name in original.namelist():
            value = original.read(name)
            if name == "xl/workbook.xml":
                value = b'<!DOCTYPE x [<!ENTITY bad SYSTEM "file:///etc/passwd">]>' + value
            archive.writestr(name, value)
    with pytest.raises(InvalidInput):
        catalog_xlsx.inspect(content.getvalue())


def rule(**changes):
    return {
        "title": "Проверенное правило",
        "product_group": "test-group",
        "marking_required": True,
        "tnved_prefixes": ["3303"],
        "valid_from": "2020-01-01",
        "source_url": "https://example.test/rules",
        "source_note": "Изолированный тест; не нормативное правило",
        **changes,
    }


def test_dated_rules_conditions_ambiguity_and_incomplete_inputs(assortment):
    state, _, seller, _ = assortment
    p = create(assortment, volume_ml="100")
    state.catalog.save_rule(
        seller, rule(conditions=[{"field": "volume_ml", "operator": "eq", "value": 100}])
    )
    assert state.catalog.classify(seller, p)["state"] == "matched"
    assert state.catalog.classify(seller, p, "2019-01-01")["state"] == "unknown"
    state.catalog.save_rule(seller, rule(title="Исключение", marking_required=False))
    assert state.catalog.classify(seller, p)["state"] == "ambiguous"
    state.catalog.save_rule(
        seller,
        rule(
            title="Нет данных",
            priority=10,
            conditions=[{"field": "composition", "operator": "eq", "value": "A"}],
        ),
    )
    assert state.catalog.classify(seller, p)["state"] == "incomplete"


def test_document_files_shared_applicability_scope_and_manual_verification(assortment):
    state, client, seller, other = assortment
    p, q = create(assortment), create(assortment, 2)
    with pytest.raises(InvalidInput):
        state.catalog.save_document(
            seller, {"kind": "declaration", "number": "001", "status": "verified"}
        )
    d = state.catalog.save_document(
        seller,
        {
            "kind": "declaration",
            "number": "001",
            "status": "verified",
            "verification_note": "Оператор проверил",
            "product_ids": [p["id"], q["id"]],
        },
    )
    attached = state.catalog.add_file(seller, d["id"], "document.pdf", b"%PDF-1.4\nfixture")
    assert state.catalog.file(seller, attached["id"])["content"].startswith(b"%PDF")
    with pytest.raises(NotFound):
        state.catalog.file(other, attached["id"])
    assert client.get(f"/api/sellers/{other}/catalog/files/{attached['id']}").status_code == 404
    plan = preview(assortment, catalog_xlsx.export(state.catalog, seller, product_ids=[p["id"]]))
    assert not plan["errors"] and plan["operations"] == []
    assert set(state.catalog.documents(seller)[0]["product_ids"]) == {p["id"], q["id"]}


def test_concurrent_edit_has_exactly_one_winner(assortment):
    state, _, seller, _ = assortment
    p = create(assortment)

    def change(title):
        try:
            state.catalog.save_product(seller, data(title=title), p["id"], p["revision"])
            return True
        except Conflict:
            return False

    with ThreadPoolExecutor(max_workers=2) as executor:
        assert sorted(executor.map(change, ["One", "Two"])) == [False, True]


def test_source_sync_never_overwrites_canonical_and_stable_link_survives(db, setup):
    first, second, connection, _, registry, _ = setup
    from fbe_flow.modules.catalog import Catalog
    from fbe_flow.modules.connections import Connections
    from fbe_flow.modules.records import Records

    catalog = Catalog(db, Connections(db, registry), registry, None)
    records = Records(db)
    records.apply(first, connection, normalized_batch("Источник"))
    source = records.list(first, "products")[0]
    p = catalog.adopt(first, data(), source["id"], "")
    records.apply(first, connection, normalized_batch("После синхронизации"))
    assert catalog.product(first, p["id"])["title"] == "Товар 1"
    assert catalog.links(first, p["id"])[0]["source_product_id"] == source["id"]
    assert records.list(first, "products")[0]["id"] == source["id"]
    with db.connection() as conn:
        check_source_binding(conn, first, source["id"], "", gtin(), "test-group")
        with pytest.raises(Conflict):
            check_source_binding(conn, first, source["id"], "", gtin(2), "test-group")
    assert catalog.list(second)["total"] == 0


def test_local_printing_and_application_do_not_change_chz_status(workspace):
    state, _, seller, connection, provider, *_ = workspace
    p = state.catalog.save_product(seller, data(product_group=provider.group))
    b = state.catalog.save_batch(seller, {"product_id": p["id"], "name": "Партия", "quantity": 1})
    code_id = str(uuid4())
    with state.database.connection() as conn:
        conn.execute(
            "INSERT INTO marking_codes(id,seller_id,connection_id,code,full_code,gtin,"
            "product_group,external_status) VALUES(?,?,?,?,?,?,?,?)",
            (
                code_id,
                seller,
                connection["id"],
                "01" + gtin() + "21SERIAL",
                "01" + gtin() + "21SERIAL\x1d91KEY\x1d92CRYPTO",
                gtin(),
                provider.group,
                "EMITTED",
            ),
        )
    calls = len(provider.calls)
    state.catalog.assign_codes(seller, b["id"], [code_id])
    state.catalog.record_code_event(seller, b["id"], [code_id], "printed", "Оператор")
    state.catalog.record_code_event(seller, b["id"], [code_id], "applied", "Оператор")
    codes = state.catalog.batch_codes(seller, b["id"])["items"]
    assert codes[0]["external_status"] == "EMITTED"
    assert {e["kind"] for e in codes[0]["local_events"]} == {"code.assigned", "printed", "applied"}
    assert len(provider.calls) == calls
    with pytest.raises(Conflict):
        state.catalog.assign_codes(seller, b["id"], [code_id])
    with pytest.raises(Conflict):
        state.catalog.save_product(seller, data(gtins=[gtin(2)]), p["id"], p["revision"])


def test_catalog_api_page_assets_and_empty_export_exist(assortment):
    state, client, seller, _ = assortment
    page = client.get(f"/sellers/{seller}/catalog")
    assert page.status_code == 200 and "Единый ассортимент" in page.text
    assert "/static/catalog.js" in page.text and "Ассортимент" in page.text
    assert client.get("/static/catalog.js").status_code == 200
    assert client.get(f"/api/sellers/{seller}/catalog/xlsx/template").status_code == 200
    with state.database.connection() as conn:
        assert conn.execute("PRAGMA foreign_key_check").fetchall() == []
