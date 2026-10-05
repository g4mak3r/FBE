"""Bounded XLSX product-runtime exchange; formulas cannot authorize external actions."""
import hashlib
import io
import json
import re
import zipfile
from copy import copy, deepcopy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from fbe_flow.core.catalog_models import (
    CatalogBatch, CatalogDocument, CatalogProduct, ClassificationRule,
)
from fbe_flow.core.errors import Conflict, InvalidInput
from fbe_flow.modules.catalog import clean_product, encode, validate

FORMAT = "fbe-assortment-1"
MAX_ROWS, MAX_COLUMNS, MAX_CELLS = 100002, 250, 5000000
MAX_BYTES = 32 * 1024 * 1024
CLEAR = "#CLEAR"
LABELS = {
    "id": "ID FBE", "revision": "Версия записи", "title": "Наименование", "sku": "Артикул FBE",
    "family": "Семейство / модель", "brand": "Бренд", "manufacturer": "Производитель",
    "country": "Страна производства", "description": "Описание", "composition": "Состав",
    "volume_ml": "Объём, мл", "net_weight_g": "Масса нетто, г", "gross_weight_g": "Масса брутто, г",
    "length_mm": "Длина упаковки, мм", "width_mm": "Ширина упаковки, мм",
    "height_mm": "Высота упаковки, мм", "package_quantity": "Количество в упаковке, шт",
    "shelf_life_days": "Срок годности, дней", "tnved": "ТН ВЭД", "okpd2": "ОКПД2", "gtins": "GTIN",
    "barcodes": "Штрихкоды", "product_group": "Группа ЧЗ",
    "marking_attestation": "Заявление о нанесении для карточки", "archived": "В архиве",
    "kind": "Тип документа", "number": "Номер документа", "issued_on": "Дата выдачи",
    "expires_on": "Действует до", "issuer": "Кем выдан", "scope": "Область применения",
    "status": "Статус проверки", "verification_note": "Основание проверки",
    "document_id": "ID документа или номер", "product_id": "ID товара или артикул FBE",
    "linked": "Связать с документом", "source_product_id": "ID исходной карточки FBE",
    "variant": "Вариант площадки", "overrides": "Данные площадки (JSON)", "name": "Название партии",
    "quantity": "Количество, шт", "manufactured_on": "Дата производства", "note": "Примечание",
    "field": "Характеристика", "value": "Значение (JSON)", "tnved_prefixes": "Префиксы ТН ВЭД",
    "okpd2_prefixes": "Префиксы ОКПД2", "conditions": "Дополнительные условия (JSON)",
    "valid_from": "Начало действия", "valid_until": "Окончание действия",
    "source_url": "Источник правила", "source_note": "Основание правила", "priority": "Приоритет",
    "enabled": "Включено", "marking_required": "Маркировка обязательна",
    "document_number": "Номер ДоС / документа (из карточки)",
    "document_kind": "Тип документа (из карточки)",
}
PRODUCT_KEYS = ["id", "revision", *[k for k in CatalogProduct.model_fields if k != "attributes"]]
SHEETS = {
    "Товары": ("product", PRODUCT_KEYS),
    "Характеристики": ("attributes", ["product_id", "field", "value"]),
    "Документы": ("document", ["id", "revision", *[
        k for k in CatalogDocument.model_fields if k != "product_ids"]]),
    "Применимость": ("document_products", ["document_id", "product_id", "linked"]),
    "Связи": ("link", ["product_id", "source_product_id", "variant", "overrides"]),
    "Партии": ("batch", ["id", "revision", *CatalogBatch.model_fields]),
    "Правила": ("rule", ["id", "revision", *ClassificationRule.model_fields]),
}
TEXT_FIELDS = {"id", "sku", "tnved", "okpd2", "product_id", "document_id", "source_product_id",
               "variant", "gtins", "barcodes", "document_number"}
MEASURES = {"volume_ml", "net_weight_g", "gross_weight_g", "length_mm", "width_mm", "height_mm"}
INT_FIELDS = {"revision", "package_quantity", "shelf_life_days", "quantity", "priority"}
BOOL_FIELDS = {"marking_attestation", "archived", "enabled", "marking_required", "linked"}
LIST_FIELDS = {"gtins", "barcodes", "tnved_prefixes", "okpd2_prefixes"}
JSON_FIELDS = {"attributes", "conditions", "overrides", "value"}
DATE_FIELDS = {"issued_on", "expires_on", "manufactured_on", "valid_from", "valid_until"}
MODELS = {"product": CatalogProduct, "document": CatalogDocument, "batch": CatalogBatch,
          "rule": ClassificationRule}

def normal(value):
    return re.sub(r"[^a-zа-я0-9]+", "", str(value).lower().replace("ё", "е"))

ALIASES = {normal(v): (k, "1") for k, v in LABELS.items()}
ALIASES.update({normal(k): (k, "1") for k in LABELS})
for target, names in {
    "sku": ["Артикул продавца", "Артикул", "offer_id", "vendorCode"],
    "title": ["Название товара", "Название", "Наименование товара", "name"],
    "tnved": ["ТНВЭД код", "ТНВЭД", "Код ТН ВЭД ЕАЭС", "Код ТНВЭД"],
    "gtins": ["Код GTIN", "Код товара GTIN"],
    "barcodes": ["Баркод", "Баркоды", "Штрихкод", "barcode"],
    "marking_attestation": ["kizMarked", "Подтверждение нанесения маркировки", "КИЗ нанесен"],
    "gross_weight_g": ["Вес с упаковкой, г", "Вес товара с упаковкой (г)", "Вес в упаковке, г"],
    "volume_ml": ["Объем товара, мл", "Объем", "Объем, мл"],
    "document_number": ["Номер декларации соответствия", "Номер ДоС", "Декларация соответствия",
                        "Номер сертификата"],
}.items():
    for label in names:
        ALIASES[normal(label)] = (target, "1")
for target, names in {
    "length_mm": ["Длина упаковки, см", "Длина упаковки", "Длина (см)"],
    "width_mm": ["Ширина упаковки, см", "Ширина упаковки", "Ширина (см)"],
    "height_mm": ["Высота упаковки, см", "Высота упаковки", "Высота (см)"],
    "gross_weight_g": ["Вес с упаковкой, кг", "Вес товара с упаковкой (кг)"],
}.items():
    for label in names:
        ALIASES[normal(label)] = (target, "1000" if target == "gross_weight_g" else "10")

def safe_workbook(content, read_only=True):
    if not isinstance(content, bytes) or not 1 <= len(content) <= MAX_BYTES:
        raise InvalidInput("XLSX: до 32 МБ")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            if len(names) != len(set(names)) or len(names) > 20000:
                raise ValueError("Повторяющиеся или избыточные части файла")
            total = 0
            for info in archive.infolist():
                lower = info.filename.lower()
                if any(key in lower for key in (
                    "vbaproject", "externallinks/", "embeddings/", "connections.xml",
                )):
                    raise ValueError("Макросы, встроенные объекты и внешние связи не поддерживаются")
                total += info.file_size
                if total > 256 * 1024 * 1024 or info.file_size > 128 * 1024 * 1024:
                    raise ValueError("Слишком большой распакованный файл")
                if lower.endswith((".xml", ".rels")):
                    with archive.open(info) as stream:
                        tail = b""
                        while block := stream.read(65536):
                            data = (tail + block).upper()
                            if b"<!DOCTYPE" in data or b"<!ENTITY" in data:
                                raise ValueError("Определения XML-сущностей не поддерживаются")
                            tail = data[-20:]
        book = load_workbook(io.BytesIO(content), read_only=read_only,
                             data_only=False, keep_links=False)
        cells = 0
        for sheet in book:
            if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
                book.close()
                raise ValueError("Слишком много строк или колонок")
            cells += sheet.max_row * sheet.max_column
        if cells > MAX_CELLS:
            book.close()
            raise ValueError("Слишком много ячеек")
        return book
    except (ValueError, KeyError, OSError, zipfile.BadZipFile, OverflowError) as exc:
        raise InvalidInput(f"Не удалось прочитать XLSX: {exc}") from exc
    except Exception as exc:
        raise InvalidInput("Повреждённый или неподдерживаемый XLSX-файл") from exc

def scalar(cell, field, scale="1"):
    if cell.data_type == "f":
        raise InvalidInput("Замените формулу её значением")
    value = cell.value
    if value is None or value == "":
        return None, False
    if value == CLEAR:
        return None, True
    if isinstance(value, str) and re.fullmatch(r"#{2,}CLEAR", value):
        value = value[1:]
    if field in TEXT_FIELDS and not isinstance(value, str):
        raise InvalidInput("Идентификатор должен быть текстом для сохранения всех цифр")
    if field in JSON_FIELDS:
        try:
            return json.loads(value) if isinstance(value, str) else value, True
        except (ValueError, TypeError) as exc:
            raise InvalidInput("Неверный JSON; текст задаётся в двойных кавычках") from exc
    if field in LIST_FIELDS:
        if not isinstance(value, str):
            raise InvalidInput("Список должен быть текстом")
        values = json.loads(value) if value.startswith("[") else re.split(r"[;\n]+", value)
        if not isinstance(values, list) or any(not isinstance(v, str) for v in values):
            raise InvalidInput("Ожидается список текстовых значений")
        return [v.strip() for v in values if v.strip()], True
    if field in BOOL_FIELDS:
        if type(value) is bool:
            return value, True
        text = str(value).strip().lower()
        if text in {"да", "true", "1"}:
            return True, True
        if text in {"нет", "false", "0"}:
            return False, True
        raise InvalidInput("Допустимо Да или Нет")
    if field in MEASURES or field in INT_FIELDS:
        if type(value) is bool:
            raise InvalidInput("Ожидается число")
        try:
            result = Decimal(str(value).replace(",", ".")) * Decimal(scale)
            if not result.is_finite():
                raise ValueError
            if field in INT_FIELDS:
                if result != result.to_integral_value():
                    raise ValueError
                return int(result), True
            return str(result), True
        except (ValueError, InvalidOperation) as exc:
            raise InvalidInput("Неверное числовое значение") from exc
    if field in DATE_FIELDS:
        if isinstance(value, datetime):
            return value.date().isoformat(), True
        if isinstance(value, date):
            return value.isoformat(), True
        if isinstance(value, str):
            for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
                try:
                    return datetime.strptime(value, fmt).date().isoformat(), True
                except ValueError:
                    pass
        raise InvalidInput("Дата: дата Excel или текст ГГГГ-ММ-ДД")
    value = str(value)
    if ILLEGAL_CHARACTERS_RE.search(value):
        raise InvalidInput("Недопустимые управляющие символы")
    return value, True

def set_cell(cell, value, key=None, escape_clear=True):
    if key == "value":
        value = encode(value)
    if value is None:
        return
    if isinstance(value, (dict, list)):
        value = "; ".join(value) if key in LIST_FIELDS else encode(value)
    if key in MEASURES:
        value = float(Decimal(str(value)))
    if key in DATE_FIELDS and value:
        value = date.fromisoformat(value)
        cell.number_format = "dd.mm.yyyy"
    if isinstance(value, str):
        if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
            raise InvalidInput("Значение нельзя без потерь поместить в ячейку XLSX")
        cell.value = "#" + value if escape_clear and re.fullmatch(r"#+CLEAR", value) else value
        cell.data_type = "s"
        if key in TEXT_FIELDS:
            cell.number_format = "@"
    else:
        cell.value = value

def sheet_table(book, name, keys, rows, blank=False):
    if len(rows) > MAX_ROWS - 2:
        raise InvalidInput("Лист превышает 100000 строк. Выберите часть ассортимента")
    sheet = book.create_sheet(name)
    sheet.freeze_panes = "C3"
    sheet.sheet_view.showGridLines = False
    for index, key in enumerate(keys, 1):
        header = sheet.cell(1, index, LABELS.get(key, key))
        header.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header.fill = PatternFill("solid", fgColor="244B5A")
        header.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.cell(2, index, key)
        sheet.column_dimensions[get_column_letter(index)].width = (
            38 if key in {"title", "description", "scope", "source_note"} else 22
        )
        header.comment = Comment(
            "Пустая ячейка сохраняет значение. #CLEAR очищает поле. "
            "При обновлении сохраняйте ID и версию.", "FBE Flow",
        )
    sheet.row_dimensions[1].height = 42
    sheet.row_dimensions[2].hidden = True
    for row_index, row in enumerate(rows, 3):
        for col_index, key in enumerate(keys, 1):
            set_cell(sheet.cell(row_index, col_index), row.get(key), key)
    if blank:
        for row_index in range(3, 103):
            for col_index, key in enumerate(keys, 1):
                if key in TEXT_FIELDS:
                    sheet.cell(row_index, col_index).number_format = "@"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{max(sheet.max_row, 3)}"
    for index, key in enumerate(keys, 1):
        if key in BOOL_FIELDS:
            validation = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
            validation.showErrorMessage = True
            validation.errorTitle, validation.error = "Неверное значение", "Выберите Да или Нет"
            sheet.add_data_validation(validation)
            validation.add(f"{get_column_letter(index)}3:{get_column_letter(index)}"
                           f"{max(sheet.max_row, 1000)}")
    return sheet

def export(catalog, seller, empty=False, product_ids=None):
    data = ({"products": [], "documents": [], "links": [], "batches": [], "rules": []}
            if empty else catalog.export_data(seller))
    if product_ids is not None:
        wanted = set(product_ids)
        if (len(wanted) != len(product_ids) or
                not wanted.issubset({v["id"] for v in data["products"]})):
            raise InvalidInput("В выгрузке выбраны неизвестные товары")
        data["products"] = [v for v in data["products"] if v["id"] in wanted]
        data["documents"] = [v for v in data["documents"]
                             if set(v["product_ids"]).intersection(wanted)]
        data["links"] = [v for v in data["links"] if v["product_id"] in wanted]
        data["batches"] = [v for v in data["batches"] if v["product_id"] in wanted]
    book = Workbook()
    book.remove(book.active)
    instructions = book.create_sheet("Как заполнить")
    instructions.append(["FBE Flow — единый ассортимент"])
    for text in [
        "Товары: одна строка на товарный вариант. GTIN, артикулы и ТН ВЭД храните текстом.",
        "Пустые ячейки сохраняют значения. #CLEAR удаляет поле; ##CLEAR задаёт текст #CLEAR.",
        "При обновлении сохраните ID и версию записи. Перед применением проверьте предпросмотр.",
        "В применимости и партиях укажите ID товара или артикул FBE; документ — ID или номер.",
        "Файлы документов находятся в программе и в резервной копии SQLite, а не внутри XLSX.",
        "Правила задаются по проверенному источнику, сроку и исключениям. Примеры не нормативны.",
        "Заявление для карточки отдельно от фактического нанесения и принятия документа ЧЗ.",
        "Статусы маркировки справочные. Изменение в Excel не меняет внешние статусы ЧЗ.",
        "Удаление строки ничего не удаляет в программе. Архив — Да; убрать применимость — Нет.",
    ]:
        instructions.append([text])
    instructions.column_dimensions["A"].width = 110
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[row[0].row].height = 34
    instructions["A1"].font = Font(size=18, bold=True, color="244B5A")
    exported = {p["id"] for p in data["products"]}
    rows = {
        "Товары": data["products"],
        "Характеристики": [{"product_id": p["id"], "field": k, "value": v}
                           for p in data["products"] for k, v in p["attributes"].items()],
        "Документы": data["documents"],
        "Применимость": [{"document_id": d["id"], "product_id": p, "linked": True}
                         for d in data["documents"] for p in d["product_ids"] if p in exported],
        "Связи": data["links"], "Партии": data["batches"], "Правила": data["rules"],
    }
    for name, (_, keys) in SHEETS.items():
        sheet_table(book, name, keys, rows[name], empty)
    statuses = []
    if not empty:
        for batch in data["batches"]:
            for offset in range(0, batch["assigned"], 200):
                for code in catalog.batch_codes(seller, batch["id"], offset, 200)["items"]:
                    statuses.append({"batch_id": batch["id"], "code_id": code["id"],
                                     "gtin": code["gtin"], "external_status": code["external_status"],
                                     "operator_events": [
                                         {"kind": e["kind"], "created_at": e["created_at"]}
                                         for e in code["local_events"]
                                     ]})
    sheet_table(book, "Статусы маркировки",
                ["batch_id", "code_id", "gtin", "external_status", "operator_events"], statuses)
    meta = book.create_sheet("_FBE")
    meta.append(["format", FORMAT])
    meta.append(["seller_id", "" if empty else seller])
    meta.append(["exported_at", datetime.now().isoformat(timespec="seconds")])
    meta.sheet_state = "veryHidden"
    stream = io.BytesIO()
    book.save(stream)
    if len(stream.getvalue()) > MAX_BYTES:
        raise InvalidInput("Выгрузка превышает 32 МБ. Выберите часть ассортимента")
    return stream.getvalue()

def inspect(content):
    book = safe_workbook(content)
    try:
        native, sheets = "_FBE" in book.sheetnames, []
        for sheet in book:
            if sheet.title in {"_FBE", "Как заполнить", "Статусы маркировки"}:
                continue
            header_row, score, headers = 1, -1, []
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30)):
                current = sum(normal(c.value) in ALIASES for c in row if c.value is not None)
                if current > score:
                    header_row, score = row[0].row, current
                    headers = [str(c.value or "") for c in row]
            if native and sheet.title in SHEETS:
                header_row = 2
                headers = [str(c.value or "") for c in next(
                    sheet.iter_rows(min_row=2, max_row=2))]
            mapping = []
            for index, header in enumerate(headers, 1):
                field, scale = ALIASES.get(
                    normal(header), ("attributes." + header[:220] if header else "", "1"))
                if native and header in SHEETS.get(sheet.title, (None, []))[1]:
                    field = header
                mapping.append({"column": index, "header": header, "field": field, "scale": scale})
            sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                           "header_row": header_row, "mapping": mapping})
        keys = [k for k in PRODUCT_KEYS if k not in {"id", "revision"}]
        return {"native": native, "sheets": sheets,
                "fields": [{"key": k, "label": LABELS.get(k, k)}
                           for k in [*keys, "document_number", "document_kind"]],
                "digest": hashlib.sha256(content).hexdigest()}
    finally:
        book.close()

def read_rows(sheet, start_row, mapping, warnings):
    if not 2 <= start_row <= max(100, sheet.max_row + 1):
        raise InvalidInput("Неверная первая строка данных")
    targets = [v["field"] for v in mapping if v.get("field")]
    columns = [v["column"] for v in mapping]
    if len(targets) != len(set(targets)):
        raise InvalidInput("Несколько колонок сопоставлены одному полю")
    if len(columns) != len(set(columns)):
        raise InvalidInput("Колонки сопоставления повторяются")
    if any(not 1 <= v <= sheet.max_column for v in columns):
        raise InvalidInput("Колонка находится за границей листа")
    result, errors, ignored = [], [], set()
    for row in sheet.iter_rows(min_row=start_row):
        if not any(v.value is not None for v in row):
            continue
        record, supplied = {}, set()
        for item in mapping:
            field, cell = item.get("field"), row[item["column"] - 1]
            if not field:
                if cell.value is not None:
                    ignored.add(item["column"])
                continue
            try:
                value, present = scalar(
                    cell, "value_text" if field.startswith("attributes.") else field,
                    item.get("scale", "1"))
                if present:
                    record[field] = value
                    supplied.add(field)
                    if cell.value == CLEAR:
                        supplied.add("!clear:" + field)
            except (InvalidInput, ValueError, TypeError) as exc:
                errors.append({"sheet": sheet.title, "row": row[0].row,
                               "column": item["column"], "field": field, "message": str(exc)})
        if record:
            result.append((row[0].row, record, supplied))
    if ignored:
        warnings.append({"sheet": sheet.title,
                         "message": "Заполненные колонки пропущены: " +
                         ", ".join(map(str, sorted(ignored)))})
    return result, errors

def preview(catalog, seller, content, profile="fbe", sheet_name=None, header_row=None,
            mapping=None, data_row=None):
    if profile not in {"fbe", "wb", "ozon", "custom"}:
        raise InvalidInput("Неизвестный формат импорта")
    book = safe_workbook(content)
    errors, warnings, operations, deferred = [], [], [], []
    old = catalog.export_data(seller)
    originals = {
        "product": {v["id"]: v for v in old["products"]},
        "document": {v["id"]: v for v in old["documents"]},
        "batch": {v["id"]: v for v in old["batches"]},
        "rule": {v["id"]: v for v in old["rules"]},
    }
    by_sku = {v["sku"]: v["id"] for v in old["products"]}
    by_gtin = {g: v["id"] for v in old["products"] for g in v["gtins"]}
    by_number, doc_numbers = {}, {}
    for doc in old["documents"]:
        by_number.setdefault(doc["number"], []).append(doc["id"])
        doc_numbers.setdefault((doc["kind"], doc["number"]), []).append(doc["id"])
    existing_links = {(v["source_product_id"], v["variant"]): v for v in old["links"]}
    product_ops, doc_ops = {}, {}
    seen_ids, seen_skus, seen_gtins = set(), set(), {}
    native = profile == "fbe"
    def error(sheet, row, exc):
        errors.append({"sheet": sheet, "row": row, "column": None, "message": str(exc)})
    def operation(kind, entity_id, before, data, sheet, row, revision=None):
        return {"kind": kind, "id": entity_id, "new": before is None, "revision": revision,
                "data": data, "sheet": sheet, "row": row, "changes": {}}
    def resolve_product(reference):
        product_id = by_sku.get(reference, reference)
        if product_id not in originals["product"] and product_id not in product_ops:
            raise InvalidInput("Товар не найден по ID или артикулу FBE")
        return product_id
    try:
        if native:
            if "_FBE" not in book.sheetnames:
                raise InvalidInput("Это не шаблон FBE. Выберите WB, Ozon или другой формат")
            meta = dict(book["_FBE"].iter_rows(values_only=True))
            if meta.get("format") != FORMAT or meta.get("seller_id") not in {"", seller}:
                raise InvalidInput("Версия шаблона или продавец не соответствует рабочей области")
            selected = [book[name] for name in SHEETS if name in book.sheetnames]
        else:
            if sheet_name not in book.sheetnames or not mapping:
                raise InvalidInput("Выберите лист и сопоставьте колонки")
            if data_row and data_row <= (header_row or 1):
                raise InvalidInput("Первая строка данных должна идти после заголовков")
            selected = [book[sheet_name]]
        for sheet in selected:
            kind, keys = SHEETS[sheet.title] if native else ("product", PRODUCT_KEYS)
            row_mapping = ([{"column": i + 1, "field": str(c.value or ""), "scale": "1"}
                            for i, c in enumerate(next(sheet.iter_rows(min_row=2, max_row=2)))]
                           if native else mapping)
            allowed = set(keys) if native else set(PRODUCT_KEYS) | {
                "attributes", "document_number", "document_kind"
            }
            if any(v.get("field") and v["field"] not in allowed and
                   not (not native and v["field"].startswith("attributes.")) for v in row_mapping):
                raise InvalidInput("Сопоставление содержит неизвестное поле")
            records, row_errors = read_rows(
                sheet, 3 if native else data_row or (header_row or 1) + 1, row_mapping, warnings)
            errors.extend(row_errors)
            if kind in {"attributes", "document_products", "link"}:
                deferred.extend((kind, sheet.title, *record) for record in records)
                continue
            for row, values, _supplied in records:
                try:
                    entity_id, revision = values.pop("id", None), values.pop("revision", None)
                    before = originals[kind].get(entity_id) if entity_id else None
                    if entity_id and not before:
                        raise InvalidInput("ID не найден у выбранного продавца")
                    if before and revision != before["revision"]:
                        raise Conflict("Версия устарела. Скачайте свежую выгрузку")
                    if entity_id and (kind, entity_id) in seen_ids:
                        raise InvalidInput("ID повторяется в файле")
                    model = MODELS[kind]
                    data = deepcopy({k: before[k] for k in model.model_fields
                                     if before and k in before})
                    extras = {k[11:]: v for k, v in values.items() if k.startswith("attributes.")}
                    values = {k: v for k, v in values.items() if not k.startswith("attributes.")}
                    doc_number = values.pop("document_number", None)
                    doc_kind = values.pop("document_kind", None) or "declaration"
                    if kind == "product":
                        data["attributes"] = {**data.get("attributes", {}), **extras}
                        if "attributes" in values:
                            data["attributes"] = values.pop("attributes") or {}
                    if kind == "batch":
                        values["product_id"] = resolve_product(
                            values.get("product_id", data.get("product_id")))
                    for key, value in values.items():
                        data[key] = [] if value is None and key in LIST_FIELDS else value
                    data = validate(model, data)
                    entity_id = entity_id or str(uuid4())
                    seen_ids.add((kind, entity_id))
                    if kind == "product":
                        if data["sku"] in seen_skus:
                            raise Conflict("Артикул повторяется; один вариант на строку")
                        if by_sku.get(data["sku"]) not in {None, entity_id}:
                            raise Conflict("Артикул существует. Для обновления используйте ID и версию")
                        for code in data["gtins"]:
                            if seen_gtins.get(code) not in {None, entity_id}:
                                raise Conflict("GTIN повторяется у разных товаров файла")
                            if by_gtin.get(code) not in {None, entity_id}:
                                raise Conflict("GTIN уже существует; свяжите с существующим товаром")
                            seen_gtins[code] = entity_id
                        seen_skus.add(data["sku"])
                        by_sku[data["sku"]] = entity_id
                    op = operation(kind, entity_id, before, data, sheet.title, row, revision)
                    operations.append(op)
                    if kind == "product":
                        product_ops[entity_id] = op
                    if kind == "document":
                        doc_ops[entity_id] = op
                        if entity_id not in by_number.get(data["number"], []):
                            by_number.setdefault(data["number"], []).append(entity_id)
                    if doc_number and kind == "product":
                        candidates = doc_numbers.get((doc_kind, doc_number), [])
                        if len(candidates) > 1:
                            raise Conflict("Номер документа неоднозначен. Внесите применимость по ID")
                        document_id = candidates[0] if candidates else str(uuid4())
                        doc_op = doc_ops.get(document_id)
                        if not doc_op:
                            prior = originals["document"].get(document_id)
                            doc_data = ({k: deepcopy(prior[k]) for k in CatalogDocument.model_fields}
                                        if prior else {"kind": doc_kind, "number": doc_number,
                                                       "product_ids": []})
                            doc_op = operation("document", document_id, prior,
                                               validate(CatalogDocument, doc_data), sheet.title, row,
                                               prior["revision"] if prior else None)
                            operations.append(doc_op)
                            doc_ops[document_id] = doc_op
                            doc_numbers[(doc_kind, doc_number)] = [document_id]
                        if entity_id not in doc_op["data"]["product_ids"]:
                            doc_op["data"]["product_ids"].append(entity_id)
                except (InvalidInput, Conflict, TypeError, ValueError) as exc:
                    error(sheet.title, row, exc)
        seen_attrs, seen_associations, seen_links = set(), set(), set()
        for kind, sheet, row, values, supplied in deferred:
            try:
                product_id = resolve_product(values.get("product_id"))
                if kind == "attributes":
                    field = values.get("field")
                    if not field:
                        raise InvalidInput("Не указано имя характеристики")
                    if (product_id, field) in seen_attrs:
                        raise InvalidInput("Характеристика повторяется в файле")
                    seen_attrs.add((product_id, field))
                    op = product_ops.get(product_id)
                    if not op:
                        raise InvalidInput("Сохраните строку товара с ID и версией в листе Товары")
                    if "value" not in values:
                        continue
                    if "!clear:value" in supplied:
                        op["data"]["attributes"].pop(field, None)
                    else:
                        op["data"]["attributes"][field] = values["value"]
                    op["data"] = validate(CatalogProduct, op["data"])
                elif kind == "document_products":
                    document_id = values.get("document_id")
                    if document_id not in doc_ops and document_id not in originals["document"]:
                        candidates = by_number.get(document_id, [])
                        if len(candidates) != 1:
                            raise InvalidInput("Документ не найден однозначно по ID или номеру")
                        document_id = candidates[0]
                    op = doc_ops.get(document_id)
                    if not op:
                        raise InvalidInput("Сохраните строку документа с ID и версией")
                    if (document_id, product_id) in seen_associations:
                        raise InvalidInput("Применимость повторяется в файле")
                    seen_associations.add((document_id, product_id))
                    ids = op["data"]["product_ids"]
                    if values.get("linked", True) and product_id not in ids:
                        ids.append(product_id)
                    elif not values.get("linked", True) and product_id in ids:
                        ids.remove(product_id)
                else:
                    source_id = values.get("source_product_id")
                    if not source_id:
                        raise InvalidInput("Не указан ID исходной карточки")
                    data = {"product_id": product_id, "source_product_id": source_id,
                            "variant": values.get("variant") or "",
                            "overrides": values.get("overrides") or {}}
                    key = (source_id, data["variant"])
                    if key in seen_links:
                        raise InvalidInput("Исходный вариант повторяется в файле")
                    seen_links.add(key)
                    prior = existing_links.get(key)
                    if prior and all(prior[k] == v for k, v in data.items()):
                        continue
                    operations.append({"kind": "link", "data": data, "before_link": prior,
                                       "sheet": sheet, "row": row,
                                       "changes": {"link": {"before": prior, "after": data}}})
            except (InvalidInput, Conflict, ValueError, TypeError) as exc:
                error(sheet, row, exc)
        operations.sort(key=lambda v: {"product": 0, "document": 1, "batch": 2,
                                       "rule": 3, "link": 4}[v["kind"]])
        for op in operations:
            if op["kind"] != "link":
                before = originals[op["kind"]].get(op["id"])
                op["changes"] = {k: {"before": before.get(k) if before else None, "after": v}
                                 for k, v in op["data"].items() if not before or before.get(k) != v}
        operations = [v for v in operations if v.get("new") or v["changes"]]
        if not errors:
            errors.extend(catalog.validate_import_operations(seller, operations))
        return catalog.prepare_import(seller, operations, errors, warnings,
                                      hashlib.sha256(content).hexdigest())
    finally:
        book.close()

def fill_template(catalog, seller, content, product_ids, sheet_name, header_row, mapping,
                  data_row=None):
    book = safe_workbook(content, read_only=False)
    try:
        if sheet_name not in book.sheetnames or not mapping or not 1 <= header_row <= 100:
            raise InvalidInput("Выберите лист шаблона и сопоставьте колонки")
        if not product_ids or len(product_ids) != len(set(product_ids)) or len(product_ids) > 50000:
            raise InvalidInput("Выберите от 1 до 50000 разных товаров")
        products = [catalog.product(seller, p) for p in product_ids]
        sheet, start = book[sheet_name], data_row or header_row + 1
        if header_row > sheet.max_row or start <= header_row or start > 100:
            raise InvalidInput("Укажите строку данных после заголовков")
        fields = [v["field"] for v in mapping if v.get("field")]
        columns = [v["column"] for v in mapping]
        if len(fields) != len(set(fields)) or len(columns) != len(set(columns)):
            raise InvalidInput("Поля или колонки шаблона повторяются")
        if any(not 1 <= v <= sheet.max_column for v in columns):
            raise InvalidInput("Колонка за границей шаблона")
        styles = {v: copy(sheet.cell(start, v)._style) for v in columns}
        for row in sheet.iter_rows(min_row=start):
            for cell in row:
                cell.value = None
        for index, product in enumerate(products, start):
            for item in mapping:
                field = item.get("field")
                if not field:
                    continue
                if field not in CatalogProduct.model_fields and not field.startswith("attributes."):
                    raise InvalidInput("Неизвестное поле шаблона")
                value = (product["attributes"].get(field[11:]) if field.startswith("attributes.")
                         else product[field])
                if field in MEASURES and value is not None:
                    scale = Decimal(item.get("scale", "1"))
                    if not scale.is_finite() or scale <= 0:
                        raise InvalidInput("Неверный перевод единиц")
                    value = str(Decimal(value) / scale)
                cell = sheet.cell(index, item["column"])
                cell._style = copy(styles[item["column"]])
                set_cell(cell, value, field, escape_clear=False)
        stream = io.BytesIO()
        book.save(stream)
        if len(stream.getvalue()) > MAX_BYTES:
            raise InvalidInput("Заполненный шаблон превышает 32 МБ")
        return stream.getvalue()
    finally:
        book.close()
