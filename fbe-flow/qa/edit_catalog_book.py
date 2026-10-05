"""Change one cell as an operator would in the exported workbook."""

import sys

from openpyxl import load_workbook

book = load_workbook(sys.argv[1])
sheet = book["Товары"]
keys = {cell.value: cell.column for cell in sheet[2]}
for row in sheet.iter_rows(min_row=3):
    if row[keys["id"] - 1].value == sys.argv[3]:
        sheet.cell(row[0].row, keys["brand"]).value = "Из XLSX"
        break
else:
    raise RuntimeError("QA product missing in export")
book.save(sys.argv[2])
