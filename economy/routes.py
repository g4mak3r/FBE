from __future__ import annotations

import json
import os
import threading
import time
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urlencode

from fastapi import APIRouter, Form, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .repository import (
    add_cost_history,
    build_dashboard,
    delete_finance_reports_for_period,
    delete_ad_stats_for_period,
    delete_ad_expenses_for_period,
    ensure_economy_schema,
    get_incremental_cursor,
    get_report_coverage,
    get_ad_coverage,
    record_sync_run,
    flatten_ad_stats,
    save_ad_campaigns,
    save_ad_expenses,
    save_ad_stats,
    save_finance_reports,
    save_orders,
    save_sales,
)
from .token_inspector import SOURCE_SPECS, inspect_tokens, normalize_token, source_for_sync
from .wb_client import EconomyWBClient, EconomyWBError

MONEY_FMT = '#,##0.00;[Red](#,##0.00);-'
INT_FMT = '#,##0;[Red](#,##0);-'
ECONOMY_SYNC_LOCK = threading.Lock()
ECONOMY_SETTINGS_LOCK = threading.RLock()


def _load_json(path: Path) -> dict[str, Any]:
    with ECONOMY_SETTINGS_LOCK:
        if not path.exists():
            return {}
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}


def _save_json(path: Path, data: dict[str, Any]) -> None:
    with ECONOMY_SETTINGS_LOCK:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
        try:
            temporary.write_text(
                json.dumps(data, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)


def _merge_json(
    path: Path,
    changes: dict[str, Any],
    *,
    remove: tuple[str, ...] = (),
) -> dict[str, Any]:
    """Atomically merge UI state so concurrent handlers cannot lose fields."""
    with ECONOMY_SETTINGS_LOCK:
        current = _load_json(path)
        for key in remove:
            current.pop(key, None)
        current.update(changes)
        _save_json(path, current)
        return current


def _safe_date(value: str, fallback: date) -> str:
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date().isoformat()
    except Exception:
        return fallback.isoformat()


def _economy_url(**params: Any) -> str:
    clean = {key: str(value) for key, value in params.items() if value not in (None, "")}
    query = urlencode(clean)
    return f"/economy?{query}" if query else "/economy"


def _cursor_with_overlap(cursor: str, fallback: str, minutes: int = 5) -> str:
    if not cursor:
        return fallback
    text = str(cursor).strip()
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return (parsed - timedelta(minutes=max(0, int(minutes)))).isoformat()
    except Exception:
        return max(fallback, text)


def _write_raw(raw_dir: Path, source: str, date_from: str, date_to: str, payload: Any) -> str:
    raw_dir.mkdir(parents=True, exist_ok=True)
    target = raw_dir / f"{source}_{date_from}_{date_to}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(target)


def _write_diagnostic(diagnostics_dir: Path, kind: str, payload: dict[str, Any]) -> str:
    diagnostics_dir.mkdir(parents=True, exist_ok=True)
    safe_kind = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in kind)
    target = diagnostics_dir / f"economy_{safe_kind}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.json"
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(target)


def _human_wb_error(source: str, exc: Exception) -> str:
    try:
        token_source = source_for_sync(source)
    except Exception:
        token_source = source if source in SOURCE_SPECS else "statistics"
    spec = SOURCE_SPECS[token_source]
    label = str(spec["label"])
    token_genitive = str(spec.get("token_genitive") or label.lower())
    required = str(spec["required_category"])
    if not isinstance(exc, EconomyWBError):
        return str(exc)

    detail = str(exc.detail or exc)
    detail_lower = detail.lower()
    if exc.status_code in {401, 403} and ("scope" in detail_lower or "category" in detail_lower):
        message = (
            f"Токен {token_genitive} не имеет категории «{required}». "
            "Создайте новый токен WB с соответствующей категорией."
        )
    elif exc.status_code == 401:
        message = f"WB отклонил токен контура «{label}». Проверьте срок действия и тип токена."
    elif exc.status_code == 403:
        message = f"Для контура «{label}» недостаточно прав доступа WB API."
    elif exc.status_code == 429:
        retry = exc.rate_limit.get("retry_seconds") or exc.rate_limit.get("reset_seconds")
        if retry is not None:
            message = f"WB ограничил запрос «{label}». Повторите обновление примерно через {int(float(retry) + 0.999)} сек."
        else:
            message = f"WB временно ограничил запросы контура «{label}»."
    else:
        message = f"Ошибка WB API в контуре «{label}»: {detail}"
    if exc.request_id:
        message += f" requestId: {exc.request_id}"
    return message


def _client(config: dict[str, Any], settings: dict[str, Any]) -> EconomyWBClient:
    return EconomyWBClient(
        default_token=str(config.get("wb_token") or ""),
        mock_mode=bool(config.get("mock_mode", False)),
        max_retries=int(config.get("economy_wb_max_retries", 4)),
        max_wait_seconds=float(config.get("economy_wb_max_wait_seconds", 190.0)),
        retry_fallback_seconds=float(config.get("economy_wb_retry_fallback_seconds", 5.0)),
        request_timeout_seconds=float(config.get("economy_wb_request_timeout_seconds", 30.0)),
    )


def _default_history_start(today: date) -> str:
    # Finance report lists are available from 2025-01-01. The current-year
    # default keeps the first load compact while remaining useful.
    return date(today.year, 1, 1).isoformat()


def _history_start(settings: dict[str, Any], today: date) -> str:
    fallback = _default_history_start(today)
    value = _safe_date(str(settings.get("history_start") or ""), datetime.strptime(fallback, "%Y-%m-%d").date())
    return max("2025-01-01", value)


def _campaign_ids(settings: dict[str, Any], config: dict[str, Any]) -> list[int]:
    raw = settings.get("promotion_campaign_ids")
    if raw in (None, ""):
        raw = config.get("economy_campaign_ids") or [37958203, 38741167]
    if isinstance(raw, str):
        values = raw.replace(";", ",").replace(" ", ",").split(",")
    elif isinstance(raw, (list, tuple, set)):
        values = list(raw)
    else:
        values = [raw]
    ids: list[int] = []
    for value in values:
        try:
            number = int(str(value).strip())
            if number > 0:
                ids.append(number)
        except Exception:
            continue
    return sorted(set(ids))


def _date_chunks(date_from: str, date_to: str, max_days: int = 31) -> list[tuple[str, str]]:
    start = datetime.strptime(date_from, "%Y-%m-%d").date()
    end = datetime.strptime(date_to, "%Y-%m-%d").date()
    chunks: list[tuple[str, str]] = []
    while start <= end:
        chunk_end = min(end, start + timedelta(days=max_days - 1))
        chunks.append((start.isoformat(), chunk_end.isoformat()))
        start = chunk_end + timedelta(days=1)
    return chunks


def _all_finance_reports(
    client: EconomyWBClient,
    date_from: str,
    date_to: str,
    *,
    period: str,
    request_gap_seconds: float,
) -> tuple[list[dict[str, Any]], int]:
    rows: list[dict[str, Any]] = []
    offset = 0
    calls = 0
    limit = 1000
    while True:
        if calls and not client.mock_mode:
            time.sleep(max(float(request_gap_seconds), 60.0))
        page = client.get_finance_reports(
            date_from,
            date_to,
            period=period,
            limit=limit,
            offset=offset,
        )
        calls += 1
        if not page:
            break
        rows.extend(page)
        if len(page) < limit:
            break
        offset += len(page)
    return rows, calls


def _record_failure(
    *,
    db_path: Path,
    diagnostics_dir: Path,
    source: str,
    date_from: str,
    date_to: str,
    started: str,
    exc: Exception,
    request_events: list[dict[str, Any]],
) -> dict[str, Any]:
    human = _human_wb_error(source, exc)
    diagnostic: dict[str, Any] = {
        "source": source,
        "date_from": date_from,
        "date_to": date_to,
        "occurred_at": datetime.now().isoformat(timespec="seconds"),
        "message": human,
        "exception_type": type(exc).__name__,
        "wb_requests": request_events,
    }
    if isinstance(exc, EconomyWBError):
        diagnostic["wb_api_error"] = exc.diagnostic()
    diagnostic_file = _write_diagnostic(diagnostics_dir, f"sync_{source}", diagnostic)
    record_sync_run(
        db_path,
        source=source,
        date_from=date_from,
        date_to=date_to,
        status="error",
        rows_read=0,
        rows_written=0,
        message=human,
        details={"diagnostic_file": diagnostic_file, "wb_requests": request_events},
        started_at=started,
    )
    return {"source": source, "ok": False, "read": 0, "written": 0, "error": human}


def _sync_operational(
    *,
    source: str,
    client: EconomyWBClient,
    db_path: Path,
    raw_dir: Path,
    diagnostics_dir: Path,
    history_start: str,
    today: str,
    full_history: bool,
) -> dict[str, Any]:
    started = datetime.now().isoformat(timespec="seconds")
    client.drain_request_events()
    try:
        cursor = "" if full_history else get_incremental_cursor(db_path, source, history_start, today)
        # WB guarantees operational history for at most 90 days. Existing older
        # local rows are preserved, but a fresh request cannot recreate them.
        guaranteed_start = (datetime.strptime(today, "%Y-%m-%d").date() - timedelta(days=89)).isoformat()
        initial_start = max(history_start, guaranteed_start)
        api_date_from = initial_start if full_history else _cursor_with_overlap(cursor, initial_start)
        if source == "orders":
            payload = client.get_orders(api_date_from)
            event_names = ("date", "orderDate")
            written = save_orders(
                db_path,
                [row for row in payload if history_start <= str(next((row.get(n) for n in event_names if row.get(n)), ""))[:10] <= today],
            )
        else:
            payload = client.get_sales(api_date_from)
            event_names = ("date", "saleDate")
            written = save_sales(
                db_path,
                [row for row in payload if history_start <= str(next((row.get(n) for n in event_names if row.get(n)), ""))[:10] <= today],
            )
        raw_file = _write_raw(raw_dir, source, history_start, today, payload)
        events = client.drain_request_events()
        details = {
            "api_date_from": api_date_from,
            "history_start": history_start,
            "full_history": full_history,
            "raw_file": raw_file,
            "wb_requests": events,
        }
        record_sync_run(
            db_path,
            source=source,
            date_from=history_start,
            date_to=today,
            status="ok",
            rows_read=len(payload),
            rows_written=written,
            message="Обновлен локальный кеш WB.",
            details=details,
            started_at=started,
        )
        return {"source": source, "ok": True, "read": len(payload), "written": written, "details": details}
    except Exception as exc:
        return _record_failure(
            db_path=db_path,
            diagnostics_dir=diagnostics_dir,
            source=source,
            date_from=history_start,
            date_to=today,
            started=started,
            exc=exc,
            request_events=client.drain_request_events(),
        )


def _sync_finance_reports(
    *,
    client: EconomyWBClient,
    db_path: Path,
    raw_dir: Path,
    diagnostics_dir: Path,
    history_start: str,
    today: str,
    period_kind: str,
    full_history: bool,
    request_gap_seconds: float,
) -> dict[str, Any]:
    started = datetime.now().isoformat(timespec="seconds")
    client.drain_request_events()
    source = "finance_weekly" if period_kind == "weekly" else "finance"
    try:
        coverage = get_report_coverage(db_path, period_kind)
        if full_history or not coverage.get("date_to"):
            fetch_from = history_start
        else:
            overlap_days = 35 if period_kind == "weekly" else 14
            last_day = datetime.strptime(str(coverage["date_to"]), "%Y-%m-%d").date()
            fetch_from = max(history_start, (last_day - timedelta(days=overlap_days)).isoformat())
        payload, calls = _all_finance_reports(
            client,
            fetch_from,
            today,
            period=period_kind,
            request_gap_seconds=request_gap_seconds,
        )
        raw_file = _write_raw(raw_dir, f"finance_{period_kind}", fetch_from, today, payload)
        # A successful 204 means that WB has no rows for the request. Preserve
        # the existing local cache instead of erasing previously downloaded
        # reports on an empty refresh response.
        removed = 0
        written = 0
        if payload:
            removed = delete_finance_reports_for_period(
                db_path,
                fetch_from,
                today,
                period_kind=period_kind,
            )
            written = save_finance_reports(db_path, payload, period_kind=period_kind)
        events = client.drain_request_events()
        details = {
            "period_kind": period_kind,
            "fetch_from": fetch_from,
            "history_start": history_start,
            "full_history": full_history,
            "removed_before_replace": removed,
            "raw_file": raw_file,
            "api_calls": calls,
            "wb_requests": events,
        }
        label = "дневные" if period_kind == "daily" else "недельные"
        record_sync_run(
            db_path,
            source=source,
            date_from=fetch_from,
            date_to=today,
            status="ok",
            rows_read=len(payload),
            rows_written=written,
            message=f"Обновлены официальные {label} отчеты WB.",
            details=details,
            started_at=started,
        )
        return {"source": source, "ok": True, "read": len(payload), "written": written, "details": details}
    except Exception as exc:
        return _record_failure(
            db_path=db_path,
            diagnostics_dir=diagnostics_dir,
            source="finance",
            date_from=history_start,
            date_to=today,
            started=started,
            exc=exc,
            request_events=client.drain_request_events(),
        )


def _sync_ads(
    *,
    client: EconomyWBClient,
    db_path: Path,
    raw_dir: Path,
    diagnostics_dir: Path,
    history_start: str,
    today: str,
    pinned_ids: list[int],
    full_history: bool,
    request_gap_seconds: float,
) -> dict[str, Any]:
    started = datetime.now().isoformat(timespec="seconds")
    client.drain_request_events()
    try:
        coverage = get_ad_coverage(db_path)
        if full_history:
            fetch_from = history_start
        elif coverage.get("date_to"):
            last_day = datetime.strptime(str(coverage["date_to"]), "%Y-%m-%d").date()
            fetch_from = max(history_start, (last_day - timedelta(days=7)).isoformat())
        else:
            fetch_from = max(history_start, (datetime.strptime(today, "%Y-%m-%d").date() - timedelta(days=30)).isoformat())

        discovered_ids = client.get_campaign_ids()
        campaign_ids = sorted(set(discovered_ids) | set(pinned_ids))
        if not campaign_ids:
            raise ValueError("WB не вернул рекламных кампаний, а ID кампаний не заданы в настройках.")

        info_rows = client.get_campaign_info(campaign_ids)
        known = {int(row.get("advertId") or 0) for row in info_rows if str(row.get("advertId") or "").isdigit()}
        info_rows.extend({"advertId": value, "name": f"Кампания {value}", "status": 0, "type": 0} for value in campaign_ids if value not in known)
        campaigns_written = save_ad_campaigns(db_path, info_rows)

        all_stats_payload: list[dict[str, Any]] = []
        all_expenses_payload: list[dict[str, Any]] = []
        stats_written = expenses_written = 0
        chunks = _date_chunks(fetch_from, today, 31)
        for index, (chunk_from, chunk_to) in enumerate(chunks):
            if index and not client.mock_mode and request_gap_seconds > 0:
                time.sleep(float(request_gap_seconds))
            stats_payload = client.get_campaign_stats(campaign_ids, chunk_from, chunk_to)
            flat_stats = flatten_ad_stats(stats_payload)
            if stats_payload:
                delete_ad_stats_for_period(db_path, chunk_from, chunk_to, campaign_ids)
                stats_written += save_ad_stats(db_path, flat_stats)
                all_stats_payload.extend(stats_payload)

            if not client.mock_mode and request_gap_seconds > 0:
                time.sleep(float(request_gap_seconds))
            expenses_payload = client.get_ad_expenses(chunk_from, chunk_to)
            if expenses_payload:
                delete_ad_expenses_for_period(db_path, chunk_from, chunk_to)
                expenses_written += save_ad_expenses(db_path, expenses_payload)
                all_expenses_payload.extend(expenses_payload)

        raw_stats = _write_raw(raw_dir, "ads_stats", fetch_from, today, all_stats_payload)
        raw_expenses = _write_raw(raw_dir, "ads_expenses", fetch_from, today, all_expenses_payload)
        events = client.drain_request_events()
        details = {
            "fetch_from": fetch_from, "history_start": history_start, "full_history": full_history,
            "campaign_ids": campaign_ids, "discovered_ids": discovered_ids, "pinned_ids": pinned_ids,
            "campaigns_written": campaigns_written, "stats_written": stats_written,
            "expenses_written": expenses_written, "chunks": chunks,
            "raw_stats": raw_stats, "raw_expenses": raw_expenses, "wb_requests": events,
        }
        record_sync_run(
            db_path, source="ads", date_from=fetch_from, date_to=today, status="ok",
            rows_read=len(all_stats_payload) + len(all_expenses_payload),
            rows_written=stats_written + expenses_written,
            message="Обновлены рекламные кампании, статистика и история списаний WB Продвижения.",
            details=details, started_at=started,
        )
        return {
            "source": "ads", "ok": True, "read": len(all_stats_payload) + len(all_expenses_payload),
            "written": stats_written + expenses_written, "campaigns": campaigns_written,
            "stats": stats_written, "expenses": expenses_written, "details": details,
        }
    except Exception as exc:
        return _record_failure(
            db_path=db_path, diagnostics_dir=diagnostics_dir, source="ads",
            date_from=history_start, date_to=today, started=started, exc=exc,
            request_events=client.drain_request_events(),
        )


def _export_economy_xlsx(db_path: Path, dashboard: dict[str, Any], target: Path) -> Path:
    wb = Workbook()
    dark = PatternFill("solid", fgColor="172033")
    white = Font(color="FFFFFF", bold=True)

    overview = wb.active
    overview.title = "Основные показатели"
    overview.merge_cells("A1:D1")
    overview["A1"] = "FBE 0.77.0 · Экономика и реклама WB"
    overview["A1"].fill = dark
    overview["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    overview.append(["Период", dashboard["date_from"], dashboard["date_to"], "Расчет из локальной SQLite-базы"])
    overview.append([])
    overview.append(["Контур", "Показатель", "Количество", "Сумма, руб."])
    for cell in overview[4]:
        cell.fill = dark
        cell.font = white
    t = dashboard["totals"]
    rows = [
        ("Заказы", "Все заказы", t["orders_qty"], t["orders_amount"]),
        ("Заказы", "Отменено", t["cancel_qty"], t["cancel_amount"]),
        ("Заказы", "Активные заказы", t["active_orders_qty"], t["active_orders_amount"]),
        ("Выкупы", "Выкуплено", t["buyout_qty"], t["buyout_amount"]),
        ("Выкупы", "Возвращено", t["return_qty"], t["return_amount"]),
        ("Выкупы", "Чистые выкупы", t["net_buyout_qty"], t["net_buyout_amount"]),
        ("Дневные отчеты WB", "Продажи", None, t["report_sales"]),
        ("Дневные отчеты WB", "К перечислению", None, t["bank_payment"]),
        ("Дневные отчеты WB", "Разница WB", None, t["wb_difference"]),
        ("Реклама", "Списано WB Продвижением", None, dashboard["advertising"]["effective_spend"]),
        ("Реклама", "Сумма заказов из рекламы", dashboard["advertising"]["orders"], dashboard["advertising"]["revenue"]),
        ("Реклама", "ДРР рекламы, %", None, dashboard["advertising"]["drr_pct"]),
        ("Себестоимость", "Чистые выкупы", t["net_buyout_qty"], t["cogs"]),
        ("Контрольный остаток", "После WB и себестоимости", None,
         t["remaining_after_wb_and_cogs"] if dashboard["quality"]["can_show_remaining"] else None),
    ]
    for item in rows:
        overview.append(list(item))
    for r in range(5, overview.max_row + 1):
        overview.cell(r, 3).number_format = INT_FMT
        overview.cell(r, 4).number_format = MONEY_FMT

    sku = wb.create_sheet("Выкупы по артикулам")
    sku.append(["Артикул продавца", "Артикул WB", "Товар", "Формат", "Выкуплено", "Возвраты", "Чистые выкупы", "Чистые выкупы, руб.", "Себестоимость единицы", "Себестоимость чистых выкупов", "Себестоимость заполнена"])
    for cell in sku[1]:
        cell.fill = dark
        cell.font = white
    for row in dashboard["rows"]:
        sku.append([
            row["seller_article"], row["nm_id"], row["title"], row["volume"],
            row["buyout_qty"], row["return_qty"], row["net_qty"], row["net_amount"],
            row["cost_price_rub"], row["cogs"], "Да" if row["cost_ready"] else "Нет",
        ])
    for r in range(2, sku.max_row + 1):
        for c in (5, 6, 7):
            sku.cell(r, c).number_format = INT_FMT
        for c in (8, 9, 10):
            sku.cell(r, c).number_format = MONEY_FMT

    def report_sheet(title: str, report_rows: list[dict[str, Any]]) -> None:
        sheet = wb.create_sheet(title)
        sheet.append(["ID отчета", "Период", "С", "По", "Создан", "Тип", "Продажи", "forPay", "Логистика", "Хранение", "Приемка", "Удержания", "Штрафы", "Доплаты", "К перечислению"])
        for cell in sheet[1]:
            cell.fill = dark
            cell.font = white
        for row in report_rows:
            sheet.append([
                row.get("report_id"), row.get("period_kind"), row.get("date_from"), row.get("date_to"),
                row.get("create_date"), row.get("report_type"), row.get("retail_amount_sum", 0),
                row.get("for_pay_sum", 0), row.get("delivery_service_sum", 0), row.get("paid_storage_sum", 0),
                row.get("paid_acceptance_sum", 0), row.get("deduction_sum", 0), row.get("penalty_sum", 0),
                row.get("additional_payment_sum", 0), row.get("bank_payment_sum", 0),
            ])
        for r in range(2, sheet.max_row + 1):
            for c in range(7, 16):
                sheet.cell(r, c).number_format = MONEY_FMT

    ads = wb.create_sheet("Реклама и ДРР")
    ads.append(["ID кампании", "Название", "Статус", "Тип", "Оплата", "Расход по статистике", "Списано WB", "Сумма заказов из рекламы", "Заказы", "Товаров в заказах", "ДРР, %", "Показы", "Клики", "CTR, %", "CPC", "Цена заказа", "Расход на товар в заказах"])
    for cell in ads[1]:
        cell.fill = dark
        cell.font = white
    for row in dashboard["advertising"]["rows"]:
        ads.append([
            row["advert_id"], row["name"], row["status_label"], row["type_label"], row["payment_type"],
            row["spend"], row["actual_spend"], row["revenue"], row["orders"], row["shks"],
            row["drr_pct"], row["views"], row["clicks"], row["ctr_pct"], row["cpc"],
            row["cpo"], row["cost_per_ordered_item"],
        ])
    for r in range(2, ads.max_row + 1):
        for c in (6, 7, 8, 15, 16, 17):
            ads.cell(r, c).number_format = MONEY_FMT
        for c in (9, 10, 12, 13):
            ads.cell(r, c).number_format = INT_FMT
        for c in (11, 14):
            ads.cell(r, c).number_format = '0.0%' if False else '0.0'

    report_sheet("Дневные отчеты WB", dashboard["report_rows"])
    report_sheet("Недельная сверка WB", dashboard["weekly_report_rows"])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for index, column in enumerate(sheet.columns, start=1):
            letter = get_column_letter(index)
            max_len = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[letter].width = min(max(max_len + 2, 11), 42)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target


def build_economy_router(
    *,
    db_path: str | Path,
    templates_dir: str | Path,
    config: dict[str, Any],
    update_wb_token: Callable[[str], None] | None = None,
    submit_job: Callable[..., str] | None = None,
    template_globals: dict[str, Any] | None = None,
) -> APIRouter:
    router = APIRouter()
    templates = Jinja2Templates(directory=str(templates_dir))
    templates.env.globals.update(template_globals or {})
    db_path = ensure_economy_schema(db_path)
    settings_path = Path(config.get("economy_settings_path") or "data/economy_settings.json")
    raw_dir = Path(config.get("economy_raw_dir") or "data/economy/raw")
    exports_dir = Path(config.get("economy_exports_dir") or "data/economy/exports")
    diagnostics_dir = Path(config.get("economy_diagnostics_dir") or "data/diagnostics")

    def _dispatch_sync_job(
        *,
        title: str,
        worker: Callable[[Callable[..., None]], dict[str, Any]],
        return_start: str,
        return_end: str,
    ) -> RedirectResponse:
        if submit_job is not None:
            job_id = submit_job(title, worker, operation_key="economy-sync")
            return RedirectResponse(
                _economy_url(
                    date_from=return_start,
                    date_to=return_end,
                    job_id=job_id,
                    msg="Обновление запущено в фоне. Эту страницу можно закрыть.",
                ),
                status_code=303,
            )
        result = worker(lambda **kwargs: None)
        return RedirectResponse(
            _economy_url(
                date_from=return_start,
                date_to=return_end,
                msg=result.get("message", "") if result.get("ok") else "",
                error=result.get("error", "") if not result.get("ok") else "",
            ),
            status_code=303,
        )

    @router.get("/economy", response_class=HTMLResponse)
    def economy_page(
        request: Request,
        date_from: str = "",
        date_to: str = "",
        msg: str = "",
        error: str = "",
        job_id: str = "",
    ):
        today = date.today()
        date_to_value = _safe_date(date_to, today)
        date_from_value = _safe_date(date_from, today - timedelta(days=29))
        if date_from_value > date_to_value:
            date_from_value, date_to_value = date_to_value, date_from_value
        dashboard = build_dashboard(db_path, date_from_value, date_to_value)
        settings = _load_json(settings_path)
        token_state = inspect_tokens(config, settings)
        month_start = today.replace(day=1)
        previous_month_end = month_start - timedelta(days=1)
        previous_month_start = previous_month_end.replace(day=1)
        return templates.TemplateResponse(
            request,
            "economy.html",
            {
                "request": request,
                "config": config,
                "dashboard": dashboard,
                "token_state": token_state,
                "overall_token": token_state["statistics"],
                "history_start": _history_start(settings, today),
                "campaign_ids": _campaign_ids(settings, config),
                "campaign_ids_text": ", ".join(str(value) for value in _campaign_ids(settings, config)),
                "quick_periods": [
                    ("Сегодня", today.isoformat(), today.isoformat()),
                    ("7 дней", (today - timedelta(days=6)).isoformat(), today.isoformat()),
                    ("30 дней", (today - timedelta(days=29)).isoformat(), today.isoformat()),
                    ("Этот месяц", month_start.isoformat(), today.isoformat()),
                    ("Прошлый месяц", previous_month_start.isoformat(), previous_month_end.isoformat()),
                ],
                "msg": msg,
                "error": error,
                "job_id": job_id,
            },
        )

    @router.post("/economy/sync")
    def economy_sync(return_date_from: str = Form(default=""), return_date_to: str = Form(default="")):
        today_obj = date.today()
        return_start = _safe_date(return_date_from, today_obj - timedelta(days=29))
        return_end = _safe_date(return_date_to, today_obj)

        def worker(report: Callable[..., None]) -> dict[str, Any]:
            if not ECONOMY_SYNC_LOCK.acquire(blocking=False):
                return {"ok": False, "error": "Обновление WB уже выполняется."}
            try:
                settings = _load_json(settings_path)
                history_start = _history_start(settings, today_obj)
                today_value = today_obj.isoformat()
                states = inspect_tokens(config, settings)
                client = _client(config, settings)
                results: list[dict[str, Any]] = []
                sources = ("orders", "sales")
                report(total=4, progress=0, message="Загружаю заказы WB…")
                for index, source in enumerate(sources, start=1):
                    state = states["statistics"]
                    if not state["can_sync"]:
                        results.append({"source": source, "ok": False, "error": state["message"], "written": 0})
                    else:
                        results.append(_sync_operational(
                            source=source, client=client, db_path=db_path, raw_dir=raw_dir,
                            diagnostics_dir=diagnostics_dir, history_start=history_start,
                            today=today_value, full_history=False,
                        ))
                    report(total=4, progress=index, message="Загружаю продажи и возвраты WB…" if index == 1 else "Загружаю финансовые отчеты WB…")
                if states["finance"]["can_sync"]:
                    results.append(_sync_finance_reports(
                        client=client, db_path=db_path, raw_dir=raw_dir, diagnostics_dir=diagnostics_dir,
                        history_start=history_start, today=today_value, period_kind="daily",
                        full_history=False,
                        request_gap_seconds=float(config.get("economy_finance_request_gap_seconds", 61.0)),
                    ))
                else:
                    results.append({"source": "finance", "ok": False, "error": states["finance"]["message"], "written": 0})
                report(total=4, progress=3, message="Загружаю статистику рекламы WB…")
                if states["promotion"]["can_sync"]:
                    results.append(_sync_ads(
                        client=client, db_path=db_path, raw_dir=raw_dir, diagnostics_dir=diagnostics_dir,
                        history_start=history_start, today=today_value,
                        pinned_ids=_campaign_ids(settings, config), full_history=False,
                        request_gap_seconds=float(config.get("economy_advert_request_gap_seconds", 0)),
                    ))
                else:
                    results.append({"source": "ads", "ok": False, "error": states["promotion"]["message"], "written": 0})

                labels = {"orders": "заказы", "sales": "выкупы и возвраты", "finance": "дневные отчеты", "ads": "реклама"}
                completed = [item for item in results if item.get("ok")]
                failed = [item for item in results if not item.get("ok")]
                parts: list[str] = []
                for item in completed:
                    if item["source"] == "orders":
                        parts.append(f"заказы - новых или измененных строк {item.get('written', 0)}")
                    elif item["source"] == "sales":
                        parts.append(f"выкупы и возвраты - новых или измененных строк {item.get('written', 0)}")
                    elif item["source"] == "finance":
                        parts.append(f"дневные финансовые отчеты - загружено {item.get('written', 0)}")
                    elif item["source"] == "ads":
                        parts.append(f"реклама - кампаний {item.get('campaigns', 0)}, строк статистики {item.get('stats', 0)}, списаний {item.get('expenses', 0)}")
                message = "Данные WB обновлены: " + "; ".join(parts) if parts else "Данные не обновлены"
                errors = "; ".join(
                    f"{labels.get(item['source'], item['source'])}: {item.get('error', 'ошибка')}"
                    for item in failed
                )
                report(total=4, progress=4, message=message)
                return {"ok": not failed, "message": message, "error": errors, "results": results}
            finally:
                ECONOMY_SYNC_LOCK.release()

        return _dispatch_sync_job(
            title="Обновление экономики WB",
            worker=worker,
            return_start=return_start,
            return_end=return_end,
        )

    @router.post("/economy/sync-history")
    def sync_history(return_date_from: str = Form(default=""), return_date_to: str = Form(default="")):
        today_obj = date.today()
        return_start = _safe_date(return_date_from, today_obj - timedelta(days=29))
        return_end = _safe_date(return_date_to, today_obj)

        def worker(report: Callable[..., None]) -> dict[str, Any]:
            if not ECONOMY_SYNC_LOCK.acquire(blocking=False):
                return {"ok": False, "error": "Обновление WB уже выполняется."}
            try:
                settings = _load_json(settings_path)
                history_start = _history_start(settings, today_obj)
                today_value = today_obj.isoformat()
                states = inspect_tokens(config, settings)
                client = _client(config, settings)
                if not states["statistics"]["can_sync"] or not states["finance"]["can_sync"]:
                    return {"ok": False, "error": "Для загрузки истории нужны категории Статистика и Финансы."}
                report(total=4, progress=0, message="Загружаю историю заказов WB…")
                results = [_sync_operational(
                    source="orders", client=client, db_path=db_path, raw_dir=raw_dir,
                    diagnostics_dir=diagnostics_dir, history_start=history_start,
                    today=today_value, full_history=True,
                )]
                report(total=4, progress=1, message="Загружаю историю продаж и возвратов WB…")
                results.append(_sync_operational(
                    source="sales", client=client, db_path=db_path, raw_dir=raw_dir,
                    diagnostics_dir=diagnostics_dir, history_start=history_start,
                    today=today_value, full_history=True,
                ))
                report(total=4, progress=2, message="Загружаю финансовую историю WB…")
                results.append(_sync_finance_reports(
                    client=client, db_path=db_path, raw_dir=raw_dir,
                    diagnostics_dir=diagnostics_dir, history_start=history_start,
                    today=today_value, period_kind="daily", full_history=True,
                    request_gap_seconds=float(config.get("economy_finance_request_gap_seconds", 61.0)),
                ))
                if states["promotion"]["can_sync"]:
                    report(total=4, progress=3, message="Загружаю историю рекламы WB…")
                    results.append(_sync_ads(
                        client=client, db_path=db_path, raw_dir=raw_dir, diagnostics_dir=diagnostics_dir,
                        history_start=history_start, today=today_value,
                        pinned_ids=_campaign_ids(settings, config), full_history=True,
                        request_gap_seconds=float(config.get("economy_advert_request_gap_seconds", 0)),
                    ))
                failed = [item for item in results if not item.get("ok")]
                ads_note = f"; реклама {results[3].get('written', 0)}" if len(results) > 3 else ""
                message = (
                    f"История обновлена с {history_start}: заказы {results[0]['written']}; "
                    f"выкупы {results[1]['written']}; дневные отчеты {results[2]['written']}{ads_note}."
                )
                error = "; ".join(item.get("error", "Ошибка") for item in failed)
                report(total=4, progress=4, message=message)
                return {"ok": not failed, "message": message, "error": error, "results": results}
            finally:
                ECONOMY_SYNC_LOCK.release()

        return _dispatch_sync_job(
            title="Полная история экономики WB",
            worker=worker,
            return_start=return_start,
            return_end=return_end,
        )

    @router.post("/economy/sync-weekly")
    def sync_weekly(return_date_from: str = Form(default=""), return_date_to: str = Form(default="")):
        today_obj = date.today()
        return_start = _safe_date(return_date_from, today_obj - timedelta(days=29))
        return_end = _safe_date(return_date_to, today_obj)

        def worker(report: Callable[..., None]) -> dict[str, Any]:
            if not ECONOMY_SYNC_LOCK.acquire(blocking=False):
                return {"ok": False, "error": "Обновление WB уже выполняется."}
            try:
                settings = _load_json(settings_path)
                state = inspect_tokens(config, settings)["finance"]
                if not state["can_sync"]:
                    return {"ok": False, "error": state["message"]}
                report(total=1, progress=0, message="Загружаю недельные отчеты WB…")
                result = _sync_finance_reports(
                    client=_client(config, settings), db_path=db_path, raw_dir=raw_dir,
                    diagnostics_dir=diagnostics_dir, history_start=_history_start(settings, today_obj),
                    today=today_obj.isoformat(), period_kind="weekly", full_history=True,
                    request_gap_seconds=float(config.get("economy_finance_request_gap_seconds", 61.0)),
                )
                if not result["ok"]:
                    return {"ok": False, "error": result["error"], "result": result}
                message = f"Недельная сверка обновлена: {result['written']} отчетов."
                report(total=1, progress=1, message=message)
                return {"ok": True, "message": message, "result": result}
            finally:
                ECONOMY_SYNC_LOCK.release()

        return _dispatch_sync_job(
            title="Недельная сверка экономики WB",
            worker=worker,
            return_start=return_start,
            return_end=return_end,
        )

    @router.post("/api/economy/tokens/check")
    def check_economy_tokens(return_date_from: str = Form(default=""), return_date_to: str = Form(default="")):
        today_obj = date.today()
        return_start = _safe_date(return_date_from, today_obj - timedelta(days=29))
        return_end = _safe_date(return_date_to, today_obj)

        def worker(report: Callable[..., None]) -> dict[str, Any]:
            if not ECONOMY_SYNC_LOCK.acquire(blocking=False):
                return {"ok": False, "error": "Обновление WB уже выполняется."}
            try:
                settings = _load_json(settings_path)
                states = inspect_tokens(config, settings)
                client = _client(config, settings)
                checks = dict(settings.get("token_checks") or {})
                connected = denied = errors = 0
                for index, source in enumerate(("statistics", "finance", "promotion"), start=1):
                    state = states[source]
                    report(total=3, progress=index - 1, message=f"Проверяю контур WB «{state['label']}»…")
                    checked_at = datetime.now().isoformat(timespec="seconds")
                    record: dict[str, Any] = {"fingerprint": state["fingerprint"], "checked_at": checked_at, "request_id": "", "http_status": None}
                    if not state["can_check"]:
                        record.update({"status": "no_access", "message": state["message"]})
                        checks[source] = record
                        denied += 1
                        continue
                    try:
                        client.check_connection(source)
                        record.update({"status": "connected", "message": f"Подключение «{state['label']}» подтверждено WB API.", "http_status": 200})
                        connected += 1
                    except Exception as exc:
                        human = _human_wb_error(source, exc)
                        is_denied = isinstance(exc, EconomyWBError) and exc.status_code in {401, 403}
                        record.update({"status": "no_access" if is_denied else "error", "message": human, "request_id": exc.request_id if isinstance(exc, EconomyWBError) else "", "http_status": exc.status_code if isinstance(exc, EconomyWBError) else None})
                        if is_denied:
                            denied += 1
                        else:
                            errors += 1
                    checks[source] = record
                _merge_json(settings_path, {"token_checks": checks})
                if connected == 3 and not denied and not errors:
                    message = "Токен WB проверен: Статистика, Финансы и Продвижение доступны."
                else:
                    message = f"Проверка токена: доступно контуров API {connected} из 3, нет доступа {denied}"
                    if errors:
                        message += f", ошибок соединения {errors}"
                report(total=3, progress=3, message=message)
                return {"ok": errors == 0, "message": message, "error": message if errors else ""}
            finally:
                ECONOMY_SYNC_LOCK.release()

        return _dispatch_sync_job(
            title="Проверка токена WB",
            worker=worker,
            return_start=return_start,
            return_end=return_end,
        )

    @router.post("/api/economy/settings")
    def save_settings(wb_token: str = Form(default=""), clear_token: str | None = Form(default=None), history_start: str = Form(default=""), campaign_ids: str = Form(default=""), return_date_from: str = Form(default=""), return_date_to: str = Form(default="")):
        changes: dict[str, Any] = {}
        if history_start:
            changes["history_start"] = max("2025-01-01", _safe_date(history_start, date(date.today().year, 1, 1)))
        if campaign_ids.strip():
            parsed = []
            for value in campaign_ids.replace(";", ",").replace(" ", ",").split(","):
                try:
                    number = int(value.strip())
                    if number > 0:
                        parsed.append(number)
                except Exception:
                    continue
            changes["promotion_campaign_ids"] = sorted(set(parsed))
        normalized = normalize_token(wb_token)
        token_changed = False
        if clear_token:
            normalized = ""
            token_changed = True
        elif normalized:
            token_changed = normalized != normalize_token(str(config.get("wb_token") or ""))
        if clear_token or normalized:
            if update_wb_token is not None:
                update_wb_token(normalized)
            else:
                config["wb_token"] = normalized
        # Credentials live only in the central FBE settings. Remove stale
        # Economy-era copies when this form is explicitly saved.
        remove: list[str] = []
        if clear_token or normalized:
            remove.extend(("unified_token", "statistics_token", "finance_token", "promotion_token"))
            if token_changed:
                remove.append("token_checks")
        _merge_json(settings_path, changes, remove=tuple(remove))
        return RedirectResponse(_economy_url(date_from=return_date_from, date_to=return_date_to, msg="Настройки экономики сохранены."), status_code=303)

    @router.post("/api/economy/cost-history")
    def cost_history(seller_article: str = Form(...), effective_from: str = Form(...), cost_price_rub: float = Form(...), notes: str = Form(default="")):
        try:
            add_cost_history(db_path, seller_article=seller_article, effective_from=effective_from, cost_price_rub=cost_price_rub, notes=notes)
            return RedirectResponse(_economy_url(msg="История себестоимости обновлена"), status_code=303)
        except Exception as exc:
            return RedirectResponse(_economy_url(error=str(exc)), status_code=303)

    @router.get("/economy/export.xlsx")
    def economy_export(date_from: str = "", date_to: str = ""):
        today_obj = date.today()
        end = _safe_date(date_to, today_obj)
        start = _safe_date(date_from, today_obj - timedelta(days=29))
        if start > end:
            start, end = end, start
        dashboard = build_dashboard(db_path, start, end)
        target = exports_dir / f"Экономика WB {start} - {end}.xlsx"
        _export_economy_xlsx(db_path, dashboard, target)
        return FileResponse(target, filename=target.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @router.get("/api/economy/summary")
    def economy_summary(date_from: str = "", date_to: str = ""):
        today_obj = date.today()
        end = _safe_date(date_to, today_obj)
        start = _safe_date(date_from, today_obj - timedelta(days=29))
        dashboard = build_dashboard(db_path, start, end)
        return JSONResponse({"ok": True, "date_from": start, "date_to": end, "totals": dashboard["totals"], "quality": dashboard["quality"], "rows": dashboard["rows"], "status": dashboard["status"]})

    return router
