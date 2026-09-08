from __future__ import annotations

import base64
import calendar
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
import threading
import uuid
import csv
import json
import os
from io import BytesIO
import re
import time
from urllib.parse import quote, urlencode
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, TypeVar

from fastapi import HTTPException, FastAPI, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field

from pdf_tools import (
    create_stickers_pdf,
    merge_alternating_pdfs,
    merge_wb_pdf_with_internal_pdf_files,
    merge_wb_pdf_with_internal_page_refs,
    merge_title_labels_with_official_kiz_pages,
    create_small_title_labels_pdf,
    merge_pdf_page_refs,
    create_native_title_labels_pdf,
    create_datamatrix_labels_pdf,
    create_supply_info_label_pdf,
    prepend_pdf_file,
)
from printers.pdf_queue import create_paired_pdf, export_internal_label_image, print_pdf_queue
from printers.bartender_label import print_internal_label, build_internal_label_row, choose_bartender_template
from printers.bullzip_bartender_pdf import render_internal_label_pdf_with_bullzip, render_internal_labels_batch_pdf_with_bullzip
from printers.manual_labels import build_manual_row, render_manual_labels_pdf_with_bullzip, run_manual_bartender_print
from printers.wb_sticker import print_wb_sticker
from printers.marking_direct import render_title_png, render_datamatrix_label_png, print_images_windows
from catalog import ProductCatalog
from catalog.bootstrap import bootstrap_catalog
from catalog.routes import build_catalog_router
from economy import build_economy_router, ensure_economy_schema
from picklist_excel import parse_wb_picklist_xlsx
from print_queue import log_print_event, print_order_pair
from settings import Settings
from demo_mode import prepare_demo_database
from local_templates import pending_template_status, import_pending_templates
from economy.token_inspector import normalize_token, decode_wb_token
from suz_client import SuzApiError, SuzClient
from datamatrix_renderer import DataMatrixRenderError, OfficialDataMatrixRenderer
from keyboard_layout import force_english_keyboard_layout
from utils import (
    age_minutes_from_created_at,
    format_age,
    format_dt,
    sort_supplies_recent,
    supply_status,
    timer_class,
)
from wb_api import WBApiError, WBClient
from marking import (
    MarkingCodeError,
    build_marking_requirements,
    canonical_gtin14,
    is_marking_required,
    parse_marking_code,
    display_marking_code,
    validate_gtin,
)
from marking_db import (
    MarkingDbError,
    assign_code,
    count_supply_suz_ordered_codes_by_gtin,
    create_suz_order,
    ensure_database,
    get_assignment,
    get_suz_order,
    list_assignments,
    list_suz_orders,
    list_open_suz_orders,
    remove_local_assignment,
    validate_assignment_candidate,
    save_suz_codes,
    update_suz_item_status,
    update_suz_order_state,
    update_wb_state,
    assign_suz_codes_to_supply_orders,
    list_supply_code_assignments,
    mark_supply_codes_printed,
    create_circulation_document,
    list_circulation_documents,
    update_circulation_document,
    update_supply_code_lifecycle,
    mark_supply_circulation_error,
    create_suz_utilisation_report,
    list_suz_utilisation_reports,
    update_suz_utilisation_report,
    list_post_sale_assignments,
    update_post_sale_wb_statuses,
    mark_post_sale_document_failed,
    update_post_sale_true_status,
    upsert_fbs_supplies,
    set_fbs_supply_order_ids,
    upsert_fbs_orders,
    update_fbs_order_statuses,
    list_fbs_lifecycle_rows,
    list_fbs_registry_rows,
    list_fbs_supply_registry_rows,
    record_fbs_sync_state,
    get_fbs_sync_state,
    ensure_fbs_registry_from_marking_codes,
    update_fbs_order_sgtin_meta,
    recover_marking_code_from_wb,
    update_marking_true_statuses,
    WB_TERMINAL_STATUSES,
    WB_CANCELED_STATUSES,
)

BASE_DIR = Path(__file__).resolve().parent
try:
    _build_meta = json.loads((BASE_DIR / "BUILD.json").read_text(encoding="utf-8"))
except Exception:
    _build_meta = {}
APP_VERSION = str(_build_meta.get("version") or "0.85")
BUILD_ID = str(_build_meta.get("build_id") or "085-release-unknown")

settings = Settings(BASE_DIR / "config.json")
config = settings.data
if config.get("mock_mode"):
    from runtime_safety import install_demo_guard
    install_demo_guard()
INSTANCE_ID = uuid.uuid4().hex
_restart_pending = False
_transitioning = False
_active_requests = 0
config.setdefault("fbs_registry_history_days", 92)
config.setdefault("fbs_operational_sync_ttl_seconds", 45)
config.setdefault("fbs_operational_retry_seconds", 30)


def _runtime_path(key: str, default: str) -> Path:
    """Resolve configured runtime paths independently of the process cwd.

    FBE is commonly launched from a desktop shortcut or an installer helper,
    so a relative ``Path('data/...')`` must never silently point at the
    caller's working directory. ``Settings`` already normalizes configured
    paths; this helper also protects tests and legacy configs that omit a key.
    """
    raw = config.get(key, default)
    path = Path(str(raw or default))
    return path if path.is_absolute() else (BASE_DIR / path).resolve()


@asynccontextmanager
async def _app_lifespan(_: FastAPI):
    yield
    shutdown = globals().get("_shutdown_background_executors")
    if callable(shutdown):
        shutdown()


app = FastAPI(title=f"FBE {APP_VERSION}", lifespan=_app_lifespan)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")
templates.env.globals["fbe_version"] = APP_VERSION
templates.env.globals["fbe_build_id"] = BUILD_ID


@app.middleware("http")
async def _fbe_local_cache_guard(request: Request, call_next):
    """Local FBE must never keep an old UI after installing a new release."""
    global _active_requests, _transitioning
    path = request.url.path
    if request.method not in {"GET", "HEAD", "OPTIONS"}:
        origin = request.headers.get("origin")
        if request.headers.get("sec-fetch-site") == "cross-site" or (
            origin and origin.rstrip("/") != str(request.base_url).rstrip("/")
        ):
            return JSONResponse({"ok": False, "error": "Cross-origin action rejected"}, status_code=403)
    passive = path in {"/api/fbe-info", "/api/connections/status"} or path.startswith("/static/")
    transition = request.method == "POST" and path in {
        "/api/connections/wb", "/api/connections/wb/disconnect", "/api/connections/mode"}
    if not passive and (_restart_pending or _transitioning):
        return JSONResponse({"ok": False, "error": "Переключение профиля. Дождитесь перезапуска FBE."}, status_code=503)
    if transition:
        with _BACKGROUND_JOBS_LOCK:
            busy = any(job.get("state") in {"queued", "running"} for job in _BACKGROUND_JOBS.values())
        if _active_requests or busy:
            return JSONResponse({"ok": False, "error": "Дождитесь завершения текущих операций и повторите переключение."}, status_code=409)
        _transitioning = True
    if not passive:
        _active_requests += 1
    try:
        response = await call_next(request)
    finally:
        if not passive:
            _active_requests -= 1
        if transition:
            _transitioning = False

    path = request.url.path
    if path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-cache, no-store, max-age=0, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    elif request.method == "GET" and (
        response.headers.get("content-type", "").startswith("text/html")
        or path in {"/", "/marking", "/marking/global"}
    ):
        response.headers["Cache-Control"] = "no-cache, no-store, max-age=0, must-revalidate"
    return response


@app.get("/api/fbe-info")
def fbe_info():
    return {
        "app": "FBE",
        "instance_id": INSTANCE_ID,
        "restart_pending": _restart_pending,
        "version": APP_VERSION,
        "build_id": BUILD_ID,
        "mock_mode": bool(config.get("mock_mode", False)),
        "mode": "demo" if bool(config.get("mock_mode", False)) else "real",
    }


def _schedule_application_restart(exit_code: int = 75) -> None:
    global _restart_pending
    _restart_pending = True
    # Mode changes affect database paths and router dependencies. Restarting the
    # small local process is safer than hot-swapping those globals mid-request.
    def _exit() -> None:
        time.sleep(0.45)
        os._exit(int(exit_code))
    threading.Thread(target=_exit, name="fbe-restart", daemon=True).start()


@app.get("/api/connections/status")
def connections_status():
    state = settings.connection_public_state()
    state["active_mode"] = "demo" if bool(config.get("mock_mode", False)) else "real"
    state["first_run"] = bool(state["active_mode"] == "real" and not state.get("wb", {}).get("connected"))
    active_sid = str(state.get("wb", {}).get("sid") or "") if state["active_mode"] == "real" else ""
    state["local_templates"] = pending_template_status(BASE_DIR, active_sid)
    if state["active_mode"] == "demo":
        state["wb"] = {
            "connected": True,
            "name": "FBE Demo Company",
            "sid": "00000000-0000-4000-8000-000000000017",
            "tin": "000000000000",
            "tradeMark": "DEMO",
        }
        state["marking"] = {
            "configured": True,
            "inn": "000000000000",
            "oms_id_set": True,
            "connection_id_set": True,
            "certificate_set": True,
        }
        state["demo_profile"] = dict(state["wb"])
        state["local_templates"] = {"pending": False, "count": 0, "files": [], "can_import": False}
        state["first_run"] = False
    return state


def _require_real_connections() -> None:
    if config.get("mock_mode"):
        raise HTTPException(403, "DEMO MODE: сначала вернитесь в REAL.")


def _validate_wb_profile(token: str) -> dict[str, Any]:
    _require_real_connections()
    if not token:
        raise HTTPException(400, "Введите API-токен Wildberries.")
    try:
        profile = WBClient(token=token, mock_mode=False).get_seller_info()
    except Exception:
        # Never echo remote bodies or exceptions: they may contain credentials.
        raise HTTPException(400, "Не удалось проверить WB. Проверьте токен и доступность API.") from None
    if not str(profile.get("sid") or "").strip():
        raise HTTPException(400, "WB не вернул seller SID. Подключение не сохранено.")
    return profile


@app.post("/api/connections/wb")
def connect_wildberries(token: str = Form(...)):
    normalized = normalize_token(token)
    profile = _validate_wb_profile(normalized)
    settings.set_runtime_wb_connection(normalized, profile)
    _schedule_application_restart()
    return {"ok": True, "restart": True, "mode": "real", "previous_instance": INSTANCE_ID}


@app.post("/api/connections/wb/check")
def check_wildberries():
    profile = _validate_wb_profile(str(config.get("wb_token") or ""))
    if str(profile.get("sid") or "").strip() != config.get("workspace_sid"):
        raise HTTPException(409, "Seller SID изменился. Подключите аккаунт заново через профиль.")
    return {"ok": True, "message": "Соединение с Wildberries проверено."}


@app.post("/api/connections/wb/disconnect")
def disconnect_wildberries():
    _require_real_connections()
    settings.clear_runtime_wb_connection()
    _schedule_application_restart()
    return {"ok": True, "restart": True, "mode": "real", "previous_instance": INSTANCE_ID}


@app.post("/api/connections/templates/import")
def import_local_bartender_templates():
    _require_real_connections()
    sid = str(config.get("workspace_sid") or "").strip()
    if not sid:
        raise HTTPException(400, "Сначала подключите Wildberries для выбора рабочего профиля.")
    try:
        result = import_pending_templates(BASE_DIR, sid)
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from None
    return {"ok": bool(result.get("ok", False)), "result": result,
            "local_templates": pending_template_status(BASE_DIR, sid)}


@app.post("/api/connections/marking")
def connect_marking(
    inn: str = Form(default=""),
    oms_id: str = Form(default=""),
    connection_id: str = Form(default=""),
    cert_thumbprint: str = Form(default=""),
):
    _require_real_connections()
    clean_inn = re.sub(r"\D", "", str(inn or ""))
    changes = {
        "suz_auth_inn": clean_inn,
        "suz_true_participant_inn": clean_inn,
        "suz_true_producer_inn": clean_inn,
        "suz_true_owner_inn": clean_inn,
        "suz_oms_id": str(oms_id or "").strip(),
        "suz_connection_id": str(connection_id or "").strip(),
        "suz_cert_thumbprint": str(cert_thumbprint or "").replace(" ", "").strip().upper(),
    }
    settings.update_external_api_settings(changes)
    suz.reset_token()
    return {"ok": True, "marking": settings.connection_public_state().get("marking", {})}


@app.post("/api/connections/mode")
def switch_runtime_mode(mode: str = Form(...)):
    if str(mode).lower() not in {"real", "demo"}:
        raise HTTPException(400, "Режим должен быть REAL или DEMO.")
    requested = str(mode).lower()
    current = "demo" if bool(config.get("mock_mode", False)) else "real"
    if requested == current:
        return {"ok": True, "restart": False, "mode": current}
    settings.set_runtime_mode(requested)
    _schedule_application_restart()
    return {"ok": True, "restart": True, "mode": requested, "previous_instance": INSTANCE_ID}


wb = WBClient(token=config.get("wb_token", ""), mock_mode=config.get("mock_mode", False))
if bool(config.get("mock_mode", False)):
    marking_db_path = prepare_demo_database(config.get("database_path", "data/demo/fbe_demo.db"))
else:
    marking_db_path = ensure_database(config.get("database_path", "data/fbe.db"))

# 0.83.0: SQLite is the sole working product catalog. On the first start only,
# the existing products.xlsx is imported automatically when the products table is empty.
catalog_bootstrap_result = bootstrap_catalog(
    marking_db_path,
    legacy_xlsx_path=config.get("catalog_legacy_excel_path") or config.get("excel_path", "data/products.xlsx"),
    sheet_name=config.get("excel_sheet", "mainSheet"),
    columns=config.get("excel_columns") or {},
    backup_dir=config.get("catalog_backups_dir", "data/backups"),
    auto_import=bool(config.get("catalog_auto_import_legacy", True)),
)
if catalog_bootstrap_result.get("imported"):
    print(
        "[FBE] Каталог перенесен в SQLite: "
        f"товаров {catalog_bootstrap_result.get('products_written', 0)}, "
        f"переходников {catalog_bootstrap_result.get('samples_written', 0)}, "
        f"именных ароматов {catalog_bootstrap_result.get('names_written', 0)}, "
        f"форматов {catalog_bootstrap_result.get('formats_written', 0)}; "
        f"backup={catalog_bootstrap_result.get('backup', '')}"
    )
app.include_router(
    build_catalog_router(
        db_path=marking_db_path,
        templates_dir=BASE_DIR / "templates",
        config=config,
        template_globals={"fbe_version": APP_VERSION, "fbe_build_id": BUILD_ID},
    )
)

def update_global_wb_token(token: str) -> None:
    # A second credential form must not bypass seller validation/restart.
    raise HTTPException(409, "Меняйте WB-подключение через профиль компании в sidebar.")


def submit_economy_background_job(*args: Any, **kwargs: Any) -> str:
    """Late-bound bridge: the router is built before the shared job runner."""
    return submit_background_job(*args, **kwargs)


ensure_economy_schema(marking_db_path)
app.include_router(
    build_economy_router(
        db_path=marking_db_path,
        templates_dir=BASE_DIR / "templates",
        config=config,
        update_wb_token=update_global_wb_token,
        submit_job=submit_economy_background_job,
        template_globals={"fbe_version": APP_VERSION, "fbe_build_id": BUILD_ID},
    )
)


def load_suz_settings() -> dict[str, Any]:
    return settings.external_api_settings()


def save_suz_local_settings(changes: dict[str, Any]) -> dict[str, Any]:
    return settings.update_external_api_settings(changes)


suz = SuzClient(base_dir=BASE_DIR, settings_loader=load_suz_settings)
datamatrix_renderer = OfficialDataMatrixRenderer(base_dir=BASE_DIR, settings_loader=load_suz_settings)

T = TypeVar("T")
_cache: dict[str, tuple[float, Any]] = {}
_cache_refreshing: set[str] = set()
_cache_lock = threading.Lock()

# 0.83.0: network work must never block page rendering or delay browser click handlers.
# FBE is a single-user local app, so an in-process executor is enough and avoids
# introducing a broker/service dependency. Long WB/CRPT operations run here while
# HTTP handlers return immediately and the UI polls compact job state.
# Keep cache refresh traffic separate from user-triggered background jobs.
# In 0.83.11 both shared one small pool; a cold dashboard could occupy every
# worker with WB cache refreshes and make marking/operational jobs wait in line.
_EXECUTOR_LOCK = threading.Lock()
_CACHE_EXECUTOR: ThreadPoolExecutor | None = None
_JOB_EXECUTOR: ThreadPoolExecutor | None = None
_BACKGROUND_JOBS: dict[str, dict[str, Any]] = {}
_BACKGROUND_JOB_KEYS: dict[str, str] = {}
_BACKGROUND_JOBS_LOCK = threading.Lock()


def _executor(kind: str) -> ThreadPoolExecutor:
    global _CACHE_EXECUTOR, _JOB_EXECUTOR
    with _EXECUTOR_LOCK:
        if kind == "cache":
            if _CACHE_EXECUTOR is None:
                _CACHE_EXECUTOR = ThreadPoolExecutor(
                    max_workers=max(1, int(config.get("cache_worker_count", 2))),
                    thread_name_prefix="fbe-cache",
                )
            return _CACHE_EXECUTOR
        if _JOB_EXECUTOR is None:
            _JOB_EXECUTOR = ThreadPoolExecutor(
                max_workers=max(2, int(config.get("background_worker_count", 4))),
                thread_name_prefix="fbe-job",
            )
        return _JOB_EXECUTOR


def _shutdown_background_executors() -> None:
    """Cancel queued work and release pools cleanly on FBE shutdown/restart."""
    global _CACHE_EXECUTOR, _JOB_EXECUTOR
    with _EXECUTOR_LOCK:
        executors = (_CACHE_EXECUTOR, _JOB_EXECUTOR)
        _CACHE_EXECUTOR = None
        _JOB_EXECUTOR = None
    for executor in executors:
        if executor is not None:
            executor.shutdown(wait=False, cancel_futures=True)


def _job_snapshot(job_id: str) -> dict[str, Any] | None:
    with _BACKGROUND_JOBS_LOCK:
        row = _BACKGROUND_JOBS.get(str(job_id))
        return dict(row) if row else None


def _job_update(job_id: str, **changes: Any) -> None:
    with _BACKGROUND_JOBS_LOCK:
        row = _BACKGROUND_JOBS.get(str(job_id))
        if not row:
            return
        row.update(changes)
        row["updated_at"] = time.time()


def _cleanup_jobs() -> None:
    cutoff = time.time() - 3600
    with _BACKGROUND_JOBS_LOCK:
        for job_id in list(_BACKGROUND_JOBS):
            row = _BACKGROUND_JOBS[job_id]
            if float(row.get("updated_at") or row.get("created_at") or 0) < cutoff:
                _BACKGROUND_JOBS.pop(job_id, None)
                operation_key = str(row.get("operation_key") or "")
                if operation_key and _BACKGROUND_JOB_KEYS.get(operation_key) == job_id:
                    _BACKGROUND_JOB_KEYS.pop(operation_key, None)


def submit_background_job(
    title: str,
    worker: Callable[[Callable[..., None]], Any],
    *,
    operation_key: str | None = None,
) -> str:
    _cleanup_jobs()
    key = str(operation_key or "").strip()
    now = time.time()
    with _BACKGROUND_JOBS_LOCK:
        if key:
            existing_id = _BACKGROUND_JOB_KEYS.get(key)
            existing = _BACKGROUND_JOBS.get(existing_id or "")
            if existing and str(existing.get("state") or "") in {"queued", "running"}:
                return str(existing_id)
            _BACKGROUND_JOB_KEYS.pop(key, None)
        job_id = uuid.uuid4().hex[:16]
        _BACKGROUND_JOBS[job_id] = {
            "id": job_id, "title": str(title), "state": "queued",
            "progress": 0, "total": 0, "message": "В очереди…",
            "created_at": now, "updated_at": now, "result": None, "error": "",
            "operation_key": key,
        }
        if key:
            _BACKGROUND_JOB_KEYS[key] = job_id

    def report(*, message: str | None = None, progress: int | None = None, total: int | None = None, **extra: Any) -> None:
        changes: dict[str, Any] = dict(extra)
        if message is not None: changes["message"] = str(message)
        if progress is not None: changes["progress"] = int(progress)
        if total is not None: changes["total"] = int(total)
        _job_update(job_id, **changes)

    def runner() -> None:
        _job_update(job_id, state="running", message="Выполняется…")
        try:
            result = worker(report)
            _job_update(job_id, state="done", progress=int((_job_snapshot(job_id) or {}).get("total") or 1), result=result, message=(result or {}).get("message", "Готово") if isinstance(result, dict) else "Готово")
        except Exception as exc:
            _job_update(job_id, state="error", error=str(exc), message=str(exc))
        finally:
            if key:
                with _BACKGROUND_JOBS_LOCK:
                    if _BACKGROUND_JOB_KEYS.get(key) == job_id:
                        _BACKGROUND_JOB_KEYS.pop(key, None)

    _executor("job").submit(runner)
    return job_id


def _background_cache_refresh(key: str, loader: Callable[[], T]) -> None:
    # Remember when this background request started. A later explicit/manual
    # refresh must always win, even if this older request finishes afterwards.
    started_at = time.time()
    try:
        value = loader()
        with _cache_lock:
            current = _cache.get(key)
            if current is None or float(current[0]) <= started_at:
                _cache[key] = (time.time(), value)
    except Exception as exc:
        print(f"[FBE cache refresh] {key}: {exc}")
    finally:
        with _cache_lock:
            _cache_refreshing.discard(key)


def cached_stale_while_revalidate(
    key: str, ttl_seconds: int, loader: Callable[[], T], fallback: Callable[[], T]
) -> T:
    """Return local/stale data immediately and refresh network state in background."""
    now = time.time()
    with _cache_lock:
        hit = _cache.get(key)
        fresh = bool(hit and now - hit[0] < ttl_seconds)
        should_refresh = not fresh and key not in _cache_refreshing
        if should_refresh:
            _cache_refreshing.add(key)
    if fresh and hit:
        return hit[1]
    if should_refresh:
        _executor("cache").submit(_background_cache_refresh, key, loader)
    if hit:
        return hit[1]
    return fallback()


def invalidate_cache(prefix: str | None = None) -> None:
    with _cache_lock:
        if prefix is None:
            _cache.clear()
            return
        for key in list(_cache.keys()):
            if key.startswith(prefix):
                _cache.pop(key, None)



def moscow_now() -> datetime:
    return datetime.now(timezone(timedelta(hours=3)))


def get_catalog() -> ProductCatalog:
    # SQLite is the only runtime source. XLSX is import/export only.
    return ProductCatalog(
        db_path=marking_db_path,
        columns=config.get("excel_columns") or {},
    )


def build_manual_rows(label_type: str, aroma: str, volume: str | None, quantity: int) -> tuple[list[dict[str, str]], list[str]]:
    catalog = get_catalog()
    errors: list[str] = []
    quantity = max(1, min(99, int(quantity)))

    if label_type == "samples":
        record = catalog.find_sample_by_aroma(aroma)
        if not record:
            return [], [f"Аромат не найден в каталоге переходников: {aroma}"]
    elif label_type == "names":
        record = catalog.find_name_label(aroma=aroma, volume=volume)
        if not record:
            return [], [f"Аромат или формат не найден в справочнике именных этикеток: {aroma} / {volume or 'формат не выбран'}"]
    else:
        return [], [f"Неизвестный тип ручной печати: {label_type}"]

    row = build_manual_row(label_type, record, quantity)
    return [row.copy() for _ in range(quantity)], errors


def manual_context(msg: str | None = None, error: str | None = None) -> dict[str, Any]:
    catalog = get_catalog()
    sample_aromas = catalog.list_sample_aromas()
    name_aromas = catalog.list_name_aromas()
    volumes = catalog.list_name_formats()
    return {
        "catalog_error": catalog.error,
        "catalog_warnings": catalog.warnings[:20],
        # Kept for backward compatibility with older templates/plugins.
        "aromas": sample_aromas or name_aromas,
        # Separate DB-backed lists: samples for transition labels; names/formats for named labels.
        "sample_aromas": sample_aromas,
        "name_aromas": name_aromas,
        "volumes": volumes,
        "default_volume": volumes[0] if volumes else "",
        "msg": msg,
        "error": error,
        "config": config,
    }


def _local_order_snapshot(*, only_new: bool = False, days: int = 14) -> list[dict[str, Any]]:
    date_from = (moscow_now().date() - timedelta(days=max(1, int(days)))).isoformat()
    rows = list_fbs_registry_rows(marking_db_path, limit=10000, date_from=date_from)
    result: list[dict[str, Any]] = []
    for row in rows:
        if only_new:
            supplier = str(row.get("supplier_status") or "").lower()
            wb_state = str(row.get("wb_status") or "").lower()
            if str(row.get("supply_id") or "").strip():
                continue
            if supplier not in {"", "confirm"} or wb_state in WB_TERMINAL_STATUSES:
                continue
        raw: dict[str, Any] = {}
        try:
            parsed = json.loads(str(row.get("raw_json") or "{}"))
            if isinstance(parsed, dict): raw = parsed
        except Exception:
            pass
        raw.setdefault("id", int(row.get("order_id") or 0))
        raw.setdefault("article", str(row.get("seller_article") or ""))
        raw.setdefault("supplyId", str(row.get("supply_id") or ""))
        raw.setdefault("warehouseId", int(row.get("warehouse_id") or 0))
        raw.setdefault("createdAt", str(row.get("created_at_wb") or ""))
        raw.setdefault("wbStatus", str(row.get("wb_status") or ""))
        raw.setdefault("supplierStatus", str(row.get("supplier_status") or ""))
        if raw.get("id"):
            result.append(raw)
    return result


def _local_supply_snapshot() -> list[dict[str, Any]]:
    """Return the persisted supply read model without contacting WB.

    Supply rows are primary because they also represent empty/new supplies. Order
    rows are merged only for counts and backward compatibility with databases
    created before the supply-level registry existed.
    """
    date_from = _registry_history_start()
    supply_rows = list_fbs_supply_registry_rows(
        marking_db_path, limit=10000, date_from=date_from
    )
    order_rows = list_fbs_registry_rows(
        marking_db_path, limit=10000, date_from=date_from
    )

    grouped: dict[str, dict[str, Any]] = {}
    for row in supply_rows:
        sid = str(row.get("supply_id") or "").strip()
        if not sid:
            continue
        count = row.get("order_count")
        try:
            count_int = int(count)
        except (TypeError, ValueError):
            count_int = -1
        grouped[sid] = {
            "id": sid,
            "name": str(row.get("name") or ""),
            "createdAt": str(row.get("created_at_wb") or ""),
            "closedAt": str(row.get("closed_at_wb") or ""),
            "scanDt": str(row.get("scan_dt_wb") or ""),
            "done": bool(int(row.get("done") or 0)),
            "ordersCount": count_int if count_int >= 0 else 0,
            "_count_known": count_int >= 0,
        }

    observed_counts: dict[str, int] = {}
    for row in order_rows:
        sid = str(row.get("supply_id") or "").strip()
        if not sid:
            continue
        observed_counts[sid] = observed_counts.get(sid, 0) + 1
        if sid not in grouped:
            grouped[sid] = {
                "id": sid,
                "name": str(row.get("supply_name") or ""),
                "createdAt": str(row.get("supply_created_at_wb") or ""),
                "closedAt": str(row.get("closed_at_wb") or ""),
                "scanDt": str(row.get("scan_dt_wb") or ""),
                "done": bool(row.get("supply_done")),
                "ordersCount": 0,
                "_count_known": False,
            }

    result: list[dict[str, Any]] = []
    for sid, item in grouped.items():
        observed = int(observed_counts.get(sid, 0))
        if not item.pop("_count_known", False):
            item["ordersCount"] = observed
        elif observed > int(item.get("ordersCount") or 0):
            # Historical completed rows can outlive the current membership result
            # of a closed supply. Never display a misleading zero in that case.
            item["ordersCount"] = observed
        result.append(item)
    return result


def _store_new_orders_cache(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    clean = [dict(item) for item in rows if isinstance(item, dict)]
    try:
        upsert_fbs_orders(marking_db_path, clean)
    except Exception as exc:
        print(f"[FBE new orders registry] {exc}")
    with _cache_lock:
        _cache["new_orders"] = (time.time(), clean)
    return clean


def refresh_new_orders_now() -> list[dict[str, Any]]:
    """Explicit WB refresh used by the ↻ button.

    Unlike stale-while-revalidate, this waits for a real /orders/new response and
    only then replaces the visible cache. An exception leaves the previous cache
    untouched, so a transient WB/network failure cannot make the UI look empty.
    """
    rows = wb.get_new_orders()
    return _store_new_orders_cache(rows)


def get_new_orders_cached() -> list[dict[str, Any]]:
    def load() -> list[dict[str, Any]]:
        rows = wb.get_new_orders()
        try: upsert_fbs_orders(marking_db_path, rows)
        except Exception: pass
        return rows
    return cached_stale_while_revalidate(
        "new_orders", int(config.get("cache_ttl_seconds", 30)), load,
        lambda: _local_order_snapshot(only_new=True, days=14),
    )


def get_supplies_cached() -> list[dict[str, Any]]:
    max_pages = int(config.get("supplies_api_max_pages", 3))
    def load() -> list[dict[str, Any]]:
        rows = wb.get_all_supplies(max_pages=max_pages)
        try:
            prepared = prepare_supplies([dict(x) for x in rows])
            upsert_fbs_supplies(marking_db_path, prepared)
        except Exception:
            pass
        return rows
    return cached_stale_while_revalidate(
        "supplies", int(config.get("cache_ttl_seconds", 30)), load, _local_supply_snapshot
    )


def get_supply_details_cached(supply_id: str) -> dict[str, Any]:
    sid = str(supply_id)
    def load() -> dict[str, Any]:
        row = wb.get_supply_details(sid) or {"id": sid}
        try:
            prepared = prepare_supplies([dict(row)])
            if prepared: upsert_fbs_supplies(marking_db_path, prepared)
        except Exception:
            pass
        return dict(row)
    def fallback() -> dict[str, Any]:
        for row in _local_supply_snapshot():
            if str(row.get("id") or "") == sid:
                return dict(row)
        return {"id": sid}
    return cached_stale_while_revalidate(
        f"supply_details:{sid}", int(config.get("cache_ttl_seconds", 30)), load, fallback
    )


def get_seller_warehouses_cached() -> list[dict[str, Any]]:
    def fallback() -> list[dict[str, Any]]:
        found: dict[int, str] = {}
        for row in _local_order_snapshot(days=92):
            try: wid = int(row.get("warehouseId") or row.get("warehouseID") or 0)
            except Exception: wid = 0
            if not wid: continue
            name = str(row.get("sellerWarehouseName") or row.get("warehouseName") or "").strip()
            if name: found[wid] = name
        return [{"id": wid, "name": name} for wid, name in found.items()]
    return cached_stale_while_revalidate(
        "seller_warehouses", int(config.get("cache_ttl_seconds", 30)),
        wb.get_seller_warehouses, fallback,
    )


def seller_warehouse_name_map() -> dict[int, str]:
    result: dict[int, str] = {}
    try:
        rows = get_seller_warehouses_cached()
    except Exception:
        rows = []
    for row in rows:
        raw_id = row.get("id") or row.get("warehouseId") or row.get("warehouseID")
        try:
            warehouse_id = int(raw_id)
        except (TypeError, ValueError):
            continue
        name = str(row.get("name") or row.get("warehouseName") or row.get("title") or "").strip()
        if name:
            result[warehouse_id] = name
    return result


def order_warehouse_id(order: dict[str, Any]) -> int:
    try:
        return int(order.get("warehouseId") or order.get("warehouseID") or 0)
    except (TypeError, ValueError):
        return 0


def order_warehouse_name(order: dict[str, Any], names: dict[int, str] | None = None) -> str:
    warehouse_id = order_warehouse_id(order)
    mapped = (names or {}).get(warehouse_id, "") if warehouse_id else ""
    if mapped:
        return mapped
    explicit = str(
        order.get("sellerWarehouseName")
        or order.get("warehouseName")
        or order.get("warehouse")
        or ""
    ).strip()
    if explicit:
        return explicit
    return f"Склад #{warehouse_id}" if warehouse_id else "Склад не определен"


def get_all_orders_cached() -> list[dict[str, Any]]:
    def load() -> list[dict[str, Any]]:
        rows = wb.get_all_orders_recent(max_pages=int(config.get("fbs_orders_sync_max_pages", 8)))
        try:
            upsert_fbs_orders(marking_db_path, rows)
        except Exception:
            pass
        return rows
    return cached_stale_while_revalidate(
        "all_orders_recent", int(config.get("cache_ttl_seconds", 30)), load,
        lambda: _local_order_snapshot(days=92),
    )


def get_supply_order_ids_cached(supply_id: str) -> list[int]:
    sid = str(supply_id)
    def load() -> list[int]:
        ids = wb.get_supply_order_ids(sid)
        try:
            set_fbs_supply_order_ids(marking_db_path, supply_id=sid, order_ids=ids)
        except Exception:
            pass
        return ids
    def fallback() -> list[int]:
        return [
            int(r.get("order_id") or 0) for r in list_fbs_registry_rows(marking_db_path, limit=10000)
            if str(r.get("supply_id") or "") == sid and int(r.get("order_id") or 0) > 0
        ]
    return cached_stale_while_revalidate(
        f"supply_order_ids:{sid}", int(config.get("cache_ttl_seconds", 30)), load, fallback
    )


def seller_article_from_order(order: dict[str, Any]) -> str:
    return str(order.get("article") or order.get("supplierArticle") or "").strip()


def sku_from_order(order: dict[str, Any]) -> str:
    skus = order.get("skus") or []
    if skus:
        return str(skus[0])
    return str(order.get("barcode") or "")


def excel_kiz_info(product: dict[str, Any] | None) -> dict[str, str]:
    """Compatibility badge driven by the catalog marking profile.

    The legacy Excel column «КИЗ» is no longer authoritative. 0.83.0 derives
    this badge from products.kiz_required, which in turn is controlled by the
    explicit marking profile.
    """
    required = bool((product or {}).get("kiz_required"))
    if required:
        profile = str((product or {}).get("marking_profile") or "").strip()
        return {"label": "КИЗ", "class": "kiz-required", "raw": f"Профиль маркировки: {profile or 'установлен'}"}
    return {"label": "", "class": "kiz-empty", "raw": ""}


def enrich_orders(orders: list[dict[str, Any]], sort: str | None = None) -> list[dict[str, Any]]:
    catalog = get_catalog()
    warehouse_names = seller_warehouse_name_map()
    enriched = []
    for order in orders:
        article = seller_article_from_order(order)
        product = catalog.find_by_article(article)
        created_at = order.get("createdAt", "")
        age_min = age_minutes_from_created_at(created_at)
        kiz = excel_kiz_info(product)
        enriched.append(
            {
                "id": int(order["id"]),
                "seller_article": article,
                "sku": sku_from_order(order),
                "created_at": created_at,
                "created_at_text": format_dt(created_at),
                "age_minutes": age_min if age_min is not None else -1,
                "age_text": format_age(age_min),
                "timer_class": timer_class(age_min),
                "supply_id": order.get("supplyId", ""),
                "warehouse_id": order_warehouse_id(order),
                "warehouse_name": order_warehouse_name(order, warehouse_names),
                "excel_found": bool(product),
                "name": product.get("display_name", "") if product else "",
                "category": product.get("category", "") if product else "",
                "excel_kiz": product.get("excel_kiz", "") if product else "",
                "product": product or {},
                "kiz_label": kiz["label"],
                "kiz_class": kiz["class"],
                "kiz_raw": kiz["raw"],
                "raw": order,
            }
        )

    sort_orders(enriched, sort or config.get("default_order_sort", "urgent"))
    return enriched


def sort_orders(items: list[dict[str, Any]], sort: str) -> None:
    if sort in {"new", "timer_asc"}:
        items.sort(key=lambda x: x["age_minutes"] if x["age_minutes"] >= 0 else 10**12)
    else:  # urgent / timer_desc
        items.sort(key=lambda x: x["age_minutes"], reverse=True)


def prepare_supplies(raw_supplies: list[dict[str, Any]], limit: int | None = None) -> list[dict[str, Any]]:
    supplies = sort_supplies_recent(raw_supplies, limit=limit)
    for supply in supplies:
        status = supply_status(supply)
        supply["fbe_status"] = status["label"]
        supply["fbe_status_class"] = status["class"]
        supply["created_at_text"] = format_dt(supply.get("createdAt"))
        supply["closed_at_text"] = format_dt(supply.get("closedAt"))
        supply["scan_dt_text"] = format_dt(supply.get("scanDt"))
        supply["transfer_or_scan_text"] = supply["scan_dt_text"] or supply["closed_at_text"] or ""
        # Client-side sorting helpers. Kept as simple scalar values for HTML data-* attrs.
        supply["fbe_status_rank"] = {"На сборке": 0, "Ждет отгрузки": 1, "Получена WB": 2}.get(supply["fbe_status"], 9)
    return supplies


def _registry_history_start(days: int | None = None) -> str:
    days = int(days or config.get("fbs_registry_history_days", 92) or 92)
    return (moscow_now().date() - timedelta(days=max(1, days))).isoformat()


def _fbs_dashboard_supply_groups(
    rows: list[dict[str, Any]],
    supply_rows: list[dict[str, Any]] | None = None,
    authoritative_open_supply_ids: set[str] | None = None,
) -> dict[str, Any]:
    """Build operational FBS tabs from both supply state and order statuses.

    Supply-level fields answer whether a WB supply itself is still open (`done`).
    Order statuses answer where its positions are now. 0.83.11 used only the
    second source, so a newly-created/empty supply disappeared until status sync,
    while stale order links could resurrect an old supply. The merged model makes
    supply state authoritative for the assembly tab and keeps order statuses for
    handover/delivery semantics.
    """
    from utils import parse_wb_datetime

    final_wb = WB_TERMINAL_STATUSES
    assembly_cutoff = datetime.now(timezone.utc) - timedelta(
        days=max(1, int(config.get("fbs_assembly_display_days", 30) or 30))
    )
    handed_cutoff = datetime.now(timezone.utc) - timedelta(
        days=max(1, int(config.get("fbs_operational_history_days", 7) or 7))
    )

    grouped: dict[str, dict[str, Any]] = {}
    delivered_rows: list[dict[str, Any]] = []

    def ensure_item(sid: str) -> dict[str, Any]:
        return grouped.setdefault(sid, {
            "id": sid,
            "name": sid,
            "createdAt": "",
            "orders": [],
            "confirm": 0,
            "complete": 0,
            "sold": 0,
            "canceled": 0,
            "supply_known": False,
            "supply_done": None,
            "supply_order_count": None,
            "closedAt": "",
            "scanDt": "",
            "last_sync_at": "",
        })

    # First load the supply-level registry. This lets a real open supply exist in
    # the UI even when it is empty or its order-status refresh has not completed.
    for supply in supply_rows or []:
        sid = str(supply.get("supply_id") or supply.get("id") or "").strip()
        if not sid:
            continue
        item = ensure_item(sid)
        item["supply_known"] = True
        item["name"] = str(supply.get("name") or item["name"] or sid)
        item["createdAt"] = str(supply.get("created_at_wb") or supply.get("createdAt") or item["createdAt"] or "")
        item["closedAt"] = str(supply.get("closed_at_wb") or supply.get("closedAt") or "")
        item["scanDt"] = str(supply.get("scan_dt_wb") or supply.get("scanDt") or "")
        item["supply_done"] = bool(supply.get("done"))
        try:
            count = int(supply.get("order_count"))
        except (TypeError, ValueError):
            count = -1
        item["supply_order_count"] = count if count >= 0 else None
        item["last_sync_at"] = str(supply.get("last_sync_at") or "")

    # Then merge order lifecycle state.
    for row in rows:
        sid = str(row.get("supply_id") or "").strip()
        wb_state = str(row.get("wb_status") or "").strip().lower()
        supplier_state = str(row.get("supplier_status") or "").strip().lower()
        if wb_state in final_wb:
            delivered_rows.append(row)
        if not sid:
            continue

        item = ensure_item(sid)
        if not item.get("name") or item.get("name") == sid:
            item["name"] = str(row.get("supply_name") or sid)
        if not item.get("createdAt"):
            item["createdAt"] = str(row.get("supply_created_at_wb") or row.get("created_at_wb") or "")
        item["orders"].append(row)

        if supplier_state == "confirm" and wb_state not in final_wb:
            created_dt = parse_wb_datetime(row.get("created_at_wb"))
            if created_dt is None or created_dt.astimezone(timezone.utc) >= assembly_cutoff:
                item["confirm"] += 1
        if supplier_state == "complete" and wb_state not in final_wb:
            item["complete"] += 1
        if wb_state == "sold":
            item["sold"] += 1
        if wb_state in WB_CANCELED_STATUSES:
            item["canceled"] += 1

    assembling: list[dict[str, Any]] = []
    handed: list[dict[str, Any]] = []
    for item in grouped.values():
        created_dt = parse_wb_datetime(item.get("createdAt"))
        confirmed_open = (
            authoritative_open_supply_ids is not None
            and str(item.get("id") or "") in authoritative_open_supply_ids
        )
        # Age is only a stale-snapshot safety net. A fresh WB response naming a
        # supply as open is authoritative even if that supply was created more
        # than the local display window ago.
        recent_enough = confirmed_open or created_dt is None or created_dt.astimezone(timezone.utc) >= assembly_cutoff
        supply_known = bool(item.get("supply_known"))
        supply_done = item.get("supply_done")
        closed_marker = bool(str(item.get("closedAt") or "").strip() or str(item.get("scanDt") or "").strip())

        # A fresh supply-level `done=True` must veto stale confirm rows. After a
        # successful operational sync, the checkpoint also carries the exact set
        # of open supply IDs observed by WB. That set prevents an old persisted
        # done=0 row from resurrecting when it is absent from the current WB list.
        visible_by_checkpoint = (
            authoritative_open_supply_ids is None
            or str(item.get("id") or "") in authoritative_open_supply_ids
        )
        is_assembling = recent_enough and visible_by_checkpoint and (
            (supply_known and supply_done is False and not closed_marker)
            or (not supply_known and item["confirm"] > 0)
        )
        handed_recent = created_dt is None or created_dt.astimezone(timezone.utc) >= handed_cutoff
        is_handed = (
            not is_assembling
            and handed_recent
            and item["complete"] > 0
        )

        if item.get("supply_order_count") is not None:
            order_count = max(int(item["supply_order_count"]), len(item["orders"]))
        else:
            order_count = len(item["orders"])

        if is_assembling:
            item["order_count"] = order_count
            item["order_count_sort"] = order_count
            item["fbe_status"] = "На сборке"
            item["fbe_status_class"] = "status-assembling"
            item["fbe_status_rank"] = 0
            item["created_at_text"] = format_dt(item.get("createdAt"))
            assembling.append(item)
        elif is_handed:
            item["order_count"] = order_count
            item["order_count_sort"] = order_count
            item["fbe_status"] = "Передано WB"
            item["fbe_status_class"] = "status-transferred"
            item["fbe_status_rank"] = 1
            item["created_at_text"] = format_dt(item.get("createdAt"))
            handed.append(item)

    def sort_key(x: dict[str, Any]):
        return parse_wb_datetime(x.get("createdAt")) or datetime.min.replace(tzinfo=timezone.utc)

    assembling.sort(key=sort_key, reverse=True)
    handed.sort(key=sort_key, reverse=True)
    return {
        "assembling": assembling,
        "handed": handed,
        "delivered_rows": delivered_rows,
        "sold_count": sum(1 for r in delivered_rows if str(r.get("wb_status") or "").lower() == "sold"),
        "canceled_count": sum(1 for r in delivered_rows if str(r.get("wb_status") or "").lower() in WB_CANCELED_STATUSES),
        "delivered_count": len(delivered_rows),
    }


def _fbs_supply_registry_freshness() -> dict[str, Any]:
    """Describe freshness from the last *completed operational WB sync*.

    Row timestamps are deliberately not used as the primary signal: local writes
    (for example creating one supply) must never make an old persisted registry
    look globally fresh.
    """
    state = get_fbs_sync_state(marking_db_path, sync_key="operational")
    details = state.get("details") if isinstance(state.get("details"), dict) else {}
    open_supply_ids = {
        str(value).strip() for value in (details.get("open_supply_ids") or [])
        if str(value).strip()
    }
    raw = str(state.get("completed_at") or "").strip()
    if not raw:
        return {
            "fresh": False,
            "last_sync_at": "",
            "age_seconds": None,
            "sync_ok": False,
            "error_text": str(state.get("error_text") or ""),
            "open_supply_ids": open_supply_ids,
        }
    try:
        dt = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        age = max(0.0, (datetime.now(timezone.utc) - dt.astimezone(timezone.utc)).total_seconds())
    except Exception:
        return {
            "fresh": False,
            "last_sync_at": raw,
            "age_seconds": None,
            "sync_ok": bool(state.get("ok")),
            "error_text": str(state.get("error_text") or ""),
            "open_supply_ids": open_supply_ids,
        }
    max_age = max(30, int(config.get("fbs_operational_stale_ui_seconds", 300) or 300))
    sync_ok = bool(state.get("ok"))
    return {
        "fresh": sync_ok and age <= max_age,
        "last_sync_at": raw,
        "age_seconds": int(age),
        "max_age_seconds": max_age,
        "sync_ok": sync_ok,
        "error_text": str(state.get("error_text") or ""),
        "open_supply_ids": open_supply_ids,
    }


def _fbs_operational_refresh_due(freshness: dict[str, Any]) -> tuple[bool, int]:
    """Bound retries after a failed WB sync while never declaring stale data fresh."""
    if freshness.get("fresh"):
        return False, 0
    age = freshness.get("age_seconds")
    if freshness.get("sync_ok") or age is None:
        return True, 0
    retry_seconds = max(
        10, min(300, int(config.get("fbs_operational_retry_seconds", 30) or 30))
    )
    remaining = max(0, retry_seconds - int(age))
    return remaining == 0, remaining

def _active_supply_selector_from_registry(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    supply_rows = list_fbs_supply_registry_rows(
        marking_db_path, limit=10000, date_from=_registry_history_start()
    )
    freshness = _fbs_supply_registry_freshness()
    open_ids = freshness.get("open_supply_ids") if freshness.get("fresh") else None
    return _fbs_dashboard_supply_groups(
        rows, supply_rows, authoritative_open_supply_ids=open_ids
    )["assembling"]


def supply_daypart(now: datetime | None = None) -> str:
    now = now or moscow_now()
    hour = now.hour
    if 5 <= hour < 12:
        return "Утро"
    if 12 <= hour < 18:
        return "День"
    if 18 <= hour < 24:
        return "Вечер"
    return "Ночь"


def clean_supply_warehouse_name(value: str) -> str:
    name = " ".join(str(value or "").split()).strip()
    name = re.sub(r"^(?:FBS|ФБС)\s*[-—|:]?\s*", "", name, flags=re.IGNORECASE)
    name = re.sub(r"^Склад\s+", "", name, flags=re.IGNORECASE)
    return name or "Склад"


def generate_supply_name(
    existing_supplies: list[dict[str, Any]] | None = None,
    *,
    warehouse_name: str = "Поставка",
    now: datetime | None = None,
) -> str:
    """Generate the v0.79 warehouse-aware FBS supply name.

    Example: ``FBE Ростов | 08.08.26 | Утро``. WB supply IDs are already
    unique, so a human-visible #counter is no longer necessary. If a name with
    the exact same warehouse/date/daypart already exists, a short ``· 2``
    suffix is added only to avoid two indistinguishable rows in the dashboard.
    """
    now = now or moscow_now()
    warehouse = clean_supply_warehouse_name(warehouse_name)
    base = f"FBE {warehouse} | {now.strftime('%d.%m.%y')} | {supply_daypart(now)}"
    existing_names = {str(item.get("name") or "").strip() for item in (existing_supplies or [])}
    if base not in existing_names:
        return base
    index = 2
    while f"{base} · {index}" in existing_names:
        index += 1
    return f"{base} · {index}"


def preview_supply_name(existing_supplies: list[dict[str, Any]] | None = None) -> str:
    return generate_supply_name(existing_supplies, warehouse_name="Поставка")



def should_autogenerate_supply_name(name: str, auto_name_suggested: str = "") -> bool:
    """Treat untouched UI suggestions as auto-generated names.

    The popup shows the suggested name inside an editable input. If parents do
    not touch it, the backend must still generate the final number at submit
    time, otherwise a stale preview can create duplicate #1/#2 names.
    """
    raw = str(name or "").strip()
    suggested = str(auto_name_suggested or "").strip()
    return not raw or (suggested and raw == suggested)


def warehouse_for_selected_order_ids(order_ids: list[int]) -> tuple[int, str] | None:
    wanted = {int(value) for value in order_ids}
    names = seller_warehouse_name_map()
    found: dict[int, str] = {}
    for order in get_new_orders_cached():
        order_id = int(order.get("id") or 0)
        if order_id not in wanted:
            continue
        warehouse_id = order_warehouse_id(order)
        found[warehouse_id] = order_warehouse_name(order, names)
    if len(found) != 1:
        return None
    warehouse_id, warehouse_name = next(iter(found.items()))
    return warehouse_id, warehouse_name


def default_expiration_date() -> str:
    """Default for WB expiration modal: current date + 2 years - 1 day."""
    today = datetime.now().date()
    try:
        target = today.replace(year=today.year + 2) - timedelta(days=1)
    except ValueError:
        target = today + timedelta(days=365 * 2 - 1)
    return target.strftime("%d.%m.%Y")


def dashboard_redirect(**params: Any) -> RedirectResponse:
    clean = {key: str(value) for key, value in params.items() if value not in (None, "")}
    url = "/"
    if clean:
        url += "?" + urlencode(clean)
    return RedirectResponse(url=url, status_code=303)


def add_orders_to_supply_in_batches(supply_id: str, order_ids: list[int]) -> int:
    ids = [int(x) for x in order_ids if int(x) > 0]
    added = 0
    for batch in chunks(ids, 100):
        wb.add_orders_to_supply(supply_id=supply_id, order_ids=batch)
        added += len(batch)
        if len(ids) > 100:
            time.sleep(0.25)
    if ids:
        try:
            # The batch above can add to a non-empty supply. Reconcile against
            # WB's full authoritative membership before calling the exact setter;
            # passing only the newly-added IDs would incorrectly evict older rows.
            full_ids = [int(x) for x in wb.get_supply_order_ids(str(supply_id)) if int(x) > 0]
            set_fbs_supply_order_ids(marking_db_path, supply_id=str(supply_id), order_ids=full_ids)
        except Exception:
            # If the read-after-write call is temporarily unavailable, preserve
            # the known additions without pretending the partial list is exact.
            try:
                upsert_fbs_orders(marking_db_path, [
                    {"id": oid, "supplyId": str(supply_id), "supplierStatus": "confirm"}
                    for oid in ids
                ])
            except Exception:
                pass
    return added


def normalize_expiration_date(value: str) -> str:
    raw = str(value or "").strip()
    parsed: datetime | None = None
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(raw, fmt)
            break
        except ValueError:
            continue
    if not parsed:
        raise ValueError("Введите дату в формате 14.06.2028")

    min_date = datetime.now().date() + timedelta(days=30)
    if parsed.date() < min_date:
        raise ValueError("Срок годности должен быть не меньше чем через 30 дней")
    return parsed.strftime("%d.%m.%Y")


def _norm_need_text(value: Any) -> str:
    return str(value or "").strip()


def _norm_need_category(value: Any) -> str:
    return _norm_need_text(value).lower().replace("ё", "е")


def _strip_volume_package(value: str) -> str:
    """Production volume for perfume accounting.

    In the product catalog, perfume volume may include packaging words such as
    "роллер" / "атомайзер". For production accounting the operator needs the
    aroma and milliliters first, so 3мл роллер -> 3мл.
    """
    value = _norm_need_text(value)
    if not value:
        return ""
    value = re.sub(r"\b(роллер|атомайзер|спрей|флакон)\b", "", value, flags=re.IGNORECASE).strip()
    value = re.sub(r"\s+", " ", value)
    return value


def _lower_first(value: str) -> str:
    value = _norm_need_text(value)
    if not value:
        return ""
    return value[:1].lower() + value[1:]


def _join_clean(parts: list[str]) -> str:
    return " ".join([_norm_need_text(p) for p in parts if _norm_need_text(p)]).strip()


def _requirement_label(product: dict[str, Any]) -> tuple[str, str]:
    """Return (group, label) for the main production unit.

    Rules:
    - Парфюмерия: subcategory does not matter. Count by aroma + volume.
      Example: Черный Опиум 3мл.
    - СПА: subcategory matters. Count by subcategory + product.
      Example: скраб-бальзам Сливочный Пломбир.
    - Набор: count the full set separately for now.
    - Fallback: use FBE display name.
    """
    category_raw = _norm_need_text(product.get("category"))
    category = _norm_need_category(category_raw)
    subcategory = _norm_need_text(product.get("subcategory"))
    product_name = _norm_need_text(product.get("product"))
    volume = _norm_need_text(product.get("volume"))

    if category == "парфюмерия":
        label = _join_clean([product_name, _strip_volume_package(volume)])
        return "Парфюмерия", label or _norm_need_text(product.get("display_name"))

    if category == "спа":
        label = _join_clean([_lower_first(subcategory), product_name, volume])
        return "СПА", label or _norm_need_text(product.get("display_name"))

    if category in {"набор", "наборы"}:
        display = _norm_need_text(product.get("display_name"))
        # Prefix with "набор" so it is clear these are intentionally not exploded
        # into individual SPA/perfume components yet.
        if display.lower().startswith("набор"):
            label = _lower_first(display)
        else:
            label = _join_clean(["набор", product_name or display])
        return "Наборы", label or _norm_need_text(product.get("seller_article"))

    label = _norm_need_text(product.get("display_name")) or _norm_need_text(product.get("seller_article"))
    return category_raw or "Прочее", label


def _sample_requirement_label(product: dict[str, Any]) -> str:
    sample = _norm_need_text(product.get("sample"))
    sample_volume = _strip_volume_package(_norm_need_text(product.get("sample_volume")))
    if not sample:
        return ""
    return _join_clean(["пробник", sample, sample_volume])


def _sorted_requirement_items(counter: dict[str, int]) -> list[dict[str, Any]]:
    return [
        {"label": label, "count": count}
        for label, count in sorted(counter.items(), key=lambda x: (-x[1], x[0].lower()))
        if label and count > 0
    ]


def _finalize_requirements_context(
    total_orders: int,
    recognized_orders: int,
    main_by_group: dict[str, dict[str, int]],
    sample_counter: dict[str, int],
    missing: list[dict[str, str]],
    warnings: list[str],
    catalog: ProductCatalog | None = None,
) -> dict[str, Any]:
    group_order = ["Парфюмерия", "СПА", "Наборы", "Прочее"]
    groups: list[dict[str, Any]] = []
    seen = set()
    for group in group_order + sorted(main_by_group.keys()):
        if group in seen or group not in main_by_group:
            continue
        seen.add(group)
        items = _sorted_requirement_items(main_by_group[group])
        if items:
            groups.append({"name": group, "items": items, "total": sum(item["count"] for item in items)})

    if catalog is not None:
        if catalog.error:
            warnings.append(f"Каталог: {catalog.error}")
        warnings.extend(catalog.warnings[:10])

    return {
        "total_orders": total_orders,
        "recognized_orders": recognized_orders,
        "groups": groups,
        "samples": _sorted_requirement_items(sample_counter),
        "missing": missing,
        "warnings": warnings[:25],
    }


def build_supply_requirements_from_articles(articles: list[str]) -> dict[str, Any]:
    """Aggregate production needs locally from seller articles and the SQLite product catalog.

    WB is not queried here at all. The browser already has the seller articles
    in the visible order table, so the accounting button sends only that local
    list to this function. This keeps the modal fast and prevents WB/API issues
    from breaking the production-needs calculation.
    """
    normalized = [str(article or "").strip() for article in articles]
    main_by_group: dict[str, dict[str, int]] = {}
    sample_counter: dict[str, int] = {}
    missing: list[dict[str, str]] = []
    warnings: list[str] = []
    recognized_orders = 0

    try:
        catalog = get_catalog()
    except Exception as exc:
        return {
            "total_orders": len(normalized),
            "recognized_orders": 0,
            "groups": [],
            "samples": [],
            "missing": [
                {"order_id": str(i + 1), "seller_article": article or "—"}
                for i, article in enumerate(normalized)
            ],
            "warnings": [f"Не удалось открыть каталог товаров: {exc}"],
        }

    for index, article in enumerate(normalized, start=1):
        try:
            if not article:
                missing.append({"order_id": str(index), "seller_article": "—"})
                continue

            product = catalog.find_by_article(article)
            if not product:
                missing.append({"order_id": str(index), "seller_article": article})
                continue

            group, label = _requirement_label(product)
            label = label or article or "без названия"
            bucket = main_by_group.setdefault(group or "Прочее", {})
            bucket[label] = bucket.get(label, 0) + 1
            recognized_orders += 1

            sample_label = _sample_requirement_label(product)
            if sample_label:
                sample_counter[sample_label] = sample_counter.get(sample_label, 0) + 1
        except Exception as exc:
            missing.append({"order_id": str(index), "seller_article": article or "—"})
            warnings.append(f"Не удалось посчитать артикул {article or '—'}: {exc}")

    return _finalize_requirements_context(
        total_orders=len(normalized),
        recognized_orders=recognized_orders,
        main_by_group=main_by_group,
        sample_counter=sample_counter,
        missing=missing,
        warnings=warnings,
        catalog=catalog,
    )


def build_supply_requirements_from_orders(orders: list[dict[str, Any]]) -> dict[str, Any]:
    """Compatibility wrapper: derive seller articles from known WB orders.

    The preferred UI path uses build_supply_requirements_from_articles(), so it
    does not make any WB calls when the accounting modal is opened.
    """
    return build_supply_requirements_from_articles([seller_article_from_order(order) for order in orders])


def requirements_context_for_order_ids(order_ids: list[int], orders_pool: list[dict[str, Any]]) -> dict[str, Any]:
    requested: list[int] = []
    for raw in order_ids:
        try:
            requested.append(int(raw))
        except Exception:
            continue

    by_id: dict[int, dict[str, Any]] = {}
    for order in orders_pool:
        raw_id = order.get("id") or order.get("orderId")
        try:
            if raw_id is not None:
                by_id[int(raw_id)] = order
        except Exception:
            continue

    orders = [by_id[oid] for oid in requested if oid in by_id]
    missing_ids = [oid for oid in requested if oid not in by_id]
    ctx = build_supply_requirements_from_orders(orders)
    for oid in missing_ids:
        ctx["missing"].append({"order_id": str(oid), "seller_article": "не найдено в текущих заказах WB"})
    ctx["total_orders"] = len(requested)
    return ctx


def sticker_ext(sticker_type: str) -> str:
    return "zpl" if sticker_type.startswith("zpl") else sticker_type


def chunks(items: list[int], size: int) -> list[list[int]]:
    return [items[i:i + size] for i in range(0, len(items), size)]


def format_sticker_number(part_a: str | None, part_b: str | None) -> dict[str, str]:
    part_a = str(part_a or "").strip()
    part_b = str(part_b or "").strip()
    if not part_a and not part_b:
        return {"part_a": "", "part_b": "", "full": "", "display": "—"}
    return {
        "part_a": part_a,
        "part_b": part_b,
        "full": f"{part_a}{part_b}",
        "display": f"{part_a} {part_b}".strip(),
    }


def sticker_metadata_path(supply_id: str, base_dir: str | Path | None = None) -> Path:
    root = (
        _runtime_path("sticker_metadata_dir", "data/sticker_meta")
        if base_dir is None
        else Path(base_dir)
    )
    if not root.is_absolute():
        root = (BASE_DIR / root).resolve()
    return root / f"{supply_id}.json"


def load_sticker_metadata(supply_id: str) -> dict[int, dict[str, str]]:
    path = sticker_metadata_path(supply_id)
    if not path.exists():
        return {}
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    result: dict[int, dict[str, str]] = {}
    for key, value in raw.items():
        try:
            result[int(key)] = value
        except Exception:
            continue
    return result


def save_sticker_metadata(supply_id: str, metadata: dict[int, dict[str, str]]) -> None:
    path = sticker_metadata_path(supply_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {str(k): v for k, v in metadata.items()}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def get_or_fetch_sticker_metadata(supply_id: str, order_ids: list[int]) -> dict[int, dict[str, str]]:
    """Return partA/partB sticker numbers for supply rows.

    WB displays the sticker number as partA + partB, for example:
    5504117 + 0266 -> 55041170266.
    We cache this in data/sticker_meta so opening the same supply does not
    repeatedly request stickers from WB.
    """
    requested_ids = [int(x) for x in order_ids]
    if not requested_ids:
        return {}

    metadata = load_sticker_metadata(supply_id)
    missing = [oid for oid in requested_ids if oid not in metadata]
    if not missing or not config.get("fetch_sticker_numbers_for_supply_view", True):
        return metadata

    sticker_type = str(config.get("sticker_number_lookup_type", config.get("pdf_sticker_type", "svg")))
    for batch in chunks(missing, 100):
        try:
            stickers = wb.get_order_stickers(
                order_ids=batch,
                sticker_type=sticker_type,
                width=int(config.get("wb_sticker_width", 58)),
                height=int(config.get("wb_sticker_height", 40)),
            )
        except Exception:
            continue

        for sticker in stickers:
            number = format_sticker_number(sticker.part_a, sticker.part_b)
            metadata[int(sticker.order_id)] = {
                "part_a": number["part_a"],
                "part_b": number["part_b"],
                "full": number["full"],
                "barcode": str(sticker.barcode or ""),
            }

    save_sticker_metadata(supply_id, metadata)
    return metadata


def save_order_stickers(supply_id: str, order_ids: list[int], sticker_type: str, base_dir: str | Path) -> list[Path]:
    """Download WB stickers in batches and return file paths in requested order.

    WB limits sticker requests to 100 order IDs. Earlier versions already passed
    lists, but this function now explicitly chunks large supplies and reorders the
    result by the original order_ids list so the final PDF keeps the WB supply
    order exactly.
    """
    requested_ids = [int(x) for x in order_ids]
    if not requested_ids:
        return []

    sticker_dir = Path(base_dir) / supply_id
    sticker_dir.mkdir(parents=True, exist_ok=True)
    ext = sticker_ext(sticker_type)

    saved_by_order_id: dict[int, Path] = {}
    metadata = load_sticker_metadata(supply_id)
    for batch in chunks(requested_ids, 100):
        stickers = wb.get_order_stickers(
            order_ids=batch,
            sticker_type=sticker_type,
            width=int(config.get("wb_sticker_width", 58)),
            height=int(config.get("wb_sticker_height", 40)),
        )
        for sticker in stickers:
            payload = wb.decode_sticker_file(sticker.file, sticker_type)
            path = sticker_dir / f"{sticker.order_id}.{ext}"
            path.write_bytes(payload)
            saved_by_order_id[int(sticker.order_id)] = path

            number = format_sticker_number(sticker.part_a, sticker.part_b)
            metadata[int(sticker.order_id)] = {
                "part_a": number["part_a"],
                "part_b": number["part_b"],
                "full": number["full"],
                "barcode": str(sticker.barcode or ""),
            }

    save_sticker_metadata(supply_id, metadata)
    return [sticker_dir / f"{oid}.{ext}" for oid in requested_ids if (sticker_dir / f"{oid}.{ext}").exists()]


def get_supply_orders(
    supply_id: str,
    *,
    known_order_ids: Iterable[int] | None = None,
) -> list[dict[str, Any]]:
    """Load supply orders, reusing a freshly fetched WB membership when given.

    Transfer preflight already asks WB for the authoritative order set. Reusing
    it avoids an immediate duplicate ``/order-ids`` request while retaining the
    cached lookup for ordinary supply-page rendering.
    """
    order_ids = (
        sorted({int(value) for value in known_order_ids if int(value) > 0})
        if known_order_ids is not None
        else get_supply_order_ids_cached(supply_id)
    )
    all_orders = get_all_orders_cached()
    order_by_id = {int(order["id"]): order for order in all_orders}
    orders = [order_by_id[oid] for oid in order_ids if oid in order_by_id]

    missing_ids = [oid for oid in order_ids if oid not in order_by_id]
    for oid in missing_ids:
        orders.append({"id": oid, "article": "", "skus": [], "createdAt": "", "supplyId": supply_id})
    save_supply_articles_snapshot(supply_id, orders)
    return orders


def supply_orders_context(supply_id: str, sort: str | None = None) -> list[dict[str, Any]]:
    orders = get_supply_orders(supply_id)
    enriched = enrich_orders(orders, sort=sort)

    direct_type = config.get("wb_sticker_type", "png")
    direct_ext = sticker_ext(direct_type)
    sticker_dir = _runtime_path("sticker_cache_dir", "data/stickers") / supply_id

    # Never download WB stickers merely because a supply page was opened.
    # Printing/PDF actions fetch missing stickers on demand.
    sticker_meta = load_sticker_metadata(supply_id)

    for item in enriched:
        path = sticker_dir / f"{item['id']}.{direct_ext}"
        meta = sticker_meta.get(int(item["id"]), {})
        item["wb_sticker_ready"] = path.exists()
        item["wb_sticker_path"] = str(path)
        item["sticker_part_a"] = meta.get("part_a", "")
        item["sticker_part_b"] = meta.get("part_b", "")
        item["sticker_number"] = meta.get("full", "")
        item["sticker_display"] = f"{meta.get('part_a', '')} {meta.get('part_b', '')}".strip()
        item["sticker_barcode"] = meta.get("barcode", "")
        item["sticker_sort"] = int(meta.get("full") or 0) if str(meta.get("full") or "").isdigit() else int(item["id"])
    return enriched




def inventory_path() -> Path:
    return Path(config.get("inventory_path", "data/inventory.xlsx"))


def _inv_clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _inv_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace("\u00a0", "").replace(",", ".")
    try:
        return float(text)
    except Exception:
        return 0.0


def _inv_header_map(ws) -> dict[str, int]:
    headers = [_inv_clean(cell.value) for cell in ws[1]]
    return {h: i for i, h in enumerate(headers) if h}


def _inv_get(row: tuple[Any, ...], idx: dict[str, int], names: list[str], default: Any = "") -> Any:
    for name in names:
        if name in idx:
            pos = idx[name]
            return row[pos] if pos < len(row) else default
    return default


def inventory_context() -> dict[str, Any]:
    """Read local inventory.xlsx and return a small stock dashboard.

    Parents can fill the `Начальный остаток` column in `materials`; FBE then shows
    current stock as: initial + purchases + movements. No WB calls are made here.
    """
    path = inventory_path()
    ctx: dict[str, Any] = {
        "path": str(path),
        "exists": path.exists(),
        "error": "",
        "materials": [],
        "summary": {"total": 0, "ok": 0, "low": 0, "zero": 0, "empty": 0},
        "warnings": [],
    }
    if not path.exists():
        ctx["error"] = f"Файл inventory.xlsx не найден: {path}"
        return ctx

    try:
        from openpyxl import load_workbook
        wb_inv = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        ctx["error"] = f"Не удалось открыть inventory.xlsx: {exc}"
        return ctx

    if "materials" not in wb_inv.sheetnames:
        ctx["error"] = "В inventory.xlsx не найден лист materials"
        return ctx

    purchases_qty: dict[str, float] = {}
    purchases_sum: dict[str, float] = {}
    if "purchases" in wb_inv.sheetnames:
        ws = wb_inv["purchases"]
        idx = _inv_header_map(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            mid = _inv_clean(_inv_get(row, idx, ["material_id", "Материал", "material id"]))
            if not mid:
                continue
            qty = _inv_float(_inv_get(row, idx, ["Количество", "qty", "quantity"]))
            total = _inv_float(_inv_get(row, idx, ["Сумма", "sum", "amount", "cost"]))
            purchases_qty[mid] = purchases_qty.get(mid, 0.0) + qty
            purchases_sum[mid] = purchases_sum.get(mid, 0.0) + total

    movements_qty: dict[str, float] = {}
    movements_sum: dict[str, float] = {}
    if "movements" in wb_inv.sheetnames:
        ws = wb_inv["movements"]
        idx = _inv_header_map(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            mid = _inv_clean(_inv_get(row, idx, ["material_id", "Материал", "material id"]))
            if not mid:
                continue
            qty = _inv_float(_inv_get(row, idx, ["Количество", "qty", "quantity"]))
            total = _inv_float(_inv_get(row, idx, ["Сумма", "sum", "amount", "cost"]))
            movements_qty[mid] = movements_qty.get(mid, 0.0) + qty
            movements_sum[mid] = movements_sum.get(mid, 0.0) + total

    ws = wb_inv["materials"]
    idx = _inv_header_map(ws)
    required = ["material_id", "Наименование", "Единица"]
    missing_headers = [h for h in required if h not in idx]
    if missing_headers:
        ctx["warnings"].append("В materials не найдены колонки: " + ", ".join(missing_headers))

    materials: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        mid = _inv_clean(_inv_get(row, idx, ["material_id", "Материал", "material id"]))
        if not mid:
            continue
        name = _inv_clean(_inv_get(row, idx, ["Наименование", "Название", "name"])) or mid
        group = _inv_clean(_inv_get(row, idx, ["Группа", "Категория", "group"]))
        unit = _inv_clean(_inv_get(row, idx, ["Единица", "unit"]))
        min_stock = _inv_float(_inv_get(row, idx, ["Мин. остаток", "Минимальный остаток", "min_stock"]))
        initial = _inv_float(_inv_get(row, idx, ["Начальный остаток", "Стартовый остаток", "initial_stock", "opening_stock"]))
        unit_price = _inv_float(_inv_get(row, idx, ["Цена за единицу", "Средняя цена", "Цена", "unit_price"]))
        initial_sum = _inv_float(_inv_get(row, idx, ["Сумма начального остатка", "Начальная сумма", "opening_value"]))
        if not unit_price and initial and initial_sum:
            unit_price = initial_sum / initial

        purchase_qty = purchases_qty.get(mid, 0.0)
        movement_qty = movements_qty.get(mid, 0.0)
        current = initial + purchase_qty + movement_qty

        cost_base_qty = initial + purchase_qty
        cost_base_sum = (initial * unit_price if unit_price else initial_sum) + purchases_sum.get(mid, 0.0)
        avg_price = (cost_base_sum / cost_base_qty) if cost_base_qty else 0.0
        stock_value = current * avg_price if avg_price else 0.0

        if current <= 0:
            status = "Нет остатка"
            status_class = "bad"
            ctx["summary"]["zero"] += 1
        elif min_stock and current < min_stock:
            status = "Заканчивается"
            status_class = "warn"
            ctx["summary"]["low"] += 1
        else:
            status = "ОК"
            status_class = "ok"
            ctx["summary"]["ok"] += 1
        if initial == 0 and purchase_qty == 0:
            ctx["summary"]["empty"] += 1

        materials.append({
            "material_id": mid,
            "group": group,
            "name": name,
            "unit": unit,
            "min_stock": min_stock,
            "initial": initial,
            "purchase_qty": purchase_qty,
            "movement_qty": movement_qty,
            "current": current,
            "unit_price": avg_price or unit_price,
            "stock_value": stock_value,
            "status": status,
            "status_class": status_class,
        })

    materials.sort(key=lambda x: (str(x.get("group") or ""), str(x.get("name") or "")))
    ctx["materials"] = materials
    ctx["summary"]["total"] = len(materials)
    ctx["total_value"] = sum(float(m.get("stock_value") or 0) for m in materials)
    return ctx


def inventory_writeoff_state_path() -> Path:
    return Path(config.get("inventory_writeoff_state_path", "data/cache/inventory_writeoffs.json"))


def read_inventory_writeoff_state() -> dict[str, Any]:
    path = inventory_writeoff_state_path()
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def write_inventory_writeoff_state(state: dict[str, Any]) -> None:
    path = inventory_writeoff_state_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def supply_articles_snapshot_path(supply_id: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(supply_id or ""))
    return Path(config.get("supply_articles_cache_dir", "data/cache/supply_articles")) / f"{safe}.json"


def save_supply_articles_snapshot(supply_id: str, orders: list[dict[str, Any]]) -> None:
    articles: list[str] = []
    order_rows: list[dict[str, Any]] = []
    for order in orders or []:
        article = seller_article_from_order(order)
        if not article:
            continue
        raw_id = order.get("id") or order.get("orderId") or ""
        articles.append(article)
        order_rows.append({"order_id": str(raw_id), "seller_article": article})
    if not articles:
        return
    payload = {
        "supply_id": str(supply_id),
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "articles": articles,
        "orders": order_rows,
    }
    path = supply_articles_snapshot_path(supply_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_supply_articles_snapshot(supply_id: str) -> list[str]:
    path = supply_articles_snapshot_path(supply_id)
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return []
    articles = data.get("articles") if isinstance(data, dict) else None
    if not isinstance(articles, list):
        return []
    return [str(x or "").strip() for x in articles if str(x or "").strip()]


def normalize_inventory_code(value: Any, width: int = 2) -> str:
    text = _inv_clean(value)
    if not text:
        return ""
    if text.endswith(".0"):
        text = text[:-2]
    if text.isdigit():
        return text.zfill(width)
    return text


def inventory_has_writeoff(supply_id: str, wb_inv=None) -> bool:
    state = read_inventory_writeoff_state()
    if str(supply_id) in state:
        return True
    close_after = False
    if wb_inv is None:
        path = inventory_path()
        if not path.exists():
            return False
        try:
            from openpyxl import load_workbook
            wb_inv = load_workbook(path, data_only=True, read_only=True)
            close_after = True
        except Exception:
            return False
    try:
        if "movements" not in wb_inv.sheetnames:
            return False
        ws = wb_inv["movements"]
        idx = _inv_header_map(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_supply = _inv_clean(_inv_get(row, idx, ["supplyId", "supply_id", "Поставка"]))
            row_type = _inv_clean(_inv_get(row, idx, ["Тип", "type"]))
            if row_supply == str(supply_id) and row_type in {"writeoff", "auto_writeoff"}:
                return True
    finally:
        if close_after:
            try:
                wb_inv.close()
            except Exception:
                pass
    return False


def inventory_table_rows(ws) -> list[dict[str, Any]]:
    idx = _inv_header_map(ws)
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue
        item = {header: (row[pos] if pos < len(row) else None) for header, pos in idx.items()}
        rows.append(item)
    return rows


def inventory_material_prices(wb_inv) -> dict[str, float]:
    """Weighted average price from initial balances and purchases."""
    base_qty: dict[str, float] = {}
    base_sum: dict[str, float] = {}
    if "materials" in wb_inv.sheetnames:
        ws = wb_inv["materials"]
        idx = _inv_header_map(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            mid = _inv_clean(_inv_get(row, idx, ["material_id", "Материал", "material id"]))
            if not mid:
                continue
            initial = _inv_float(_inv_get(row, idx, ["Начальный остаток", "Стартовый остаток", "initial_stock", "opening_stock"]))
            unit_price = _inv_float(_inv_get(row, idx, ["Цена за единицу", "Средняя цена", "Цена", "unit_price"]))
            initial_sum = _inv_float(_inv_get(row, idx, ["Сумма начального остатка", "Начальная сумма", "opening_value"]))
            if initial and not initial_sum and unit_price:
                initial_sum = initial * unit_price
            if initial:
                base_qty[mid] = base_qty.get(mid, 0.0) + initial
                base_sum[mid] = base_sum.get(mid, 0.0) + initial_sum
    if "purchases" in wb_inv.sheetnames:
        ws = wb_inv["purchases"]
        idx = _inv_header_map(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            mid = _inv_clean(_inv_get(row, idx, ["material_id", "Материал", "material id"]))
            if not mid:
                continue
            qty = _inv_float(_inv_get(row, idx, ["Количество", "qty", "quantity"]))
            total = _inv_float(_inv_get(row, idx, ["Сумма", "sum", "amount", "cost"]))
            if qty:
                base_qty[mid] = base_qty.get(mid, 0.0) + qty
                base_sum[mid] = base_sum.get(mid, 0.0) + total
    prices: dict[str, float] = {}
    for mid, qty in base_qty.items():
        prices[mid] = (base_sum.get(mid, 0.0) / qty) if qty else 0.0
    return prices


def inventory_product_key_from_article(article: str, product_keys: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    article = str(article or "").strip()
    if not article:
        return None
    found = product_keys.get(article)
    if found:
        return found
    # Fallback for the user's PRF system, e.g. PRF-R0308P0101.
    m = re.search(r"^(?:PRF|PRW|TOW)-([RSF]\d{2})(\d{2})(?:P(\d{2})(\d{2}))?", article, flags=re.IGNORECASE)
    if not m:
        return None
    fmt, aroma, sample_fmt_suffix, sample_aroma = m.groups()
    sample_format = f"P{sample_fmt_suffix}" if sample_fmt_suffix else ""
    return {
        "seller_article": article,
        "Категория": "Парфюмерия",
        "format_code": fmt.upper(),
        "aroma_code": normalize_inventory_code(aroma),
        "production_key": f"perfume::{fmt.upper()}::{normalize_inventory_code(aroma)}",
        "sample_volume_code": sample_fmt_suffix or "",
        "sample_aroma_code": normalize_inventory_code(sample_aroma) if sample_aroma else "",
        "sample_production_key": f"sample::{sample_format}::{normalize_inventory_code(sample_aroma)}" if sample_format and sample_aroma else "",
    }


def build_inventory_writeoff_movements(articles: list[str], supply_id: str) -> tuple[list[dict[str, Any]], list[str]]:
    path = inventory_path()
    if not path.exists():
        return [], [f"Файл inventory.xlsx не найден: {path}"]

    from openpyxl import load_workbook
    wb_inv = load_workbook(path, data_only=True, read_only=True)
    warnings: list[str] = []
    try:
        if "product_keys" not in wb_inv.sheetnames:
            return [], ["В inventory.xlsx не найден лист product_keys"]
        if "format_recipes" not in wb_inv.sheetnames:
            return [], ["В inventory.xlsx не найден лист format_recipes"]

        product_keys: dict[str, dict[str, Any]] = {}
        for row in inventory_table_rows(wb_inv["product_keys"]):
            article = _inv_clean(row.get("seller_article") or row.get("Артикул продавца"))
            if article:
                product_keys[article] = row

        recipes_by_format: dict[str, list[dict[str, Any]]] = {}
        for row in inventory_table_rows(wb_inv["format_recipes"]):
            fmt = _inv_clean(row.get("format_code"))
            if not fmt:
                continue
            recipes_by_format.setdefault(fmt, []).append(row)

        prices = inventory_material_prices(wb_inv)
    finally:
        try:
            wb_inv.close()
        except Exception:
            pass

    aggregated: dict[tuple[str, str, str, str], dict[str, Any]] = {}

    def add_recipe(seller_article: str, production_key: str, format_code: str, aroma_code: str, multiplier: int = 1) -> None:
        format_code = _inv_clean(format_code)
        aroma_code = normalize_inventory_code(aroma_code)
        if not format_code or not aroma_code:
            return
        recipe_rows = recipes_by_format.get(format_code) or []
        if not recipe_rows:
            warnings.append(f"Нет рецептуры format_code={format_code} для {seller_article}")
            return
        for recipe in recipe_rows:
            raw_ref = _inv_clean(recipe.get("material_ref") or recipe.get("material_id"))
            if not raw_ref:
                continue
            material_id = raw_ref.replace("{aroma_code}", aroma_code)
            qty_per_unit = _inv_float(recipe.get("qty_per_unit") or recipe.get("Количество на 1 шт"))
            loss_pct = _inv_float(recipe.get("loss_pct") or recipe.get("Потери %"))
            qty = qty_per_unit * (1 + loss_pct / 100.0) * multiplier
            if not qty:
                continue
            unit = _inv_clean(recipe.get("unit") or recipe.get("Единица"))
            key = (seller_article, production_key, material_id, unit)
            item = aggregated.setdefault(key, {
                "seller_article": seller_article,
                "production_key": production_key,
                "material_id": material_id,
                "unit": unit,
                "qty": 0.0,
                "sum": 0.0,
            })
            item["qty"] += qty
            item["sum"] += qty * prices.get(material_id, 0.0)

    for article in [str(a or "").strip() for a in articles if str(a or "").strip()]:
        key_row = inventory_product_key_from_article(article, product_keys)
        if not key_row:
            warnings.append(f"Артикул не найден в inventory product_keys: {article}")
            continue
        category = _norm_need_category(key_row.get("Категория") or key_row.get("category"))
        if category and category != "парфюмерия":
            warnings.append(f"Пропущен не парфюмерный артикул: {article}")
            continue
        fmt = _inv_clean(key_row.get("format_code"))
        aroma = normalize_inventory_code(key_row.get("aroma_code"))
        production_key = _inv_clean(key_row.get("production_key")) or f"perfume::{fmt}::{aroma}"
        add_recipe(article, production_key, fmt, aroma)

        sample_production_key = _inv_clean(key_row.get("sample_production_key"))
        sample_aroma = normalize_inventory_code(key_row.get("sample_aroma_code"))
        sample_format = ""
        if sample_production_key:
            parts = sample_production_key.split("::")
            if len(parts) >= 3:
                sample_format = parts[1]
                sample_aroma = sample_aroma or normalize_inventory_code(parts[2])
        if not sample_format:
            sample_suffix = normalize_inventory_code(key_row.get("sample_volume_code"))
            sample_format = f"P{sample_suffix}" if sample_suffix else ""
        if sample_format and sample_aroma:
            add_recipe(article, sample_production_key or f"sample::{sample_format}::{sample_aroma}", sample_format, sample_aroma)

    movement_rows: list[dict[str, Any]] = []
    for item in aggregated.values():
        qty = round(float(item["qty"]), 6)
        total = round(float(item["sum"]), 2)
        movement_rows.append({
            "Дата": datetime.now().strftime("%d.%m.%Y %H:%M:%S"),
            "Тип": "writeoff",
            "supplyId": str(supply_id),
            "seller_article": item["seller_article"],
            "production_key": item["production_key"],
            "material_id": item["material_id"],
            "Количество": -qty,
            "Единица": item["unit"],
            "Сумма": -total if total else 0,
            "Комментарий": "автосписание после выхода поставки из активных",
        })
    return movement_rows, warnings[:100]


def first_empty_movements_row(ws) -> int:
    for r in range(2, max(ws.max_row, 2) + 1):
        if all(ws.cell(r, c).value in (None, "") for c in range(1, 11)):
            return r
    return ws.max_row + 1


def append_inventory_movements(rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    path = inventory_path()
    from openpyxl import load_workbook
    wb_inv = load_workbook(path)
    try:
        if "movements" not in wb_inv.sheetnames:
            ws = wb_inv.create_sheet("movements")
            headers = ["Дата", "Тип", "supplyId", "seller_article", "production_key", "material_id", "Количество", "Единица", "Сумма", "Комментарий"]
            for col, header in enumerate(headers, start=1):
                ws.cell(row=1, column=col, value=header)
        ws = wb_inv["movements"]
        headers = [_inv_clean(cell.value) for cell in ws[1]]
        if not headers or not headers[0]:
            headers = ["Дата", "Тип", "supplyId", "seller_article", "production_key", "material_id", "Количество", "Единица", "Сумма", "Комментарий"]
            for col, header in enumerate(headers, start=1):
                ws.cell(row=1, column=col, value=header)
        idx = {header: i + 1 for i, header in enumerate(headers) if header}
        row_num = first_empty_movements_row(ws)
        for row in rows:
            for key, value in row.items():
                col = idx.get(key)
                if col:
                    ws.cell(row=row_num, column=col, value=value)
            row_num += 1
        wb_inv.save(path)
    finally:
        try:
            wb_inv.close()
        except Exception:
            pass


def inventory_writeoff_supply_if_needed(supply: dict[str, Any]) -> dict[str, Any] | None:
    if not bool(config.get("inventory_auto_writeoff_enabled", True)):
        return None
    supply_id = str(supply.get("id") or "").strip()
    if not supply_id:
        return None
    if supply.get("fbe_status") == "На сборке":
        return None
    if inventory_has_writeoff(supply_id):
        return None

    articles: list[str] = []
    try:
        orders = get_supply_orders(supply_id)
        articles = [seller_article_from_order(order) for order in orders if seller_article_from_order(order)]
        if articles:
            save_supply_articles_snapshot(supply_id, orders)
    except Exception:
        articles = []

    if not articles:
        articles = load_supply_articles_snapshot(supply_id)
    if not articles:
        return {"supply_id": supply_id, "ok": False, "message": "нет сохраненных артикулов для списания"}

    rows, warnings = build_inventory_writeoff_movements(articles, supply_id)
    if not rows:
        return {"supply_id": supply_id, "ok": False, "message": "нет строк списания", "warnings": warnings}
    append_inventory_movements(rows)

    state = read_inventory_writeoff_state()
    state[supply_id] = {
        "written_at": datetime.now().isoformat(timespec="seconds"),
        "supply_name": str(supply.get("name") or ""),
        "status": str(supply.get("fbe_status") or ""),
        "articles_count": len(articles),
        "movement_rows": len(rows),
        "warnings": warnings[:20],
    }
    write_inventory_writeoff_state(state)
    invalidate_cache("inventory")
    return {"supply_id": supply_id, "ok": True, "rows": len(rows), "warnings": warnings}


def auto_writeoff_inactive_supplies(supplies: list[dict[str, Any]]) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    if not bool(config.get("inventory_auto_writeoff_enabled", True)):
        return results
    for supply in supplies:
        try:
            result = inventory_writeoff_supply_if_needed(supply)
            if result:
                results.append(result)
        except Exception as exc:
            results.append({"supply_id": str(supply.get("id") or ""), "ok": False, "message": str(exc)})
    return results


def get_fbs_operational_sync_cached() -> dict[str, Any]:
    # Operational state is intentionally fresher than analytics/history. Cap old
    # configs at 90 s so an upgraded installation cannot keep a 3-minute stale
    # view of supplies. The request itself still runs stale-while-revalidate.
    configured_ttl = int(config.get("fbs_operational_sync_ttl_seconds", 60) or 60)
    ttl = min(90, max(30, configured_ttl))
    return cached_stale_while_revalidate(
        "fbs_operational_registry",
        ttl,
        lambda: _sync_fbs_lifecycle_registry(include_supplies=True, deep=False),
        lambda: {"local": True, "errors": []},
    )


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, msg: str | None = None, error: str | None = None, sort: str | None = None, refresh: int | None = None, open_supply: str | None = None):
    if refresh:
        # Manual refresh is intentionally synchronous and narrowly scoped to new
        # orders. The old implementation cleared the entire dashboard cache first,
        # so the stale-while-revalidate fallback could temporarily render 0 orders.
        # We now keep the last good cache until WB has returned a valid response.
        try:
            fresh_orders = refresh_new_orders_now()
            return RedirectResponse(
                url="/?" + urlencode({"msg": f"Новые заказы WB обновлены: {len(fresh_orders)}"}) + "#wb",
                status_code=303,
            )
        except Exception as exc:
            return RedirectResponse(
                url="/?" + urlencode({"error": f"Не удалось обновить новые заказы WB. Последние данные сохранены: {exc}"}) + "#wb",
                status_code=303,
            )
    catalog = get_catalog()
    try:
        new_orders = enrich_orders(get_new_orders_cached(), sort=sort)
    except Exception as exc:
        new_orders = []
        error = error or str(exc)

    # Do not start a second /supplies request merely to generate a suggested name.
    # The operational synchronizer below is the single owner of dashboard supply
    # refresh; use the local snapshot for immediate naming/render fallback.
    try:
        supplies = prepare_supplies(_local_supply_snapshot())
    except Exception:
        supplies = []

    freshness_before = _fbs_supply_registry_freshness()
    refresh_due, _ = _fbs_operational_refresh_due(freshness_before)
    if refresh_due:
        # A persisted cache from an older run must not keep an in-memory sync cache
        # fresh. Force one background operational refresh on this page load.
        invalidate_cache("fbs_operational_registry")

    sync_info: dict[str, Any] = {}
    if freshness_before.get("fresh") or refresh_due:
        try:
            sync_info = get_fbs_operational_sync_cached()
        except Exception as exc:
            sync_info = {"errors": [str(exc)]}
    else:
        sync_info = {
            "local": True,
            "backoff": True,
            "errors": [str(freshness_before.get("error_text") or "WB временно недоступен")],
        }

    history_start = _registry_history_start()
    registry_rows = list_fbs_registry_rows(marking_db_path, limit=10000, date_from=history_start)
    supply_registry_rows = list_fbs_supply_registry_rows(
        marking_db_path, limit=10000, date_from=history_start
    )
    registry_freshness = _fbs_supply_registry_freshness()
    authoritative_open_ids = (
        registry_freshness.get("open_supply_ids")
        if registry_freshness.get("fresh") else None
    )
    supply_groups = _fbs_dashboard_supply_groups(
        registry_rows, supply_registry_rows,
        authoritative_open_supply_ids=authoritative_open_ids,
    )
    active_dashboard_supplies = supply_groups["assembling"]
    transferred_dashboard_supplies = supply_groups["handed"]
    fbs_sync_pending = not bool(registry_freshness.get("fresh"))
    if fbs_sync_pending:
        # Showing known-stale operational rows is worse than showing an explicit
        # sync state: this is exactly what caused old supplies to flash on startup.
        active_dashboard_supplies = []
        transferred_dashboard_supplies = []
    delivered_count = int(supply_groups["delivered_count"] or 0)
    delivered_sold_count = int(supply_groups["sold_count"] or 0)
    delivered_canceled_count = int(supply_groups["canceled_count"] or 0)

    lifecycle_rows = list_fbs_lifecycle_rows(marking_db_path, limit=10000, date_from=history_start)
    post_sale_summary = {
        "at_wb": sum(1 for r in lifecycle_rows if str(r.get("supplier_status") or "").lower() == "complete" and str(r.get("wb_status") or "").lower() not in WB_TERMINAL_STATUSES),
        "ready": sum(1 for r in lifecycle_rows if str(r.get("wb_status") or "").lower() == "sold" and str(r.get("circulation_status") or "").upper() == "INTRODUCED" and str(r.get("post_sale_status") or "") not in {"retired", "return_submitted", "returned"}),
        "retired": sum(1 for r in lifecycle_rows if str(r.get("post_sale_status") or "") in {"retired", "retirement_submitted"}),
        "returns": sum(1 for r in lifecycle_rows if str(r.get("post_sale_status") or "") in {"return_submitted", "returned"}),
        "sold": delivered_sold_count,
        "canceled": delivered_canceled_count,
    }

    sample_aromas = catalog.list_sample_aromas() or []
    name_aromas = catalog.list_name_aromas() or []
    volumes = catalog.list_name_formats() or []

    return templates.TemplateResponse(
        request,
        "dashboard.html",
        {
            "request": request,
            "new_orders": new_orders,
            "supplies": supplies,
            "active_supplies": active_dashboard_supplies,
            "active_dashboard_supplies": active_dashboard_supplies,
            "transferred_dashboard_supplies": transferred_dashboard_supplies,
            "delivered_count": delivered_count,
            "delivered_sold_count": delivered_sold_count,
            "delivered_canceled_count": delivered_canceled_count,
            "history_start": history_start,
            "fbs_sync_info": sync_info,
            "fbs_sync_pending": fbs_sync_pending,
            "fbs_registry_freshness": registry_freshness,
            "dashboard_rendered_at": int(time.time()),
            "catalog_error": catalog.error,
            "catalog_warnings": catalog.warnings[:20],
            "msg": msg,
            "error": error,
            "default_expiration": default_expiration_date(),
            "config": config,
            "sort": sort or config.get("default_order_sort", "urgent"),
            "sample_aromas": sample_aromas,
            "name_aromas": name_aromas,
            "aromas": sample_aromas or name_aromas,
            "volumes": volumes,
            "default_volume": (volumes[0] if volumes else ""),
            "inventory_writeoff_results": [],
            "suggested_supply_name": preview_supply_name(supplies),
            "open_supply": open_supply or "",
            "post_sale_summary": post_sale_summary,
        },
    )


@app.get("/api/fbs/operational-state")
def fbs_operational_state():
    """Lightweight polling endpoint used by the dashboard after a cold start."""
    freshness = _fbs_supply_registry_freshness()
    refresh_due, retry_in = _fbs_operational_refresh_due(freshness)
    if refresh_due:
        # Ensure a refresh exists even if the first page request was interrupted.
        with _cache_lock:
            refreshing = "fbs_operational_registry" in _cache_refreshing
        if not refreshing:
            invalidate_cache("fbs_operational_registry")
            try:
                get_fbs_operational_sync_cached()
            except Exception:
                pass
    with _cache_lock:
        refreshing = "fbs_operational_registry" in _cache_refreshing
    freshness = _fbs_supply_registry_freshness()
    _, retry_in = _fbs_operational_refresh_due(freshness)
    return {
        "ok": True,
        "fresh": bool(freshness.get("fresh")),
        "refreshing": refreshing,
        "last_sync_at": freshness.get("last_sync_at") or "",
        "age_seconds": freshness.get("age_seconds"),
        "retry_in_seconds": retry_in,
        "error": str(freshness.get("error_text") or ""),
    }


@app.get("/marking", response_class=HTMLResponse)
def marking_global_page(request: Request, refresh: int | None = None):
    if refresh:
        invalidate_cache()
    try:
        get_fbs_operational_sync_cached()
        history_start = _registry_history_start()
        rows = list_fbs_registry_rows(marking_db_path, limit=10000, date_from=history_start)
        supply_rows = list_fbs_supply_registry_rows(
            marking_db_path, limit=10000, date_from=history_start
        )
        freshness = _fbs_supply_registry_freshness()
        open_ids = freshness.get("open_supply_ids") if freshness.get("fresh") else None
        active = _fbs_dashboard_supply_groups(
            rows, supply_rows, authoritative_open_supply_ids=open_ids
        )["assembling"]
    except Exception as exc:
        active = []
        error = str(exc)
    else:
        error = ""
    return templates.TemplateResponse(
        request,
        "marking_global.html",
        {
            "request": request,
            "supplies": active,
            "error": error,
            "config": config,
            "suz_status": suz_readiness(),
            "legal_settings": _circulation_defaults(),
        },
    )


class GlobalMarkingRefreshPayload(BaseModel):
    supply_ids: list[str] = Field(default_factory=list)


def _marking_supply_refresh_needed(supply_id: str) -> dict[str, bool]:
    open_orders = bool(list_open_suz_orders(marking_db_path, supply_id))
    assignments = list_supply_code_assignments(marking_db_path, supply_id)
    if not assignments:
        return {"suz": open_orders, "true": False}
    reports = list_suz_utilisation_reports(marking_db_path, supply_id)
    has_pending_report = any(
        _utilisation_report_blocks_application(r, assignments)
        for r in reports
    )
    terminal_docs = TRUE_DOCUMENT_SUCCESS_STATUSES | TRUE_DOCUMENT_FAILURE_STATUSES
    has_pending_doc = any(
        str(r.get("status") or "").upper() not in terminal_docs
        for r in list_circulation_documents(marking_db_path, supply_id)
    )
    not_introduced = any(
        not _circulation_code_terminal(r.get("circulation_status"))
        for r in assignments
    )
    return {"suz": open_orders, "true": bool(has_pending_report or has_pending_doc or not_introduced)}


def _global_marking_refresh_worker(supply_ids: list[str], report: Callable[..., None]) -> dict[str, Any]:
    ids = [str(x).strip() for x in supply_ids if str(x).strip()]
    ids = list(dict.fromkeys(ids))
    errors: list[str] = []
    refreshed = 0
    skipped = 0
    report(total=len(ids), progress=0, message=f"Проверяю {len(ids)} поставок…")
    for index, supply_id in enumerate(ids, start=1):
        needs = _marking_supply_refresh_needed(supply_id)
        if not needs.get("suz") and not needs.get("true"):
            skipped += 1
            report(progress=index, total=len(ids), message=f"{index}/{len(ids)} · {supply_id}: уже актуально")
            continue
        report(progress=index - 1, total=len(ids), message=f"{index}/{len(ids)} · {supply_id}: запрос серверов…")
        if needs.get("suz"):
            try:
                suz_result = _sync_all_supply_orders_impl(supply_id)
                errors.extend([f"{supply_id}: {e}" for e in (suz_result.get("errors") or [])])
            except Exception as exc:
                errors.append(f"{supply_id} · СУЗ: {exc}")
        if needs.get("true"):
            try:
                true_result = _refresh_circulation_status_impl(supply_id, allow_empty=True)
                errors.extend([f"{supply_id}: {e}" for e in (true_result.get("errors") or [])])
            except Exception as exc:
                errors.append(f"{supply_id} · ЧЗ: {exc}")
        refreshed += 1
        report(progress=index, total=len(ids), message=f"{index}/{len(ids)} · {supply_id}: готово")
    return {
        "ok": True, "refreshed": refreshed, "skipped": skipped, "errors": errors,
        "message": (
            f"Статусы актуализированы · проверено {refreshed}, без запросов {skipped}"
            + (f" · временных ошибок {len(errors)}" if errors else "")
        ),
    }


@app.get("/api/jobs/{job_id}")
def background_job_status(job_id: str):
    row = _job_snapshot(job_id)
    if not row:
        return JSONResponse({"ok": False, "error": "Задача не найдена"}, status_code=404)
    return {"ok": True, **row}


@app.post("/api/marking/global/refresh")
def start_global_marking_refresh(payload: GlobalMarkingRefreshPayload):
    ids = [str(x).strip() for x in payload.supply_ids if str(x).strip()]
    if not ids:
        return {"ok": True, "job_id": "", "message": "Активных поставок для проверки нет"}
    job_id = submit_background_job(
        "Обновление Честного Знака",
        lambda report: _global_marking_refresh_worker(ids, report),
        operation_key="marking-refresh:" + ",".join(sorted(set(ids))),
    )
    return {"ok": True, "job_id": job_id, "message": "Проверка запущена"}


@app.get("/api/analytics/today")
def api_analytics_today(force: int | None = None):
    return JSONResponse({
        "ok": False,
        "enabled": False,
        "updated_label": "аналитика выключена",
        "orders_sum_label": "—",
        "orders_count": 0,
        "buyouts_sum_label": "—",
        "buyouts_count": 0,
        "message": "Аналитика WB временно отключена, чтобы не получать 429.",
    })


@app.get("/manual-print", response_class=HTMLResponse)
def manual_print_page(request: Request, msg: str | None = None, error: str | None = None):
    ctx = manual_context(msg=msg, error=error)
    ctx["request"] = request
    return templates.TemplateResponse(request, "manual_print.html", ctx)


@app.post("/manual-print/pdf")
def manual_print_pdf(
    label_type: str = Form(...),
    aroma: str = Form(...),
    volume: str = Form(default=""),
    quantity: int = Form(...),
):
    rows, errors = build_manual_rows(label_type, aroma, volume, quantity)
    if errors or not rows:
        return RedirectResponse(url=f"/?error={'; '.join(errors) or 'Нет строк для печати'}#manual", status_code=303)

    kind_slug = "samples" if label_type == "samples" else "names"
    by_type = config.get("manual_pdf_output_by_type") or {}
    out_pdf = Path(by_type.get(kind_slug) or Path(config.get("manual_pdf_output_dir", "data/manual_pdf")) / f"FBE_MANUAL_{kind_slug}.pdf")

    try:
        pdf_path = render_manual_labels_pdf_with_bullzip(config, label_type, rows, out_pdf)
    except Exception as exc:
        return RedirectResponse(url=f"/?error=Ошибка BarTender/Bullzip: {exc}#manual", status_code=303)

    # Open in the browser instead of silently downloading. The file path is fixed
    # and overwritten on each run, so the manual PDF folder does not accumulate copies.
    return FileResponse(
        str(pdf_path),
        media_type="application/pdf",
        filename=pdf_path.name,
        headers={"Content-Disposition": f'inline; filename="{pdf_path.name}"'},
    )


@app.post("/manual-print/direct")
def manual_print_direct(
    label_type: str = Form(...),
    aroma: str = Form(...),
    volume: str = Form(default=""),
    quantity: int = Form(...),
):
    rows, errors = build_manual_rows(label_type, aroma, volume, quantity)
    if errors or not rows:
        return RedirectResponse(url=f"/?error={'; '.join(errors) or 'Нет строк для печати'}#manual", status_code=303)

    try:
        run_manual_bartender_print(
            config=config,
            label_type=label_type,
            rows=rows,
            printer_name=str(config.get("small_label_printer_name") or config.get("internal_label_printer_name") or config.get("wb_printer_name")),
        )
    except Exception as exc:
        return RedirectResponse(url=f"/?error=Ошибка прямой печати: {exc}#manual", status_code=303)

    label = "переходников" if label_type == "samples" else "именных этикеток"
    return RedirectResponse(url=f"/?msg=Отправлено на печать: {len(rows)} {label} ({aroma})#manual", status_code=303)


@app.post("/supplies/create")
def create_supply(
    name: str = Form(default=""),
    auto_name_suggested: str = Form(default=""),
):
    custom_name = name.strip()
    final_name = generate_supply_name(get_supplies_cached()) if should_autogenerate_supply_name(custom_name, auto_name_suggested) else custom_name
    supply_id = wb.create_supply(final_name)
    try:
        upsert_fbs_supplies(marking_db_path, [{
            "id": supply_id, "name": final_name, "done": False,
            "createdAt": datetime.now(timezone.utc).isoformat(),
            "fbe_status": "На сборке",
        }])
    except Exception:
        pass
    invalidate_cache("supplies")
    invalidate_cache("fbs_operational_registry")
    return dashboard_redirect(msg=f"Создана поставка {supply_id}: {final_name}", open_supply=supply_id)


@app.post("/supplies/create-by-warehouses")
def create_supplies_by_warehouses():
    """Create one FBS supply per seller virtual warehouse for all new orders."""
    try:
        invalidate_cache("new_orders")
        fresh_orders = wb.get_new_orders()
        names = seller_warehouse_name_map()
        grouped: dict[int, dict[str, Any]] = {}
        for order in fresh_orders:
            order_id = int(order.get("id") or 0)
            if not order_id:
                continue
            warehouse_id = order_warehouse_id(order)
            if not warehouse_id:
                return dashboard_redirect(
                    error=f"Задание {order_id}: WB не вернул warehouseId. Автоматическое распределение остановлено."
                )
            bucket = grouped.setdefault(
                warehouse_id,
                {
                    "warehouse_id": warehouse_id,
                    "warehouse_name": order_warehouse_name(order, names),
                    "order_ids": [],
                },
            )
            bucket["order_ids"].append(order_id)

        if not grouped:
            return dashboard_redirect(msg="Новых заказов для распределения по складам нет")

        existing = list(get_supplies_cached())
        created: list[dict[str, Any]] = []
        errors: list[str] = []
        for warehouse_id, bucket in sorted(grouped.items(), key=lambda pair: str(pair[1]["warehouse_name"]).lower()):
            warehouse_name = str(bucket["warehouse_name"] or f"Склад #{warehouse_id}")
            supply_name = generate_supply_name(existing, warehouse_name=warehouse_name)
            supply_id = ""
            try:
                supply_id = wb.create_supply(supply_name)
                try:
                    upsert_fbs_supplies(marking_db_path, [{
                        "id": supply_id, "name": supply_name, "done": False,
                        "createdAt": datetime.now(timezone.utc).isoformat(),
                        "fbe_status": "На сборке",
                    }])
                except Exception:
                    pass
                count = add_orders_to_supply_in_batches(supply_id, list(bucket["order_ids"]))
                created.append({
                    "supply_id": supply_id,
                    "name": supply_name,
                    "warehouse_name": warehouse_name,
                    "count": count,
                })
                existing.append({"id": supply_id, "name": supply_name})
            except Exception as exc:
                errors.append(
                    f"{warehouse_name}: "
                    + (f"поставка {supply_id} создана, но задания не добавлены: {exc}" if supply_id else str(exc))
                )

        invalidate_cache()
        summary = "; ".join(
            f"{item['warehouse_name']} — {item['count']} → {item['name']}" for item in created
        )
        if errors:
            return dashboard_redirect(
                error=("Создано: " + summary + ". " if summary else "") + "Ошибки: " + "; ".join(errors),
                open_supply=created[0]["supply_id"] if len(created) == 1 else "",
            )
        return dashboard_redirect(
            msg=f"Создано поставок: {len(created)}. {summary}",
            open_supply=created[0]["supply_id"] if len(created) == 1 else "",
        )
    except Exception as exc:
        return dashboard_redirect(error=f"Не удалось создать поставки по складам: {exc}")


@app.post("/supplies/add-orders")
def add_orders_to_supply(
    supply_id: str = Form(...),
    order_ids: list[int] = Form(default=[]),
):
    if not order_ids:
        return dashboard_redirect(error="Не выбраны заказы")
    selected_ids = [int(x) for x in order_ids]
    selected_warehouse = warehouse_for_selected_order_ids(selected_ids)
    if selected_warehouse is None and len(selected_ids) > 1:
        return dashboard_redirect(
            error="Выбраны задания разных FBS-складов. Используйте «Создать поставки по складам» или выбирайте один склад."
        )
    added = add_orders_to_supply_in_batches(supply_id=supply_id, order_ids=selected_ids)
    invalidate_cache("new_orders")
    invalidate_cache("all_orders")
    invalidate_cache("all_orders_recent")
    invalidate_cache(f"supply_order_ids:{supply_id}")
    invalidate_cache("supplies")
    return dashboard_redirect(msg=f"Добавлено в поставку: {added}", open_supply=supply_id)


@app.post("/supplies/{source_supply_id}/move-orders")
def move_orders_to_supply(
    source_supply_id: str,
    target_supply_id: str = Form(...),
    order_ids: list[int] = Form(default=[]),
):
    if not order_ids:
        return dashboard_redirect(error="Не выбраны задания", open_supply=source_supply_id)
    target = str(target_supply_id or "").strip()
    if not target:
        return dashboard_redirect(error="Не выбрана поставка для переноса", open_supply=source_supply_id)
    if target == source_supply_id:
        return dashboard_redirect(error="Выберите другую поставку", open_supply=source_supply_id)

    moved = add_orders_to_supply_in_batches(supply_id=target, order_ids=[int(x) for x in order_ids])
    try:
        source_ids = [int(x) for x in wb.get_supply_order_ids(str(source_supply_id)) if int(x) > 0]
        set_fbs_supply_order_ids(
            marking_db_path, supply_id=str(source_supply_id), order_ids=source_ids
        )
    except Exception:
        pass
    invalidate_cache("new_orders")
    invalidate_cache("all_orders")
    invalidate_cache("all_orders_recent")
    invalidate_cache(f"supply_order_ids:{source_supply_id}")
    invalidate_cache(f"supply_order_ids:{target}")
    invalidate_cache("supplies")
    return dashboard_redirect(msg=f"Перенесено в поставку {target}: {moved}", open_supply=target)


@app.post("/supplies/create-and-add")
def create_supply_and_add_orders(
    name: str = Form(default=""),
    auto_name_suggested: str = Form(default=""),
    order_ids: list[int] = Form(default=[]),
):
    if not order_ids:
        return dashboard_redirect(error="Не выбраны заказы")
    selected_ids = [int(x) for x in order_ids]
    selected_warehouse = warehouse_for_selected_order_ids(selected_ids)
    if selected_warehouse is None:
        return dashboard_redirect(
            error="Одну поставку нельзя создать из заданий разных FBS-складов. Используйте «Создать поставки по складам»."
        )
    _warehouse_id, warehouse_name = selected_warehouse
    custom_name = name.strip()
    final_name = (
        generate_supply_name(get_supplies_cached(), warehouse_name=warehouse_name)
        if should_autogenerate_supply_name(custom_name, auto_name_suggested)
        else custom_name
    )
    supply_id = wb.create_supply(final_name)
    try:
        upsert_fbs_supplies(marking_db_path, [{
            "id": supply_id, "name": final_name, "done": False,
            "createdAt": datetime.now(timezone.utc).isoformat(),
            "fbe_status": "На сборке",
        }])
    except Exception:
        pass
    added = add_orders_to_supply_in_batches(supply_id=supply_id, order_ids=selected_ids)
    invalidate_cache()
    return dashboard_redirect(msg=f"Создана поставка {supply_id}: {final_name}. Добавлено заказов: {added}", open_supply=supply_id)


@app.get("/supplies/{supply_id}/inline", response_class=HTMLResponse)
def supply_inline_panel(
    request: Request,
    supply_id: str,
    sort: str | None = None,
):
    supply = get_supply_details_cached(supply_id)

    enriched = supply_orders_context(supply_id, sort=sort)
    order_ids = [int(o.get("id") or 0) for o in enriched if int(o.get("id") or 0) > 0]
    shipping = _supply_shipping_context(supply_id, order_ids)

    if supply:
        supply["fbe_status"] = shipping["label"]
        supply["fbe_status_class"] = shipping["class"]
        supply["created_at_text"] = format_dt(supply.get("createdAt"))
        supply["closed_at_text"] = format_dt(supply.get("closedAt"))
        supply["scan_dt_text"] = format_dt(supply.get("scanDt"))
        supply["transfer_or_scan_text"] = supply["scan_dt_text"] or supply["closed_at_text"] or ""

    marking = load_marking_workspace(supply_id, enriched_orders=enriched, fetch_wb=False)
    registry_rows = list_fbs_registry_rows(marking_db_path, limit=10000, date_from=_registry_history_start())

    return templates.TemplateResponse(
        request,
        "supply_panel.html",
        {
            "request": request,
            "supply_id": supply_id,
            "supply": supply,
            "orders": enriched,
            "active_supplies": _active_supply_selector_from_registry(registry_rows),
            "default_expiration": default_expiration_date(),
            "marking": marking,
            "shipping": shipping,
            "config": config,
            "sort": sort or config.get("default_order_sort", "urgent"),
        },
    )


def build_marking_report_xlsx_bytes(marking: dict[str, Any], supply_id: str) -> bytes:
    """Build one immutable audit row per marked WB assembly order.

    XLSX/XML cannot contain ASCII 29 directly. The exact code therefore remains
    in SQLite and is exported losslessly as Base64, while a readable column uses
    the explicit <GS> marker.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "КИЗы поставки"

    green = PatternFill("solid", fgColor="22C55E")
    green_dark = PatternFill("solid", fgColor="166534")
    green_light = PatternFill("solid", fgColor="DCFCE7")
    red_light = PatternFill("solid", fgColor="FEE2E2")
    gray = PatternFill("solid", fgColor="F1F5F9")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "№",
        "Код поставки",
        "ID задания WB",
        "Стикер WB",
        "Артикул продавца",
        "Наименование",
        "GTIN-14",
        "Серийный номер",
        "КИЗ (читаемый, <GS>)",
        "КИЗ Base64 (точная копия)",
        "orderId СУЗ",
        "blockId СУЗ",
        "Товарная группа",
        "Статус FBE",
        "Статус нанесения",
        "ID отчета о нанесении",
        "Статус ввода в оборот",
        "UUID документа ввода",
        "Ошибка обработки",
        "Статус WB",
        "Получен",
        "Закреплен",
        "Напечатан",
        "Передан в WB",
    ]
    last_col = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, "FBE - отчет по кодам маркировки")
    ws.cell(1, 1).font = Font(bold=True, size=16, color="FFFFFF")
    ws.cell(1, 1).fill = green_dark
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.cell(3, 1, "Поставка").font = Font(bold=True)
    ws.cell(3, 2, supply_id)
    ws.cell(4, 1, "Сформирован").font = Font(bold=True)
    ws.cell(4, 2, datetime.now().strftime("%d.%m.%Y %H:%M"))
    ws.cell(5, 1, "Маркируемых позиций").font = Font(bold=True)
    ws.cell(5, 2, int(marking.get("marking_orders") or 0))
    ws.cell(6, 1, "КИЗ закреплено").font = Font(bold=True)
    ws.cell(6, 2, int(marking.get("assigned_orders") or 0))

    header_row = 8
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = green
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 34

    stored = {
        int(item.get("order_id") or 0): item
        for item in list_supply_code_assignments(marking_db_path, supply_id)
        if int(item.get("order_id") or 0)
    }
    row_number = header_row + 1
    ordered_rows = sorted(
        list(marking.get("orders") or []),
        key=lambda item: (str(item.get("name") or "").lower(), int(item.get("order_id") or 0)),
    )
    for index, item in enumerate(ordered_rows, start=1):
        order_id = int(item.get("order_id") or 0)
        assignment = stored.get(order_id) or item.get("assignment") or {}
        raw_code = str(assignment.get("raw_code") or "")
        readable = display_marking_code(raw_code, max_length=10000) if raw_code else ""
        exact_b64 = base64.b64encode(raw_code.encode("utf-8")).decode("ascii") if raw_code else ""
        values = [
            index,
            supply_id,
            order_id,
            str(item.get("sticker_display") or ""),
            str(item.get("seller_article") or ""),
            str(item.get("name") or ""),
            str(item.get("gtin14") or item.get("gtin") or ""),
            str(assignment.get("serial") or ""),
            readable,
            exact_b64,
            str(assignment.get("oms_order_id") or ""),
            str(assignment.get("suz_block_id") or ""),
            str(assignment.get("product_group") or ""),
            str(assignment.get("status") or "не закреплен"),
            str(assignment.get("utilization_status") or ""),
            str(assignment.get("utilization_report_id") or ""),
            str(assignment.get("circulation_status") or ""),
            str(assignment.get("circulation_document_id") or ""),
            str(assignment.get("error_text") or ""),
            str(item.get("wb_decision_label") or ""),
            str(assignment.get("created_at") or ""),
            str(assignment.get("assigned_at") or ""),
            str(assignment.get("printed_at") or ""),
            str(assignment.get("wb_sent_at") or ""),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_number, col, value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if raw_code:
                cell.fill = green_light
            else:
                cell.fill = red_light
        for col in (3, 7, 8, 9, 10, 11, 12, 15, 16, 17, 18):
            ws.cell(row_number, col).number_format = "@"
        row_number += 1

    if not ordered_rows:
        ws.merge_cells(start_row=header_row + 1, start_column=1, end_row=header_row + 2, end_column=last_col)
        cell = ws.cell(header_row + 1, 1, "В поставке нет маркируемых позиций")
        cell.fill = gray
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        1: 6, 2: 22, 3: 17, 4: 17, 5: 20, 6: 44, 7: 18, 8: 20,
        9: 70, 10: 78, 11: 38, 12: 38, 13: 16, 14: 18, 15: 20,
        16: 38, 17: 22, 18: 40, 19: 44, 20: 24, 21: 20, 22: 20,
        23: 20, 24: 20,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{max(header_row, row_number - 1)}"

    legend = wb_out.create_sheet("Хранение КИЗ")
    legend.append(["Поле", "Назначение"])
    legend.append(["КИЗ (читаемый, <GS>)", "Для визуальной проверки. Каждый <GS> означает исходный ASCII 29."])
    legend.append(["КИЗ Base64 (точная копия)", "Полная строка КИЗ в UTF-8 без потери управляющих символов. Декодирование Base64 восстанавливает исходную строку байт-в-байт."])
    legend.append(["Основное хранилище", "Оригинальный КИЗ с настоящими GS-разделителями хранится в data/fbe.db. XLSX не используется для передачи кода в WB."])
    legend.append(["Ввод в оборот", "True API получает только идентификационную часть кода без криптохвоста. Полный raw_code в SQLite не меняется."])
    legend.append(["Передача в WB", "Доступна только после статуса INTRODUCED. FBE берет полный raw_code напрямую из SQLite и передает его JSON-запросом. Excel в этом контуре не участвует."])
    for cell in legend[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = green_dark
    for row in legend.iter_rows(min_row=2, max_row=legend.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    legend.column_dimensions["A"].width = 34
    legend.column_dimensions["B"].width = 105
    legend.freeze_panes = "A2"

    bio = BytesIO()
    wb_out.save(bio)
    return bio.getvalue()


class MarkingScanPayload(BaseModel):
    order_id: int
    code: str


class SuzCertificatePayload(BaseModel):
    thumbprint: str


class SuzOrderPayload(BaseModel):
    reserve_per_gtin: int = 0
    # Used only for the chemistry group. Perfumery has no paymentType field.
    payment_type: int = 1


class MarkingLegalSettingsPayload(BaseModel):
    participant_inn: str
    oms_id: str = ""
    connection_id: str = ""


class CirculationIntroducePayload(BaseModel):
    confirmed: bool = False
    # Internal guard used by the step-specific background jobs. It prevents
    # a race where the application-report button could accidentally advance
    # to introduction if CRPT changed state between render and click.
    requested_stage: str = ""
    # Explicit one-time acknowledgement. The default remains safe: conflicting
    # catalog rows for one GTIN block the legally significant operation.
    ignore_compliance_conflicts: bool = False
    participant_inn: str
    # Kept optional for backward compatibility with v0.66 browser payloads.
    # 0.83.0 intentionally uses one INN for participant, producer and owner.
    producer_inn: str = ""
    owner_inn: str = ""
    production_date: str
    compliance_by_gtin: dict[str, dict[str, str]] = Field(default_factory=dict)
    # Legacy v0.66 fields are accepted as a fallback, but the current UI sends
    # compliance_by_gtin so every GTIN can have its own declaration details.
    tnved_by_gtin: dict[str, str] = Field(default_factory=dict)
    certificate_type: str = ""
    certificate_number: str = ""
    certificate_date: str = ""
    certificate_valid_until: str = ""


def suz_readiness() -> dict[str, Any]:
    settings_now = load_suz_settings()
    spec_raw = str(settings_now.get("suz_api_spec_path") or config.get("suz_api_spec_path") or "").strip()
    spec_path = Path(spec_raw) if spec_raw else None
    if spec_path is not None and not spec_path.is_absolute():
        spec_path = Path.cwd() / spec_path
    cert = str(settings_now.get("suz_cert_thumbprint") or "").replace(" ", "").upper()

    # Runtime ordering only requires the OMS identifiers and a signing certificate.
    # The bundled API guide is useful for audit/reference, but its absence must not
    # disable a production action after a successful live connection test.
    checks = [
        {
            "label": "OMS ID",
            "ready": bool(str(settings_now.get("suz_oms_id") or "").strip()),
            "value": str(settings_now.get("suz_oms_id") or "").strip(),
            "required": True,
        },
        {
            "label": "Идентификатор соединения СУЗ",
            "ready": bool(str(settings_now.get("suz_connection_id") or "").strip()),
            "value": str(settings_now.get("suz_connection_id") or "").strip(),
            "required": True,
        },
        {
            "label": "Сертификат УКЭП",
            "ready": bool(cert),
            "value": cert[-12:] if cert else "",
            "required": True,
        },
        {
            "label": "Документация API СУЗ 3.0.37",
            "ready": bool(spec_path and spec_path.exists()),
            "value": str(spec_path) if spec_path else spec_raw,
            "required": False,
        },
    ]
    required_items = [item for item in checks if item["required"]]
    return {
        "items": checks,
        "ready": all(item["ready"] for item in required_items),
        "configured": sum(1 for item in checks if item["ready"]),
        "total": len(checks),
        "oms_id": str(settings_now.get("suz_oms_id") or ""),
        "connection_id": str(settings_now.get("suz_connection_id") or ""),
        "cert_thumbprint": cert,
        "template_id": int(settings_now.get("suz_template_id") or 46),
        "payment_type": int(settings_now.get("suz_payment_type") or 1),
        "serial_number_type": str(settings_now.get("suz_serial_number_type") or "OPERATOR"),
        "release_method_type": str(settings_now.get("suz_release_method_type") or "PRODUCTION"),
    }


def _meta_order_id(item: dict[str, Any]) -> int:
    return int(item.get("id") or item.get("orderId") or item.get("orderID") or 0)


def _sgtin_meta_info(item: dict[str, Any] | None) -> dict[str, Any]:
    """Normalize all known WB FBS metadata response shapes for SGTIN.

    WB has changed the metadata response model over time. The current endpoint
    can expose concrete SGTIN values as an ``sgtins`` collection and validation
    state separately. Older FBE builds only inspected ``meta.sgtin.value`` and
    therefore missed codes that had been attached manually in the WB cabinet.
    Only strings that successfully parse as GS1 marking codes are accepted as
    concrete KIZ values; words such as ``filled`` are kept as state only.
    """
    result = {
        "available": False,
        "decision": "",
        "values": [],
        "detail_value": None,
        "raw_candidates": [],
    }
    if not isinstance(item, dict):
        return result

    candidates: list[Any] = []
    decisions: list[str] = []

    def add_decision(value: Any) -> None:
        text = str(value or "").strip()
        if text and text not in decisions:
            decisions.append(text)

    def harvest(value: Any, *, key_hint: str = "") -> None:
        if value is None:
            return
        key = str(key_hint or "").lower()
        if isinstance(value, str):
            if key in {"decision", "status", "state", "validationstatus", "validation_status"}:
                add_decision(value)
            elif key in {"sgtin", "sgtins", "value", "code", "identifier", "identificationcode", "identification_code"}:
                candidates.append(value)
            return
        if isinstance(value, (int, float, bool)):
            return
        if isinstance(value, list):
            if key in {"sgtin", "sgtins", "value", "codes", "identifiers"}:
                result["available"] = True
            for child in value:
                harvest(child, key_hint=key_hint)
            return
        if isinstance(value, dict):
            if key in {"sgtin", "sgtins"}:
                result["available"] = True
            for child_key, child_value in value.items():
                ck = str(child_key or "").lower()
                if ck in {"decision", "status", "state", "validationstatus", "validation_status"}:
                    add_decision(child_value)
                if ck in {"sgtin", "sgtins"}:
                    result["available"] = True
                harvest(child_value, key_hint=ck)

    # Prefer explicitly named SGTIN branches, but keep legacy response support.
    for key in ("sgtins", "sgtin"):
        if key in item:
            harvest(item.get(key), key_hint=key)
    meta = item.get("meta")
    if isinstance(meta, dict) and "sgtin" in meta:
        harvest(meta.get("sgtin"), key_hint="sgtin")
    details = item.get("metaDetails")
    if isinstance(details, list):
        for detail in details:
            if not isinstance(detail, dict) or str(detail.get("key") or "").lower() != "sgtin":
                continue
            result["available"] = True
            add_decision(detail.get("decision"))
            result["detail_value"] = detail.get("value")
            harvest(detail.get("value"), key_hint="value")
            harvest(detail.get("sgtins"), key_hint="sgtins")

    # Some API revisions return a nested metadata object under another key.
    for key in ("metadata", "identifiers", "labeling", "marking"):
        branch = item.get(key)
        if isinstance(branch, (dict, list)):
            harvest(branch, key_hint=key)

    valid_values: list[str] = []
    raw_candidates: list[str] = []
    for candidate in candidates:
        text = str(candidate or "")
        if not text or text in raw_candidates:
            continue
        raw_candidates.append(text)
        try:
            parsed = parse_marking_code(text)
        except Exception:
            continue
        raw = str(parsed.get("raw_code") or "")
        if raw and raw not in valid_values:
            valid_values.append(raw)

    result["raw_candidates"] = raw_candidates
    result["values"] = valid_values
    if decisions:
        result["decision"] = decisions[0]
    if valid_values:
        result["available"] = True
    return result


def _status_from_wb_meta(info: dict[str, Any], current_status: str) -> str:
    decision = str(info.get("decision") or "")
    if decision in {"valid", "sgtinMaySell"}:
        return "accepted_wb"
    if decision in {"filled", "pending"} or info.get("values"):
        return "sent_wb"
    return current_status



INTRODUCED_CODE_STATUSES = {"INTRODUCED"}
APPLIED_CODE_STATUSES = {"APPLIED", "APPLIED_NOT_PAID"}
TRUE_DOCUMENT_SUCCESS_STATUSES = {"CHECKED_OK"}
TRUE_DOCUMENT_FAILURE_STATUSES = {
    "CHECKED_NOT_OK", "PROCESSING_ERROR", "ERROR", "REJECTED", "CANCELLED", "FAILED"
}


def _circulation_code_terminal(value: Any) -> bool:
    """A terminal CIS must not be polled automatically.

    This is deliberately broader than ``INTRODUCED``. True API can report a
    terminal rejection/withdrawal too, and those responses must not create an
    endless status loop. A terminal error remains eligible for an explicit new
    legal operation; callers that decide whether a code already passed
    introduction use ``_circulation_code_success_terminal`` below.
    """
    status = str(value or "").strip().upper()
    return (
        _circulation_code_success_terminal(status)
        or status in {
            "REJECTED", "ERROR", "FAILED", "CANCELLED", "DISPOSED",
            "WRITTEN_OFF", "EXPIRED", "DISAGGREGATED",
        }
        or status.startswith((
            "REJECTED_", "ERROR_", "FAILED_", "CANCELLED_",
            "DISPOSED_", "WRITTEN_OFF_", "EXPIRED_", "DISAGGREGATED_",
        ))
    )


def _circulation_code_success_terminal(value: Any) -> bool:
    """Return whether a CIS has reached the successful legal lifecycle branch."""
    status = str(value or "").strip().upper()
    return (
        status in INTRODUCED_CODE_STATUSES
        or status == "WITHDRAWN"
        or status.startswith("RETIRED")
    )


def _circulation_code_failure_terminal(value: Any) -> bool:
    """Return whether a CIS reached a terminal failure branch.

    A failed CIS linked to a document is retryable only for that code. This
    distinction prevents a successful sibling group from being resubmitted and
    also prevents a stale ``submitted`` document row from blocking recovery.
    """
    return _circulation_code_terminal(value) and not _circulation_code_success_terminal(value)


def _true_doc_status(payload: Any) -> str:
    if isinstance(payload, dict):
        for key in ("status", "document_status", "documentStatus", "downloadStatus"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip().upper()
        for key in ("result", "document", "body", "data"):
            nested = payload.get(key)
            status = _true_doc_status(nested)
            if status:
                return status
    elif isinstance(payload, list):
        for item in payload:
            status = _true_doc_status(item)
            if status:
                return status
    return ""


def _true_doc_error(payload: Any) -> str:
    if isinstance(payload, dict):
        for key in ("error_message", "errorMessage", "description", "error"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()
        for key in ("errors", "processingErrors", "validationErrors"):
            value = payload.get(key)
            if value:
                try:
                    return json.dumps(value, ensure_ascii=False)
                except Exception:
                    return str(value)
    return ""


def _extract_cis_statuses(rows: list[dict[str, Any]]) -> dict[str, str]:
    result: dict[str, str] = {}

    def visit(value: Any) -> None:
        if isinstance(value, list):
            for item in value:
                visit(item)
            return
        if not isinstance(value, dict):
            return
        requested = str(
            value.get("requestedCis")
            or value.get("cis")
            or value.get("identificationCode")
            or value.get("uitCode")
            or ""
        ).strip()
        status = str(
            value.get("status") or value.get("cisStatus") or value.get("cis_status") or ""
        ).strip().upper()
        if requested and status:
            result[requested] = status
        for key in ("cisInfo", "results", "items", "data", "result", "cises"):
            nested = value.get(key)
            if isinstance(nested, (dict, list)):
                visit(nested)

    visit(rows)
    return result


def _assignment_product_group(assignment: dict[str, Any]) -> str:
    product_group = str(assignment.get("product_group") or "").strip()
    if product_group:
        return product_group
    by_article = str(_suz_product_group_for_article(str(assignment.get("seller_article") or "")) or "")
    if by_article:
        return by_article

    # Historical WB-only assignments may have lost the seller article locally.
    # If the recovered KIZ gives us a GTIN and that GTIN maps unambiguously to
    # one marking group in the current Catalog, use it for the True API lookup.
    gtin = canonical_gtin14(assignment.get("gtin")) or str(assignment.get("gtin") or "")
    if gtin:
        groups: set[str] = set()
        catalog_now = get_catalog()
        for article, product in catalog_now.by_article.items():
            product_gtin = canonical_gtin14(product.get("gtin")) or str(product.get("gtin") or "")
            if product_gtin != gtin:
                continue
            inferred = _infer_suz_product_group(product, article)
            if inferred:
                groups.add(str(inferred))
        if len(groups) == 1:
            return next(iter(groups))
    return ""


UTILISATION_TERMINAL_STATUSES = {
    "SUCCESS", "PARTIALLY", "REJECTED", "ERROR", "FAILED", "CANCELLED"
}


def _chemistry_application_complete(assignment: dict[str, Any]) -> bool:
    """True when a chemistry code no longer needs the manual application report."""
    return (
        str(assignment.get("utilization_status") or "").strip().upper() in APPLIED_CODE_STATUSES
        or _circulation_code_success_terminal(assignment.get("circulation_status"))
    )


def _utilisation_report_blocks_application(
    report_row: dict[str, Any], assignments: list[dict[str, Any]]
) -> bool:
    """Return whether this OMS report still blocks chemistry application.

    Report rows are transport/audit state; the code lifecycle is authoritative for
    advancing the workflow. A stale SUBMITTED/PROCESSING report must not keep step
    3 blocked once all codes linked to that report are APPLIED/INTRODUCED.
    """
    status = str(report_row.get("status") or "").strip().upper()
    if status in UTILISATION_TERMINAL_STATUSES:
        return False
    report_id = str(report_row.get("report_id") or "").strip()
    if not report_id:
        return False
    linked = [
        item for item in assignments
        if str(item.get("utilization_report_id") or "").strip() == report_id
        and _assignment_product_group(item) == "chemistry"
    ]
    if not linked:
        # Orphaned historical reports are diagnostic history, not a workflow gate.
        return False
    return any(not _chemistry_application_complete(item) for item in linked)


def _refresh_supply_true_statuses(supply_id: str) -> dict[str, Any]:
    """Refresh CIS lifecycle independently for each product group.

    Mixed supplies must not be all-or-nothing. If CRPT answers for chemistry but
    the perfumery request fails (or vice versa), the successful statuses are
    committed immediately so one remote group cannot keep another group blocked.
    """
    assignments = list_supply_code_assignments(marking_db_path, supply_id)
    if not assignments:
        raise MarkingDbError("В поставке еще нет закрепленных КИЗов")

    grouped: dict[str, list[dict[str, Any]]] = {}
    group_errors: list[str] = []
    refresh_rows = [
        assignment for assignment in assignments
        if not _circulation_code_terminal(assignment.get("circulation_status"))
    ]
    for assignment in refresh_rows:
        group = _assignment_product_group(assignment)
        if not group:
            group_errors.append(
                f"Не удалось определить товарную группу для артикула {assignment.get('seller_article') or '—'}"
            )
            continue
        grouped.setdefault(group, []).append(assignment)

    all_statuses: dict[str, str] = {}
    raw_responses: dict[str, list[dict[str, Any]]] = {}
    aggregate_counts = {"updated": 0, "applied": 0, "introduced": 0, "errors": 0}
    for product_group, group_rows in grouped.items():
        raw_codes = [str(item.get("raw_code") or "") for item in group_rows if str(item.get("raw_code") or "")]
        try:
            response_rows: list[dict[str, Any]] = []
            for start in range(0, len(raw_codes), 1000):
                response_rows.extend(
                    suz.get_cises_info(raw_codes[start : start + 1000], product_group=product_group)
                )
            raw_responses[product_group] = response_rows
            group_statuses = _extract_cis_statuses(response_rows)
            all_statuses.update(group_statuses)
            counts = update_supply_code_lifecycle(
                marking_db_path,
                supply_id=supply_id,
                statuses_by_identification_code=group_statuses,
            )
            for key in aggregate_counts:
                aggregate_counts[key] += int(counts.get(key) or 0)
        except Exception as exc:
            group_errors.append(f"{product_group}: {exc}")

    unresolved = max(0, len(refresh_rows) - len(all_statuses))
    return {
        "counts": aggregate_counts,
        "unresolved": unresolved,
        "statuses": all_statuses,
        "responses": raw_responses,
        "errors": group_errors,
    }


def _circulation_defaults() -> dict[str, Any]:
    settings_now = load_suz_settings()
    participant = str(
        settings_now.get("suz_true_participant_inn")
        or settings_now.get("suz_auth_inn")
        or ""
    ).strip()
    return {
        "participant_inn": participant,
        "oms_id": str(settings_now.get("suz_oms_id") or "").strip(),
        "connection_id": str(settings_now.get("suz_connection_id") or "").strip(),
        # Compatibility keys for older templates; all three roles are the same
        # for the user's own-production scenario.
        "producer_inn": participant,
        "owner_inn": participant,
        "production_date": moscow_now().date().isoformat(),
        "certificate_type": str(settings_now.get("suz_true_certificate_type") or "").strip(),
        "certificate_number": str(settings_now.get("suz_true_certificate_number") or "").strip(),
        "certificate_date": str(settings_now.get("suz_true_certificate_date") or "").strip(),
        "tnved_by_gtin": dict(settings_now.get("suz_true_tnved_by_gtin") or {}),
        "compliance_by_gtin": dict(settings_now.get("suz_true_compliance_by_gtin") or {}),
    }


def _normalize_certificate_type(value: Any) -> str:
    normalized = str(value or "").strip().upper()
    aliases = {
        "ДЕКЛАРАЦИЯ О СООТВЕТСТВИИ": "CONFORMITY_DECLARATION",
        "ДЕКЛАРАЦИЯ": "CONFORMITY_DECLARATION",
        "СЕРТИФИКАТ СООТВЕТСТВИЯ": "CONFORMITY_CERTIFICATE",
        "СЕРТИФИКАТ": "CONFORMITY_CERTIFICATE",
    }
    return aliases.get(normalized, normalized)


def _normalize_excel_digits(value: Any) -> str:
    text = str(value or "").strip().replace(" ", "").replace("\u00a0", "")
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return re.sub(r"\D", "", text)


def _normalize_catalog_date(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    # Excel date cells are often returned as `YYYY-MM-DD 00:00:00`; text cells
    # may be maintained by operators as DD.MM.YYYY.
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def _catalog_compliance(product: dict[str, Any] | None) -> dict[str, str]:
    product = product or {}
    return {
        "tnved_code": _normalize_excel_digits(product.get("tnved_code")),
        "certificate_type": _normalize_certificate_type(product.get("certificate_type")),
        "certificate_number": str(product.get("certificate_number") or "").strip(),
        "certificate_date": _normalize_catalog_date(product.get("certificate_date")),
        "certificate_valid_until": _normalize_catalog_date(product.get("certificate_valid_until")),
    }


def _catalog_compliance_for_identity(
    catalog: Any,
    *,
    articles: list[str] | tuple[str, ...] | None = None,
    gtin: str = "",
) -> tuple[dict[str, str], list[str], list[str]]:
    """Resolve compliance from the real catalog rows linked to an article/GTIN.

    Marking demand rows are grouped by GTIN and contain ``articles`` rather than
    ``seller_article``. Older code looked only for ``seller_article`` and therefore
    silently lost the catalog declaration data in the Chestny ZNAK panel.
    """
    requested_articles = []
    for value in articles or []:
        article = str(value or "").strip()
        if article and article not in requested_articles:
            requested_articles.append(article)

    products: list[dict[str, Any]] = []
    matched_articles: list[str] = []
    for article in requested_articles:
        product = catalog.find_by_article(article)
        if product and product not in products:
            products.append(product)
            matched_articles.append(str(product.get("seller_article") or article).strip())

    # Use GTIN only as a fallback when no requested article was found. This is
    # essential for repeated GTINs: each assigned KIZ must inherit the legal
    # details of its exact seller article rather than a random sibling row that
    # happens to share the same GTIN.
    target_gtin = canonical_gtin14(gtin)
    if target_gtin and not products:
        for product in getattr(catalog, "by_article", {}).values():
            if canonical_gtin14(product.get("gtin")) != target_gtin:
                continue
            if product not in products:
                products.append(product)
                matched_articles.append(str(product.get("seller_article") or "").strip())

    result = {
        "tnved_code": "",
        "certificate_type": "",
        "certificate_number": "",
        "certificate_date": "",
        "certificate_valid_until": "",
    }
    conflicts: list[str] = []
    for product in products:
        candidate = _catalog_compliance(product)
        for key, value in candidate.items():
            if not value:
                continue
            if result[key] and result[key] != value:
                conflicts.append(key)
                continue
            result[key] = value

    matched_articles = [value for value in dict.fromkeys(matched_articles) if value]
    return result, matched_articles, sorted(set(conflicts))


def _merge_compliance(*sources: dict[str, Any] | None) -> dict[str, str]:
    result = {
        "tnved_code": "",
        "certificate_type": "",
        "certificate_number": "",
        "certificate_date": "",
        "certificate_valid_until": "",
    }
    for source in sources:
        if not isinstance(source, dict):
            continue
        for key in result:
            value = source.get(key)
            if value not in (None, ""):
                result[key] = str(value).strip()
    result["certificate_type"] = _normalize_certificate_type(result["certificate_type"])
    result["tnved_code"] = _normalize_excel_digits(result["tnved_code"])
    return result


def _validate_inn(value: str, label: str) -> str:
    normalized = re.sub(r"\D", "", str(value or ""))
    if len(normalized) not in {10, 12}:
        raise MarkingDbError(f"{label}: ИНН должен содержать 10 или 12 цифр")
    digits = [int(char) for char in normalized]
    if len(digits) == 10:
        weights = [2, 4, 10, 3, 5, 9, 4, 6, 8]
        valid = (sum(value * weight for value, weight in zip(digits[:9], weights)) % 11) % 10 == digits[9]
    else:
        weights_11 = [7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
        weights_12 = [3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
        check_11 = (sum(value * weight for value, weight in zip(digits[:10], weights_11)) % 11) % 10
        check_12 = (sum(value * weight for value, weight in zip(digits[:11], weights_12)) % 11) % 10
        valid = check_11 == digits[10] and check_12 == digits[11]
    if not valid:
        raise MarkingDbError(f"{label}: контрольные цифры ИНН не совпадают")
    return normalized


def _validate_iso_date(value: str, label: str, *, allow_future: bool = False) -> str:
    try:
        parsed = datetime.strptime(str(value or ""), "%Y-%m-%d").date()
    except ValueError as exc:
        raise MarkingDbError(f"{label}: используйте формат ГГГГ-ММ-ДД") from exc
    if not allow_future and parsed > moscow_now().date():
        raise MarkingDbError(f"{label} не может быть позднее текущей даты")
    return parsed.isoformat()


def _add_calendar_months(value: str, months: int) -> str:
    base = datetime.strptime(str(value), "%Y-%m-%d").date()
    total = base.year * 12 + (base.month - 1) + int(months)
    year, month_index = divmod(total, 12)
    month = month_index + 1
    day = min(base.day, calendar.monthrange(year, month)[1])
    return base.replace(year=year, month=month, day=day).isoformat()


def _chemistry_application_attributes(product: dict[str, Any], production_date: str) -> dict[str, str]:
    article = str(product.get("seller_article") or "товар").strip()
    try:
        shelf_life_months = int(product.get("shelf_life_months") or 0)
    except Exception:
        shelf_life_months = 0
    if shelf_life_months <= 0:
        raise MarkingDbError(
            f"{article}: для профиля «Дезодоранты и косметика» заполните «Срок годности, мес.» в Каталоге"
        )
    alcohol_raw = str(product.get("alcohol_volume_pct") if product.get("alcohol_volume_pct") is not None else "").strip()
    if alcohol_raw == "":
        raise MarkingDbError(
            f"{article}: заполните «Этиловый спирт, %» в Каталоге; если спирта нет, укажите 0"
        )
    try:
        alcohol = float(alcohol_raw.replace(",", "."))
    except ValueError as exc:
        raise MarkingDbError(f"{article}: некорректное значение этилового спирта") from exc
    if not 0 <= alcohol <= 99.9:
        raise MarkingDbError(f"{article}: этиловый спирт должен быть от 0 до 99,9%")
    expiration_date = _add_calendar_months(production_date, shelf_life_months)
    if expiration_date <= moscow_now().date().isoformat():
        raise MarkingDbError(f"{article}: рассчитанный срок годности уже истек")
    alcohol_text = (f"{alcohol:.1f}").rstrip("0").rstrip(".")
    return {
        "productionDate": production_date,
        "expirationDate": expiration_date,
        "alcoholVolume": alcohol_text,
    }


def _utilisation_report_status(payload: Any) -> str:
    if isinstance(payload, dict):
        for key in ("reportStatus", "report_status", "status"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip().upper()
        for key in ("result", "data", "report", "reports"):
            status = _utilisation_report_status(payload.get(key))
            if status:
                return status
    elif isinstance(payload, list):
        for item in payload:
            status = _utilisation_report_status(item)
            if status:
                return status
    return ""


def _utilisation_report_error(payload: Any) -> str:
    if isinstance(payload, dict):
        for key in ("errorReason", "error_reason", "error", "description", "message"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()
        for key in ("result", "data", "report", "reports"):
            error = _utilisation_report_error(payload.get(key))
            if error:
                return error
    elif isinstance(payload, list):
        for item in payload:
            error = _utilisation_report_error(item)
            if error:
                return error
    return ""


def _marking_human_status(assignment: dict[str, Any] | None) -> dict[str, str]:
    if not assignment:
        return {"key": "waiting", "label": "Код ожидается", "technical": "", "class": "muted"}
    utilization = str(assignment.get("utilization_status") or "").strip().upper()
    circulation = str(assignment.get("circulation_status") or "").strip().upper()
    error_text = str(assignment.get("error_text") or "").strip()
    if error_text or circulation in TRUE_DOCUMENT_FAILURE_STATUSES:
        return {"key": "error", "label": "Ошибка", "technical": circulation or utilization, "class": "error"}
    if circulation == "INTRODUCED":
        return {"key": "introduced", "label": "В обороте", "technical": "INTRODUCED", "class": "ok"}
    if circulation == "SUBMITTED":
        return {"key": "processing", "label": "Ввод в оборот обрабатывается", "technical": "SUBMITTED", "class": "warn"}
    if utilization == "SUBMITTED":
        return {"key": "processing", "label": "Отчет о нанесении обрабатывается", "technical": "SUBMITTED", "class": "warn"}
    if utilization in APPLIED_CODE_STATUSES:
        return {"key": "applied", "label": "Нанесен", "technical": utilization, "class": "ok"}
    technical = circulation or utilization or "EMITTED"
    if technical == "EMITTED":
        return {"key": "emitted", "label": "Эмитирован", "technical": technical, "class": "info"}
    return {"key": "other", "label": technical.replace("_", " ").title(), "technical": technical, "class": "warn"}


def _build_marking_status_cards(marking: dict[str, Any], lifecycle_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    counts = {"emitted": 0, "applied": 0, "introduced": 0, "processing": 0, "error": 0, "other": 0}
    for row in lifecycle_rows:
        status = _marking_human_status(row)
        key = status["key"]
        if key in counts:
            counts[key] += 1
    cards: list[dict[str, Any]] = []
    pending = int(marking.get("pending_codes_total") or 0)
    if pending:
        cards.append({"key": "ordered", "label": "Заказано, ожидается", "count": pending, "class": "warn"})
    labels = {
        "emitted": ("Эмитировано", "info"),
        "applied": ("Нанесено", "ok"),
        "introduced": ("В обороте", "ok"),
        "processing": ("Обрабатывается", "warn"),
        "error": ("Ошибка", "error"),
        "other": ("Другой статус", "warn"),
    }
    for key in ("emitted", "applied", "introduced", "processing", "error", "other"):
        count = int(counts[key])
        if count:
            label, css_class = labels[key]
            cards.append({"key": key, "label": label, "count": count, "class": css_class})
    if not cards:
        cards.append({"key": "empty", "label": "Коды еще не заказаны", "count": 0, "class": "muted"})
    return cards


@app.post("/api/supplies/{supply_id}/marking/legal-settings")
def marking_save_legal_settings(supply_id: str, payload: MarkingLegalSettingsPayload):
    try:
        participant_inn = _validate_inn(payload.participant_inn, "ИНН производителя")
        settings_before = load_suz_settings()
        oms_id = str(payload.oms_id or settings_before.get("suz_oms_id") or "").strip()
        connection_id = str(payload.connection_id or settings_before.get("suz_connection_id") or "").strip()
        if not oms_id:
            raise MarkingDbError("OMS ID не заполнен")
        if not connection_id:
            raise MarkingDbError("Идентификатор соединения СУЗ не заполнен")
        save_suz_local_settings({
            "suz_oms_id": oms_id,
            "suz_connection_id": connection_id,
            "suz_auth_inn": participant_inn,
            "suz_true_participant_inn": participant_inn,
            "suz_true_producer_inn": participant_inn,
            "suz_true_owner_inn": participant_inn,
        })
        if (
            oms_id != str(settings_before.get("suz_oms_id") or "").strip()
            or connection_id != str(settings_before.get("suz_connection_id") or "").strip()
        ):
            suz.reset_token()
        return {
            "ok": True,
            "message": "OMS ID, идентификатор соединения СУЗ и ИНН сохранены. Они меняются только вручную в этом блоке.",
            "oms_id": oms_id,
            "connection_id": connection_id,
            "participant_inn": participant_inn,
            "production_date": moscow_now().date().isoformat(),
        }
    except MarkingDbError as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


@app.post("/api/marking/legal-settings")
def marking_save_global_legal_settings(payload: MarkingLegalSettingsPayload):
    return marking_save_legal_settings("global", payload)


def load_marking_workspace(
    supply_id: str,
    *,
    enriched_orders: list[dict[str, Any]] | None = None,
    fetch_wb: bool = True,
) -> dict[str, Any]:
    """Build the supply marking workspace with permanent code assignments."""
    orders = enriched_orders if enriched_orders is not None else supply_orders_context(supply_id)
    order_ids = [int(item.get("id") or item.get("order_id") or 0) for item in orders]
    order_ids = [value for value in order_ids if value]
    assignments = list_assignments(marking_db_path, order_ids)

    wb_meta: dict[int, dict[str, Any]] = {}
    if fetch_wb and order_ids:
        try:
            for start in range(0, len(order_ids), 100):
                for item in wb.get_orders_meta(order_ids[start : start + 100]):
                    item_id = _meta_order_id(item)
                    if item_id:
                        wb_meta[item_id] = item
        except Exception:
            # A temporary WB metadata error must not hide locally saved KIZs.
            wb_meta = {}

    ordered_for_supply = count_supply_suz_ordered_codes_by_gtin(marking_db_path, supply_id)
    marking = build_marking_requirements(
        orders,
        get_catalog(),
        assignments=assignments,
        wb_meta=wb_meta,
        free_codes={},
        pending_codes=ordered_for_supply,
    )
    marking["suz"] = suz_readiness()
    marking["datamatrix"] = datamatrix_renderer.readiness()
    suz_orders = list_suz_orders(marking_db_path, supply_id)
    suz_pending_runtime = 0
    for order in suz_orders:
        requested_total = 0
        received_total = 0
        runtime_items: list[dict[str, Any]] = []
        for item in order.get("items") or []:
            requested = int(item.get("requested_quantity") or 0)
            received = int(item.get("received_quantity") or 0)
            remaining = max(0, requested - received)
            requested_total += requested
            received_total += received
            suz_pending_runtime += remaining
            buffer_status = str(item.get("buffer_status") or "PENDING").upper()
            runtime_items.append({
                "gtin": str(item.get("gtin") or ""),
                "requested": requested,
                "received": received,
                "remaining": remaining,
                "buffer_status": buffer_status,
                "rejection_reason": str(item.get("rejection_reason") or ""),
            })
        order["requested_total"] = requested_total
        order["received_total"] = received_total
        order["remaining_total"] = max(0, requested_total - received_total)
        order["runtime_items"] = runtime_items
        buffer_states = {str(item.get("buffer_status") or "PENDING").upper() for item in runtime_items}
        order_status_live = str(order.get("status") or "PENDING").replace("_", " ").upper()
        if requested_total and received_total >= requested_total:
            runtime_label = "RECEIVED · КИЗы получены"
        elif "REJECTED" in buffer_states:
            runtime_label = "REJECTED · ошибка СУЗ"
        elif "ACTIVE" in buffer_states:
            runtime_label = "ACTIVE · получаем КИЗы"
        elif "PENDING" in buffer_states or order_status_live in {"CREATED", "PENDING"}:
            runtime_label = "PENDING · СУЗ формирует КИЗы"
        elif received_total:
            runtime_label = f"{order_status_live} · получаем КИЗы"
        else:
            runtime_label = order_status_live
        order["runtime_label"] = runtime_label
    marking["suz_orders"] = suz_orders
    marking["suz_pending_runtime"] = suz_pending_runtime

    order_plan = {"perfumery": 0, "chemistry": 0}
    order_plan_gtins = {"perfumery": 0, "chemistry": 0}
    catalog_for_plan = get_catalog()
    for row in marking.get("rows") or []:
        quantity = int(row.get("to_order") or 0)
        if quantity <= 0:
            continue
        groups: set[str] = set()
        for article in row.get("articles") or []:
            product = catalog_for_plan.find_by_article(str(article or "").strip())
            if product:
                group = _infer_suz_product_group(product, str(product.get("seller_article") or article or ""))
                if group:
                    groups.add(group)
        if len(groups) == 1:
            group = next(iter(groups))
            if group in order_plan:
                order_plan[group] += quantity
                order_plan_gtins[group] += 1
    marking["suz_order_plan"] = order_plan
    marking["suz_order_plan_gtins"] = order_plan_gtins
    marking["stored_codes"] = len(list_supply_code_assignments(marking_db_path, supply_id))
    marking["all_marking_assigned"] = bool(
        int(marking.get("marking_orders") or 0) > 0
        and int(marking.get("assigned_orders") or 0) == int(marking.get("marking_orders") or 0)
    )

    lifecycle_rows = list_supply_code_assignments(marking_db_path, supply_id)
    lifecycle_by_order = {
        int(item.get("order_id") or 0): item
        for item in lifecycle_rows
        if int(item.get("order_id") or 0)
    }
    for order_row in marking.get("orders") or []:
        assignment = lifecycle_by_order.get(int(order_row.get("order_id") or 0)) or order_row.get("assignment")
        order_row["human_code_status"] = _marking_human_status(assignment)

    article_status_groups: dict[tuple[str, str], dict[str, Any]] = {}
    for order_row in marking.get("orders") or []:
        article = str(order_row.get("seller_article") or "—").strip() or "—"
        gtin = str(order_row.get("gtin14") or order_row.get("gtin") or "—").strip() or "—"
        key = (article, gtin)
        group = article_status_groups.setdefault(
            key,
            {
                "name": str(order_row.get("name") or article),
                "seller_article": article,
                "gtin": gtin,
                "count": 0,
                "status_counts": Counter(),
                "status_classes": {},
                "wb_counts": Counter(),
                "wb_classes": {},
            },
        )
        group["count"] += 1
        code_status = order_row.get("human_code_status") or {}
        code_label = str(code_status.get("label") or "Код ожидается")
        group["status_counts"][code_label] += 1
        group["status_classes"][code_label] = str(code_status.get("class") or "muted")
        wb_label = str(order_row.get("wb_decision_label") or "Статус WB не получен")
        group["wb_counts"][wb_label] += 1
        group["wb_classes"][wb_label] = str(order_row.get("wb_decision_class") or "unknown")

    status_article_rows: list[dict[str, Any]] = []
    for group in article_status_groups.values():
        status_parts = [
            f"{label}: {count}" for label, count in group.pop("status_counts").most_common()
        ]
        status_classes = group.pop("status_classes")
        wb_parts = [f"{label}: {count}" for label, count in group.pop("wb_counts").most_common()]
        wb_classes = group.pop("wb_classes")
        group["status_text"] = " · ".join(status_parts)
        group["status_class"] = (
            next(iter(status_classes.values())) if len(status_classes) == 1 else "warn"
        )
        group["wb_text"] = " · ".join(wb_parts)
        group["wb_class"] = next(iter(wb_classes.values())) if len(wb_classes) == 1 else "pending"
        status_article_rows.append(group)
    status_article_rows.sort(key=lambda item: (str(item["name"]).lower(), str(item["seller_article"])))
    marking["status_article_rows"] = status_article_rows
    marking["status_cards"] = _build_marking_status_cards(marking, lifecycle_rows)
    marking["processing_codes"] = (
        sum(int(item.get("count") or 0) for item in marking["status_cards"] if item.get("key") == "processing")
        + int(marking.get("suz_pending_runtime") or 0)
    )
    marking["printed_codes"] = sum(1 for item in lifecycle_rows if str(item.get("printed_at") or "").strip())
    circulation_documents = list_circulation_documents(marking_db_path, supply_id)
    utilisation_reports = list_suz_utilisation_reports(marking_db_path, supply_id)
    marking["utilisation_reports"] = utilisation_reports
    marking["latest_utilisation_report"] = utilisation_reports[0] if utilisation_reports else None
    introduced_rows = [
        item for item in lifecycle_rows
        if _circulation_code_success_terminal(item.get("circulation_status"))
    ]
    applied_rows = [
        item for item in lifecycle_rows
        if str(item.get("utilization_status") or "").upper() in APPLIED_CODE_STATUSES
        or _circulation_code_success_terminal(item.get("circulation_status"))
    ]
    lifecycle_groups = sorted({
        _assignment_product_group(item) for item in lifecycle_rows if _assignment_product_group(item)
    })
    marking["circulation_documents"] = circulation_documents
    marking["latest_circulation_document"] = circulation_documents[0] if circulation_documents else None
    marking["applied_codes"] = len(applied_rows)
    marking["introduced_codes"] = len(introduced_rows)
    marking["all_marking_introduced"] = bool(
        marking["all_marking_assigned"]
        and len(lifecycle_rows) == int(marking.get("marking_orders") or 0)
        and len(introduced_rows) == len(lifecycle_rows)
    )
    marking["circulation_groups"] = lifecycle_groups
    circulation_defaults = _circulation_defaults()
    marking["circulation_defaults"] = circulation_defaults
    stored_compliance = circulation_defaults.get("compliance_by_gtin", {})
    legacy_tnved = circulation_defaults.get("tnved_by_gtin", {})
    seen_circulation_gtins: set[str] = set()
    circulation_gtins: list[dict[str, str]] = []
    catalog = get_catalog()
    compliance_conflicts: list[str] = []
    unmatched_compliance_gtins: list[str] = []
    for item in (marking.get("rows") or []):
        gtin = canonical_gtin14(item.get("gtin"))
        if not gtin or gtin in seen_circulation_gtins:
            continue
        seen_circulation_gtins.add(gtin)
        articles = [str(value or "").strip() for value in (item.get("articles") or [])]
        legacy_article = str(item.get("seller_article") or "").strip()
        if legacy_article and legacy_article not in articles:
            articles.append(legacy_article)
        catalog_values, matched_articles, conflicts = _catalog_compliance_for_identity(
            catalog, articles=articles, gtin=gtin
        )
        if not matched_articles:
            unmatched_compliance_gtins.append(gtin)
        if conflicts:
            compliance_conflicts.append(f"GTIN {gtin}: {', '.join(conflicts)}")
        local_fallback = dict(stored_compliance.get(gtin) or {})
        if legacy_tnved.get(gtin) and not local_fallback.get("tnved_code"):
            local_fallback["tnved_code"] = legacy_tnved.get(gtin)
        global_fallback = {
            "certificate_type": circulation_defaults.get("certificate_type", ""),
            "certificate_number": circulation_defaults.get("certificate_number", ""),
            "certificate_date": circulation_defaults.get("certificate_date", ""),
        }
        # SQLite catalog is authoritative when a field is filled; local settings
        # remain a convenient fallback for older catalog files.
        compliance = _merge_compliance(global_fallback, local_fallback, catalog_values)
        circulation_gtins.append({
            "gtin": gtin,
            "article": ", ".join(matched_articles or [value for value in articles if value]),
            "name": str(item.get("name") or ""),
            "tnved": compliance["tnved_code"],
            "certificate_type": compliance["certificate_type"],
            "certificate_number": compliance["certificate_number"],
            "certificate_date": compliance["certificate_date"],
            "certificate_valid_until": compliance["certificate_valid_until"],
        })
    marking["compliance_conflicts"] = compliance_conflicts
    marking["unmatched_compliance_gtins"] = unmatched_compliance_gtins
    marking["circulation_gtins"] = circulation_gtins
    marking["compliance_missing_count"] = sum(
        1 for item in circulation_gtins
        if len(_normalize_excel_digits(item.get("tnved"))) != 10
        or not str(item.get("certificate_type") or "").strip()
        or not str(item.get("certificate_number") or "").strip()
        or not str(item.get("certificate_date") or "").strip()
    )
    marking["legal_settings_ready"] = bool(
        str(circulation_defaults.get("oms_id") or "").strip()
        and str(circulation_defaults.get("participant_inn") or "").strip()
        and int(marking["compliance_missing_count"]) == 0
        and not compliance_conflicts
        and not unmatched_compliance_gtins
    )

    circulation_blockers: list[str] = []
    if not str(circulation_defaults.get("oms_id") or "").strip():
        circulation_blockers.append("Не заполнен OMS ID")
    if not str(circulation_defaults.get("participant_inn") or "").strip():
        circulation_blockers.append("Не заполнен ИНН производителя")
    if int(marking["compliance_missing_count"]) > 0:
        circulation_blockers.append(
            f"Не заполнены ТН ВЭД или документы соответствия: {marking['compliance_missing_count']}"
        )
    if unmatched_compliance_gtins:
        circulation_blockers.append("Часть GTIN не сопоставлена с Каталогом")
    if compliance_conflicts:
        circulation_blockers.append("Для одного GTIN в Каталоге указаны разные реквизиты")
    if not marking["all_marking_assigned"]:
        circulation_blockers.append(
            f"Закреплено КИЗ: {marking.get('assigned_orders', 0)} из {marking.get('marking_orders', 0)}"
        )
    chemistry_rows = [item for item in lifecycle_rows if _assignment_product_group(item) == "chemistry"]
    chemistry_pending_application = [
        item for item in chemistry_rows
        if str(item.get("utilization_status") or "").strip().upper() not in APPLIED_CODE_STATUSES
        and not _circulation_code_success_terminal(item.get("circulation_status"))
    ]
    chemistry_metadata_blockers: list[str] = []
    if chemistry_pending_application:
        catalog_for_chemistry = get_catalog()
        missing_print = sum(1 for item in chemistry_pending_application if not str(item.get("printed_at") or "").strip())
        if missing_print:
            chemistry_metadata_blockers.append(
                f"Для дезодорантов сначала напечатайте и нанесите КИЗы: {missing_print}"
            )
        checked_articles: set[str] = set()
        for code_row in chemistry_pending_application:
            article = str(code_row.get("seller_article") or "").strip()
            if not article or article in checked_articles:
                continue
            checked_articles.add(article)
            product = catalog_for_chemistry.find_by_article(article)
            if not product:
                chemistry_metadata_blockers.append(f"{article}: товар не найден в Каталоге")
                continue
            try:
                _chemistry_application_attributes(product, str(circulation_defaults.get("production_date") or ""))
            except MarkingDbError as exc:
                chemistry_metadata_blockers.append(str(exc))

    active_utilisation = next(
        (item for item in utilisation_reports
         if _utilisation_report_blocks_application(item, chemistry_rows)),
        None,
    )
    marking["chemistry_codes"] = len(chemistry_rows)
    marking["chemistry_printed_codes"] = sum(
        1 for item in chemistry_rows if str(item.get("printed_at") or "").strip()
    )
    marking["chemistry_codes_fully_printed"] = bool(
        chemistry_rows and marking["chemistry_printed_codes"] >= len(chemistry_rows)
    )
    marking["chemistry_applied_codes"] = max(0, len(chemistry_rows) - len(chemistry_pending_application))
    marking["chemistry_pending_application_count"] = len(chemistry_pending_application)
    marking["chemistry_needs_utilisation"] = bool(chemistry_pending_application)
    marking["chemistry_application_in_progress"] = bool(active_utilisation)
    marking["chemistry_metadata_blockers"] = chemistry_metadata_blockers
    if chemistry_pending_application and active_utilisation:
        circulation_blockers.append(
            f"Отчет о нанесении {active_utilisation.get('report_id')} обрабатывается; сначала обновите статусы"
        )
    elif chemistry_pending_application:
        circulation_blockers.extend(chemistry_metadata_blockers)
    terminal_document_statuses = TRUE_DOCUMENT_SUCCESS_STATUSES | TRUE_DOCUMENT_FAILURE_STATUSES
    active_document = next(
        (
            item
            for item in circulation_documents
            if str(item.get("status") or "").upper() not in terminal_document_statuses
        ),
        None,
    )
    if active_document:
        circulation_blockers.append(
            f"Документ {active_document.get('document_uuid')} уже отправлен; сначала обновите его статус"
        )
    successful_document = next(
        (
            item
            for item in circulation_documents
            if str(item.get("status") or "").upper() in TRUE_DOCUMENT_SUCCESS_STATUSES
        ),
        None,
    )
    if successful_document and not marking["all_marking_introduced"]:
        circulation_blockers.append(
            f"Документ {successful_document.get('document_uuid')} уже обработан успешно; "
            "обновите статусы КИЗ и не отправляйте дубликат"
        )
    if marking["all_marking_introduced"]:
        circulation_blockers.append("Все КИЗы поставки уже введены в оборот")
    conflict_blocker = "Для одного GTIN в Каталоге указаны разные реквизиты"
    hard_circulation_blockers = [
        item for item in circulation_blockers if item != conflict_blocker
    ]
    # The chemistry application report precedes introduction into circulation and
    # does not use TN VED/declaration fields. Do not block that first stage because
    # of compliance data that becomes relevant only to LP_INTRODUCE_GOODS.
    if marking.get("chemistry_needs_utilisation") and not marking.get("chemistry_application_in_progress"):
        hard_circulation_blockers = [
            item for item in hard_circulation_blockers
            if not item.startswith("Не заполнены ТН ВЭД или документы соответствия")
            and item != "Часть GTIN не сопоставлена с Каталогом"
        ]
        requires_conflict_override = False
    else:
        requires_conflict_override = bool(compliance_conflicts)
    marking["circulation_blockers"] = circulation_blockers
    marking["circulation_hard_blockers"] = hard_circulation_blockers
    marking["requires_compliance_conflict_override"] = requires_conflict_override
    marking["can_introduce"] = not circulation_blockers
    # The button may remain available when the only blocker is a repeated-GTIN
    # conflict. The browser must collect an explicit one-time acknowledgement,
    # and the API verifies the same flag independently.
    marking["can_introduce_with_override"] = not hard_circulation_blockers
    marking["circulation_action_label"] = (
        "Подтвердить нанесение" if marking.get("chemistry_needs_utilisation")
        else "Ввести КИЗы в оборот"
    )
    marking["can_refresh_circulation"] = bool(lifecycle_rows)

    blockers: list[str] = []
    if not marking["suz"].get("ready"):
        missing_runtime = [
            str(item.get("label"))
            for item in marking["suz"].get("items") or []
            if item.get("required") and not item.get("ready")
        ]
        blockers.append("Не настроено подключение СУЗ: " + ", ".join(missing_runtime))
    if int(marking.get("marking_orders") or 0) <= 0:
        blockers.append("В поставке нет товаров с профилем, требующим маркировку")
    if int(marking.get("missing_orders") or 0) > 0:
        blockers.append("Есть маркируемые товары без корректного GTIN")
    if int(marking.get("unresolved_orders") or 0) > 0:
        blockers.append("Есть артикулы, не найденные в каталоге товаров")
    if int(marking.get("to_order_total") or 0) <= 0:
        blockers.append("Все коды для этой поставки уже заказаны в СУЗ")

    marking["suz_order_blockers"] = blockers
    marking["suz_can_create"] = not blockers
    marking["print_blockers"] = []
    if not marking["all_marking_assigned"]:
        marking["print_blockers"].append(
            f"Закреплено КИЗ: {marking.get('assigned_orders', 0)} из {marking.get('marking_orders', 0)}"
        )
    if not marking["datamatrix"].get("ready"):
        marking["print_blockers"].append(
            str(marking["datamatrix"].get("error") or "Не готов официальный рендерер DataMatrix")
        )
    marking["can_print_native"] = not marking["print_blockers"]
    marking["direct_print_blockers"] = list(marking["print_blockers"])
    if not str(config.get("small_label_printer_name") or "").strip():
        marking["direct_print_blockers"].append("Не указан принтер маленьких этикеток")
    marking["can_print_direct"] = not marking["direct_print_blockers"]
    marking["direct_printer_name"] = str(config.get("small_label_printer_name") or "")
    # Explicit sequential workflow gates used by the UI. Perfume application reporting
    # is automatic in the official perfume group; chemistry/deodorants require the
    # separate SUZ utilisation report before introduction into circulation.
    marking["codes_fully_received"] = bool(
        int(marking.get("marking_orders") or 0) > 0
        and int(marking.get("assigned_orders") or 0) >= int(marking.get("marking_orders") or 0)
    )
    marking["codes_fully_printed"] = bool(
        marking["codes_fully_received"]
        and int(marking.get("printed_codes") or 0) >= int(marking.get("marking_orders") or 0)
    )
    marking["application_report_required"] = bool(marking.get("chemistry_codes"))
    if marking["application_report_required"]:
        # The manual OMS utilisation report belongs only to chemistry/deodorants.
        # Perfume codes in the same WB supply must not participate in step-3 gates.
        marking["application_report_complete"] = bool(
            marking.get("all_marking_introduced")
            or (
                marking.get("chemistry_codes_fully_printed")
                and int(marking.get("chemistry_pending_application_count") or 0) == 0
            )
        )
    else:
        # Perfumery has no manual report in FBE; preserve the sequential UI by
        # completing the visual step after its codes have physically been printed.
        marking["application_report_complete"] = bool(
            marking.get("all_marking_introduced") or marking["codes_fully_printed"]
        )
    marking["can_submit_application_report"] = bool(
        marking["application_report_required"]
        and marking.get("chemistry_codes_fully_printed")
        and int(marking.get("chemistry_pending_application_count") or 0) > 0
        and not marking.get("chemistry_application_in_progress")
        and not marking.get("chemistry_metadata_blockers")
    )
    marking["can_start_introduction"] = bool(
        marking["codes_fully_printed"]
        and marking["application_report_complete"]
        and marking.get("can_introduce_with_override")
        and not marking.get("all_marking_introduced")
    )
    marking["processing_codes"] = int(marking.get("processing_codes") or 0)
    if marking.get("chemistry_application_in_progress") and not marking.get("all_marking_introduced"):
        marking["processing_codes"] += 1
    if active_document and not marking.get("all_marking_introduced"):
        marking["processing_codes"] += 1
    marking["has_live_work"] = bool(marking.get("processing_codes"))
    marking["all_wb_sent"] = bool(marking.get("marking_orders") and int(marking.get("sent_wb_orders") or 0) >= int(marking.get("marking_orders") or 0))

    marking["can_send_wb"] = marking["all_marking_introduced"]
    if not marking["all_marking_introduced"]:
        marking["wb_blocker"] = (
            f"Сначала введите КИЗы в оборот: {marking.get('introduced_codes', 0)} "
            f"из {marking.get('marking_orders', 0)} имеют статус INTRODUCED"
        )
    else:
        marking["wb_blocker"] = ""
    return marking


def _marking_error_payload(exc: Exception, status_code: int = 400) -> JSONResponse:
    upstream_status = getattr(exc, "status_code", None)
    actual_status = int(upstream_status) if isinstance(upstream_status, int) and 400 <= upstream_status <= 599 else status_code
    payload: dict[str, Any] = {
        "ok": False,
        "error": str(exc),
        "status_code": actual_status,
    }
    request_id = str(getattr(exc, "request_id", "") or "").strip()
    response_body = getattr(exc, "response_body", None)
    stage = str(getattr(exc, "stage", "") or "").strip()
    if request_id:
        payload["request_id"] = request_id
    if response_body not in (None, ""):
        payload["details"] = response_body
    if stage:
        payload["stage"] = stage
    print(f"[FBE marking error] HTTP {actual_status}: {exc}")
    if request_id:
        print(f"[FBE marking error] requestId: {request_id}")
    if response_body not in (None, ""):
        try:
            rendered = json.dumps(response_body, ensure_ascii=False) if not isinstance(response_body, str) else response_body
        except Exception:
            rendered = str(response_body)
        print(f"[FBE marking error] True API response: {rendered}")
    return JSONResponse(payload, status_code=actual_status)


def _save_true_api_diagnostic(
    *,
    supply_id: str,
    document_payload: dict[str, Any] | None,
    exc: Exception,
) -> str:
    try:
        folder = _runtime_path("economy_diagnostics_dir", "data/diagnostics")
        folder.mkdir(parents=True, exist_ok=True)
        stamp = moscow_now().strftime("%Y%m%d_%H%M%S_%f")
        safe_supply = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(supply_id or "supply"))
        path = folder / f"true_api_{safe_supply}_{stamp}.json"
        body = {
            "created_at": moscow_now().isoformat(),
            "supply_id": supply_id,
            "document_type": "LP_INTRODUCE_GOODS",
            "document_payload": document_payload,
            "error": str(exc),
            "status_code": getattr(exc, "status_code", None),
            "request_id": getattr(exc, "request_id", ""),
            "stage": getattr(exc, "stage", ""),
            "response_body": getattr(exc, "response_body", None),
        }
        path.write_text(json.dumps(body, ensure_ascii=False, indent=2), encoding="utf-8")
        return str(path)
    except Exception as diagnostic_exc:
        print(f"[FBE diagnostics] Не удалось сохранить диагностику True API: {diagnostic_exc}")
        return ""


def _find_marking_order(supply_id: str, order_id: int) -> tuple[dict[str, Any], dict[str, Any]]:
    orders = supply_orders_context(supply_id)
    order = next((item for item in orders if int(item.get("id") or 0) == int(order_id)), None)
    if order is None:
        raise MarkingDbError("Сборочное задание не найдено в этой поставке")
    product = get_catalog().find_by_article(str(order.get("seller_article") or ""))
    if not product:
        raise MarkingDbError("Артикул заказа не найден в каталоге товаров")
    if not is_marking_required(product.get("kiz_required")):
        raise MarkingDbError("Профиль каталога указывает, что товар не требует маркировки")
    gtin = str(product.get("gtin") or "").strip()
    reason = validate_gtin(gtin)
    if reason:
        raise MarkingDbError(reason)
    return order, product


@app.get("/supplies/{supply_id}/marking-needs", response_class=HTMLResponse)
def supply_marking_needs(request: Request, supply_id: str, remote: int = 0):
    try:
        marking = load_marking_workspace(supply_id, fetch_wb=bool(remote))
    except Exception as exc:
        marking = {
            "total_orders": 0,
            "marking_orders": 0,
            "ready_orders": 0,
            "assigned_orders": 0,
            "sent_wb_orders": 0,
            "gtin_count": 0,
            "free_codes_total": 0,
            "pending_codes_total": 0,
            "to_order_total": 0,
            "missing_orders": 0,
            "unresolved_orders": 0,
            "rows": [],
            "orders": [],
            "missing": [],
            "unresolved": [],
            "category_counts": [],
            "circulation_gtins": [],
            "compliance_missing_count": 0,
            "status_cards": [],
            "status_article_rows": [],
            "suz_orders": [],
            "suz_pending_runtime": 0,
            "suz_order_plan": {"perfumery": 0, "chemistry": 0},
            "suz_order_plan_gtins": {"perfumery": 0, "chemistry": 0},
            "processing_codes": 0,
            "has_live_work": False,
            "printed_codes": 0,
            "applied_codes": 0,
            "introduced_codes": 0,
            "all_marking_introduced": False,
            "all_wb_sent": False,
            "codes_fully_received": False,
            "codes_fully_printed": False,
            "application_report_required": False,
            "application_report_complete": False,
            "can_submit_application_report": False,
            "can_start_introduction": False,
            "latest_circulation_document": None,
            "latest_utilisation_report": None,
            "utilisation_reports": [],
            "chemistry_codes": 0,
            "chemistry_needs_utilisation": False,
            "chemistry_application_in_progress": False,
            "chemistry_metadata_blockers": [],
            "circulation_action_label": "Ввести КИЗы в оборот",
            "circulation_groups": [],
            "can_order": False,
            "suz_can_create": False,
            "can_print_native": False,
            "can_print_direct": False,
            "can_introduce": False,
            "can_introduce_with_override": False,
            "requires_compliance_conflict_override": False,
            "compliance_conflicts": [],
            "circulation_hard_blockers": ["Сначала устраните ошибку загрузки поставки"],
            "can_send_wb": False,
            "suz_order_blockers": ["Сначала устраните ошибку загрузки поставки"],
            "print_blockers": ["Сначала устраните ошибку загрузки поставки"],
            "direct_print_blockers": ["Сначала устраните ошибку загрузки поставки"],
            "circulation_blockers": ["Сначала устраните ошибку загрузки поставки"],
            "wb_blocker": "Сначала устраните ошибку загрузки поставки",
            "direct_printer_name": str(config.get("marking_direct_printer_name") or ""),
            "circulation_defaults": _circulation_defaults(),
            "warnings": [f"Не удалось рассчитать потребность: {exc}"],
            "suz": suz_readiness(),
        }
    return templates.TemplateResponse(
        request,
        "marking_needs_panel.html",
        {"request": request, "supply_id": supply_id, "marking": marking},
    )


@app.post("/api/input-layout/english")
def input_layout_english():
    if not bool(config.get("marking_scanner_enabled", False)):
        return {"ok": False, "disabled": True, "message": "Сканирование КИЗ в FBE отключено"}
    result = force_english_keyboard_layout()
    # Layout switching is a convenience. The parser has a server-side RU->EN
    # recovery fallback, so a Windows refusal must not make the input unusable.
    return result


@app.post("/api/supplies/{supply_id}/marking/validate")
def marking_validate_code(supply_id: str, payload: MarkingScanPayload):
    if not bool(config.get("marking_scanner_enabled", False)):
        return _marking_error_payload(ValueError("Сканирование КИЗ в FBE отключено. Используйте официальный PDF Честного Знака."), 410)
    try:
        order, product = _find_marking_order(supply_id, payload.order_id)
        parsed = parse_marking_code(payload.code)
        expected_gtin = canonical_gtin14(product.get("gtin"))
        if parsed["gtin"] != expected_gtin:
            raise MarkingCodeError(
                f"GTIN не совпадает: в КИЗ {parsed['gtin']}, для товара ожидается {expected_gtin}"
            )
        availability = validate_assignment_candidate(
            marking_db_path,
            raw_code=parsed["raw_code"],
            order_id=int(payload.order_id),
        )
        layout_note = " Раскладка RU была автоматически исправлена на EN." if parsed.get("layout_corrected") else ""
        return {
            "ok": True,
            "message": f"КИЗ проверен: GTIN совпадает.{layout_note}",
            "gtin": parsed["gtin"],
            "serial": parsed["serial"],
            "display_code": parsed["display_code"],
            "normalized_code": parsed["raw_code"],
            "layout_corrected": bool(parsed.get("layout_corrected")),
            "already_saved": bool(availability.get("already_saved_for_order")),
            "source": availability.get("source"),
        }
    except (MarkingCodeError, MarkingDbError, ValueError) as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


@app.post("/api/supplies/{supply_id}/marking/scan")
def marking_scan_code(supply_id: str, payload: MarkingScanPayload):
    if not bool(config.get("marking_scanner_enabled", False)):
        return _marking_error_payload(ValueError("Ручное закрепление КИЗ отключено: коды автоматически назначаются из СУЗ конкретным заданиям WB."), 410)
    try:
        order, product = _find_marking_order(supply_id, payload.order_id)
        parsed = parse_marking_code(payload.code)
        expected_gtin = canonical_gtin14(product.get("gtin"))
        if parsed["gtin"] != expected_gtin:
            raise MarkingCodeError(
                f"GTIN не совпадает: в КИЗ {parsed['gtin']}, для товара ожидается {expected_gtin}"
            )
        assignment = assign_code(
            marking_db_path,
            raw_code=parsed["raw_code"],
            gtin=parsed["gtin"],
            serial=parsed["serial"],
            seller_article=str(order.get("seller_article") or ""),
            supply_id=supply_id,
            order_id=int(payload.order_id),
        )
        return {
            "ok": True,
            "message": f"КИЗ распознан и локально закреплен за заданием {payload.order_id}",
            "gtin": parsed["gtin"],
            "serial": parsed["serial"],
            "assignment_id": assignment.get("id"),
        }
    except (MarkingCodeError, MarkingDbError, ValueError) as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


@app.post("/api/supplies/{supply_id}/marking/{order_id}/remove")
def marking_remove_assignment(supply_id: str, order_id: int):
    try:
        # Deletion deliberately does not require a fresh WB request or a valid
        # current SQLite catalog row. Even a malformed accidental scan must always
        # be removable while it has not been sent to WB.
        remove_local_assignment(marking_db_path, order_id, supply_id=supply_id)
        return {"ok": True, "message": "Локальная привязка удалена"}
    except MarkingDbError as exc:
        return _marking_error_payload(exc, 409)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


def _refresh_circulation_status_impl(supply_id: str, *, allow_empty: bool = True) -> dict[str, Any]:
    """Refresh only remote lifecycle state; return partial success on transient CRPT errors."""
    assignments = list_supply_code_assignments(marking_db_path, supply_id)
    if not assignments:
        if allow_empty:
            return {
                "ok": True, "skipped": True, "partial": False,
                "message": "Закрепленных КИЗов пока нет — проверка ЧЗ не требуется.",
                "applied": 0, "introduced": 0, "total": 0, "documents": [], "errors": [],
            }
        raise MarkingDbError("В поставке еще нет закрепленных КИЗов")

    document_updates: list[dict[str, Any]] = []
    document_errors: list[str] = []

    for report in list_suz_utilisation_reports(marking_db_path, supply_id):
        current_status = str(report.get("status") or "").strip().upper()
        # Keep the transport/document status for audit, but stop polling a report
        # once it is terminal or all of its own linked chemistry codes have
        # advanced. This is essential for mixed supplies: perfume rows must not
        # keep a chemistry report alive, and a stale remote SENT value must not
        # generate a request forever after /cises/info says APPLIED.
        if (
            current_status in UTILISATION_TERMINAL_STATUSES
            or not _utilisation_report_blocks_application(report, assignments)
        ):
            continue
        report_id = str(report.get("report_id") or "").strip()
        if not report_id:
            continue
        try:
            response = suz.get_report_info(report_id)
            status = _utilisation_report_status(response) or current_status or "PROCESSING"
            error_text = (
                _utilisation_report_error(response)
                if status in {"PARTIALLY", "REJECTED", "ERROR", "FAILED", "CANCELLED"}
                else ""
            )
            updated_report = update_suz_utilisation_report(
                marking_db_path, report_id, status=status, response=response, error_text=error_text
            )
            document_updates.append({"report_id": report_id, "kind": "utilisation", "status": status, "row": updated_report})
            if error_text:
                document_errors.append(f"Отчет о нанесении {report_id}: {error_text}")
        except Exception as exc:
            document_errors.append(f"Отчет о нанесении {report_id}: {exc}")

    terminal_statuses = TRUE_DOCUMENT_SUCCESS_STATUSES | TRUE_DOCUMENT_FAILURE_STATUSES
    for document in list_circulation_documents(marking_db_path, supply_id):
        current_status = str(document.get("status") or "").strip().upper()
        if current_status in terminal_statuses:
            continue
        document_uuid = str(document.get("document_uuid") or "").strip()
        if not document_uuid:
            continue
        linked_codes = [
            item for item in assignments
            if str(item.get("circulation_document_id") or "").strip() == document_uuid
        ]
        if linked_codes and all(
            _circulation_code_terminal(item.get("circulation_status"))
            for item in linked_codes
        ):
            continue
        try:
            response = suz.get_true_document_info(document_uuid)
            status = _true_doc_status(response) or current_status or "PROCESSING"
            error_text = _true_doc_error(response) if status in TRUE_DOCUMENT_FAILURE_STATUSES else ""
            updated = update_circulation_document(
                marking_db_path, document_uuid, status=status, response=response, error_text=error_text
            )
            document_updates.append(updated or {"document_uuid": document_uuid, "status": status})
            if status in TRUE_DOCUMENT_FAILURE_STATUSES:
                reason = error_text or f"Документ обработан со статусом {status}"
                mark_supply_circulation_error(
                    marking_db_path, supply_id=supply_id, document_uuid=document_uuid, error_text=reason
                )
                document_errors.append(f"{document_uuid}: {reason}")
        except Exception as exc:
            document_errors.append(f"{document_uuid}: {exc}")

    lifecycle: dict[str, Any] = {}
    try:
        lifecycle = _refresh_supply_true_statuses(supply_id)
        document_errors.extend([f"Статусы КИ: {error}" for error in (lifecycle.get("errors") or [])])
    except Exception as exc:
        # A transient /cises/info disconnect is not a business conflict. Keep the
        # last local state and let auto-refresh retry later instead of surfacing HTTP 409.
        document_errors.append(f"Статусы КИ: {exc}")

    refreshed = load_marking_workspace(supply_id, fetch_wb=False)
    message = (
        f"Статусы обновлены: нанесено {refreshed.get('applied_codes', 0)}, "
        f"в обороте {refreshed.get('introduced_codes', 0)} из {refreshed.get('marking_orders', 0)}."
    )
    if lifecycle.get("unresolved"):
        message += f" Не сопоставлено: {lifecycle['unresolved']}."
    if document_errors:
        message += " Часть серверов временно не ответила — FBE повторит проверку автоматически."
    return {
        "ok": True,
        "partial": bool(document_errors),
        "message": message,
        "applied": int(refreshed.get("applied_codes") or 0),
        "introduced": int(refreshed.get("introduced_codes") or 0),
        "total": int(refreshed.get("marking_orders") or 0),
        "documents": document_updates,
        "errors": document_errors,
    }


@app.post("/api/supplies/{supply_id}/marking/circulation/status")
def marking_refresh_circulation_status(supply_id: str):
    job_id = submit_background_job(
        "Проверка статусов Честного Знака",
        lambda report: _refresh_circulation_status_impl(supply_id, allow_empty=True),
        operation_key=f"marking-refresh:{supply_id}",
    )
    return {"ok": True, "job_id": job_id, "message": "Проверка статусов запущена"}


def marking_introduce_into_circulation(supply_id: str, payload: CirculationIntroducePayload):
    """Run the legal marking lifecycle for perfumery and chemistry.

    Chemistry has one extra mandatory stage: after the physical application of
    the DataMatrix, FBE submits an OMS utilisation/application report. Only after
    that report succeeds are the codes eligible for LP_INTRODUCE_GOODS.
    """
    document_payload: dict[str, Any] | None = None
    try:
        if not payload.confirmed:
            raise MarkingDbError("Требуется явное подтверждение юридически значимой операции")
        assignments = list_supply_code_assignments(marking_db_path, supply_id)
        if not assignments:
            raise MarkingDbError("В поставке еще нет закрепленных КИЗов")

        workspace = load_marking_workspace(supply_id, fetch_wb=False)
        if not workspace.get("all_marking_assigned"):
            raise MarkingDbError(
                f"Сначала закрепите все КИЗы: {workspace.get('assigned_orders', 0)} "
                f"из {workspace.get('marking_orders', 0)}"
            )

        participant_inn = _validate_inn(
            payload.participant_inn, "ИНН участника / производителя / собственника"
        )
        production_date = _validate_iso_date(payload.production_date, "Дата производства")

        # Always refresh before deciding which legal stage is next.
        try:
            _refresh_supply_true_statuses(supply_id)
        except Exception:
            # Newly emitted chemistry codes can be absent/lagging in True API;
            # this must not prevent the mandatory application report.
            pass
        assignments = list_supply_code_assignments(marking_db_path, supply_id)
        pending_rows = [
            item for item in assignments
            if not _circulation_code_success_terminal(item.get("circulation_status"))
        ]
        if not pending_rows:
            raise MarkingDbError("Все КИЗы этой поставки уже введены в оборот")

        # ---- Chemistry: mandatory report about physical application ----
        chemistry_unapplied = [
            item for item in pending_rows
            if _assignment_product_group(item) == "chemistry"
            and str(item.get("utilization_status") or "").strip().upper() not in APPLIED_CODE_STATUSES
        ]
        requested_stage = str(payload.requested_stage or "").strip().lower()
        if requested_stage == "introduction" and chemistry_unapplied:
            raise MarkingDbError("Сначала завершите шаг 3 — отчет о нанесении")
        if chemistry_unapplied:
            reports_now = list_suz_utilisation_reports(marking_db_path, supply_id)
            active_report_ids = {
                str(report.get("report_id") or "").strip()
                for report in reports_now
                if str(report.get("status") or "").strip().upper() not in UTILISATION_TERMINAL_STATUSES
                and _utilisation_report_blocks_application(report, assignments)
            }
            submitted = [
                item for item in chemistry_unapplied
                if (
                    str(item.get("utilization_report_id") or "").strip() in active_report_ids
                    or (
                        str(item.get("utilization_status") or "").strip().upper() == "SUBMITTED"
                        and not str(item.get("utilization_report_id") or "").strip()
                    )
                )
            ]
            chemistry_unapplied = [item for item in chemistry_unapplied if item not in submitted]
            if submitted and not chemistry_unapplied:
                report_ids = sorted({str(item.get("utilization_report_id") or "") for item in submitted if item.get("utilization_report_id")})
                raise MarkingDbError(
                    "Отчет о нанесении для дезодорантов уже отправлен"
                    + (f" ({', '.join(report_ids)})" if report_ids else "")
                    + ". Нажмите «Проверить статусы»."
                )

            not_printed = [item for item in chemistry_unapplied if not str(item.get("printed_at") or "").strip()]
            if not_printed:
                raise MarkingDbError(
                    f"Перед отчетом о нанесении физически напечатайте и нанесите КИЗы: {len(not_printed)}"
                )

            catalog = get_catalog()
            grouped_reports: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
            for item in chemistry_unapplied:
                article = str(item.get("seller_article") or "").strip()
                product = catalog.find_by_article(article) if article else None
                if not product:
                    raise MarkingDbError(f"{article or 'Без артикула'}: товар не найден в Каталоге")
                attrs = _chemistry_application_attributes(product, production_date)
                gtin = canonical_gtin14(item.get("gtin"))
                key = (gtin, attrs["productionDate"], attrs["expirationDate"], attrs["alcoholVolume"])
                grouped_reports.setdefault(key, []).append(item)

            report_ids: list[str] = []
            for (gtin, production, expiration, alcohol), rows in grouped_reports.items():
                attrs = {
                    "productionDate": production,
                    "expirationDate": expiration,
                    "alcoholVolume": alcohol,
                }
                response = suz.send_utilisation_report(
                    [str(item.get("raw_code") or "") for item in rows],
                    product_group="chemistry",
                    attributes=attrs,
                )
                report_id = str(response.get("reportId") or "").strip()
                request_payload = dict(response.get("requestPayload") or {})
                create_suz_utilisation_report(
                    marking_db_path,
                    supply_id=supply_id,
                    report_id=report_id,
                    product_group="chemistry",
                    gtin=gtin,
                    payload=request_payload,
                    response=response,
                    order_ids=[int(item.get("order_id") or 0) for item in rows],
                )
                report_ids.append(report_id)

            return {
                "ok": True,
                "stage": "utilisation",
                "message": (
                    f"Отчет о нанесении отправлен: {len(report_ids)}. "
                    "FBE автоматически отслеживает обработку. После подтверждения статус станет «Нанесено», "
                    "и действие «Ввести в оборот» станет доступно без ручной перезагрузки страницы."
                ),
                "report_ids": report_ids,
                "codes": len(chemistry_unapplied),
            }

        if requested_stage == "application":
            return {
                "ok": True,
                "stage": "application_complete",
                "message": "Отчет о нанесении уже подтвержден; ввод в оборот из шага 3 не выполнялся.",
            }

        # ---- All groups: True API introduction into circulation ----
        documents_by_id = {
            str(item.get("document_uuid") or "").strip(): item
            for item in list_circulation_documents(marking_db_path, supply_id)
            if str(item.get("document_uuid") or "").strip()
        }
        failed_document_ids = {
            document_uuid
            for document_uuid, item in documents_by_id.items()
            if str(item.get("status") or "").strip().upper() in TRUE_DOCUMENT_FAILURE_STATUSES
        }
        # True API can expose a terminal CIS rejection before the document
        # status endpoint catches up (or when that later poll was unavailable).
        # Treat only the rejected linked code as retryable; successful sibling
        # codes remain protected by their own INTRODUCED/RETIRED state.
        failed_document_ids.update(
            str(item.get("circulation_document_id") or "").strip()
            for item in assignments
            if str(item.get("circulation_document_id") or "").strip()
            and _circulation_code_failure_terminal(item.get("circulation_status"))
        )
        # A known external UUID remains an idempotency barrier after CHECKED_OK:
        # /cises/info may expose INTRODUCED a little later. Only an explicit
        # terminal failure makes those linked codes eligible for replacement.
        already_submitted = [
            item for item in pending_rows
            if str(item.get("circulation_document_id") or "").strip()
            and str(item.get("circulation_document_id") or "").strip() not in failed_document_ids
        ]
        pending_rows = [
            item for item in pending_rows
            if not str(item.get("circulation_document_id") or "").strip()
            or str(item.get("circulation_document_id") or "").strip() in failed_document_ids
        ]
        if not pending_rows:
            ids = ", ".join(sorted({
                str(item.get("circulation_document_id") or "").strip()
                for item in already_submitted
                if str(item.get("circulation_document_id") or "").strip()
            }))
            raise MarkingDbError(
                f"Документ ввода в оборот уже отправлен ({ids}). Повторная отправка заблокирована; "
                "обновите статусы позже."
            )

        not_applied = [
            item for item in pending_rows
            if str(item.get("utilization_status") or "").strip().upper() not in APPLIED_CODE_STATUSES
        ]
        if not_applied:
            statuses = sorted({
                str(item.get("utilization_status") or item.get("circulation_status") or "не определен")
                for item in not_applied
            })
            raise MarkingDbError(
                "Ввод в оборот доступен только после статуса «Нанесен» (APPLIED). "
                f"Текущие статусы: {', '.join(statuses)}"
            )

        workspace = load_marking_workspace(supply_id, fetch_wb=False)
        workspace_conflicts = list(workspace.get("compliance_conflicts") or [])
        if workspace_conflicts and not payload.ignore_compliance_conflicts:
            raise MarkingDbError(
                "Для одного GTIN в Каталоге указаны разные реквизиты. "
                "Проверьте карточки или установите явную галочку «Игнорировать предупреждение»."
            )

        legacy_global = {
            "certificate_type": payload.certificate_type,
            "certificate_number": payload.certificate_number,
            "certificate_date": payload.certificate_date,
            "certificate_valid_until": payload.certificate_valid_until,
        }
        validated_rows: list[tuple[dict[str, Any], dict[str, str]]] = []
        catalog = get_catalog()
        for item in pending_rows:
            gtin = canonical_gtin14(item.get("gtin"))
            supplied = dict(payload.compliance_by_gtin.get(gtin) or {})
            if not supplied.get("tnved_code"):
                supplied["tnved_code"] = payload.tnved_by_gtin.get(gtin, "")

            article = str(item.get("seller_article") or "").strip()
            catalog_values, matched_articles, conflicts = _catalog_compliance_for_identity(
                catalog, articles=[article] if article else [], gtin=gtin
            )
            if conflicts:
                if not payload.ignore_compliance_conflicts:
                    raise MarkingDbError(
                        f"GTIN {gtin}: в Каталоге найдены противоречивые реквизиты ({', '.join(conflicts)})"
                    )
                if not article or article not in matched_articles:
                    raise MarkingDbError(
                        f"GTIN {gtin}: предупреждение подтверждено, но у КИЗа нет однозначного артикула"
                    )

            compliance = _merge_compliance(legacy_global, catalog_values, supplied)
            tnved = _normalize_excel_digits(compliance["tnved_code"])
            item_label = f"артикул {article}, GTIN {gtin}" if article else f"GTIN {gtin}"
            if len(tnved) != 10:
                raise MarkingDbError(f"{item_label}: код ТН ВЭД должен содержать ровно 10 цифр")

            certificate_type = _normalize_certificate_type(compliance["certificate_type"])
            certificate_number = compliance["certificate_number"].strip()
            certificate_date_raw = compliance["certificate_date"].strip()
            certificate_valid_until_raw = compliance["certificate_valid_until"].strip()
            certificate_values = [certificate_type, certificate_number, certificate_date_raw]
            if any(certificate_values) and not all(certificate_values):
                raise MarkingDbError(f"{item_label}: заполните одновременно тип, номер и дату начала документа")
            if all(certificate_values):
                if certificate_type not in {"CONFORMITY_DECLARATION", "CONFORMITY_CERTIFICATE"}:
                    raise MarkingDbError(f"{item_label}: неподдерживаемый тип документа о соответствии")
                certificate_date = _validate_iso_date(
                    certificate_date_raw, f"{item_label}: дата начала действия документа"
                )
                if certificate_valid_until_raw:
                    certificate_valid_until = _validate_iso_date(
                        certificate_valid_until_raw,
                        f"{item_label}: дата окончания действия документа",
                        allow_future=True,
                    )
                    if certificate_valid_until < certificate_date:
                        raise MarkingDbError(f"{item_label}: дата окончания документа раньше даты начала")
                    if production_date > certificate_valid_until:
                        raise MarkingDbError(f"{item_label}: документ соответствия истек {certificate_valid_until}")
                else:
                    certificate_valid_until = ""
                if production_date < certificate_date:
                    raise MarkingDbError(
                        f"{item_label}: дата производства раньше начала действия документа {certificate_date}"
                    )
            else:
                certificate_date = ""
                certificate_valid_until = ""

            validated_rows.append((item, {
                "tnved_code": tnved,
                "certificate_type": certificate_type,
                "certificate_number": certificate_number,
                "certificate_date": certificate_date,
                "certificate_valid_until": certificate_valid_until,
            }))

        rows_by_group: dict[str, list[tuple[dict[str, Any], dict[str, str]]]] = {}
        for row in validated_rows:
            group = _assignment_product_group(row[0])
            if group not in {"perfumery", "chemistry"}:
                raise MarkingDbError(f"Товарная группа {group or 'не определена'} пока не поддерживается")
            rows_by_group.setdefault(group, []).append(row)

        created_documents: list[str] = []
        for product_group, group_rows in rows_by_group.items():
            products: list[dict[str, Any]] = []
            order_ids: list[int] = []
            for item, compliance in group_rows:
                product: dict[str, Any] = {
                    "uit_code": suz.identification_code(str(item.get("raw_code") or "")),
                    "tnved_code": compliance["tnved_code"],
                }
                if compliance["certificate_type"]:
                    product["certificate_document"] = compliance["certificate_type"]
                    product["certificate_document_number"] = compliance["certificate_number"]
                    product["certificate_document_date"] = compliance["certificate_date"]
                products.append(product)
                order_ids.append(int(item.get("order_id") or 0))

            document_payload = {
                "participant_inn": participant_inn,
                "producer_inn": participant_inn,
                "owner_inn": participant_inn,
                "production_date": production_date,
                "production_type": "OWN_PRODUCTION",
                "products": products,
            }
            response = suz.create_true_document(
                product_group=product_group,
                document_type="LP_INTRODUCE_GOODS",
                document_payload=document_payload,
            )
            document_uuid = str(response.get("uuid") or "").strip()
            if not document_uuid:
                raise MarkingDbError(f"True API не вернул UUID документа для {product_group}")
            create_circulation_document(
                marking_db_path,
                supply_id=supply_id,
                document_uuid=document_uuid,
                document_type="LP_INTRODUCE_GOODS",
                product_group=product_group,
                payload=document_payload,
                response=response,
                order_ids=order_ids,
                submission_kind="introduction",
            )
            created_documents.append(document_uuid)

        # Keep only unambiguous GTIN fallbacks; exact article data remains authoritative.
        fallback_compliance_by_gtin: dict[str, dict[str, str]] = {}
        conflicting_fallback_gtins: set[str] = set()
        for assigned_item, compliance in validated_rows:
            gtin = canonical_gtin14(assigned_item.get("gtin"))
            if gtin in conflicting_fallback_gtins:
                continue
            previous = fallback_compliance_by_gtin.get(gtin)
            if previous is not None and previous != compliance:
                fallback_compliance_by_gtin.pop(gtin, None)
                conflicting_fallback_gtins.add(gtin)
                continue
            fallback_compliance_by_gtin[gtin] = dict(compliance)
        save_suz_local_settings({
            "suz_auth_inn": participant_inn,
            "suz_true_participant_inn": participant_inn,
            "suz_true_producer_inn": participant_inn,
            "suz_true_owner_inn": participant_inn,
            "suz_true_tnved_by_gtin": {
                gtin: item["tnved_code"] for gtin, item in fallback_compliance_by_gtin.items()
            },
            "suz_true_compliance_by_gtin": fallback_compliance_by_gtin,
        })
        override_note = (
            " Различия реквизитов повторяющегося GTIN подтверждены; для каждого КИЗа использован его артикул."
            if workspace_conflicts and payload.ignore_compliance_conflicts else ""
        )
        return {
            "ok": True,
            "stage": "introduction",
            "message": (
                f"Документы ввода в оборот отправлены: {', '.join(created_documents)}."
                f"{override_note} Нажмите «Проверить статусы»."
            ),
            "document_uuids": created_documents,
            "document_uuid": created_documents[0] if len(created_documents) == 1 else "",
            "codes": len(validated_rows),
            "already_submitted_codes": len(already_submitted),
        }
    except (MarkingDbError, SuzApiError) as exc:
        diagnostic_path = _save_true_api_diagnostic(
            supply_id=supply_id,
            document_payload=document_payload,
            exc=exc,
        )
        response = _marking_error_payload(exc, 409)
        if diagnostic_path:
            try:
                payload_body = json.loads(response.body.decode("utf-8"))
                payload_body["diagnostic_file"] = diagnostic_path
                return JSONResponse(payload_body, status_code=response.status_code)
            except Exception:
                pass
        return response
    except Exception as exc:
        _save_true_api_diagnostic(supply_id=supply_id, document_payload=document_payload, exc=exc)
        return _marking_error_payload(exc, 500)


@app.post("/api/supplies/{supply_id}/marking/circulation/introduce")
def marking_introduce_compatibility_job(supply_id: str, payload: CirculationIntroducePayload):
    """Compatibility URL: execute the legal operation outside the request thread."""
    requested_stage = str(payload.requested_stage or "introduction").strip().lower()
    if requested_stage == "application":
        worker = lambda report: _application_report_worker(supply_id, payload, report)
        title = "Отчет о нанесении"
        operation_key = f"marking-application:{supply_id}"
    else:
        worker = lambda report: _introduction_worker(supply_id, payload, report)
        title = "Ввод КИЗов в оборот"
        operation_key = f"marking-introduction:{supply_id}"
    job_id = submit_background_job(title, worker, operation_key=operation_key)
    return {"ok": True, "job_id": job_id, "message": f"{title} запущен"}


def _marking_lifecycle_payload_call(supply_id: str, payload: CirculationIntroducePayload) -> dict[str, Any]:
    return _json_response_payload(marking_introduce_into_circulation(supply_id, payload))


def _payload_for_stage(payload: CirculationIntroducePayload, stage: str) -> CirculationIntroducePayload:
    data = payload.model_dump() if hasattr(payload, "model_dump") else payload.dict()
    data["requested_stage"] = stage
    return CirculationIntroducePayload(**data)


def _application_report_worker(supply_id: str, payload: CirculationIntroducePayload, report: Callable[..., None]) -> dict[str, Any]:
    """Submit the chemistry application report and release the UI immediately.

    The remote OMS report is asynchronous. Waiting for its terminal status inside
    this background job made the browser follow the job for up to three minutes,
    which looked like a blocked workflow. The live marking watcher already owns
    status polling, so this worker must stop once CRPT accepted the report.
    """
    report(total=100, progress=5, message="Проверяю напечатанные КИЗы и данные дезодорантов…", phase="preflight")
    payload = _payload_for_stage(payload, "application")
    result = _marking_lifecycle_payload_call(supply_id, payload)
    if result.get("stage") != "utilisation":
        # If the report became complete between rendering and click, this is not an
        # error. Do not submit introduction from the report button.
        workspace = load_marking_workspace(supply_id, fetch_wb=False)
        if workspace.get("application_report_complete"):
            return {"ok": True, "message": "Отчет о нанесении уже подтвержден Честным Знаком."}
        raise MarkingDbError("Честный Знак не подтвердил создание отчета о нанесении")

    report_ids = [str(value) for value in (result.get("report_ids") or []) if str(value)]
    codes = int(result.get("codes") or 0)
    suffix = f" ({', '.join(report_ids)})" if report_ids else ""
    message = (
        f"Отчет о нанесении отправлен{suffix}: {codes} КИЗ. "
        "FBE отпустил интерфейс и продолжает проверять статус автоматически."
    )
    report(progress=100, message=message, phase="submitted")
    return {"ok": True, "processing": True, "report_ids": report_ids, "codes": codes, "message": message}


def _introduction_worker(supply_id: str, payload: CirculationIntroducePayload, report: Callable[..., None]) -> dict[str, Any]:
    """Submit introduction documents; the live watcher polls their remote state."""
    report(total=100, progress=5, message="Проверяю статус «Нанесен» и юридические реквизиты…", phase="preflight")
    payload = _payload_for_stage(payload, "introduction")
    result = _marking_lifecycle_payload_call(supply_id, payload)
    if result.get("stage") == "utilisation":
        raise MarkingDbError("Сначала завершите шаг 3 — отчет о нанесении")
    document_ids = [str(value) for value in (result.get("document_uuids") or []) if str(value)]
    codes = int(result.get("codes") or 0)
    message = (
        f"Документы ввода в оборот отправлены: {len(document_ids) or 1}; КИЗов: {codes}. "
        "FBE продолжает проверку Честного Знака автоматически — интерфейс не заблокирован."
    )
    report(progress=100, message=message, phase="submitted")
    return {"ok": True, "processing": True, "document_uuids": document_ids, "codes": codes, "message": message}


@app.post("/api/supplies/{supply_id}/marking/application/start")
def marking_start_application_report_job(supply_id: str, payload: CirculationIntroducePayload):
    job_id = submit_background_job(
        "Отчет о нанесении",
        lambda report: _application_report_worker(supply_id, payload, report),
        operation_key=f"marking-application:{supply_id}",
    )
    return {"ok": True, "job_id": job_id, "message": "Отчет о нанесении запущен"}


@app.post("/api/supplies/{supply_id}/marking/circulation/introduce/start")
def marking_start_introduction_job(supply_id: str, payload: CirculationIntroducePayload):
    job_id = submit_background_job(
        "Ввод КИЗов в оборот",
        lambda report: _introduction_worker(supply_id, payload, report),
        operation_key=f"marking-introduction:{supply_id}",
    )
    return {"ok": True, "job_id": job_id, "message": "Ввод в оборот запущен"}


@app.post("/api/supplies/{supply_id}/marking/{order_id}/send-wb")
def marking_send_to_wb(supply_id: str, order_id: int):
    try:
        _find_marking_order(supply_id, order_id)
        assignment = get_assignment(marking_db_path, order_id)
        if not assignment:
            raise MarkingDbError("Сначала получите КИЗ из СУЗ и закрепите его за заданием WB")
        if str(assignment.get("circulation_status") or "").strip().upper() not in INTRODUCED_CODE_STATUSES:
            raise MarkingDbError(
                "КИЗ еще не введен в оборот. Сначала завершите документ LP_INTRODUCE_GOODS "
                "и дождитесь статуса INTRODUCED."
            )
        raw_code = str(assignment.get("raw_code") or "")
        if not 16 <= len(raw_code) <= 135:
            raise MarkingDbError(f"WB принимает КИЗ длиной 16–135 символов, получено: {len(raw_code)}")

        meta_rows = wb.get_orders_meta([int(order_id)])
        meta_item = next((item for item in meta_rows if _meta_order_id(item) == int(order_id)), None)
        meta_info = _sgtin_meta_info(meta_item)
        if not meta_info["available"]:
            raise MarkingDbError("WB не вернул поле sgtin для этого задания, поэтому добавить КИЗ нельзя")

        wb.set_order_sgtin(int(order_id), [raw_code])
        update_wb_state(
            marking_db_path,
            int(order_id),
            status="sent_wb",
            decision=str(meta_info.get("decision") or ""),
            details="КИЗ отправлен методом PUT /api/v3/orders/{orderId}/meta/sgtin",
            event_type="sent_wb",
        )

        # WB validation may be asynchronous. One immediate read usually returns
        # filled/pending; the normal workspace refresh will continue checking.
        time.sleep(0.35)
        try:
            refreshed = wb.get_orders_meta([int(order_id)])
            refreshed_item = next((item for item in refreshed if _meta_order_id(item) == int(order_id)), None)
            refreshed_info = _sgtin_meta_info(refreshed_item)
            status = _status_from_wb_meta(refreshed_info, "sent_wb")
            update_wb_state(
                marking_db_path,
                int(order_id),
                status=status,
                decision=str(refreshed_info.get("decision") or ""),
                details=json.dumps(refreshed_info, ensure_ascii=False),
                event_type="wb_status_after_send",
            )
        except Exception:
            pass

        return {"ok": True, "message": "КИЗ отправлен в WB. Обновите статусы для результата проверки."}
    except MarkingDbError as exc:
        return _marking_error_payload(exc, 409)
    except WBApiError as exc:
        update_wb_state(
            marking_db_path,
            int(order_id),
            status="wb_error",
            details=str(exc),
            event_type="wb_send_error",
        )
        return _marking_error_payload(exc, 502)
    except Exception as exc:
        return _marking_error_payload(exc, 500)




@app.post("/api/supplies/{supply_id}/marking/send-wb-all")
def marking_send_all_to_wb(supply_id: str):
    """Send persisted raw KIZ strings directly from SQLite to WB.

    No XLSX/text export participates in this operation. ASCII 29 separators are
    serialized by JSON as \u001d and decoded by WB back to the original control
    character.
    """
    try:
        marking = load_marking_workspace(supply_id, fetch_wb=True)
        if int(marking.get("marking_orders") or 0) <= 0:
            raise MarkingDbError("В поставке нет маркируемых заданий")
        if not marking.get("all_marking_assigned"):
            raise MarkingDbError(
                f"Сначала закрепите все КИЗы: {marking.get('assigned_orders', 0)} из {marking.get('marking_orders', 0)}"
            )
        if not marking.get("all_marking_introduced"):
            raise MarkingDbError(
                f"Сначала введите все КИЗы в оборот: {marking.get('introduced_codes', 0)} "
                f"из {marking.get('marking_orders', 0)}"
            )

        sent = 0
        skipped = 0
        errors: list[str] = []
        order_ids: list[int] = []
        for item in marking.get("orders") or []:
            order_id = int(item.get("order_id") or 0)
            assignment = item.get("assignment") or {}
            raw_code = str(assignment.get("raw_code") or "")
            if not order_id or not raw_code:
                errors.append(f"{order_id or '?'}: отсутствует raw_code")
                continue
            current_status = str(assignment.get("status") or "")
            if current_status == "accepted_wb":
                skipped += 1
                continue
            if str(assignment.get("circulation_status") or "").strip().upper() not in INTRODUCED_CODE_STATUSES:
                errors.append(f"{order_id}: КИЗ не имеет статуса INTRODUCED")
                continue
            if not 16 <= len(raw_code) <= 135:
                errors.append(f"{order_id}: длина КИЗ {len(raw_code)} вне диапазона WB 16-135")
                continue
            if not item.get("wb_meta_available"):
                errors.append(f"{order_id}: WB не разрешил метаданные sgtin")
                continue
            try:
                wb.set_order_sgtin(order_id, [raw_code])
                update_wb_state(
                    marking_db_path,
                    order_id,
                    status="sent_wb",
                    decision=str(item.get("wb_decision") or ""),
                    details="КИЗ передан напрямую из SQLite через PUT /api/v3/orders/{orderId}/meta/sgtin",
                    event_type="sent_wb_bulk",
                )
                order_ids.append(order_id)
                sent += 1
                time.sleep(0.08)
            except Exception as exc:
                update_wb_state(
                    marking_db_path,
                    order_id,
                    status="wb_error",
                    details=str(exc),
                    event_type="wb_send_error",
                )
                errors.append(f"{order_id}: {exc}")

        # Refresh asynchronous WB validation state in batches.
        if order_ids:
            time.sleep(0.4)
            for start in range(0, len(order_ids), 100):
                try:
                    meta_rows = wb.get_orders_meta(order_ids[start : start + 100])
                except Exception:
                    continue
                for meta_item in meta_rows:
                    order_id = _meta_order_id(meta_item)
                    if not order_id:
                        continue
                    info = _sgtin_meta_info(meta_item)
                    existing = get_assignment(marking_db_path, order_id) or {}
                    status = _status_from_wb_meta(info, str(existing.get("status") or "sent_wb"))
                    update_wb_state(
                        marking_db_path,
                        order_id,
                        status=status,
                        decision=str(info.get("decision") or ""),
                        details=json.dumps(info, ensure_ascii=False),
                        event_type="wb_status_after_bulk_send",
                    )

        message = f"Передано в WB: {sent}. Уже принято ранее: {skipped}."
        if errors:
            message += " Ошибки: " + "; ".join(errors)
        return {
            "ok": not errors,
            "message": message,
            "sent": sent,
            "skipped": skipped,
            "errors": errors,
        }
    except (MarkingDbError, WBApiError) as exc:
        return _marking_error_payload(exc, 409)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


def _marking_send_all_to_wb_worker(supply_id: str, report: Callable[..., None]) -> dict[str, Any]:
    """Run the bulk WB KIZ upload outside the request thread.

    The legacy synchronous endpoint is intentionally retained for compatibility,
    while first-party FBE UI calls this worker through the background-job API.
    """
    report(total=100, progress=5, message="Проверяю готовность КИЗов к передаче WB…", phase="preflight")
    response = marking_send_all_to_wb(supply_id)
    if isinstance(response, JSONResponse):
        try:
            payload = json.loads(response.body.decode("utf-8"))
        except Exception:
            payload = {"ok": False, "error": f"HTTP {response.status_code}"}
        if response.status_code >= 400 or payload.get("ok") is False:
            raise MarkingDbError(str(payload.get("error") or payload.get("message") or f"HTTP {response.status_code}"))
    elif isinstance(response, dict):
        payload = dict(response)
        if payload.get("ok") is False and not payload.get("sent"):
            raise MarkingDbError(str(payload.get("error") or payload.get("message") or "Передача КИЗов в WB завершилась ошибкой"))
    else:
        raise MarkingDbError("FBE получил неизвестный ответ при передаче КИЗов в WB")

    sent = int(payload.get("sent") or 0)
    skipped = int(payload.get("skipped") or 0)
    errors = list(payload.get("errors") or [])
    report(
        progress=100,
        message=str(payload.get("message") or f"Передано в WB: {sent}. Уже принято ранее: {skipped}."),
        phase="complete" if not errors else "partial",
    )
    return payload


@app.post("/api/supplies/{supply_id}/marking/send-wb-all/start")
def marking_start_send_all_to_wb_job(supply_id: str):
    job_id = submit_background_job(
        "Передача КИЗов в WB",
        lambda report: _marking_send_all_to_wb_worker(supply_id, report),
        operation_key=f"marking-send-wb:{supply_id}",
    )
    return {"ok": True, "job_id": job_id, "message": "Передача КИЗов в WB запущена"}


@app.get("/api/suz/certificates")
def suz_certificates():
    try:
        certificates = suz.list_certificates()
        return {"ok": True, "certificates": certificates}
    except SuzApiError as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


@app.post("/api/suz/certificate")
def suz_select_certificate(payload: SuzCertificatePayload):
    thumbprint = str(payload.thumbprint or "").replace(" ", "").upper()
    if not thumbprint:
        return _marking_error_payload(ValueError("Не выбран сертификат УКЭП"), 400)
    save_suz_local_settings({"suz_cert_thumbprint": thumbprint})
    suz.reset_token()
    return {"ok": True, "message": "Сертификат сохранен локально", "thumbprint": thumbprint}


@app.post("/api/suz/test")
def suz_test_connection():
    try:
        suz.get_true_api_token(force=True)
        result = suz.ping(force_token=True)
        return {
            "ok": True,
            "message": "Подключение к True API и СУЗ подтверждено",
            "oms_id": result.get("omsId"),
            "api_version": result.get("apiVersion"),
            "oms_version": result.get("omsVersion"),
        }
    except SuzApiError as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


def _chunks(values: list[dict[str, Any]], size: int) -> list[list[dict[str, Any]]]:
    return [values[index:index + size] for index in range(0, len(values), size)]


def _infer_suz_product_group(product: dict[str, Any], article: str) -> str | None:
    explicit_group = str(product.get("product_group") or "").strip()
    if explicit_group:
        return explicit_group

    article_upper = str(article or "").strip().upper()
    category = str(product.get("category") or "").strip().lower()
    subcategory = str(product.get("subcategory") or "").strip().lower()
    if (
        category == "спа"
        or article_upper.startswith("SPA-")
        or "дезодоран" in subcategory
        or "космет" in category
    ):
        return "chemistry"

    perfume_words = ("духи", "парфюм", "туалетная вода", "парфюмерная вода")
    if (
        category == "парфюмерия"
        or article_upper.startswith(("PRF-", "PRW-", "TOW-"))
        or any(word in subcategory for word in perfume_words)
    ):
        return "perfumery"
    return None


def _suz_product_group_for_article(article: str) -> str | None:
    """Return the OMS goods group stored in the catalog, with legacy inference fallback."""
    article_text = str(article or "").strip()
    product = get_catalog().find_by_article(article_text)
    if not product:
        return None
    return _infer_suz_product_group(product, article_text)


def _suz_template_for_group(product_group: str) -> int:
    if product_group == "chemistry":
        return 46
    if product_group == "perfumery":
        return 9
    raise SuzApiError(f"Для товарной группы {product_group!r} не настроен шаблон КМ")


@app.post("/api/supplies/{supply_id}/suz/order")
def suz_create_supply_order(supply_id: str, payload: SuzOrderPayload):
    try:
        readiness = suz_readiness()
        if not readiness.get("ready"):
            missing = ", ".join(
                str(item.get("label"))
                for item in readiness.get("items") or []
                if not item.get("ready")
            )
            raise SuzApiError(f"Подключение СУЗ не настроено: {missing or 'неизвестная причина'}")

        reserve = max(0, min(10, int(payload.reserve_per_gtin or 0)))
        chemistry_payment_type = int(payload.payment_type or 1)
        if chemistry_payment_type not in {1, 2}:
            raise SuzApiError("Тип оплаты для chemistry должен быть 1 или 2")

        marking = load_marking_workspace(supply_id, fetch_wb=False)
        if marking.get("missing_orders") or marking.get("unresolved_orders"):
            raise SuzApiError("Нельзя создавать заказ: в поставке есть товары без GTIN или неизвестные артикулы")

        settings_now = load_suz_settings()
        serial_number_type = str(settings_now.get("suz_serial_number_type") or "OPERATOR")
        release_method_type = str(settings_now.get("suz_release_method_type") or "PRODUCTION")

        grouped_products: dict[str, list[dict[str, Any]]] = {"chemistry": [], "perfumery": []}
        unsupported_articles: list[str] = []
        conflicting_gtins: list[str] = []
        catalog = get_catalog()

        for row in marking.get("rows") or []:
            quantity = int(row.get("to_order") or 0)
            if quantity <= 0:
                continue
            articles = [str(value) for value in (row.get("articles") or []) if str(value).strip()]
            products_for_articles = [catalog.find_by_article(article) for article in articles]
            products_for_articles = [item for item in products_for_articles if item]
            groups = {
                _infer_suz_product_group(item, str(item.get("seller_article") or ""))
                for item in products_for_articles
            }
            groups.discard(None)
            groups.discard("")
            if not groups:
                unsupported_articles.extend(articles)
                continue
            if len(groups) != 1:
                conflicting_gtins.append(str(row.get("gtin") or ""))
                continue
            product_group = next(iter(groups))
            if product_group not in grouped_products:
                unsupported_articles.extend(articles)
                continue

            template_ids = {
                int(item.get("template_id") or 0)
                for item in products_for_articles
                if int(item.get("template_id") or 0) > 0
            }
            if len(template_ids) > 1:
                conflicting_gtins.append(str(row.get("gtin") or ""))
                continue
            cis_types = {
                str(item.get("cis_type") or "").strip()
                for item in products_for_articles
                if str(item.get("cis_type") or "").strip()
            }
            if len(cis_types) > 1:
                conflicting_gtins.append(str(row.get("gtin") or ""))
                continue
            product_row = {
                "gtin": str(row.get("gtin") or ""),
                "quantity": quantity + reserve,
                "serialNumberType": serial_number_type,
                "templateId": next(iter(template_ids)) if template_ids else _suz_template_for_group(product_group),
                "cisType": next(iter(cis_types)) if cis_types else str(settings_now.get("suz_cis_type") or "UNIT"),
            }
            grouped_products[product_group].append(product_row)

        if conflicting_gtins:
            raise SuzApiError(
                "Один GTIN связан одновременно с товарами разных групп СУЗ: "
                + ", ".join(sorted(set(conflicting_gtins)))
                + ". Проверьте вкладку «Каталог товаров»."
            )
        if unsupported_articles:
            articles_text = ", ".join(sorted(set(unsupported_articles)))
            raise SuzApiError(
                "Не удалось определить товарную группу СУЗ для артикулов: "
                f"{articles_text}. Для них нужно проверить Категорию и Подкатегорию в каталоге товаров."
            )
        if not any(grouped_products.values()):
            raise SuzApiError("Недостающих кодов для заказа нет")

        contact_person = str(settings_now.get("suz_contact_person") or "").strip()
        created: list[dict[str, Any]] = []

        # A single OMS order has exactly one productGroup. A mixed WB supply is split
        # automatically into independent chemistry and perfumery orders.
        for product_group in ("chemistry", "perfumery"):
            products = grouped_products[product_group]
            if not products:
                continue

            attributes: dict[str, Any] = {"releaseMethodType": release_method_type}
            if contact_person:
                attributes["contactPerson"] = contact_person
            # paymentType exists in the chemistry order schema, but is absent from
            # the perfumery order schema. Sending it to perfumery may cause HTTP 400.
            if product_group == "chemistry":
                attributes["paymentType"] = chemistry_payment_type

            for batch in _chunks(products, 10):
                try:
                    result = suz.create_order(
                        batch,
                        attributes=attributes,
                        product_group=product_group,
                    )
                    oms_order_id = str(result.get("orderId") or "").strip()
                    if not oms_order_id:
                        raise SuzApiError(f"СУЗ не вернул orderId: {result}")
                    request_payload = dict(result.pop("requestPayload", {}))
                    record = create_suz_order(
                        marking_db_path,
                        supply_id=supply_id,
                        oms_order_id=oms_order_id,
                        expected_complete_ms=int(result.get("expectedCompleteTimestamp") or 0),
                        payload=request_payload,
                        response=result,
                    )
                    created.append(
                        {
                            "local_id": record.get("id"),
                            "order_id": oms_order_id,
                            "product_group": product_group,
                            "gtin_count": len(batch),
                            "code_count": sum(int(item.get("quantity") or 0) for item in batch),
                        }
                    )
                except Exception as exc:
                    if product_group == "chemistry" and chemistry_payment_type == 2 and "HTTP 400" in str(exc):
                        exc = SuzApiError(
                            f"{exc}. Для chemistry режим оплаты 2 разрешен только отдельным "
                            "пилотным GTIN. Выберите оплату при эмиссии и повторите заказ."
                        )
                    if created:
                        identifiers = ", ".join(
                            f"{item['product_group']}:{item['order_id']}" for item in created
                        )
                        raise SuzApiError(
                            f"Часть заказов уже создана в СУЗ ({identifiers}), затем группа "
                            f"{product_group} завершилась ошибкой: {exc}. Не нажимайте кнопку повторно — "
                            "обновите окно и проверьте созданные заказы."
                        ) from exc
                    raise

        group_labels = {
            "chemistry": "косметика/дезодоранты",
            "perfumery": "парфюмерия",
        }
        summary = ", ".join(
            f"{group_labels.get(item['product_group'], item['product_group'])}: {item['code_count']} КМ"
            for item in created
        )
        return {
            "ok": True,
            "message": f"Создано заказов СУЗ: {len(created)} ({summary})",
            "orders": created,
        }
    except SuzApiError as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


def _json_response_payload(value: Any) -> dict[str, Any]:
    if isinstance(value, JSONResponse):
        try:
            payload = json.loads(value.body.decode("utf-8"))
        except Exception:
            payload = {"ok": False, "error": f"HTTP {value.status_code}"}
        if value.status_code >= 400 or not payload.get("ok", False):
            raise SuzApiError(str(payload.get("error") or payload.get("message") or f"HTTP {value.status_code}"))
        return dict(payload)
    if isinstance(value, dict):
        if value.get("ok") is False:
            raise SuzApiError(str(value.get("error") or value.get("message") or "Операция СУЗ завершилась ошибкой"))
        return dict(value)
    raise SuzApiError("FBE получил неизвестный ответ внутреннего обработчика СУЗ")


def _suz_order_live_summary(supply_id: str, local_ids: list[int]) -> tuple[str, bool, int, int]:
    parts: list[str] = []
    requested_all = 0
    received_all = 0
    complete = True
    for local_id in local_ids:
        order = get_suz_order(marking_db_path, int(local_id)) or {}
        order_requested = 0
        order_received = 0
        states: list[str] = []
        for item in order.get("items") or []:
            requested = int(item.get("requested_quantity") or 0)
            received = int(item.get("received_quantity") or 0)
            order_requested += requested
            order_received += received
            states.append(str(item.get("buffer_status") or "PENDING").upper())
        requested_all += order_requested
        received_all += order_received
        if order_requested <= 0 or order_received < order_requested:
            complete = False
        state = "/".join(sorted(set(states))) if states else str(order.get("status") or "PENDING").upper()
        parts.append(f"{str(order.get('oms_order_id') or local_id)[:8]}… {state}: {order_received}/{order_requested}")
    return " · ".join(parts), complete, received_all, requested_all


def _suz_create_and_watch_worker(
    supply_id: str,
    reserve_per_gtin: int,
    payment_type: int,
    report: Callable[..., None],
) -> dict[str, Any]:
    report(total=100, progress=3, message="Проверяю подключение OMS перед заказом…", phase="preflight")
    # GET ping is idempotent: retry it before the non-idempotent POST /order. We do
    # NOT blindly replay POST /order after an ambiguous transport failure because
    # that could create a duplicate paid code order in CRPT.
    ping_error: Exception | None = None
    for ping_attempt in range(1, 4):
        try:
            suz.ping()
            ping_error = None
            break
        except SuzApiError as exc:
            ping_error = exc
            report(
                progress=3 + ping_attempt,
                message=f"OMS временно не ответил ({ping_attempt}/3). Повторяю проверку автоматически…",
                phase="preflight",
            )
            if ping_attempt < 3:
                time.sleep(1.2 * ping_attempt)
    if ping_error is not None:
        raise SuzApiError(f"OMS недоступен после 3 проверок: {ping_error}") from ping_error
    report(progress=9, message="Подписываю заявку УКЭП и отправляю заказ КИЗ…", phase="ordering")
    created_payload = _json_response_payload(
        suz_create_supply_order(
            supply_id,
            SuzOrderPayload(reserve_per_gtin=reserve_per_gtin, payment_type=payment_type),
        )
    )
    local_ids = [int(item.get("local_id") or 0) for item in created_payload.get("orders") or [] if int(item.get("local_id") or 0)]
    if not local_ids:
        return created_payload

    report(progress=18, message="СУЗ принял заказ. Выполняю первую короткую проверку кодов…", phase="pending")
    # Do not occupy a scarce user-job worker for the full remote processing
    # window. The accepted OMS order is already durable in SQLite; after this
    # short initial check the independent status refresh owns further polling.
    initial_watch_seconds = min(
        30.0, max(5.0, float(config.get("suz_order_initial_watch_seconds", 15) or 15))
    )
    deadline = time.time() + initial_watch_seconds
    attempt = 0
    last_errors: list[str] = []
    while time.time() < deadline:
        attempt += 1
        transient_errors: list[str] = []
        for local_id in local_ids:
            order = get_suz_order(marking_db_path, local_id) or {}
            if int(order.get("closed") or 0):
                continue
            try:
                _sync_suz_order_record(supply_id, local_id)
            except (SuzApiError, MarkingCodeError) as exc:
                transient_errors.append(str(exc))
                update_suz_order_state(marking_db_path, local_id, error_text=str(exc))
        summary, complete, received, requested = _suz_order_live_summary(supply_id, local_ids)
        progress_value = 20 + int(75 * (received / requested)) if requested else min(92, 20 + attempt * 3)
        if complete:
            report(progress=100, message=f"КИЗы получены и закреплены: {received}/{requested}. {summary}", phase="complete")
            return {
                "ok": True,
                "message": f"КИЗы получены и закреплены: {received}/{requested}.",
                "orders": created_payload.get("orders") or [],
                "received": received,
                "requested": requested,
            }
        last_errors = transient_errors or last_errors
        message = f"СУЗ обрабатывает заказ · получено {received}/{requested}. {summary}"
        if transient_errors:
            message += " · сервер временно не ответил, повторяю автоматически"
        report(progress=min(96, progress_value), message=message, phase="receiving")
        time.sleep(min(6.0, 1.2 + attempt * 0.55))

    summary, complete, received, requested = _suz_order_live_summary(supply_id, local_ids)
    # Do not convert a slow CRPT response into a failed order: keep the persisted OMS
    # order and let the normal live watcher continue later.
    return {
        "ok": True,
        "partial": True,
        "message": (
            f"Заказ принят СУЗ, получено {received}/{requested}. FBE продолжит проверять его автоматически. {summary}"
            + (f" Последний временный ответ: {last_errors[-1]}" if last_errors else "")
        ),
        "orders": created_payload.get("orders") or [],
        "received": received,
        "requested": requested,
    }


@app.post("/api/supplies/{supply_id}/suz/order/start")
def suz_start_supply_order_job(supply_id: str, payload: SuzOrderPayload):
    job_id = submit_background_job(
        "Заказ КИЗ",
        lambda report: _suz_create_and_watch_worker(
            supply_id,
            int(payload.reserve_per_gtin or 0),
            int(payload.payment_type or 1),
            report,
        ),
        operation_key=f"suz-order:{supply_id}",
    )
    return {"ok": True, "job_id": job_id, "message": "Заказ КИЗ запущен"}


def _save_suz_code_response(
    *,
    local_order_id: int,
    gtin: str,
    response: dict[str, Any],
) -> int:
    codes = response.get("codes") or []
    parsed_codes: list[dict[str, Any]] = []
    for raw_code in codes:
        parsed = parse_marking_code(raw_code)
        if parsed.get("gtin") != gtin:
            raise SuzApiError(
                f"СУЗ вернул код с GTIN {parsed.get('gtin')}, ожидался {gtin}"
            )
        parsed_codes.append(parsed)
    return save_suz_codes(
        marking_db_path,
        local_order_id=local_order_id,
        gtin=gtin,
        block_id=str(response.get("blockId") or ""),
        parsed_codes=parsed_codes,
    )


def _retry_suz_blocks(
    *,
    local_order_id: int,
    oms_order_id: str,
    gtin: str,
    missing: int,
) -> int:
    """Recover code blocks previously issued through API but not saved locally."""
    if missing <= 0:
        return 0
    try:
        blocks = suz.get_code_blocks(order_id=oms_order_id, gtin=gtin)
    except SuzApiError:
        return 0

    inserted_total = 0
    for block in sorted(blocks, key=lambda item: int(item.get("blockDateTime") or 0)):
        if inserted_total >= missing:
            break
        block_id = str(block.get("blockId") or "").strip()
        if not block_id:
            continue
        response = suz.retry_codes(block_id=block_id)
        inserted_total += _save_suz_code_response(
            local_order_id=local_order_id,
            gtin=gtin,
            response=response,
        )
    return inserted_total


def _sync_suz_order_record(supply_id: str, local_order_id: int) -> dict[str, Any]:
    """Refresh OMS state, download full code strings and persist them losslessly."""
    order = get_suz_order(marking_db_path, local_order_id)
    if not order or str(order.get("supply_id") or "") != str(supply_id):
        raise SuzApiError("Заказ СУЗ не найден для этой поставки")

    oms_order_id = str(order.get("oms_order_id") or "")
    statuses = suz.order_status(oms_order_id)
    status_by_gtin = {str(item.get("gtin") or ""): item for item in statuses}
    ready_states = {"ACTIVE", "EXHAUSTED", "CLOSED", "COMPLETED"}
    received_total = 0
    rejected = False
    item_states: list[str] = []

    for item in order.get("items") or []:
        gtin = str(item.get("gtin") or "")
        status_data = status_by_gtin.get(gtin, {})
        if status_data:
            update_suz_item_status(
                marking_db_path,
                local_order_id=local_order_id,
                gtin=gtin,
                status_data=status_data,
            )
        buffer_status = str(
            status_data.get("bufferStatus") or item.get("buffer_status") or "PENDING"
        ).upper()
        item_states.append(buffer_status)
        if buffer_status == "REJECTED":
            rejected = True
            continue
        if buffer_status not in ready_states:
            continue

        # Re-read local quantities because a previous attempt may have saved only
        # part of a block before interruption.
        fresh_order = get_suz_order(marking_db_path, local_order_id) or order
        fresh_item = next(
            (row for row in fresh_order.get("items") or [] if str(row.get("gtin") or "") == gtin),
            item,
        )
        requested = int(fresh_item.get("requested_quantity") or 0)
        received = int(fresh_item.get("received_quantity") or 0)
        missing = max(0, requested - received)
        if missing <= 0:
            continue

        # First recover any already issued block. This prevents consuming a new
        # block after a local write/network failure.
        recovered = _retry_suz_blocks(
            local_order_id=local_order_id,
            oms_order_id=oms_order_id,
            gtin=gtin,
            missing=missing,
        )
        received_total += recovered
        missing = max(0, missing - recovered)

        if missing > 0 and buffer_status == "ACTIVE":
            response = suz.get_codes(
                order_id=oms_order_id,
                gtin=gtin,
                quantity=missing,
            )
            received_total += _save_suz_code_response(
                local_order_id=local_order_id,
                gtin=gtin,
                response=response,
            )

    if rejected:
        state = "rejected"
    else:
        fresh_order = get_suz_order(marking_db_path, local_order_id) or order
        complete = bool(fresh_order.get("items")) and all(
            int(item.get("received_quantity") or 0) >= int(item.get("requested_quantity") or 0)
            for item in fresh_order.get("items") or []
        )
        if complete:
            state = "codes_received"
        elif any(status in ready_states for status in item_states):
            state = "partly_received"
        else:
            state = "pending"

    update_suz_order_state(
        marking_db_path,
        local_order_id,
        status=state,
        error_text="",
        closed=(state == "codes_received"),
    )

    # Assign every newly received code to a concrete WB order immediately.
    workspace = load_marking_workspace(supply_id, fetch_wb=False)
    assignment_result = assign_suz_codes_to_supply_orders(
        marking_db_path,
        supply_id=supply_id,
        order_rows=list(workspace.get("orders") or []),
    )
    labels = {
        "codes_received": "КИЗы получены, сохранены и закреплены за заданиями WB.",
        "partly_received": "Часть КИЗов получена; остальные буферы еще обрабатываются.",
        "pending": "Заказ еще обрабатывается СУЗ.",
        "rejected": "Заказ или один из GTIN отклонен СУЗ.",
    }
    message = labels.get(state, "Статус заказа обновлен.")
    if assignment_result.get("assigned"):
        message += f" Новых привязок: {assignment_result['assigned']}."
    return {
        "ok": True,
        "order_id": local_order_id,
        "oms_order_id": oms_order_id,
        "message": message,
        "status": state,
        "received": received_total,
        "assigned": int(assignment_result.get("assigned") or 0),
        "shortages": int(assignment_result.get("shortages") or 0),
    }


def _sync_all_supply_orders_impl(supply_id: str) -> dict[str, Any]:
    orders = [item for item in list_open_suz_orders(marking_db_path, supply_id) if item]
    if not orders:
        return {"ok": True, "message": "Открытых заказов СУЗ для проверки нет", "orders": [], "received": 0, "errors": []}

    results: list[dict[str, Any]] = []
    errors: list[str] = []
    for order in orders:
        local_id = int(order.get("id") or 0)
        try:
            results.append(_sync_suz_order_record(supply_id, local_id))
        except (SuzApiError, MarkingCodeError) as exc:
            update_suz_order_state(marking_db_path, local_id, error_text=str(exc))
            errors.append(f"{order.get('oms_order_id')}: {exc}")

    received_total = sum(int(item.get("received") or 0) for item in results)
    message = f"Проверено заказов: {len(orders)}. Получено новых КИЗов: {received_total}."
    if errors:
        message += " Часть запросов временно не ответила: " + "; ".join(errors)
    return {
        "ok": True,
        "partial": bool(errors),
        "message": message,
        "received": received_total,
        "orders": results,
        "errors": errors,
    }


@app.post("/api/supplies/{supply_id}/suz/orders/sync-all")
def suz_sync_all_supply_orders(supply_id: str):
    try:
        return _sync_all_supply_orders_impl(supply_id)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


@app.post("/api/supplies/{supply_id}/suz/orders/{local_order_id}/sync")
def suz_sync_supply_order(supply_id: str, local_order_id: int):
    try:
        return _sync_suz_order_record(supply_id, local_order_id)
    except (SuzApiError, MarkingCodeError) as exc:
        update_suz_order_state(marking_db_path, local_order_id, error_text=str(exc))
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        update_suz_order_state(marking_db_path, local_order_id, error_text=str(exc))
        return _marking_error_payload(exc, 500)


def _build_marking_print_pdf(
    supply_id: str,
    uploaded_by_gtin: dict[str, tuple[Path, int]],
    *,
    print_mode: str = "paired",
) -> tuple[Path, int, list[str], str]:
    """Build an official-size PDF lane from untouched KIZ pages.

    legacy paired mode: retained only for backward compatibility; 0.83.0 uses native 30x20 mm pages.
    kiz_only: concatenate official KIZ pages in supply order.

    The title pages are generated locally as vector PDF text; BarTender is not
    involved, so a stalled BarTender/Bullzip request cannot block the result.
    If title-page generation unexpectedly fails, FBE automatically returns
    a KIZ-only PDF instead of losing the uploaded official files.
    """
    marking = load_marking_workspace(supply_id, fetch_wb=False)
    rows = [row for row in (marking.get("print_rows") or []) if int(row.get("count") or 0) > 0]
    if not rows:
        raise ValueError("В поставке нет готовых маркируемых позиций")

    expected_gtins = {str(row.get("gtin") or "") for row in rows}
    supplied_gtins = set(uploaded_by_gtin)
    missing = sorted(expected_gtins - supplied_gtins)
    unknown = sorted(supplied_gtins - expected_gtins)
    if missing:
        raise ValueError("Не загружен официальный PDF для GTIN: " + ", ".join(missing))
    if unknown:
        raise ValueError("Загружен PDF для GTIN, которого нет в поставке: " + ", ".join(unknown))

    from pypdf import PdfReader

    catalog = get_catalog()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    work_dir = Path(config.get("marking_print_output_dir", "data/marking_print")) / supply_id / ts
    work_dir.mkdir(parents=True, exist_ok=True)

    items: list[dict[str, Any]] = []
    warnings: list[str] = []
    log_lines = [
        "FBE official-size KIZ print package",
        f"supply_id={supply_id}",
        "Official KIZ pages are copied without decoding, scaling or redrawing Data Matrix.",
        "sequence;gtin;seller_article;product;kiz_pdf;kiz_page;width_mm;height_mm",
    ]

    sequence = 0
    reader_cache: dict[str, Any] = {}
    for row in rows:
        gtin = str(row.get("gtin") or "")
        official_pdf, page_count = uploaded_by_gtin[gtin]
        expected = int(row.get("count") or 0)
        if page_count != expected:
            raise ValueError(
                f"GTIN {gtin}: в поставке нужно {expected} КИЗ, а в PDF {page_count} страниц. "
                "Скачайте в Честном Знаке стандартный PDF ровно на нужное количество."
            )

        articles = [str(value) for value in (row.get("articles") or []) if str(value).strip()]
        product = catalog.find_by_article(articles[0]) if articles else None
        if not product:
            raise ValueError(f"GTIN {gtin}: не удалось найти товар в каталоге товаров")
        if len(articles) > 1:
            warnings.append(
                f"GTIN {gtin} связан с несколькими артикулами ({', '.join(articles)}); "
                f"для титульной страницы использован {articles[0]}."
            )

        key = str(Path(official_pdf).resolve())
        reader = reader_cache.get(key)
        if reader is None:
            reader = PdfReader(key)
            reader_cache[key] = reader

        for page_index in range(page_count):
            page = reader.pages[page_index]
            width_pt = float(page.mediabox.width)
            height_pt = float(page.mediabox.height)
            sequence += 1
            items.append(
                {
                    "sequence": sequence,
                    "gtin": gtin,
                    "product": product,
                    "kiz_ref": (official_pdf, page_index),
                    "width_pt": width_pt,
                    "height_pt": height_pt,
                }
            )
            log_lines.append(
                f"{sequence};{gtin};{product.get('seller_article','')};"
                f"{product.get('display_name','')};{official_pdf};{page_index + 1};"
                f"{width_pt * 25.4 / 72.0:.2f};{height_pt * 25.4 / 72.0:.2f}"
            )

    items.sort(key=lambda value: int(value["sequence"]))
    kiz_refs = [item["kiz_ref"] for item in items]
    requested_mode = str(print_mode or "paired").strip().lower()
    if requested_mode not in {"paired", "kiz_only"}:
        requested_mode = "paired"

    actual_mode = requested_mode
    if requested_mode == "kiz_only":
        final_pdf = work_dir / f"FBE_KIZ_ONLY_{_safe_filename_part(supply_id)}_{ts}.pdf"
        merge_pdf_page_refs(kiz_refs, final_pdf)
    else:
        try:
            title_pdf = work_dir / "title_pages.pdf"
            print_scale_percent = float(config.get("marking_print_scale_percent", 55) or 55)
            title_reference_percent = float(
                config.get("marking_title_reference_scale_percent", 85) or 85
            )
            if print_scale_percent <= 0:
                print_scale_percent = 55.0
            title_layout_scale = title_reference_percent / print_scale_percent

            create_small_title_labels_pdf(
                [
                    {
                        "seller_article": item["product"].get("seller_article", ""),
                        "display_name": item["product"].get("display_name", "")
                        or item["product"].get("product", ""),
                        "width_pt": item["width_pt"],
                        "height_pt": item["height_pt"],
                    }
                    for item in items
                ],
                title_pdf,
                layout_scale=title_layout_scale,
            )
            title_refs = [(title_pdf, index) for index in range(len(items))]
            final_pdf = work_dir / f"FBE_KIZ_TITLES_{_safe_filename_part(supply_id)}_{ts}.pdf"
            merge_title_labels_with_official_kiz_pages(title_refs, kiz_refs, final_pdf)
        except Exception as exc:
            actual_mode = "kiz_only_fallback"
            warnings.append(
                "Титульные страницы не созданы; FBE автоматически собрал "
                f"официальные КИЗы без титулов. Причина: {exc}"
            )
            final_pdf = work_dir / f"FBE_KIZ_ONLY_FALLBACK_{_safe_filename_part(supply_id)}_{ts}.pdf"
            merge_pdf_page_refs(kiz_refs, final_pdf)

    if warnings:
        log_lines.extend(["", "Warnings:", *warnings])
    log_lines.extend(["", f"MODE={actual_mode}", f"FINAL={final_pdf}", f"KIZ_PAGES={len(items)}"])
    (work_dir / "print_order_log.txt").write_text("\n".join(log_lines), encoding="utf-8")

    # Local UI state only: an official print PDF has been assembled.
    for order in list_suz_orders(marking_db_path, supply_id, limit=100):
        if str(order.get("status") or "") != "rejected":
            update_suz_order_state(
                marking_db_path,
                int(order.get("id") or 0),
                status="print_pdf_ready",
                error_text="",
                closed=True,
            )

    return final_pdf, len(items), warnings, actual_mode


@app.post("/supplies/{supply_id}/marking/official-print-pdf")
async def marking_official_print_pdf(
    supply_id: str,
    gtins: list[str] = Form(...),
    kiz_pdfs: list[UploadFile] = File(...),
    print_mode: str = Form("paired"),
):
    if len(gtins) != len(kiz_pdfs):
        return _marking_error_payload(
            ValueError("Количество GTIN и загруженных PDF не совпадает"), 400
        )

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    upload_dir = Path(config.get("marking_print_output_dir", "data/marking_print")) / supply_id / ts / "official_pdf"
    upload_dir.mkdir(parents=True, exist_ok=True)
    uploaded_by_gtin: dict[str, tuple[Path, int]] = {}

    try:
        from pypdf import PdfReader

        for index, (gtin_raw, upload) in enumerate(zip(gtins, kiz_pdfs), start=1):
            gtin = canonical_gtin14(gtin_raw)
            error = validate_gtin(gtin)
            if error:
                raise ValueError(f"Строка {index}: {error}")
            if gtin in uploaded_by_gtin:
                raise ValueError(f"GTIN {gtin}: выбран более одного PDF")
            filename = str(upload.filename or "").lower()
            if filename and not filename.endswith(".pdf"):
                raise ValueError(f"GTIN {gtin}: нужен PDF из Честного Знака")
            content = await upload.read()
            if not content:
                raise ValueError(f"GTIN {gtin}: загружен пустой файл")
            path = upload_dir / f"{index:02d}_{gtin}.pdf"
            path.write_bytes(content)
            try:
                reader = PdfReader(str(path))
                page_count = len(reader.pages)
            except Exception as exc:
                raise ValueError(f"GTIN {gtin}: PDF не читается: {exc}") from exc
            if page_count <= 0:
                raise ValueError(f"GTIN {gtin}: в PDF нет страниц")
            uploaded_by_gtin[gtin] = (path, page_count)

        final_pdf, pair_count, warnings, actual_mode = _build_marking_print_pdf(
            supply_id, uploaded_by_gtin, print_mode=print_mode
        )
    except Exception as exc:
        return _marking_error_payload(exc, 400)

    safe_supply = _safe_filename_part(supply_id, "supply")
    download_prefix = "FBE_KIZ_TITLES" if actual_mode == "paired" else "FBE_KIZ_ONLY"
    download_name = f"{download_prefix}_{safe_supply}.pdf"
    return FileResponse(
        str(final_pdf),
        media_type="application/pdf",
        filename=download_name,
        headers={
            "Content-Disposition": f'inline; filename="{download_name}"',
            "X-FBE-KIZ-Pages": str(pair_count),
            "X-FBE-Print-Mode": actual_mode,
            "X-FBE-Warnings": str(len(warnings)),
        },
    )


def _marking_sticker_suffixes(supply_id: str, rows: list[dict[str, Any]]) -> dict[int, str]:
    """Fetch/cache WB stickers and return the last four digits for KIZ title labels.

    A WB order id / assembly task id is NOT a sticker number and must never be used as
    a visual fallback: it can point the operator to the wrong physical order.
    """
    order_ids = [int(item.get("order_id") or item.get("id") or 0) for item in rows]
    order_ids = [value for value in order_ids if value]
    metadata = get_or_fetch_sticker_metadata(supply_id, order_ids)
    result: dict[int, str] = {}
    missing: list[int] = []
    for order_id in order_ids:
        meta = metadata.get(order_id) or {}
        full = str(meta.get("full") or "").strip()
        digits = "".join(ch for ch in full if ch.isdigit())
        if len(digits) < 4:
            missing.append(order_id)
            continue
        result[order_id] = digits[-4:]
    if missing:
        raise ValueError(
            "WB не вернул номер стикера для заданий: " + ", ".join(str(value) for value in missing)
            + ". FBE не будет подменять его номером сборочного задания."
        )
    return result


def _build_native_marking_print_pdf(supply_id: str) -> tuple[Path, int]:
    """Build native 30x20 mm title/DataMatrix pairs from persisted raw KIZs."""
    marking = load_marking_workspace(supply_id, fetch_wb=False)
    blockers = list(marking.get("print_blockers") or [])
    if blockers:
        raise ValueError("; ".join(str(value) for value in blockers))

    rows = sorted(
        list(marking.get("orders") or []),
        key=lambda item: (str(item.get("name") or "").lower(), int(item.get("order_id") or 0)),
    )
    if not rows:
        raise ValueError("В поставке нет маркируемых позиций")
    sticker_suffixes = _marking_sticker_suffixes(supply_id, rows)

    stored = {
        int(item.get("order_id") or 0): item
        for item in list_supply_code_assignments(marking_db_path, supply_id)
        if int(item.get("order_id") or 0)
    }
    missing_orders = [
        int(item.get("order_id") or 0)
        for item in rows
        if not str((stored.get(int(item.get("order_id") or 0)) or item.get("assignment") or {}).get("raw_code") or "")
    ]
    if missing_orders:
        raise ValueError(
            "Не закреплен КИЗ за заданиями WB: " + ", ".join(str(value) for value in missing_orders)
        )

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    work_dir = Path(config.get("marking_print_output_dir", "data/marking_print")) / supply_id / ts
    png_dir = work_dir / "datamatrix"
    png_dir.mkdir(parents=True, exist_ok=True)

    title_labels: list[dict[str, object]] = []
    matrix_labels: list[dict[str, object]] = []
    order_ids: list[int] = []
    log_lines = [
        f"FBE {APP_VERSION} native marking print",
        f"supply_id={supply_id}",
        "page_size_mm=30x20",
        "raw KIZs are read directly from SQLite; XLSX is not used",
        "sequence;wb_order_id;seller_article;gtin;serial;png",
    ]

    for sequence, item in enumerate(rows, start=1):
        order_id = int(item.get("order_id") or 0)
        assignment = stored.get(order_id) or item.get("assignment") or {}
        raw_code = str(assignment.get("raw_code") or "")
        gtin = str(assignment.get("gtin") or item.get("gtin14") or item.get("gtin") or "")
        serial = str(assignment.get("serial") or "")
        png_path = png_dir / f"{sequence:04d}_{order_id}.png"
        datamatrix_renderer.render_png(
            raw_code,
            png_path,
            dpi=int(config.get("marking_datamatrix_dpi", 300) or 300),
        )

        title_labels.append(
            {
                "seller_article": str(item.get("seller_article") or assignment.get("seller_article") or ""),
                "display_name": str(item.get("name") or ""),
                "assembly_task": sticker_suffixes[order_id],
            }
        )
        matrix_labels.append(
            {
                "png_path": png_path,
                "gtin": gtin,
                "serial": serial,
            }
        )
        order_ids.append(order_id)
        log_lines.append(
            f"{sequence};{order_id};{item.get('seller_article','')};{gtin};{serial};{png_path.name}"
        )

    title_pdf = work_dir / "titles_30x20.pdf"
    matrix_pdf = work_dir / "datamatrix_30x20.pdf"
    final_pdf = work_dir / f"FBE_KIZ_30x20_{_safe_filename_part(supply_id)}_{ts}.pdf"
    label_width_mm = float(config.get("marking_label_width_mm", 30) or 30)
    label_height_mm = float(config.get("marking_label_height_mm", 20) or 20)
    if abs(label_width_mm - 30.0) > 0.01 or abs(label_height_mm - 20.0) > 0.01:
        raise ValueError(f"Для версии {APP_VERSION} размер маркировочных этикеток должен быть ровно 30×20 мм")
    create_native_title_labels_pdf(
        title_labels, title_pdf, width_mm=label_width_mm, height_mm=label_height_mm
    )
    create_datamatrix_labels_pdf(
        matrix_labels, matrix_pdf, width_mm=label_width_mm, height_mm=label_height_mm
    )
    title_refs = [(title_pdf, index) for index in range(len(rows))]
    matrix_refs = [(matrix_pdf, index) for index in range(len(rows))]
    merge_title_labels_with_official_kiz_pages(title_refs, matrix_refs, final_pdf)
    mark_supply_codes_printed(marking_db_path, supply_id, order_ids)

    log_lines.extend(["", f"FINAL={final_pdf}", f"PAIRS={len(rows)}"])
    (work_dir / "print_order_log.txt").write_text("\n".join(log_lines), encoding="utf-8")
    return final_pdf, len(rows)


def _build_native_marking_direct_images(supply_id: str) -> tuple[list[Path], list[int]]:
    """Render exact 30x20 title/DataMatrix pages as PNGs for silent Windows printing."""
    marking = load_marking_workspace(supply_id, fetch_wb=False)
    blockers = list(marking.get("direct_print_blockers") or [])
    if blockers:
        raise ValueError("; ".join(str(value) for value in blockers))

    rows = sorted(
        list(marking.get("orders") or []),
        key=lambda item: (str(item.get("name") or "").lower(), int(item.get("order_id") or 0)),
    )
    stored = {
        int(item.get("order_id") or 0): item
        for item in list_supply_code_assignments(marking_db_path, supply_id)
        if int(item.get("order_id") or 0)
    }
    if not rows:
        raise ValueError("В поставке нет маркируемых позиций")
    sticker_suffixes = _marking_sticker_suffixes(supply_id, rows)

    width_mm = float(config.get("marking_label_width_mm", 30) or 30)
    height_mm = float(config.get("marking_label_height_mm", 20) or 20)
    if abs(width_mm - 30.0) > 0.01 or abs(height_mm - 20.0) > 0.01:
        raise ValueError("Для моментальной печати размер этикеток должен быть ровно 30×20 мм")
    dpi = int(config.get("marking_datamatrix_dpi", 300) or 300)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    work_dir = Path(config.get("marking_print_output_dir", "data/marking_print")) / supply_id / ts / "direct"
    matrix_dir = work_dir / "matrix_source"
    page_dir = work_dir / "pages"
    matrix_dir.mkdir(parents=True, exist_ok=True)
    page_dir.mkdir(parents=True, exist_ok=True)

    pages: list[Path] = []
    order_ids: list[int] = []
    for sequence, item in enumerate(rows, start=1):
        order_id = int(item.get("order_id") or 0)
        assignment = stored.get(order_id) or item.get("assignment") or {}
        raw_code = str(assignment.get("raw_code") or "")
        if not raw_code:
            raise ValueError(f"Задание {order_id}: КИЗ еще не закреплен")
        gtin = str(assignment.get("gtin") or item.get("gtin14") or item.get("gtin") or "")
        serial = str(assignment.get("serial") or "")
        matrix_source = matrix_dir / f"{sequence:04d}_{order_id}.png"
        datamatrix_renderer.render_png(raw_code, matrix_source, dpi=dpi)

        title_page = page_dir / f"{sequence:04d}_{order_id}_title.png"
        matrix_page = page_dir / f"{sequence:04d}_{order_id}_kiz.png"
        render_title_png(
            seller_article=str(item.get("seller_article") or assignment.get("seller_article") or ""),
            display_name=str(item.get("name") or ""),
            assembly_task=sticker_suffixes[order_id],
            output_path=title_page,
            width_mm=width_mm,
            height_mm=height_mm,
            dpi=dpi,
        )
        render_datamatrix_label_png(
            datamatrix_png=matrix_source,
            gtin=gtin,
            serial=serial,
            output_path=matrix_page,
            width_mm=width_mm,
            height_mm=height_mm,
            dpi=dpi,
        )
        pages.extend([title_page, matrix_page])
        order_ids.append(order_id)
    return pages, order_ids


@app.post("/api/supplies/{supply_id}/marking/print-direct")
def marking_print_native_direct(supply_id: str):
    try:
        pages, order_ids = _build_native_marking_direct_images(supply_id)
        printer_name = str(config.get("small_label_printer_name") or "").strip()
        mode = print_images_windows(
            printer_name=printer_name,
            image_paths=pages,
            width_mm=30.0,
            height_mm=20.0,
            job_name=f"FBE KIZ {supply_id}",
            dry_run=bool(config.get("mock_mode") or config.get("marking_direct_print_dry_run", False)),
            safe_margin_mm=float(config.get("marking_direct_print_safe_margin_mm", 0.4) or 0.4),
        )
        if mode != "dry-run":
            mark_supply_codes_printed(marking_db_path, supply_id, order_ids)
        message = (
            f"На {printer_name} отправлено {len(order_ids)} комплектов: название + КИЗ."
            if mode != "dry-run"
            else f"Тестовый режим: подготовлено {len(order_ids)} комплектов, печать не отправлена."
        )
        return {"ok": True, "message": message, "pairs": len(order_ids), "pages": len(pages), "mode": mode}
    except (ValueError, DataMatrixRenderError, MarkingDbError, RuntimeError, FileNotFoundError) as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


@app.post("/supplies/{supply_id}/marking/print-30x20")
def marking_print_native_30x20(supply_id: str):
    try:
        final_pdf, pair_count = _build_native_marking_print_pdf(supply_id)
    except (ValueError, DataMatrixRenderError, MarkingDbError) as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)

    download_name = final_pdf.name
    return FileResponse(
        str(final_pdf),
        media_type="application/pdf",
        filename=download_name,
        headers={
            "Content-Disposition": f'inline; filename="{download_name}"',
            "X-FBE-Label-Size": "30x20mm",
            "X-FBE-Pairs": str(pair_count),
        },
    )


@app.post("/marking/print-30x20")
def marking_print_all_native_30x20():
    try:
        from pypdf import PdfReader, PdfWriter
        supplies = prepare_supplies(get_supplies_cached())
        active = [item for item in supplies if item.get("fbe_status") == "На сборке"]
        writer = PdfWriter()
        total_pairs = 0
        included = 0
        skipped: list[str] = []
        for supply in active:
            supply_id = str(supply.get("id") or "").strip()
            if not supply_id:
                continue
            try:
                workspace = load_marking_workspace(supply_id, fetch_wb=False)
                if not workspace.get("can_print_native"):
                    skipped.append(f"{supply_id}: не готово к печати")
                    continue
                pdf_path, pair_count = _build_native_marking_print_pdf(supply_id)
                reader = PdfReader(str(pdf_path))
                for page in reader.pages:
                    writer.add_page(page)
                total_pairs += int(pair_count or 0)
                included += 1
            except Exception as exc:
                skipped.append(f"{supply_id}: {exc}")
        if included <= 0:
            raise MarkingDbError("Нет поставок, готовых к формированию PDF")
        output_dir = Path(config.get("marking_print_output_dir", "data/marking_print")) / "global"
        output_dir.mkdir(parents=True, exist_ok=True)
        stamp = moscow_now().strftime("%Y%m%d_%H%M%S")
        output = output_dir / f"FBE_KIZ_ALL_30x20_{stamp}.pdf"
        with output.open("wb") as fh:
            writer.write(fh)
        return FileResponse(
            str(output),
            media_type="application/pdf",
            filename=output.name,
            headers={
                "Content-Disposition": f'inline; filename="{output.name}"',
                "X-FBE-Label-Size": "30x20mm",
                "X-FBE-Pairs": str(total_pairs),
                "X-FBE-Supplies": str(included),
                "X-FBE-Skipped": str(len(skipped)),
            },
        )
    except (ValueError, DataMatrixRenderError, MarkingDbError) as exc:
        return _marking_error_payload(exc, 400)
    except Exception as exc:
        return _marking_error_payload(exc, 500)


@app.get("/supplies/{supply_id}/marking-needs.xlsx")
def supply_marking_needs_xlsx(supply_id: str):
    marking = load_marking_workspace(supply_id, fetch_wb=False)
    content = build_marking_report_xlsx_bytes(marking, supply_id)
    safe_id = re.sub(r'[\\/:*?"<>|]+', "-", str(supply_id)).strip() or "Поставка"
    filename = f"{safe_id} {datetime.now().strftime('%d.%m')}.xlsx"
    ascii_fallback = re.sub(r"[^A-Za-z0-9_.-]+", "-", filename).strip("-") or "FBE-report.xlsx"
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{ascii_fallback}"; '
                f"filename*=UTF-8''{quote(filename)}"
            )
        },
    )


@app.get("/supplies/{supply_id}/needs", response_class=HTMLResponse)
def supply_needs_inline(request: Request, supply_id: str):
    try:
        orders = get_supply_orders(supply_id)
        needs = build_supply_requirements_from_orders(orders)
    except Exception as exc:
        needs = {
            "total_orders": 0,
            "recognized_orders": 0,
            "groups": [],
            "samples": [],
            "missing": [],
            "warnings": [f"Не удалось посчитать учет поставки: {exc}"],
        }
    return templates.TemplateResponse(
        request,
        "needs_panel.html",
        {
            "request": request,
            "title": "Учет: чего и сколько потребуется",
            "subtitle": f"Поставка {supply_id}",
            "needs": needs,
            "articles": [seller_article_from_order(order) for order in orders],
        },
    )


@app.post("/needs/articles", response_class=HTMLResponse)
def needs_from_articles(
    request: Request,
    articles: list[str] = Form(default=[]),
    subtitle: str = Form(default="Выбранные позиции"),
):
    try:
        safe_articles = [str(article or "").strip() for article in articles if str(article or "").strip()]
        needs = build_supply_requirements_from_articles(safe_articles)
    except Exception as exc:
        needs = {
            "total_orders": len(articles),
            "recognized_orders": 0,
            "groups": [],
            "samples": [],
            "missing": [{"order_id": str(i + 1), "seller_article": str(article or "—")} for i, article in enumerate(articles)],
            "warnings": [f"Не удалось посчитать учет локально: {exc}"],
        }
    return templates.TemplateResponse(
        request,
        "needs_panel.html",
        {
            "request": request,
            "title": "Учет: чего и сколько потребуется",
            "subtitle": subtitle or "Выбранные позиции",
            "needs": needs,
            "articles": safe_articles,
        },
    )


def build_requirements_xlsx_bytes(needs: dict[str, Any], subtitle: str = "Потребность") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "Потребность"

    title_fill = PatternFill("solid", fgColor="146C43")
    header_fill = PatternFill("solid", fgColor="E7F6EC")
    sample_fill = PatternFill("solid", fgColor="FFF4CC")
    thin = Side(style="thin", color="D6DEE6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:C1")
    ws["A1"] = "FBE — учет: чего и сколько потребуется"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Источник"
    ws["B2"] = subtitle or "Потребность"
    ws["A3"] = "Заказов"
    ws["B3"] = int(needs.get("total_orders") or 0)
    ws["A4"] = "Распознано"
    ws["B4"] = int(needs.get("recognized_orders") or 0)

    row = 6
    ws.cell(row=row, column=1, value="Блок")
    ws.cell(row=row, column=2, value="Позиция")
    ws.cell(row=row, column=3, value="Количество")
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    row += 1

    def write_item(block: str, label: str, count: int, fill=None):
        nonlocal row
        ws.cell(row=row, column=1, value=block)
        ws.cell(row=row, column=2, value=label)
        ws.cell(row=row, column=3, value=int(count or 0))
        for c in range(1,4):
            cell = ws.cell(row=row, column=c)
            cell.border = border
            if fill:
                cell.fill = fill
            if c == 3:
                cell.alignment = Alignment(horizontal="center")
        row += 1

    for group in needs.get("groups") or []:
        for item in group.get("items") or []:
            write_item(str(group.get("name") or ""), str(item.get("label") or ""), int(item.get("count") or 0))

    for item in needs.get("samples") or []:
        write_item("Пробники", str(item.get("label") or ""), int(item.get("count") or 0), sample_fill)

    missing = needs.get("missing") or []
    if missing:
        row += 1
        ws.cell(row=row, column=1, value="Не распознано")
        ws.cell(row=row, column=1).font = Font(bold=True, color="B91C1C")
        row += 1
        ws.cell(row=row, column=1, value="№")
        ws.cell(row=row, column=2, value="Артикул продавца")
        for cell in ws[row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FEE2E2")
            cell.border = border
        row += 1
        for idx, item in enumerate(missing, start=1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=str(item.get("seller_article") or "—"))
            for c in range(1,3):
                ws.cell(row=row, column=c).border = border
            row += 1

    warnings = needs.get("warnings") or []
    if warnings:
        row += 1
        ws.cell(row=row, column=1, value="Предупреждения")
        ws.cell(row=row, column=1).font = Font(bold=True)
        row += 1
        for w in warnings:
            ws.cell(row=row, column=1, value=str(w))
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            row += 1

    widths = {1: 18, 2: 56, 3: 14}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for r in range(1, ws.max_row + 1):
        ws.row_dimensions[r].height = 22
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A6:C{max(6, row-1)}"

    bio = BytesIO()
    wb_out.save(bio)
    return bio.getvalue()


@app.post("/needs/articles.xlsx")
def needs_articles_xlsx(
    articles: list[str] = Form(default=[]),
    subtitle: str = Form(default="Потребность"),
):
    safe_articles = [str(article or "").strip() for article in articles if str(article or "").strip()]
    needs = build_supply_requirements_from_articles(safe_articles)
    content = build_requirements_xlsx_bytes(needs, subtitle=subtitle or "Потребность")
    filename = f"fbe-needs-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/orders/needs-preview", response_class=HTMLResponse)
def selected_orders_needs_preview(request: Request, order_ids: list[int] = Form(default=[])):
    try:
        safe_ids = []
        for x in order_ids:
            try:
                safe_ids.append(int(x))
            except Exception:
                continue
        needs = requirements_context_for_order_ids(safe_ids, get_new_orders_cached())
    except Exception as exc:
        needs = {
            "total_orders": len(order_ids),
            "recognized_orders": 0,
            "groups": [],
            "samples": [],
            "missing": [],
            "warnings": [f"Не удалось посчитать учет выбранных заказов: {exc}"],
        }
    return templates.TemplateResponse(
        request,
        "needs_panel.html",
        {
            "request": request,
            "title": "Учет: чего и сколько потребуется",
            "subtitle": "Выбранные новые заказы",
            "needs": needs,
            "articles": [seller_article_from_order(order) for order in get_new_orders_cached() if int(order.get("id") or order.get("orderId") or 0) in safe_ids],
        },
    )


def _supply_shipping_context(supply_id: str, order_ids: list[int] | None = None, *, remote: bool = False) -> dict[str, Any]:
    supplied_ids = [int(x) for x in (order_ids or []) if int(x) > 0]
    if not supplied_ids:
        try:
            supplied_ids = [int(x) for x in get_supply_order_ids_cached(supply_id) if int(x) > 0]
        except Exception:
            supplied_ids = []

    local_rows = [
        r for r in list_fbs_registry_rows(marking_db_path, limit=10000)
        if str(r.get("supply_id") or "") == str(supply_id)
    ]
    if not supplied_ids:
        supplied_ids = [int(r.get("order_id") or 0) for r in local_rows if int(r.get("order_id") or 0) > 0]
    order_ids = sorted(set(supplied_ids))

    statuses: list[dict[str, Any]] = []
    status_error = ""
    if remote and order_ids:
        try:
            statuses = wb.get_order_statuses(order_ids)
            update_fbs_order_statuses(marking_db_path, statuses)
            update_post_sale_wb_statuses(marking_db_path, statuses)
        except Exception as exc:
            status_error = str(exc)

    if not statuses and local_rows:
        statuses = [
            {
                "id": int(r.get("order_id") or 0),
                "supplierStatus": str(r.get("supplier_status") or ""),
                "wbStatus": str(r.get("wb_status") or ""),
            }
            for r in local_rows if int(r.get("order_id") or 0) > 0
        ]

    supplier_states = [str(x.get("supplierStatus") or "").strip().lower() for x in statuses]
    wb_states = [str(x.get("wbStatus") or "").strip().lower() for x in statuses]
    final_wb = WB_TERMINAL_STATUSES
    active_pairs = [(a,b) for a,b in zip(supplier_states,wb_states) if b not in final_wb]
    if any(a == "confirm" for a,b in active_pairs):
        stage = "assembly"
        label = "На сборке"
        css = "status-assembling"
    elif any(a == "complete" for a,b in active_pairs):
        stage = "wb"
        label = "Передано WB"
        css = "status-transferred"
    elif statuses and all(b in final_wb for b in wb_states):
        stage = "delivered"
        label = "Доставлено"
        css = "status-delivered"
    else:
        stage = "unknown"
        label = "Статус WB"
        css = "status-unknown"

    trbx_ids: list[str] = []
    trbx_error = ""
    try:
        if remote:
            trbx_ids = wb.get_supply_shipping_units(supply_id)
            with _cache_lock:
                _cache[f"shipping_units:{supply_id}"] = (time.time(), list(trbx_ids))
        else:
            trbx_ids = cached_stale_while_revalidate(
                f"shipping_units:{supply_id}", 60,
                lambda: wb.get_supply_shipping_units(supply_id), lambda: [],
            )
    except Exception as exc:
        trbx_error = str(exc)

    return {
        "stage": stage,
        "label": label,
        "class": css,
        "order_ids": order_ids,
        "order_count": len(order_ids),
        "statuses": statuses,
        "status_error": status_error,
        "trbx_ids": trbx_ids,
        "trbx_count": len(trbx_ids),
        "trbx_max": max(1, len(order_ids) + 1),
        "trbx_error": trbx_error,
        "can_deliver": stage == "assembly",
        "supply_qr_ready": stage in {"wb", "delivered"},
    }




class SupplyTransferPreflightPayload(BaseModel):
    supply_ids: list[str] = Field(default_factory=list)
    confirmed_assembled: bool = False


class SupplyTransferCargoItem(BaseModel):
    supply_id: str
    amount: int = Field(ge=1, le=1000)


class SupplyTransferCargoPayload(BaseModel):
    items: list[SupplyTransferCargoItem] = Field(default_factory=list)


class SupplyTransferDeliverPayload(BaseModel):
    supply_ids: list[str] = Field(default_factory=list)
    confirmed_labels_applied: bool = False


class SupplyTransferQrPayload(BaseModel):
    supply_ids: list[str] = Field(default_factory=list)


def _unique_supply_ids(values: list[str]) -> list[str]:
    result: list[str] = []
    for value in values or []:
        sid = str(value or "").strip()
        if sid and sid not in result:
            result.append(sid)
    return result


def _metadata_values(item: dict[str, Any] | None, key: str) -> list[str]:
    """Read one FBS metadata key across legacy meta and current metaDetails shapes."""
    if not isinstance(item, dict):
        return []
    wanted = str(key or "").strip().lower()
    raw_values: list[Any] = []

    def harvest(value: Any) -> None:
        if value is None:
            return
        if isinstance(value, list):
            for child in value:
                harvest(child)
            return
        if isinstance(value, dict):
            if "value" in value:
                harvest(value.get("value"))
            else:
                for child in value.values():
                    harvest(child)
            return
        text = str(value).strip()
        if text:
            raw_values.append(text)

    meta = item.get("meta")
    if isinstance(meta, dict) and wanted in meta:
        harvest(meta.get(wanted))
    if wanted in item:
        harvest(item.get(wanted))
    details = item.get("metaDetails")
    if isinstance(details, list):
        for detail in details:
            if isinstance(detail, dict) and str(detail.get("key") or "").strip().lower() == wanted:
                harvest(detail.get("value"))

    result: list[str] = []
    for value in raw_values:
        text = str(value).strip()
        if text and text not in result:
            result.append(text)
    return result


def _metadata_negative_decisions(item: dict[str, Any] | None) -> list[str]:
    if not isinstance(item, dict):
        return []
    errors: list[str] = []
    if item.get("isError"):
        raw_errors = item.get("errors")
        if isinstance(raw_errors, list):
            errors.extend(str(value) for value in raw_errors if str(value).strip())
        elif raw_errors:
            errors.append(str(raw_errors))
        else:
            errors.append("WB пометил метаданные задания как ошибочные")
    details = item.get("metaDetails")
    bad_tokens = ("invalid", "missing", "required", "error", "rejected", "expired", "notintroduced", "notfound")
    if isinstance(details, list):
        for detail in details:
            if not isinstance(detail, dict):
                continue
            decision = str(detail.get("decision") or "").strip()
            compact = re.sub(r"[^a-z0-9]+", "", decision.lower())
            if decision and any(token in compact for token in bad_tokens):
                key = str(detail.get("key") or "metadata")
                errors.append(f"{key}: {decision}")
    return errors


def _required_meta_keys(order: dict[str, Any] | None) -> set[str]:
    if not isinstance(order, dict):
        return set()
    raw = order.get("requiredMeta") or order.get("required_meta") or []
    if isinstance(raw, str):
        raw = [raw]
    if not isinstance(raw, list):
        return set()
    return {str(value).strip().lower() for value in raw if str(value).strip()}


def _supply_transfer_preflight_one(supply_id: str) -> dict[str, Any]:
    sid = str(supply_id).strip()
    blockers: list[str] = []
    warnings: list[str] = []

    try:
        order_ids = [int(value) for value in wb.get_supply_order_ids(sid) if int(value) > 0]
        try:
            set_fbs_supply_order_ids(marking_db_path, supply_id=sid, order_ids=order_ids)
        except Exception:
            pass
    except Exception as exc:
        return {
            "supply_id": sid, "name": sid, "ready": False, "order_count": 0,
            "blockers": [f"Не удалось получить состав поставки у WB: {exc}"], "warnings": [],
            "existing_cargo_count": 0, "max_cargo": 0, "order_ids": [],
        }

    if not order_ids:
        blockers.append("В поставке нет сборочных заданий")

    shipping = _supply_shipping_context(sid, order_ids, remote=True)
    if shipping.get("status_error"):
        blockers.append(f"Не удалось проверить статусы заданий у WB: {shipping['status_error']}")
    if shipping.get("trbx_error"):
        warnings.append(f"Не удалось проверить существующие грузоместа: {shipping['trbx_error']}")
    if shipping.get("stage") != "assembly":
        blockers.append(f"Поставка не на сборке: {shipping.get('label') or shipping.get('stage')}")

    supply_details: dict[str, Any] = {}
    try:
        supply_details = wb.get_supply_details(sid) or {}
    except Exception as exc:
        warnings.append(f"Карточка поставки WB недоступна: {exc}")
    name = str(supply_details.get("name") or sid)

    # Fresh metadata is the safety source before PATCH /deliver. WB explicitly
    # recommends checking it because deliver now validates assembly-order metadata.
    meta_rows: list[dict[str, Any]] = []
    if order_ids:
        for batch in chunks(order_ids, 100):
            try:
                meta_rows.extend(wb.get_orders_meta(batch))
            except Exception as exc:
                blockers.append(f"Не удалось проверить сроки/КИЗы у WB: {exc}")
                meta_rows = []
                break
    meta_by_id = {_meta_order_id(row): row for row in meta_rows if _meta_order_id(row)}

    order_details: dict[int, dict[str, Any]] = {}
    try:
        current_orders = get_supply_orders(sid, known_order_ids=order_ids)
        order_details = {int(row.get("id") or 0): row for row in current_orders if int(row.get("id") or 0) > 0}
    except Exception as exc:
        warnings.append(f"Не удалось получить расширенные данные заданий: {exc}")

    catalog = get_catalog()
    missing_expiration: list[int] = []
    missing_sgtin: list[int] = []
    invalid_meta: list[str] = []
    for oid in order_ids:
        order = order_details.get(oid, {})
        required = _required_meta_keys(order)
        meta = meta_by_id.get(oid)

        for issue in _metadata_negative_decisions(meta):
            invalid_meta.append(f"{oid}: {issue}")

        if "expiration" in required and not _metadata_values(meta, "expiration"):
            missing_expiration.append(oid)

        article = seller_article_from_order(order)
        product = catalog.find_by_article(article) if article else None
        catalog_requires_marking = bool(product and is_marking_required(product.get("kiz_required")))
        wb_requires_sgtin = "sgtin" in required
        if catalog_requires_marking or wb_requires_sgtin:
            sgtin_info = _sgtin_meta_info(meta)
            if not list(sgtin_info.get("values") or []):
                missing_sgtin.append(oid)

    if invalid_meta:
        preview = "; ".join(invalid_meta[:6])
        if len(invalid_meta) > 6:
            preview += f"; еще {len(invalid_meta) - 6}"
        blockers.append("WB сообщает ошибки метаданных: " + preview)
    if missing_expiration:
        preview = ", ".join(str(value) for value in missing_expiration[:8])
        suffix = f" и еще {len(missing_expiration) - 8}" if len(missing_expiration) > 8 else ""
        blockers.append(f"Не указан обязательный срок годности у заданий: {preview}{suffix}")
    if missing_sgtin:
        preview = ", ".join(str(value) for value in missing_sgtin[:8])
        suffix = f" и еще {len(missing_sgtin) - 8}" if len(missing_sgtin) > 8 else ""
        blockers.append(f"WB не видит КИЗ у маркируемых заданий: {preview}{suffix}")

    return {
        "supply_id": sid,
        "name": name,
        "ready": not blockers,
        "order_count": len(order_ids),
        "order_ids": order_ids,
        "blockers": blockers,
        "warnings": warnings,
        "existing_cargo_count": int(shipping.get("trbx_count") or 0),
        "max_cargo": max(1, len(order_ids) + 1),
        "checks": {
            "statuses": "ok" if not shipping.get("status_error") and shipping.get("stage") == "assembly" else "error",
            "metadata_rows": len(meta_rows),
            "missing_expiration": len(missing_expiration),
            "missing_sgtin": len(missing_sgtin),
        },
    }


def _persist_supply_delivered_local(supply_id: str, name: str, order_ids: list[int]) -> None:
    """Commit a successful WB deliver transition to the local operational registry.

    The WB mutation already succeeded at this point, so local persistence is a
    read-model update rather than a second source of truth. Keeping it in one
    helper prevents the legacy single-supply route and the bulk wizard from
    drifting apart again.
    """
    sid = str(supply_id or "").strip()
    ids = sorted({int(x) for x in order_ids if int(x) > 0})
    if not sid:
        return
    upsert_fbs_supplies(marking_db_path, [{
        "id": sid,
        "name": str(name or sid),
        "done": True,
        "fbe_status": "Передано WB",
    }])
    if ids:
        update_fbs_order_statuses(marking_db_path, [
            {"id": oid, "supplierStatus": "complete"} for oid in ids
        ])
        set_fbs_supply_order_ids(marking_db_path, supply_id=sid, order_ids=ids)


def _shipping_archive_dir(supply_id: str) -> Path:
    return Path(config.get("shipping_unit_qr_output_dir", "data/shipping_qr")) / str(supply_id)


def _cargo_manifest_matches(supply_id: str, trbx_ids: list[str]) -> bool:
    """Return True only when the saved archive represents the exact current WB cargo-place set."""
    path = _shipping_archive_dir(supply_id) / "manifest.json"
    if not path.exists():
        return False
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        saved = [str(row.get("trbx_id") or "") for row in (payload.get("units") or [])]
    except Exception:
        return False
    return saved == [str(value) for value in trbx_ids]


def _archive_cargo_qr(supply_id: str, trbx_ids: list[str], *, order_count: int) -> dict[str, Any]:
    """Download cargo labels deterministically, one exact trbxId per request.

    A newly-created shipping unit can briefly exist before its sticker is readable.
    Retrying the read-only sticker request is safe; creating the unit itself is never
    repeated blindly.
    """
    sid = str(supply_id)
    clean_ids = []
    for value in trbx_ids or []:
        text = str(value).strip()
        if text and text not in clean_ids:
            clean_ids.append(text)
    if not clean_ids:
        raise WBApiError("WB не вернул ID грузомест")

    out_dir = _shipping_archive_dir(sid)
    out_dir.mkdir(parents=True, exist_ok=True)
    units: list[dict[str, Any]] = []
    paths: list[Path] = []

    for index, trbx_id in enumerate(clean_ids, start=1):
        sticker: dict[str, Any] | None = None
        last_error: Exception | None = None
        for attempt in range(6):
            try:
                stickers = wb.get_supply_shipping_unit_stickers(
                    sid, [trbx_id], sticker_type="png", width=58, height=40
                )
                if len(stickers) == 1 and str(stickers[0].get("file") or ""):
                    sticker = stickers[0]
                    break
                last_error = WBApiError(
                    f"WB пока не вернул QR для грузоместа {trbx_id}"
                )
            except Exception as exc:
                last_error = exc
            if attempt < 5:
                time.sleep(0.45 + attempt * 0.25)
        if not sticker:
            raise WBApiError(
                f"Грузоместо {trbx_id} создано, но QR пока не получен: {last_error or 'пустой ответ WB'}"
            )

        raw_file = str(sticker.get("file") or "")
        payload = wb.decode_sticker_file(raw_file, "png")
        safe_id = _safe_filename_part(trbx_id, f"cargo_{index}")
        path = out_dir / f"{index:02d}_{safe_id}.png"
        path.write_bytes(payload)
        paths.append(path)
        units.append({
            "index": index,
            "trbx_id": trbx_id,
            "barcode": str(sticker.get("barcode") or ""),
            "file": path.name,
        })
        if index < len(clean_ids):
            time.sleep(0.21)

    body_pdf = out_dir / "_cargo_qr_body.pdf"
    final_pdf = out_dir / "cargo_qr_saved.pdf"
    create_stickers_pdf(paths, body_pdf, width_mm=58, height_mm=40)
    _prepend_supply_print_label(
        supply_id=sid,
        order_count=order_count,
        body_pdf=body_pdf,
        final_pdf=final_pdf,
        title="ГРУЗОМЕСТА",
        footer=f"ГРУЗОМЕСТ: {len(paths)} · QR сохранены локально до передачи WB",
    )
    identity = _supply_print_identity(sid, order_count)
    manifest = {
        "supply_id": sid,
        "supply_name": str(identity.get("supply_name") or sid),
        "saved_at": moscow_now().isoformat(),
        "order_count": int(order_count),
        "cargo_count": len(paths),
        "units": units,
        "pdf": final_pdf.name,
    }
    (out_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def _merge_trbx_ids(*groups: list[str]) -> list[str]:
    result: list[str] = []
    for group in groups:
        for value in group or []:
            text = str(value).strip()
            if text and text not in result:
                result.append(text)
    return result


def _reconcile_supply_trbx_ids(supply_id: str, known_ids: list[str], expected: int) -> list[str]:
    """Reconcile POST-returned IDs with GET /trbx without requiring read-after-write consistency."""
    ids = _merge_trbx_ids(known_ids)
    if len(ids) >= expected:
        return ids
    for attempt in range(6):
        try:
            current = wb.get_supply_shipping_units(supply_id)
            ids = _merge_trbx_ids(ids, current)
        except Exception:
            current = []
        if len(ids) >= expected:
            break
        if attempt < 5:
            time.sleep(0.35 + attempt * 0.25)
    return ids

def _archive_supply_qr(supply_id: str, *, order_count: int) -> Path:
    sid = str(supply_id)
    last_error: Exception | None = None
    data: dict[str, Any] | None = None
    for attempt in range(3):
        try:
            data = wb.get_supply_barcode(sid, barcode_type="png", width=58, height=40)
            if str((data or {}).get("file") or ""):
                break
        except Exception as exc:
            last_error = exc
        if attempt < 2:
            time.sleep(0.5)
    raw_file = str((data or {}).get("file") or "")
    if not raw_file:
        if last_error:
            raise WBApiError(f"Поставка передана, но QR поставки пока не получен: {last_error}")
        raise WBApiError("Поставка передана, но WB пока не вернул QR поставки")

    payload = wb.decode_sticker_file(raw_file, "png")
    out_dir = Path(config.get("supply_qr_output_dir", "data/supply_qr")) / sid
    out_dir.mkdir(parents=True, exist_ok=True)
    img_path = out_dir / f"supply_qr_{_safe_filename_part(sid)}.png"
    img_path.write_bytes(payload)
    body_pdf = out_dir / "_supply_qr_body.pdf"
    final_pdf = out_dir / "supply_qr_saved.pdf"
    create_stickers_pdf([img_path], body_pdf, width_mm=58, height_mm=40)
    _prepend_supply_print_label(
        supply_id=sid,
        order_count=order_count,
        body_pdf=body_pdf,
        final_pdf=final_pdf,
        title="QR ПОСТАВКИ",
        footer="Сохранен FBE сразу после передачи в доставку WB",
    )
    return final_pdf


@app.post("/api/supplies/transfer/preflight")
def transfer_supplies_preflight(payload: SupplyTransferPreflightPayload):
    supply_ids = _unique_supply_ids(payload.supply_ids)
    if not supply_ids:
        return JSONResponse({"ok": False, "error": "Не выбраны поставки"}, status_code=400)
    if not payload.confirmed_assembled:
        return JSONResponse(
            {"ok": False, "error": "Подтвердите, что выбранные поставки физически собраны"},
            status_code=409,
        )
    rows = [_supply_transfer_preflight_one(sid) for sid in supply_ids]
    return {
        "ok": True,
        "all_ready": all(bool(row.get("ready")) for row in rows),
        "supplies": rows,
    }


@app.post("/api/supplies/transfer/cargo")
def transfer_supplies_create_cargo(payload: SupplyTransferCargoPayload):
    if not payload.items:
        return JSONResponse({"ok": False, "error": "Не передано количество грузомест"}, status_code=400)
    results: list[dict[str, Any]] = []
    for item in payload.items:
        sid = str(item.supply_id).strip()
        desired = int(item.amount)
        check = _supply_transfer_preflight_one(sid)
        max_cargo = int(check.get("max_cargo") or 1)
        if desired > max_cargo:
            results.append({
                "supply_id": sid, "name": check.get("name") or sid, "ok": False,
                "error": f"Для этой поставки WB допускает максимум {max_cargo} грузомест",
            })
            continue
        if not check.get("ready"):
            results.append({
                "supply_id": sid, "name": check.get("name") or sid, "ok": False,
                "error": "; ".join(str(value) for value in check.get("blockers") or []) or "Проверка готовности не пройдена",
            })
            continue
        try:
            # Read what already exists first. For a unit created by an earlier FBE
            # attempt this is how we recover it instead of creating a duplicate.
            existing = wb.get_supply_shipping_units(sid)
            if len(existing) > desired:
                raise ValueError(
                    f"У поставки уже создано {len(existing)} грузомест; нельзя уменьшить до {desired} без удаления в WB"
                )

            created_ids: list[str] = []
            if len(existing) < desired:
                missing = desired - len(existing)
                # IMPORTANT: POST /trbx is executed exactly once. Current WB API
                # returns the created trbxIds; use those IDs immediately for QR.
                created_ids = wb.add_supply_shipping_units(sid, missing)

            trbx_ids = _merge_trbx_ids(existing, created_ids)
            trbx_ids = _reconcile_supply_trbx_ids(sid, trbx_ids, desired)
            if len(trbx_ids) != desired:
                raise WBApiError(
                    f"Грузоместа созданы в WB, но FBE получил ID только {len(trbx_ids)} из {desired}. "
                    "Повторно грузоместа не создавались. Нажмите кнопку еще раз: FBE сначала перечитает уже существующие грузоместа."
                )

            # QR is requested from the exact IDs received from WB, not from array
            # positions and not from a mandatory second GET.
            manifest = _archive_cargo_qr(sid, trbx_ids, order_count=int(check.get("order_count") or 0))
            with _cache_lock:
                _cache[f"shipping_units:{sid}"] = (time.time(), list(trbx_ids))
            results.append({
                "supply_id": sid, "name": check.get("name") or sid, "ok": True,
                "cargo_count": len(trbx_ids),
                "trbx_ids": trbx_ids,
                "created_count": len(created_ids),
                "pdf_url": f"/supplies/{quote(sid)}/shipping-units-qr-saved.pdf",
                "manifest": manifest,
            })
        except Exception as exc:
            results.append({"supply_id": sid, "name": check.get("name") or sid, "ok": False, "error": str(exc)})
    ready_ids = [str(row.get("supply_id") or "") for row in results if row.get("ok")]
    batch_url = ""
    if ready_ids:
        batch_url = "/supplies/transfer/cargo-batch.pdf?" + urlencode({"supply_ids": ",".join(ready_ids)})
    return {
        "ok": any(bool(row.get("ok")) for row in results),
        "all_ok": all(bool(row.get("ok")) for row in results),
        "results": results,
        "batch_pdf_url": batch_url,
    }


@app.post("/api/supplies/transfer/deliver")
def transfer_supplies_deliver(payload: SupplyTransferDeliverPayload):
    supply_ids = _unique_supply_ids(payload.supply_ids)
    if not supply_ids:
        return JSONResponse({"ok": False, "error": "Не выбраны поставки"}, status_code=400)
    if not payload.confirmed_labels_applied:
        return JSONResponse(
            {"ok": False, "error": "Подтвердите, что QR грузомест распечатаны и наклеены"},
            status_code=409,
        )

    results: list[dict[str, Any]] = []
    for sid in supply_ids:
        check = _supply_transfer_preflight_one(sid)
        if not check.get("ready"):
            results.append({
                "supply_id": sid, "name": check.get("name") or sid, "ok": False,
                "error": "; ".join(str(value) for value in check.get("blockers") or []) or "Проверка готовности не пройдена",
            })
            continue
        try:
            # The archived manifest is the safety record of the exact trbx IDs for
            # which QR labels were obtained. Prefer it, then reconcile with WB.
            manifest_path = _shipping_archive_dir(sid) / "manifest.json"
            manifest_ids: list[str] = []
            if manifest_path.exists():
                try:
                    manifest_data = json.loads(manifest_path.read_text(encoding="utf-8"))
                    manifest_ids = [
                        str(row.get("trbx_id") or "").strip()
                        for row in (manifest_data.get("units") or [])
                        if str(row.get("trbx_id") or "").strip()
                    ]
                except Exception:
                    manifest_ids = []
            remote_ids = wb.get_supply_shipping_units(sid)
            trbx_ids = _merge_trbx_ids(manifest_ids, remote_ids)
            if not trbx_ids:
                raise ValueError("У поставки нет сохраненных грузомест. Сначала получите QR грузомест")
            if not manifest_ids:
                _archive_cargo_qr(sid, trbx_ids, order_count=int(check.get("order_count") or 0))

            wb.deliver_supply(sid)
            # PATCH /deliver is authoritative: WB moves every order in the supply
            # to supplierStatus=complete. Persist that transition immediately so
            # the dashboard moves the row to «Передано WB» without waiting for a
            # later background status refresh.
            order_ids = [int(x) for x in (check.get("order_ids") or []) if int(x) > 0]
            try:
                _persist_supply_delivered_local(
                    sid, str(check.get("name") or sid), order_ids
                )
            except Exception:
                pass
            invalidate_cache("supplies")
            invalidate_cache("fbs_operational_registry")
            invalidate_cache(f"supply_order_ids:{sid}")
            results.append({
                "supply_id": sid,
                "name": check.get("name") or sid,
                "ok": True,
                "delivered": True,
            })
        except Exception as exc:
            results.append({"supply_id": sid, "name": check.get("name") or sid, "ok": False, "error": str(exc)})

    return {
        "ok": any(bool(row.get("ok")) for row in results),
        "all_ok": all(bool(row.get("ok")) for row in results),
        "results": results,
    }


@app.post("/api/supplies/transfer/supply-qr")
def transfer_supplies_supply_qr(payload: SupplyTransferQrPayload):
    supply_ids = _unique_supply_ids(payload.supply_ids)
    if not supply_ids:
        return JSONResponse({"ok": False, "error": "Не выбраны переданные поставки"}, status_code=400)

    results: list[dict[str, Any]] = []
    for sid in supply_ids:
        name = sid
        try:
            try:
                details = wb.get_supply_details(sid) or {}
                name = str(details.get("name") or sid)
            except Exception:
                pass
            try:
                order_count = len(wb.get_supply_order_ids(sid))
            except Exception:
                order_count = 0
            # QR may become readable shortly after PATCH /deliver. This helper has
            # its own safe GET retries and persists the file locally.
            _archive_supply_qr(sid, order_count=order_count)
            results.append({
                "supply_id": sid, "name": name, "ok": True,
                "qr_url": f"/supplies/{quote(sid)}/supply-qr-saved.pdf",
            })
        except Exception as exc:
            results.append({"supply_id": sid, "name": name, "ok": False, "error": str(exc)})

    qr_ids = [str(row.get("supply_id") or "") for row in results if row.get("ok")]
    batch_url = ""
    if qr_ids:
        batch_url = "/supplies/transfer/supply-qr-batch.pdf?" + urlencode({"supply_ids": ",".join(qr_ids)})
    return {
        "ok": any(bool(row.get("ok")) for row in results),
        "all_ok": all(bool(row.get("ok")) for row in results),
        "results": results,
        "batch_pdf_url": batch_url,
    }


def _transfer_background_payload(response: Any) -> dict[str, Any]:
    """Normalize a legacy transfer handler result for the shared job runner."""
    if isinstance(response, JSONResponse):
        try:
            payload = json.loads(response.body.decode("utf-8"))
        except Exception:
            payload = {"ok": False, "error": f"HTTP {response.status_code}"}
        if response.status_code >= 400 or payload.get("ok") is False:
            raise MarkingDbError(
                str(payload.get("error") or payload.get("message") or f"HTTP {response.status_code}")
            )
        return dict(payload)
    if isinstance(response, dict):
        return dict(response)
    raise MarkingDbError("FBE получил неизвестный ответ от операции передачи поставки")


def _transfer_background_message(payload: dict[str, Any]) -> str:
    results = list(payload.get("results") or [])
    successful = sum(1 for row in results if isinstance(row, dict) and row.get("ok"))
    failed = sum(1 for row in results if isinstance(row, dict) and not row.get("ok"))
    if failed and successful:
        return f"Операция завершена частично: успешно {successful}, ошибок {failed}."
    if failed:
        return f"Операция завершена с ошибками: {failed}."
    return "Операция передачи поставки завершена."


def _transfer_preflight_worker(
    payload: SupplyTransferPreflightPayload, report: Callable[..., None]
) -> dict[str, Any]:
    report(total=100, progress=5, message="Проверяю поставки непосредственно у WB…", phase="preflight")
    result = _transfer_background_payload(transfer_supplies_preflight(payload))
    report(progress=100, message="Проверка поставок завершена.", phase="complete")
    return result


def _transfer_cargo_worker(
    payload: SupplyTransferCargoPayload, report: Callable[..., None]
) -> dict[str, Any]:
    report(total=100, progress=5, message="Создаю грузоместа и сохраняю QR…", phase="cargo")
    result = _transfer_background_payload(transfer_supplies_create_cargo(payload))
    report(progress=100, message=_transfer_background_message(result), phase="complete")
    return result


def _transfer_deliver_worker(
    payload: SupplyTransferDeliverPayload, report: Callable[..., None]
) -> dict[str, Any]:
    report(total=100, progress=5, message="Передаю поставки WB…", phase="deliver")
    result = _transfer_background_payload(transfer_supplies_deliver(payload))
    report(progress=100, message=_transfer_background_message(result), phase="complete")
    return result


def _transfer_supply_qr_worker(
    payload: SupplyTransferQrPayload, report: Callable[..., None]
) -> dict[str, Any]:
    report(total=100, progress=5, message="Получаю и сохраняю QR поставок…", phase="supply-qr")
    result = _transfer_background_payload(transfer_supplies_supply_qr(payload))
    report(progress=100, message=_transfer_background_message(result), phase="complete")
    return result


@app.post("/api/supplies/transfer/preflight/start")
def transfer_supplies_preflight_job(payload: SupplyTransferPreflightPayload):
    ids = _unique_supply_ids(payload.supply_ids)
    if not ids:
        return JSONResponse({"ok": False, "error": "Не выбраны поставки"}, status_code=400)
    if not payload.confirmed_assembled:
        return JSONResponse(
            {"ok": False, "error": "Подтвердите, что выбранные поставки физически собраны"},
            status_code=409,
        )
    key = "transfer-preflight:" + ",".join(sorted(ids))
    job_id = submit_background_job(
        "Проверка поставок WB",
        lambda report: _transfer_preflight_worker(payload, report),
        operation_key=key if ids else None,
    )
    return {"ok": True, "job_id": job_id, "message": "Проверка поставок запущена"}


@app.post("/api/supplies/transfer/cargo/start")
def transfer_supplies_cargo_job(payload: SupplyTransferCargoPayload):
    if not payload.items:
        return JSONResponse({"ok": False, "error": "Не передано количество грузомест"}, status_code=400)
    items_key = ",".join(
        f"{str(item.supply_id).strip()}={int(item.amount)}"
        for item in sorted(payload.items, key=lambda row: str(row.supply_id))
    )
    job_id = submit_background_job(
        "Создание грузомест WB",
        lambda report: _transfer_cargo_worker(payload, report),
        operation_key="transfer-cargo:" + items_key if items_key else None,
    )
    return {"ok": True, "job_id": job_id, "message": "Создание грузомест запущено"}


@app.post("/api/supplies/transfer/deliver/start")
def transfer_supplies_deliver_job(payload: SupplyTransferDeliverPayload):
    ids = _unique_supply_ids(payload.supply_ids)
    if not ids:
        return JSONResponse({"ok": False, "error": "Не выбраны поставки"}, status_code=400)
    if not payload.confirmed_labels_applied:
        return JSONResponse(
            {"ok": False, "error": "Подтвердите, что QR грузомест распечатаны и наклеены"},
            status_code=409,
        )
    job_id = submit_background_job(
        "Передача поставок WB",
        lambda report: _transfer_deliver_worker(payload, report),
        operation_key="transfer-deliver:" + ",".join(sorted(ids)) if ids else None,
    )
    return {"ok": True, "job_id": job_id, "message": "Передача поставок запущена"}


@app.post("/api/supplies/transfer/supply-qr/start")
def transfer_supplies_supply_qr_job(payload: SupplyTransferQrPayload):
    ids = _unique_supply_ids(payload.supply_ids)
    if not ids:
        return JSONResponse({"ok": False, "error": "Не выбраны переданные поставки"}, status_code=400)
    job_id = submit_background_job(
        "Получение QR поставок WB",
        lambda report: _transfer_supply_qr_worker(payload, report),
        operation_key="transfer-supply-qr:" + ",".join(sorted(ids)) if ids else None,
    )
    return {"ok": True, "job_id": job_id, "message": "Получение QR поставок запущено"}


def _saved_pdf_batch(supply_ids: list[str], *, kind: str) -> Path:
    ids = _unique_supply_ids(supply_ids)
    if not ids:
        raise ValueError("Не выбраны поставки")
    page_refs: list[tuple[Path, int]] = []
    if kind == "cargo":
        root = Path(config.get("shipping_unit_qr_output_dir", "data/shipping_qr"))
        source_name = "cargo_qr_saved.pdf"
        batch_prefix = "cargo_qr_batch"
    elif kind == "supply":
        root = Path(config.get("supply_qr_output_dir", "data/supply_qr"))
        source_name = "supply_qr_saved.pdf"
        batch_prefix = "supply_qr_batch"
    else:
        raise ValueError("Неизвестный тип QR")
    for sid in ids:
        source = root / sid / source_name
        if not source.exists():
            raise ValueError(f"Нет сохраненного PDF для поставки {sid}")
        for page_index in range(_pdf_page_count(source)):
            page_refs.append((source, page_index))
    batch_dir = root / "_batches"
    batch_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output = batch_dir / f"{batch_prefix}_{stamp}.pdf"
    merge_pdf_page_refs(page_refs, output)
    return output


@app.get("/supplies/transfer/cargo-batch.pdf")
def saved_cargo_qr_batch_pdf(supply_ids: str = ""):
    try:
        ids = [value.strip() for value in str(supply_ids).split(",") if value.strip()]
        path = _saved_pdf_batch(ids, kind="cargo")
        return FileResponse(str(path), media_type="application/pdf", filename="FBE_cargo_QR_batch.pdf")
    except ValueError as exc:
        return JSONResponse({"ok": False, "error": str(exc)}, status_code=404)


@app.get("/supplies/transfer/supply-qr-batch.pdf")
def saved_supply_qr_batch_pdf(supply_ids: str = ""):
    try:
        ids = [value.strip() for value in str(supply_ids).split(",") if value.strip()]
        path = _saved_pdf_batch(ids, kind="supply")
        return FileResponse(str(path), media_type="application/pdf", filename="FBE_supply_QR_batch.pdf")
    except ValueError as exc:
        return JSONResponse({"ok": False, "error": str(exc)}, status_code=404)


@app.get("/supplies/{supply_id}/shipping-units-qr-saved.pdf")
def saved_supply_shipping_units_qr_pdf(supply_id: str):
    path = _shipping_archive_dir(supply_id) / "cargo_qr_saved.pdf"
    if not path.exists():
        return RedirectResponse(
            url=f"/supplies/{supply_id}?error=" + quote("Сохраненный PDF грузомест еще не создан"),
            status_code=303,
        )
    return FileResponse(str(path), media_type="application/pdf", filename=f"cargo_qr_{supply_id}.pdf")


@app.get("/supplies/{supply_id}/supply-qr-saved.pdf")
def saved_supply_qr_pdf(supply_id: str):
    path = Path(config.get("supply_qr_output_dir", "data/supply_qr")) / str(supply_id) / "supply_qr_saved.pdf"
    if not path.exists():
        return RedirectResponse(
            url=f"/supplies/{supply_id}?error=" + quote("Сохраненный QR поставки еще не создан"),
            status_code=303,
        )
    return FileResponse(str(path), media_type="application/pdf", filename=f"supply_qr_{supply_id}.pdf")


@app.get("/supplies/{supply_id}", response_class=HTMLResponse)
def supply_page(
    request: Request,
    supply_id: str,
    msg: str | None = None,
    error: str | None = None,
    sort: str | None = None,
):
    supply = get_supply_details_cached(supply_id)

    enriched = supply_orders_context(supply_id, sort=sort)
    order_ids = [int(o.get("id") or 0) for o in enriched if int(o.get("id") or 0) > 0]
    shipping = _supply_shipping_context(supply_id, order_ids)

    if supply:
        supply["fbe_status"] = shipping["label"]
        supply["fbe_status_class"] = shipping["class"]
        supply["created_at_text"] = format_dt(supply.get("createdAt"))
        supply["closed_at_text"] = format_dt(supply.get("closedAt"))
        supply["scan_dt_text"] = format_dt(supply.get("scanDt"))
        supply["transfer_or_scan_text"] = supply["scan_dt_text"] or supply["closed_at_text"] or ""

    registry_rows = list_fbs_registry_rows(marking_db_path, limit=10000, date_from=_registry_history_start())

    return templates.TemplateResponse(
        request,
        "supply.html",
        {
            "request": request,
            "supply_id": supply_id,
            "supply": supply,
            "orders": enriched,
            "active_supplies": _active_supply_selector_from_registry(registry_rows),
            "shipping": shipping,
            "msg": msg,
            "error": error,
            "default_expiration": default_expiration_date(),
            "config": config,
            "sort": sort or config.get("default_order_sort", "urgent"),
        },
    )


@app.post("/supplies/{supply_id}/download-stickers")
def download_stickers(supply_id: str, order_ids: list[int] = Form(default=[])):
    if not order_ids:
        order_ids = get_supply_order_ids_cached(supply_id)

    sticker_type = config.get("wb_sticker_type", "png")
    paths = save_order_stickers(
        supply_id, [int(x) for x in order_ids], sticker_type,
        _runtime_path("sticker_cache_dir", "data/stickers"),
    )

    return RedirectResponse(
        url=f"/supplies/{supply_id}?msg=WB-стикеры для печати получены: {len(paths)}",
        status_code=303,
    )


@app.post("/supplies/{supply_id}/stickers-pdf")
def stickers_pdf(supply_id: str, order_ids: list[int] = Form(default=[])):
    # v21: this action is explicitly for selected rows only.
    # The full-supply WB PDF button was removed from the UI.
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Не выбраны задания", status_code=303)

    pdf_sticker_type = config.get("pdf_sticker_type", "svg")
    paths = save_order_stickers(supply_id, [int(x) for x in order_ids], pdf_sticker_type, Path(config.get("pdf_output_dir", "data/stickers_pdf")) / "stickers")
    pdf_path = Path(config.get("pdf_output_dir", "data/stickers_pdf")) / f"wb_stickers_{supply_id}.pdf"
    create_stickers_pdf(paths, pdf_path, width_mm=int(config.get("wb_sticker_width", 58)), height_mm=int(config.get("wb_sticker_height", 40)))

    return FileResponse(
        str(pdf_path),
        media_type="application/pdf",
        filename=f"wb_stickers_{supply_id}.pdf",
    )


@app.post("/supplies/{supply_id}/print-wb")
def print_wb_only(supply_id: str, order_ids: list[int] = Form(default=[])):
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Не выбраны задания", status_code=303)

    sticker_type = config.get("wb_sticker_type", "png")
    ext = sticker_ext(sticker_type)
    sticker_dir = _runtime_path("sticker_cache_dir", "data/stickers") / supply_id

    missing = [int(x) for x in order_ids if not (sticker_dir / f"{int(x)}.{ext}").exists()]
    if missing:
        save_order_stickers(
            supply_id, missing, sticker_type,
            _runtime_path("sticker_cache_dir", "data/stickers"),
        )

    printed = 0
    errors: list[str] = []
    for order_id in order_ids:
        path = sticker_dir / f"{int(order_id)}.{ext}"
        try:
            print_wb_sticker(
                printer_name=config["wb_printer_name"],
                sticker_path=str(path),
                sticker_type=sticker_type,
                dry_run=config.get("dry_run_print", True),
                width_mm=int(config.get("wb_sticker_width", 58)),
                height_mm=int(config.get("wb_sticker_height", 40)),
            )
            printed += 1
        except Exception as exc:
            errors.append(f"{order_id}: {exc}")

    if errors:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=WB напечатано {printed}. Ошибки: {'; '.join(errors)}", status_code=303)
    return RedirectResponse(url=f"/supplies/{supply_id}?msg=WB-стикеры напечатаны: {printed}", status_code=303)


@app.post("/supplies/{supply_id}/print-internal")
def print_internal_only(supply_id: str, order_ids: list[int] = Form(default=[])):
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Не выбраны задания", status_code=303)

    catalog = get_catalog()
    order_id_set = {int(x) for x in order_ids}
    orders = get_supply_orders(supply_id)
    order_by_id = {int(order["id"]): order for order in orders if int(order["id"]) in order_id_set}

    printed = 0
    errors: list[str] = []
    for order_id in order_ids:
        order_id = int(order_id)
        order = order_by_id.get(order_id)
        if not order:
            errors.append(f"{order_id}: нет деталей заказа")
            continue

        article = seller_article_from_order(order)
        product = catalog.find_by_article(article)
        if not product:
            errors.append(f"{order_id}: артикул {article} не найден в каталоге")
            continue

        try:
            print_internal_label(config, product)
            log_print_event({
                "type": "internal_label_only",
                "order_id": order_id,
                "seller_article": article,
                "product": product,
                "dry_run": config.get("dry_run_print", True),
            }, config=config)
            printed += 1
        except Exception as exc:
            errors.append(f"{order_id}: {exc}")

    if errors:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Внутренних этикеток напечатано {printed}. Ошибки: {'; '.join(errors)}", status_code=303)
    return RedirectResponse(url=f"/supplies/{supply_id}?msg=Внутренние этикетки напечатаны: {printed}", status_code=303)


def _supply_operation_redirect(supply_id: str, *, msg: str = "", error: str = "", return_to: str = "") -> RedirectResponse:
    if str(return_to or "").strip().lower() == "dashboard":
        return dashboard_redirect(msg=msg, error=error, open_supply=supply_id)
    params = {}
    if msg:
        params["msg"] = msg
    if error:
        params["error"] = error
    suffix = ("?" + urlencode(params)) if params else ""
    return RedirectResponse(url=f"/supplies/{supply_id}{suffix}", status_code=303)


@app.post("/supplies/{supply_id}/shipping-units")
def create_supply_shipping_units(
    supply_id: str, amount: int = Form(default=1), return_to: str = Form(default="")
):
    try:
        amount = int(amount)
        if amount < 1 or amount > 1000:
            raise ValueError("Количество грузомест должно быть от 1 до 1000")
        current_ids = get_supply_order_ids_cached(supply_id)
        shipping = _supply_shipping_context(supply_id, current_ids, remote=True)
        if shipping.get("stage") != "assembly":
            return _supply_operation_redirect(
                supply_id, error="Грузоместа можно создавать, пока поставка находится на сборке", return_to=return_to
            )
        max_total = int(shipping.get("trbx_max") or max(1, len(current_ids) + 1))
        existing = int(shipping.get("trbx_count") or 0)
        if existing + amount > max_total:
            return _supply_operation_redirect(
                supply_id,
                error=f"Для этой поставки можно создать максимум {max_total} грузомест",
                return_to=return_to,
            )
        created_ids = wb.add_supply_shipping_units(supply_id, amount)
        all_ids = wb.get_supply_shipping_units(supply_id)
        with _cache_lock:
            _cache[f"shipping_units:{supply_id}"] = (time.time(), list(all_ids))
        invalidate_cache(f"supply_order_ids:{supply_id}")
        return _supply_operation_redirect(
            supply_id, msg=f"Грузомест создано: {len(created_ids)}; всего: {len(all_ids)}", return_to=return_to
        )
    except (WBApiError, ValueError) as exc:
        return _supply_operation_redirect(supply_id, error=str(exc), return_to=return_to)


@app.post("/supplies/{supply_id}/shipping-units/print")
def print_supply_shipping_units(supply_id: str, return_to: str = Form(default="")):
    try:
        trbx_ids = wb.get_supply_shipping_units(supply_id)
        if not trbx_ids:
            return _supply_operation_redirect(supply_id, error="Сначала создайте грузоместо", return_to=return_to)
        order_count = len(get_supply_order_ids_cached(supply_id))
        manifest = _archive_cargo_qr(supply_id, trbx_ids, order_count=order_count)
        out_dir = _shipping_archive_dir(supply_id)
        printed = 0
        # The manifest was built from one WB request per trbxId, therefore every
        # local file below is already deterministically bound to that cargo place.
        for unit in manifest.get("units") or []:
            path = out_dir / str(unit.get("file") or "")
            if not path.exists():
                raise WBApiError(f"Локальный QR грузоместа {unit.get('trbx_id') or '?'} не найден")
            print_wb_sticker(
                printer_name=config["wb_printer_name"], sticker_path=str(path), sticker_type="png",
                dry_run=config.get("dry_run_print", True), width_mm=58, height_mm=40,
            )
            printed += 1
        return _supply_operation_redirect(
            supply_id, msg=f"QR грузомест сохранены и отправлены на печать: {printed}", return_to=return_to
        )
    except Exception as exc:
        return _supply_operation_redirect(supply_id, error=f"Печать QR грузомест: {exc}", return_to=return_to)


@app.get("/supplies/{supply_id}/shipping-units-qr.pdf")
def supply_shipping_units_qr_pdf(supply_id: str):
    try:
        trbx_ids = wb.get_supply_shipping_units(supply_id)
        if not trbx_ids:
            return RedirectResponse(
                url=f"/supplies/{supply_id}?error=" + quote("Сначала создайте грузоместо"),
                status_code=303,
            )
        order_count = len(get_supply_order_ids_cached(supply_id))
        _archive_cargo_qr(supply_id, trbx_ids, order_count=order_count)
        pdf_path = _shipping_archive_dir(supply_id) / "cargo_qr_saved.pdf"
        return FileResponse(
            str(pdf_path), media_type="application/pdf", filename=f"cargo_qr_{supply_id}.pdf"
        )
    except WBApiError as exc:
        return RedirectResponse(
            url=f"/supplies/{supply_id}?error=" + quote(f"QR грузомест недоступен: {exc}"),
            status_code=303,
        )


@app.post("/supplies/{supply_id}/deliver")
def transfer_supply_to_delivery(supply_id: str, return_to: str = Form(default="")):
    """Legacy single-supply action, kept safe by the same preflight/archive rules as the bulk wizard."""
    try:
        check = _supply_transfer_preflight_one(supply_id)
        if not check.get("ready"):
            raise WBApiError(
                "; ".join(str(value) for value in check.get("blockers") or [])
                or "Проверка готовности поставки не пройдена"
            )
        order_ids = [int(x) for x in (check.get("order_ids") or []) if int(x) > 0]
        trbx_ids = wb.get_supply_shipping_units(supply_id)
        if not trbx_ids:
            raise WBApiError("У поставки нет грузомест. Сначала создайте и распечатайте QR грузомест")
        if not _cargo_manifest_matches(supply_id, trbx_ids):
            _archive_cargo_qr(supply_id, trbx_ids, order_count=len(order_ids))

        wb.deliver_supply(supply_id)
        try:
            _persist_supply_delivered_local(
                supply_id, str(check.get("name") or supply_id), order_ids
            )
        except Exception:
            pass
        invalidate_cache("supplies")
        invalidate_cache("fbs_operational_registry")
        invalidate_cache(f"supply_order_ids:{supply_id}")
        qr_message = ""
        try:
            _archive_supply_qr(supply_id, order_count=len(order_ids))
            qr_message = " QR поставки получен и сохранен локально."
        except Exception as exc:
            qr_message = f" Поставка передана, но QR пока не удалось сохранить: {exc}"
        if order_ids:
            try:
                _sync_fbs_lifecycle_registry(selected_order_ids=order_ids, include_supplies=False)
            except Exception:
                pass
        return _supply_operation_redirect(
            supply_id, msg="Поставка передана WB." + qr_message, return_to=return_to
        )
    except (WBApiError, ValueError) as exc:
        return _supply_operation_redirect(supply_id, error=str(exc), return_to=return_to)


@app.get("/supplies/{supply_id}/supply-qr.pdf")
def supply_qr_pdf_get(supply_id: str):
    try:
        saved = Path(config.get("supply_qr_output_dir", "data/supply_qr")) / supply_id / "supply_qr_saved.pdf"
        if not saved.exists():
            _archive_supply_qr(supply_id, order_count=len(get_supply_order_ids_cached(supply_id)))
        return FileResponse(
            str(saved), media_type="application/pdf", filename=f"supply_qr_{supply_id}.pdf"
        )
    except WBApiError as exc:
        return RedirectResponse(
            url=f"/supplies/{supply_id}?error=" + quote(f"QR поставки недоступен: {exc}"),
            status_code=303,
        )


@app.post("/supplies/{supply_id}/supply-qr-pdf")
def supply_qr_pdf(supply_id: str):
    try:
        pdf_path = _archive_supply_qr(supply_id, order_count=len(get_supply_order_ids_cached(supply_id)))
        return FileResponse(str(pdf_path), media_type="application/pdf", filename=f"supply_qr_{supply_id}.pdf")
    except WBApiError as exc:
        return RedirectResponse(
            url=f"/supplies/{supply_id}?error=QR поставки недоступен. Обычно WB отдает его только после перевода поставки в доставку. Детали: {exc}",
            status_code=303,
        )


@app.post("/supplies/{supply_id}/print-supply-qr")
def print_supply_qr(supply_id: str, return_to: str = Form(default="")):
    try:
        qr_type = config.get("supply_qr_type", "png")
        data = wb.get_supply_barcode(supply_id, barcode_type=qr_type)
        payload = wb.decode_sticker_file(data.get("file", ""), qr_type)
    except WBApiError as exc:
        return _supply_operation_redirect(
            supply_id, error=f"QR поставки недоступен: {exc}", return_to=return_to
        )

    out_dir = Path(config.get("supply_qr_output_dir", "data/supply_qr")) / supply_id
    out_dir.mkdir(parents=True, exist_ok=True)
    ext = sticker_ext(qr_type)
    qr_path = out_dir / f"supply_qr_{supply_id}.{ext}"
    qr_path.write_bytes(payload)

    try:
        print_wb_sticker(
            printer_name=config["wb_printer_name"],
            sticker_path=str(qr_path),
            sticker_type=qr_type,
            dry_run=config.get("dry_run_print", True),
            width_mm=int(config.get("wb_sticker_width", 58)),
            height_mm=int(config.get("wb_sticker_height", 40)),
        )
    except Exception as exc:
        return _supply_operation_redirect(supply_id, error=f"Ошибка печати QR поставки: {exc}", return_to=return_to)

    return _supply_operation_redirect(supply_id, msg="QR поставки отправлен на печать", return_to=return_to)



def build_internal_label_rows_for_orders(supply_id: str, order_ids: list[int]) -> tuple[list[dict[str, str]], list[str]]:
    """Create BarTender CSV rows in the exact selected/supply order.

    This is the stable workflow: FBE prepares the BarTender input file, BarTender
    generates the internal-label PDF, then FBE mixes WB PDF + BarTender PDF.
    """
    from printers.bartender_label import INTERNAL_LABEL_FIELDS, build_internal_label_row

    catalog = get_catalog()
    order_id_set = {int(x) for x in order_ids}
    orders = get_supply_orders(supply_id)
    order_by_id = {int(order["id"]): order for order in orders if int(order["id"]) in order_id_set}

    rows: list[dict[str, str]] = []
    errors: list[str] = []
    order_log: list[str] = []

    for order_id in order_ids:
        oid = int(order_id)
        order = order_by_id.get(oid)
        if not order:
            errors.append(f"{oid}: нет деталей заказа")
            continue
        article = seller_article_from_order(order)
        product = catalog.find_by_article(article)
        if not product:
            errors.append(f"{oid}: артикул {article} не найден в каталоге")
            continue
        row = build_internal_label_row(product)
        # Keep order_id for the human-readable debug log, but not in the CSV contract.
        rows.append(row)
        order_log.append(f"{oid};{article};{row.get('category','')};{row.get('product_name','')}")

    # Validate row fields are exactly what BarTender expects.
    normalized = [{field: row.get(field, "") for field in INTERNAL_LABEL_FIELDS} for row in rows]
    return normalized, errors + ([] if normalized else ["Нет строк для BarTender CSV"])


def write_bartender_batch_csv(rows: list[dict[str, str]]) -> Path:
    from printers.bartender_label import INTERNAL_LABEL_FIELDS

    csv_path = Path(config["bartender_csv"])
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=INTERNAL_LABEL_FIELDS,
            delimiter=";",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)
    return csv_path


def write_bartender_batch_log(supply_id: str, order_ids: list[int], rows: list[dict[str, str]]) -> Path:
    out = _runtime_path("print_debug_dir", "data/print") / f"bartender_batch_order_{supply_id}.txt"
    out.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "Порядок строк в internal_label.csv для BarTender",
        "order_id;seller_article;category;product_name",
    ]
    for oid, row in zip(order_ids, rows):
        lines.append(
            f"{oid};{row.get('seller_article','')};{row.get('category','')};{row.get('product_name','')}"
        )
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


@app.post("/supplies/{supply_id}/prepare-bartender-csv")
def prepare_bartender_csv_selected(supply_id: str, order_ids: list[int] = Form(default=[])):
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Не выбраны задания", status_code=303)
    rows, errors = build_internal_label_rows_for_orders(supply_id, [int(x) for x in order_ids])
    if not rows:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={'; '.join(errors)}", status_code=303)
    csv_path = write_bartender_batch_csv(rows)
    log_path = write_bartender_batch_log(supply_id, [int(x) for x in order_ids], rows)
    msg = f"CSV для BarTender подготовлен: {len(rows)} строк. Файл: {csv_path}. Контроль порядка: {log_path}"
    if errors:
        msg += f". Пропущено: {'; '.join(errors)}"
    return RedirectResponse(url=f"/supplies/{supply_id}?msg={msg}", status_code=303)


@app.post("/supplies/{supply_id}/prepare-bartender-csv-all")
def prepare_bartender_csv_all(supply_id: str):
    order_ids = get_supply_order_ids_cached(supply_id)
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=В поставке нет заданий", status_code=303)
    rows, errors = build_internal_label_rows_for_orders(supply_id, [int(x) for x in order_ids])
    if not rows:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={'; '.join(errors)}", status_code=303)
    csv_path = write_bartender_batch_csv(rows)
    log_path = write_bartender_batch_log(supply_id, [int(x) for x in order_ids], rows)
    msg = f"CSV для BarTender подготовлен по всей поставке: {len(rows)} строк. Файл: {csv_path}. Контроль порядка: {log_path}"
    if errors:
        msg += f". Пропущено: {'; '.join(errors)}"
    return RedirectResponse(url=f"/supplies/{supply_id}?msg={msg}", status_code=303)




def _safe_filename_part(value: str, fallback: str = "item") -> str:
    value = str(value or "").strip()
    value = re.sub(r"[^0-9A-Za-zА-Яа-я._-]+", "_", value)
    return value[:80] or fallback


def _pdf_page_count(path: str | Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(path)).pages)


def _supply_print_identity(supply_id: str, order_count: int) -> dict[str, Any]:
    """Resolve stable human-facing supply data for the separator label."""
    supply: dict[str, Any] = {}
    try:
        supply = get_supply_details_cached(str(supply_id)) or {}
    except Exception:
        supply = {}
    if not supply:
        try:
            supply = next(
                (row for row in _local_supply_snapshot() if str(row.get("id") or "") == str(supply_id)),
                {},
            )
        except Exception:
            supply = {}

    name = str(supply.get("name") or supply_id or "Поставка").strip()
    date_text = ""
    try:
        from utils import parse_wb_datetime
        dt = parse_wb_datetime(supply.get("createdAt") or supply.get("created_at_wb"))
        if dt is not None:
            date_text = dt.astimezone(moscow_now().tzinfo).strftime("%d.%m.%Y")
    except Exception:
        date_text = ""
    if not date_text:
        date_text = moscow_now().strftime("%d.%m.%Y")
    return {
        "supply_name": name,
        "supply_id": str(supply_id),
        "date_text": date_text,
        "order_count": int(order_count),
    }


def _prepend_supply_print_label(
    *,
    supply_id: str,
    order_count: int,
    body_pdf: str | Path,
    final_pdf: str | Path,
    title: str = "ПОСТАВКА",
    footer: str = "",
) -> Path:
    body = Path(body_pdf)
    final = Path(final_pdf)
    identity = _supply_print_identity(supply_id, order_count)
    header = body.parent / f"_supply_header_{_safe_filename_part(supply_id)}.pdf"
    create_supply_info_label_pdf(
        **identity,
        output_path=header,
        width_mm=float(config.get("wb_sticker_width", 58) or 58),
        height_mm=float(config.get("wb_sticker_height", 40) or 40),
        title=title,
        footer=footer,
    )
    prepend_pdf_file(header, body, final)
    return final


def build_final_pdf_from_original_wb_files(
    supply_id: str,
    wb_pdf_path: str | Path,
    picklist_xlsx_path: str | Path,
) -> tuple[Path | None, int, list[str], list[str]]:
    """Battle-safe workflow based on the proven manual process.

    v31: fast BarTender workload.
    - WB stickers remain the original WB portal PDF, untouched.
    - The WB picklist gives the exact order.
    - FBE groups internal labels by .btw template and starts BarTender once per
      used template, not once per item.
    - After BarTender returns one multi-page PDF per template, FBE maps those
      pages back to the original WB/picklist order and alternates pages.
    """
    errors: list[str] = []
    warnings: list[str] = []

    wb_pdf_path = Path(wb_pdf_path)
    picklist_xlsx_path = Path(picklist_xlsx_path)
    pick_rows, parse_warnings = parse_wb_picklist_xlsx(picklist_xlsx_path)
    warnings.extend(parse_warnings)

    if not pick_rows:
        return None, 0, errors + ["Лист подбора не содержит распознанных строк с артикулами продавца"], warnings

    try:
        wb_pages = _pdf_page_count(wb_pdf_path)
    except Exception as exc:
        return None, 0, errors + [f"Не удалось прочитать PDF WB: {exc}"], warnings

    if wb_pages != len(pick_rows):
        return None, 0, errors + [
            f"Стоп: количество страниц WB PDF ({wb_pages}) не совпадает с количеством строк Листа подбора ({len(pick_rows)}). "
            "Печать заблокирована, чтобы не перепутать QR-коды и этикетки."
        ], warnings

    catalog = get_catalog()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    work_dir = Path(config.get("portal_pdf_output_dir", "data/portal_pdf")) / supply_id / ts
    internal_pdf_dir = work_dir / "internal_pdf_batches"
    work_dir.mkdir(parents=True, exist_ok=True)
    internal_pdf_dir.mkdir(parents=True, exist_ok=True)

    log_lines = [
        "Оригинальный WB PDF + Лист подбора",
        "FBE v0.39: BarTender batch mode",
        f"supply_id={supply_id}",
        f"wb_pdf={wb_pdf_path}",
        f"picklist={picklist_xlsx_path}",
        "",
        "index;sticker_number;seller_article;template;template_page_index;excel_product;internal_pdf",
    ]

    # Build all rows first. Do not start BarTender until we know the whole
    # shipment is resolvable; otherwise we may waste time and still block merge.
    items: list[dict[str, Any]] = []
    for index, row in enumerate(pick_rows, start=1):
        article = row["seller_article"]
        product = catalog.find_by_article(article)
        if not product:
            errors.append(f"Строка {index}: артикул {article} не найден в каталоге товаров")
            continue

        try:
            template_path = str(Path(choose_bartender_template(config, product)).resolve())
        except Exception as exc:
            errors.append(f"Строка {index}, артикул {article}: не выбран шаблон BarTender: {exc}")
            continue

        csv_row = build_internal_label_row(product)
        items.append({
            "index": index,
            "pick_row": row,
            "article": article,
            "product": product,
            "template_path": template_path,
            "csv_row": csv_row,
        })

    if errors:
        log_lines.extend(["", "Ошибки до запуска BarTender:", *errors])
        (work_dir / "portal_order_log.txt").write_text("\n".join(log_lines), encoding="utf-8")
        return None, 0, errors, warnings

    # Group by template while preserving picklist order inside each template.
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        grouped.setdefault(item["template_path"], []).append(item)

    if len(grouped) > 1:
        warnings.append(
            f"BarTender запущен пакетно по шаблонам: {len(grouped)} запуска вместо {len(items)}. "
            "Порядок в финальной PDF восстановлен по Листу подбора."
        )
    else:
        warnings.append(f"BarTender запущен одним пакетным заданием на {len(items)} этикеток.")

    internal_page_refs_by_index: dict[int, tuple[Path, int]] = {}
    rendered_count = 0

    for group_no, (template_path, group_items) in enumerate(grouped.items(), start=1):
        template_stem = _safe_filename_part(Path(template_path).stem, f"template_{group_no}")
        batch_name = f"{supply_id}_{ts}_{group_no:02d}_{template_stem}"
        out_pdf = internal_pdf_dir / f"{group_no:02d}_{template_stem}_{len(group_items)}labels.pdf"
        rows = [item["csv_row"] for item in group_items]

        try:
            batch_pdf = render_internal_labels_batch_pdf_with_bullzip(
                config=config,
                rows=rows,
                template_path=template_path,
                output_pdf=out_pdf,
                batch_name=batch_name,
            )
        except Exception as exc:
            errors.append(
                f"Шаблон {Path(template_path).name}: ошибка BarTender/Bullzip при пакетной печати "
                f"{len(group_items)} этикеток: {exc}"
            )
            continue

        try:
            batch_pages = _pdf_page_count(batch_pdf)
        except Exception as exc:
            errors.append(f"Шаблон {Path(template_path).name}: не удалось прочитать PDF BarTender: {exc}")
            continue

        if batch_pages != len(group_items):
            errors.append(
                f"Шаблон {Path(template_path).name}: BarTender создал {batch_pages} страниц, "
                f"а CSV содержит {len(group_items)} строк. Склейка заблокирована. "
                "Проверь, что в .btw не включены фильтры/record selection и печатаются все записи."
            )
            continue

        for page_index, item in enumerate(group_items):
            original_index = int(item["index"])
            internal_page_refs_by_index[original_index] = (Path(batch_pdf), page_index)
            rendered_count += 1
            sticker_number = item["pick_row"].get("sticker_number", "")
            article = item["article"]
            product = item["product"]
            log_lines.append(
                f"{original_index};{sticker_number};{article};{Path(template_path).name};"
                f"{page_index + 1};{product.get('display_name','')};{batch_pdf}"
            )

    if errors:
        log_lines.extend(["", "Ошибки:", *errors])
        if warnings:
            log_lines.extend(["", "Предупреждения:", *warnings])
        (work_dir / "portal_order_log.txt").write_text("\n".join(log_lines), encoding="utf-8")
        return None, rendered_count, errors, warnings

    if rendered_count != wb_pages:
        return None, rendered_count, [
            f"Стоп: BarTender создал {rendered_count} этикеток, а страниц WB PDF {wb_pages}. "
            "Печать заблокирована, чтобы не нарушить порядок."
        ], warnings

    internal_page_refs: list[tuple[Path, int]] = []
    missing: list[str] = []
    for index in range(1, wb_pages + 1):
        ref = internal_page_refs_by_index.get(index)
        if not ref:
            missing.append(str(index))
        else:
            internal_page_refs.append(ref)

    if missing:
        return None, rendered_count, [
            "Стоп: не хватает внутренних этикеток для строк Листа подбора: " + ", ".join(missing)
        ], warnings

    body_pdf = work_dir / f"_FBE_ORIGINAL_WB_BODY_{supply_id}_{ts}.pdf"
    final_pdf = work_dir / f"FBE_ORIGINAL_WB_FINAL_{supply_id}_{ts}.pdf"
    merge_wb_pdf_with_internal_page_refs(wb_pdf_path, internal_page_refs, body_pdf)
    _prepend_supply_print_label(
        supply_id=supply_id,
        order_count=wb_pages,
        body_pdf=body_pdf,
        final_pdf=final_pdf,
        title="ПОСТАВКА / ЗАКАЗЫ + ЭТИКЕТКИ",
    )

    log_lines.extend(["", f"FIRST_LABEL=supply name + supply id + date + orders={wb_pages}", f"FINAL={final_pdf}"])
    if warnings:
        log_lines.extend(["", "Предупреждения:", *warnings])
    (work_dir / "portal_order_log.txt").write_text("\n".join(log_lines), encoding="utf-8")
    return final_pdf, rendered_count, errors, warnings


@app.post("/supplies/{supply_id}/set-expiration")
def set_expiration_for_supply(
    supply_id: str,
    expiration: str = Form(...),
):
    try:
        normalized_expiration = normalize_expiration_date(expiration)
    except ValueError as exc:
        return dashboard_redirect(error=str(exc), open_supply=supply_id)

    order_ids = get_supply_order_ids_cached(supply_id)
    if not order_ids:
        return dashboard_redirect(error="В поставке нет сборочных заданий", open_supply=supply_id)

    updated = 0
    errors: list[str] = []
    gap = float(config.get("metadata_api_gap_seconds", 0.07))
    for order_id in order_ids:
        try:
            wb.set_order_expiration(int(order_id), normalized_expiration)
            updated += 1
        except Exception as exc:
            errors.append(f"{order_id}: {exc}")
        if gap > 0:
            time.sleep(gap)

    if errors:
        short_errors = "; ".join(errors[:5])
        if len(errors) > 5:
            short_errors += f"; ещё ошибок: {len(errors) - 5}"
        return dashboard_redirect(
            error=f"Срок {normalized_expiration} указан для {updated} из {len(order_ids)}. Ошибки: {short_errors}",
            open_supply=supply_id,
        )

    return dashboard_redirect(
        msg=f"Срок годности {normalized_expiration} указан для всех заданий в поставке: {updated}",
        open_supply=supply_id,
    )


@app.post("/supplies/{supply_id}/original-wb-final-pdf")
async def original_wb_final_pdf(
    supply_id: str,
    wb_pdf: UploadFile = File(...),
    picklist_excel: UploadFile = File(...),
):
    work_dir = Path(config.get("portal_pdf_output_dir", "data/portal_pdf")) / supply_id / "uploads"
    work_dir.mkdir(parents=True, exist_ok=True)

    # Fixed filenames are intentional: each run replaces the previous uploaded
    # pair, so the folder does not accumulate confusing copies.
    wb_pdf_path = work_dir / "wb_original_stickers.pdf"
    picklist_path = work_dir / "wb_picklist.xlsx"
    wb_pdf_path.write_bytes(await wb_pdf.read())
    picklist_path.write_bytes(await picklist_excel.read())

    final_pdf, count, errors, warnings = build_final_pdf_from_original_wb_files(
        supply_id=supply_id,
        wb_pdf_path=wb_pdf_path,
        picklist_xlsx_path=picklist_path,
    )

    if not final_pdf:
        text = "; ".join(errors or ["Не удалось сформировать итоговую ленту"])
        if warnings:
            text += ". Предупреждения: " + "; ".join(warnings)
        return RedirectResponse(url=f"/supplies/{supply_id}?error={text}", status_code=303)

    return FileResponse(
        str(final_pdf),
        media_type="application/pdf",
        filename=f"FBE_ORIGINAL_WB_{supply_id}.pdf",
        headers={"Content-Disposition": f'inline; filename="FBE_ORIGINAL_WB_{supply_id}.pdf"'},
    )


@app.get("/supplies/{supply_id}/pdf-mixer", response_class=HTMLResponse)
def pdf_mixer_page(request: Request, supply_id: str, msg: str | None = None, error: str | None = None):
    return templates.TemplateResponse(
        request,
        "pdf_mixer.html",
        {
            "request": request,
            "supply_id": supply_id,
            "msg": msg,
            "error": error,
            "default_expiration": default_expiration_date(),
            "config": config,
        },
    )


@app.post("/supplies/{supply_id}/mix-pdf")
async def mix_pdf_uploads(
    supply_id: str,
    wb_pdf: UploadFile = File(...),
    bartender_pdf: UploadFile = File(...),
):
    work_dir = Path(config.get("mixed_pdf_output_dir", "data/mixed_pdf")) / supply_id
    work_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    wb_path = work_dir / f"wb_{ts}.pdf"
    bt_path = work_dir / f"bartender_{ts}.pdf"
    out_path = work_dir / f"fbe_mixed_{supply_id}_{ts}.pdf"

    wb_path.write_bytes(await wb_pdf.read())
    bt_path.write_bytes(await bartender_pdf.read())

    try:
        merge_alternating_pdfs(wb_path, bt_path, out_path)
    except Exception as exc:
        return RedirectResponse(url=f"/supplies/{supply_id}/pdf-mixer?error=Ошибка склейки PDF: {exc}", status_code=303)

    return FileResponse(
        str(out_path),
        media_type="application/pdf",
        filename=f"FBE_ЛЕНТА_{supply_id}.pdf",
    )




def build_automated_final_pdf_for_orders(supply_id: str, order_ids: list[int]) -> tuple[Path | None, int, list[str]]:
    """One-button workflow:

    1. Save WB sticker PDF in the exact selected/supply order.
    2. For each order, write one-row CSV and let BarTender print to Bullzip PDF.
    3. Merge WB PDF + internal PDFs page-by-page.
    4. Return final PDF to the browser.

    This mirrors the proven manual workflow while removing manual copy/paste.
    """
    order_ids = [int(x) for x in order_ids]
    if not order_ids:
        return None, 0, ["Нет выбранных заданий"]

    catalog = get_catalog()
    orders = get_supply_orders(supply_id)
    order_by_id = {int(order["id"]): order for order in orders if int(order["id"]) in set(order_ids)}

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    work_dir = Path(config.get("auto_pdf_output_dir", "data/auto_pdf")) / supply_id / ts
    wb_img_dir = work_dir / "wb_images"
    internal_pdf_dir = work_dir / "internal_pdf"
    work_dir.mkdir(parents=True, exist_ok=True)

    # v26: use SVG by default for high-quality vector WB PDF composition.
    sticker_type = config.get("pdf_sticker_type", "svg")
    ext = sticker_ext(sticker_type)
    save_order_stickers(supply_id, order_ids, sticker_type, wb_img_dir.parent)
    sticker_dir = wb_img_dir.parent / supply_id
    wb_paths = [sticker_dir / f"{oid}.{ext}" for oid in order_ids]
    missing_wb = [str(oid) for oid, path in zip(order_ids, wb_paths) if not path.exists()]
    if missing_wb:
        return None, 0, [f"Не удалось получить WB-стикеры: {', '.join(missing_wb)}"]

    wb_pdf = work_dir / f"wb_stickers_{supply_id}_{ts}.pdf"
    create_stickers_pdf(
        wb_paths,
        wb_pdf,
        width_mm=int(config.get("wb_sticker_width", 58)),
        height_mm=int(config.get("wb_sticker_height", 40)),
    )

    internal_pdfs: list[Path] = []
    errors: list[str] = []
    log_lines = [
        "Автоматическая PDF-лента FBE",
        "Порядок:",
        "order_id;seller_article;category;internal_pdf",
    ]

    for index, order_id in enumerate(order_ids, start=1):
        order = order_by_id.get(order_id)
        if not order:
            errors.append(f"{order_id}: нет деталей заказа")
            continue

        article = seller_article_from_order(order)
        product = catalog.find_by_article(article)
        if not product:
            errors.append(f"{order_id}: артикул {article} не найден в каталоге")
            continue

        out_pdf = internal_pdf_dir / f"{index:03d}_{order_id}_{article}.pdf"
        try:
            internal_pdf = render_internal_label_pdf_with_bullzip(
                config=config,
                product=product,
                order_id=order_id,
                output_pdf=out_pdf,
            )
            internal_pdfs.append(internal_pdf)
            log_lines.append(f"{order_id};{article};{product.get('category','')};{internal_pdf}")
        except Exception as exc:
            errors.append(f"{order_id}: ошибка BarTender/Bullzip: {exc}")

    if not internal_pdfs:
        return None, 0, errors or ["BarTender не сформировал ни одной внутренней этикетки"]

    if len(internal_pdfs) != len(wb_paths):
        errors.append(
            f"Количество WB-страниц ({len(wb_paths)}) и внутренних этикеток ({len(internal_pdfs)}) различается. "
            "Итоговая лента может не совпасть 1 к 1."
        )

    body_pdf = work_dir / f"_FBE_FINAL_BODY_{supply_id}_{ts}.pdf"
    final_pdf = work_dir / f"FBE_FINAL_{supply_id}_{ts}.pdf"
    merge_wb_pdf_with_internal_pdf_files(wb_pdf, internal_pdfs, body_pdf)
    _prepend_supply_print_label(
        supply_id=supply_id,
        order_count=len(order_ids),
        body_pdf=body_pdf,
        final_pdf=final_pdf,
        title="ПОСТАВКА / ЗАКАЗЫ + ЭТИКЕТКИ",
    )

    log_lines.extend([
        "",
        f"WB PDF: {wb_pdf}",
        f"FINAL PDF: {final_pdf}",
    ])
    if errors:
        log_lines.extend(["", "Ошибки:", *errors])
    (work_dir / "order_log.txt").write_text("\n".join(log_lines), encoding="utf-8")

    return final_pdf, len(internal_pdfs), errors


@app.post("/supplies/{supply_id}/auto-final-pdf")
def auto_final_pdf_selected(supply_id: str, order_ids: list[int] = Form(default=[])):
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Не выбраны задания", status_code=303)

    pdf_path, count, errors = build_automated_final_pdf_for_orders(supply_id, [int(x) for x in order_ids])
    if not pdf_path:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={'; '.join(errors)}", status_code=303)

    filename = f"FBE_FINAL_{supply_id}.pdf"
    return FileResponse(str(pdf_path), media_type="application/pdf", filename=filename)


@app.post("/supplies/{supply_id}/auto-final-pdf-all")
def auto_final_pdf_all(supply_id: str):
    order_ids = get_supply_order_ids_cached(supply_id)
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=В поставке нет заданий", status_code=303)

    pdf_path, count, errors = build_automated_final_pdf_for_orders(supply_id, order_ids)
    if not pdf_path:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={'; '.join(errors)}", status_code=303)

    return FileResponse(str(pdf_path), media_type="application/pdf", filename=f"FBE_FINAL_{supply_id}.pdf")

def build_pair_pdf_for_orders(supply_id: str, order_ids: list[int]) -> tuple[Path | None, int, list[str]]:
    """Build one ordered PDF: WB sticker → internal label → WB sticker → internal label.

    This is the preferred fast mode for Xprinter through Windows driver because the printer
    receives one PDF queue instead of many tiny raster jobs.
    """
    catalog = get_catalog()
    order_id_set = {int(x) for x in order_ids}
    orders = get_supply_orders(supply_id)
    order_by_id = {int(order["id"]): order for order in orders if int(order["id"]) in order_id_set}

    # v26: PDF mode uses vector SVG by default for WB sticker quality.
    sticker_type = config.get("pdf_sticker_type", "svg")
    ext = sticker_ext(sticker_type)
    sticker_dir = _runtime_path("sticker_cache_dir", "data/stickers") / supply_id

    missing = [oid for oid in order_id_set if not (sticker_dir / f"{oid}.{ext}").exists()]
    if missing:
        save_order_stickers(
            supply_id, missing, sticker_type,
            _runtime_path("sticker_cache_dir", "data/stickers"),
        )

    out_root = Path(config.get("paired_pdf_output_dir", "data/paired_pdf")) / supply_id
    internal_dir = out_root / "internal_images"
    pdf_dir = out_root / "pdf"
    pairs: list[tuple[Path, Path]] = []
    errors: list[str] = []

    for order_id in order_ids:
        order_id = int(order_id)
        order = order_by_id.get(order_id)
        if not order:
            errors.append(f"{order_id}: нет деталей заказа")
            continue

        article = seller_article_from_order(order)
        product = catalog.find_by_article(article)
        if not product:
            errors.append(f"{order_id}: артикул {article} не найден в каталоге")
            continue

        wb_path = sticker_dir / f"{order_id}.{ext}"
        if not wb_path.exists():
            errors.append(f"{order_id}: WB-стикер не скачан")
            continue

        try:
            internal_img = export_internal_label_image(
                config=config,
                product=product,
                order_id=order_id,
                output_dir=internal_dir,
            )
            pairs.append((wb_path, internal_img))
            log_print_event({
                "type": "paired_pdf_prepare",
                "order_id": order_id,
                "seller_article": article,
                "wb_sticker_path": str(wb_path),
                "internal_image_path": str(internal_img),
                "product": product,
                "dry_run": config.get("dry_run_print", True),
            }, config=config)
        except Exception as exc:
            errors.append(f"{order_id}: {exc}")

    if not pairs:
        return None, 0, errors or ["Нет готовых пар для PDF"]

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    body_pdf = pdf_dir / f"_paired_queue_body_{supply_id}_{ts}.pdf"
    pdf_path = pdf_dir / f"paired_queue_{supply_id}_{ts}.pdf"
    create_paired_pdf(
        pairs,
        body_pdf,
        width_mm=int(config.get("wb_sticker_width", 58)),
        height_mm=int(config.get("wb_sticker_height", 40)),
    )
    _prepend_supply_print_label(
        supply_id=supply_id,
        order_count=len(pairs),
        body_pdf=body_pdf,
        final_pdf=pdf_path,
        title="ПОСТАВКА / ЗАКАЗЫ + ЭТИКЕТКИ",
    )
    return pdf_path, len(pairs), errors


@app.post("/supplies/{supply_id}/paired-pdf")
def paired_pdf_selected(supply_id: str, order_ids: list[int] = Form(default=[])):
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Не выбраны задания", status_code=303)

    pdf_path, count, errors = build_pair_pdf_for_orders(supply_id, [int(x) for x in order_ids])
    if not pdf_path:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={'; '.join(errors)}", status_code=303)

    mode = print_pdf_queue(config, pdf_path)
    msg = f"PDF-лента создана: {count} пар. Режим: {mode}. Файл: {pdf_path}"
    if errors:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={msg}. Ошибки: {'; '.join(errors)}", status_code=303)
    return RedirectResponse(url=f"/supplies/{supply_id}?msg={msg}", status_code=303)


@app.post("/supplies/{supply_id}/paired-pdf-all")
def paired_pdf_all(supply_id: str):
    order_ids = get_supply_order_ids_cached(supply_id)
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=В поставке нет заданий", status_code=303)

    pdf_path, count, errors = build_pair_pdf_for_orders(supply_id, order_ids)
    if not pdf_path:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={'; '.join(errors)}", status_code=303)

    mode = print_pdf_queue(config, pdf_path)
    msg = f"PDF-лента создана: {count} пар. Режим: {mode}. Файл: {pdf_path}"
    if errors:
        return RedirectResponse(url=f"/supplies/{supply_id}?error={msg}. Ошибки: {'; '.join(errors)}", status_code=303)
    return RedirectResponse(url=f"/supplies/{supply_id}?msg={msg}", status_code=303)


def print_pairs_for_orders(supply_id: str, order_ids: list[int]) -> tuple[int, list[str]]:
    catalog = get_catalog()
    order_id_set = {int(x) for x in order_ids}
    orders = get_supply_orders(supply_id)
    order_by_id = {int(order["id"]): order for order in orders if int(order["id"]) in order_id_set}

    sticker_type = config.get("wb_sticker_type", "png")
    ext = sticker_ext(sticker_type)
    sticker_dir = _runtime_path("sticker_cache_dir", "data/stickers") / supply_id

    missing = [oid for oid in order_id_set if not (sticker_dir / f"{oid}.{ext}").exists()]
    if missing:
        save_order_stickers(
            supply_id, missing, sticker_type,
            _runtime_path("sticker_cache_dir", "data/stickers"),
        )

    printed = 0
    errors: list[str] = []

    for order_id in order_ids:
        order_id = int(order_id)
        order = order_by_id.get(order_id)
        if not order:
            errors.append(f"{order_id}: нет деталей заказа")
            continue

        article = seller_article_from_order(order)
        product = catalog.find_by_article(article)
        if not product:
            errors.append(f"{order_id}: артикул {article} не найден в каталоге")
            continue

        sticker_path = sticker_dir / f"{order_id}.{ext}"
        if not sticker_path.exists():
            errors.append(f"{order_id}: WB-стикер не скачан")
            continue

        try:
            print_order_pair(
                config=config,
                order_id=order_id,
                seller_article=article,
                wb_sticker_path=str(sticker_path),
                product=product,
            )
            printed += 1
        except Exception as exc:
            errors.append(f"{order_id}: {exc}")
    return printed, errors


@app.post("/supplies/{supply_id}/print")
def print_supply_orders(supply_id: str, order_ids: list[int] = Form(default=[])):
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=Не выбраны задания", status_code=303)

    printed, errors = print_pairs_for_orders(supply_id, [int(x) for x in order_ids])
    if errors:
        return RedirectResponse(
            url=f"/supplies/{supply_id}?error=Напечатано {printed}. Ошибки: {'; '.join(errors)}",
            status_code=303,
        )
    return RedirectResponse(url=f"/supplies/{supply_id}?msg=Напечатано пар: {printed}", status_code=303)


@app.post("/supplies/{supply_id}/print-all-pairs")
def print_all_pairs(supply_id: str):
    order_ids = get_supply_order_ids_cached(supply_id)
    if not order_ids:
        return RedirectResponse(url=f"/supplies/{supply_id}?error=В поставке нет заданий", status_code=303)

    printed, errors = print_pairs_for_orders(supply_id, order_ids)
    if errors:
        return RedirectResponse(
            url=f"/supplies/{supply_id}?error=Напечатано {printed}. Ошибки: {'; '.join(errors)}",
            status_code=303,
        )
    return RedirectResponse(url=f"/supplies/{supply_id}?msg=Напечатано всех пар: {printed}", status_code=303)


class PostSaleIdsPayload(BaseModel):
    order_ids: list[int] = Field(default_factory=list)
    confirmed: bool = False


class PostSaleLookupPayload(BaseModel):
    order_ids: list[int] = Field(default_factory=list)


def _post_sale_group_rows(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        group = _assignment_product_group(row)
        if group not in {"perfumery", "chemistry"}:
            raise MarkingDbError(
                f"Задание {row.get('order_id') or '—'}: товарная группа КИЗ не определена; "
                "юридическая операция заблокирована"
            )
        result.setdefault(group, []).append(row)
    return result


def _recover_wb_sgtin_metadata(order_ids: list[int]) -> dict[str, Any]:
    """Fetch FBS SGTIN metadata from WB and restore missing local KIZ links.

    This is deliberately independent from the FBE assignment history. A seller
    may have attached the KIZ manually in the WB cabinet; in that case WB is the
    authoritative source for the assembly-order association.
    """
    ids = sorted({int(x) for x in order_ids if int(x) > 0})
    stats = {
        "requested": len(ids), "meta_rows": 0, "with_code": 0,
        "recovered": 0, "already_linked": 0, "meta_only": 0,
        "conflicts": 0, "errors": [],
    }
    if not ids:
        return stats

    registry_rows = list_fbs_lifecycle_rows(marking_db_path, limit=max(10000, len(ids) * 2))
    by_id = {int(r.get("order_id") or 0): r for r in registry_rows if int(r.get("order_id") or 0)}

    for start in range(0, len(ids), 100):
        batch = ids[start:start + 100]
        try:
            meta_rows = wb.get_orders_meta(batch)
        except Exception as exc:
            stats["errors"].append(f"WB metadata {batch[0]}…: {exc}")
            continue
        meta_by_id = {_meta_order_id(item): item for item in meta_rows if _meta_order_id(item)}
        stats["meta_rows"] += len(meta_by_id)

        if start + 100 < len(ids):
            time.sleep(0.22)
        for oid in batch:
            item = meta_by_id.get(oid)
            info = _sgtin_meta_info(item)
            values = [str(v) for v in info.get("values") or [] if str(v)]
            decision = str(info.get("decision") or "").strip()
            state = decision or ("filled" if values else ("reported" if info.get("available") else ""))
            update_fbs_order_sgtin_meta(
                marking_db_path, order_id=oid, state=state, values=values
            )
            if values:
                stats["with_code"] += 1
            elif info.get("available") or decision:
                stats["meta_only"] += 1
                continue
            else:
                continue

            row = by_id.get(oid) or {}
            for raw_code in values:
                try:
                    parsed = parse_marking_code(raw_code)
                except Exception as exc:
                    stats["errors"].append(f"{oid}: WB вернул КИЗ, но FBE не смог разобрать его: {exc}")
                    continue
                outcome = recover_marking_code_from_wb(
                    marking_db_path,
                    order_id=oid,
                    raw_code=str(parsed.get("raw_code") or raw_code),
                    gtin=str(parsed.get("gtin") or ""),
                    serial=str(parsed.get("serial") or ""),
                    seller_article=str(row.get("seller_article") or ""),
                    supply_id=str(row.get("supply_id") or ""),
                    wb_decision=decision,
                )
                status = str(outcome.get("status") or "")
                if status == "recovered":
                    stats["recovered"] += 1
                    wb_state = str(row.get("wb_status") or "").lower()
                    if wb_state:
                        update_post_sale_wb_statuses(
                            marking_db_path, [{"id": oid, "wbStatus": wb_state}]
                        )
                    break
                if status == "already_linked":
                    stats["already_linked"] += 1
                    break
                if status in {"order_conflict", "code_conflict"}:
                    stats["conflicts"] += 1
                    update_fbs_order_sgtin_meta(
                        marking_db_path, order_id=oid,
                        state=f"conflict:{state or 'wb'}", values=values
                    )
                    stats["errors"].append(f"{oid}: конфликт КИЗ — значение в WB отличается от локальной привязки FBE")
                    break
    return stats


def _reconcile_post_sale_true_statuses(order_ids: list[int] | None = None) -> dict[str, Any]:
    """Refresh True API state for locally linked KIZ, including WB-recovered codes."""
    wanted = {int(x) for x in (order_ids or []) if int(x) > 0}
    rows = list_fbs_lifecycle_rows(marking_db_path, limit=10000)
    rows = [r for r in rows if r.get("raw_code") and (not wanted or int(r.get("order_id") or 0) in wanted)]
    stats = {"checked": 0, "updated": 0, "documents_checked": 0, "documents_updated": 0, "errors": []}
    if not rows:
        return stats

    # Document transport state and the KIZ lifecycle are deliberately stored and
    # checked separately. Once every linked code reached its expected terminal
    # state, an old document status is audit history and is no longer polled.
    terminal_documents = TRUE_DOCUMENT_SUCCESS_STATUSES | TRUE_DOCUMENT_FAILURE_STATUSES
    for document in list_circulation_documents(marking_db_path, "POSTSALE", limit=500):
        document_uuid = str(document.get("document_uuid") or "").strip()
        document_type = str(document.get("document_type") or "").strip().upper()
        if not document_uuid or str(document.get("status") or "").strip().upper() in terminal_documents:
            continue
        link_field = "retire_document_id" if document_type == "LK_RECEIPT" else "return_document_id"
        linked = [row for row in rows if str(row.get(link_field) or "").strip() == document_uuid]
        if not linked:
            continue
        if document_type == "LK_RECEIPT" and all(
            str(row.get("post_sale_status") or "") in {"retired", "return_submitted", "returned"}
            for row in linked
        ):
            continue
        if document_type == "LP_RETURN" and all(
            str(row.get("post_sale_status") or "") == "returned" for row in linked
        ):
            continue
        try:
            response = suz.get_true_document_info(document_uuid)
            status = _true_doc_status(response) or str(document.get("status") or "PROCESSING")
            error_text = _true_doc_error(response) if status in TRUE_DOCUMENT_FAILURE_STATUSES else ""
            update_circulation_document(
                marking_db_path,
                document_uuid,
                status=status,
                response=response,
                error_text=error_text,
            )
            stats["documents_checked"] += 1
            stats["documents_updated"] += 1
            if status in TRUE_DOCUMENT_FAILURE_STATUSES:
                reason = error_text or f"Документ обработан со статусом {status}"
                reopened = mark_post_sale_document_failed(
                    marking_db_path,
                    document_uuid=document_uuid,
                    document_type=document_type,
                    error_text=reason,
                )
                suffix = f"; к повторной отправке: {reopened}" if reopened else ""
                stats["errors"].append(f"Документ {document_uuid}: {reason}{suffix}")
            elif error_text:
                stats["errors"].append(f"Документ {document_uuid}: {error_text}")
        except Exception as exc:
            stats["errors"].append(f"Документ {document_uuid}: {exc}")

    if not wanted:
        rows = [
            row for row in rows
            if str(row.get("post_sale_status") or "") in {"retirement_submitted", "return_submitted"}
            or (
                str(row.get("wb_status") or "").lower() == "sold"
                and str(row.get("post_sale_status") or "") not in {"retired", "returned"}
            )
        ]

    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        product_group = _assignment_product_group(row)
        if product_group not in {"perfumery", "chemistry"}:
            stats["errors"].append(
                f"Задание {row.get('order_id') or '—'}: не определена товарная группа КИЗ"
            )
            continue
        grouped.setdefault(product_group, []).append(row)

    for product_group, group_rows in grouped.items():
        raw_codes = [str(r.get("raw_code") or "") for r in group_rows if r.get("raw_code")]
        status_map: dict[str, str] = {}
        try:
            for start in range(0, len(raw_codes), 1000):
                response_rows = suz.get_cises_info(raw_codes[start:start + 1000], product_group=product_group)
                status_map.update(_extract_cis_statuses(response_rows))
                stats["checked"] += len(raw_codes[start:start + 1000])
        except Exception as exc:
            stats["errors"].append(f"ЧЗ {product_group}: {exc}")
            continue
        counts = update_marking_true_statuses(
            marking_db_path, statuses_by_identification_code=status_map
        )
        post_sale_counts = update_post_sale_true_status(
            marking_db_path, statuses_by_identification_code=status_map
        )
        stats["updated"] += int(counts.get("updated") or 0)
        stats["updated"] += int(post_sale_counts.get("updated") or 0)
    return stats


def _sync_fbs_lifecycle_registry(*, selected_order_ids: list[int] | None = None, include_supplies: bool = True, deep: bool = True) -> dict[str, Any]:
    """Synchronise the local FBS registry with WB without requiring a permanently running FBE.

    Supply/order relationships are persisted in SQLite. This is intentionally
    independent of the marking lifecycle: the registry can tell us that an
    assembly task exists and was sold even when its KIZ row is missing, which
    makes operational gaps visible instead of showing a misleading zero.
    """
    stats = {
        'supplies': 0,
        'supply_order_links': 0,
        'orders': 0,
        'statuses': 0,
        'status_updates': 0,
        'metadata': {},
        'errors': [],
        'open_supply_ids': [],
        'core': {
            'supplies_ok': not include_supplies,
            'membership_ok': not include_supplies,
            'orders_ok': not include_supplies,
            'statuses_ok': False,
        },
    }
    ensure_fbs_registry_from_marking_codes(marking_db_path)

    if include_supplies:
        try:
            raw_supplies = wb.get_all_supplies(
                max_pages=max(1, int(config.get('supplies_api_max_pages', 3) or 3))
            )
            supplies = prepare_supplies([dict(item) for item in raw_supplies])
            stats['supplies'] = upsert_fbs_supplies(marking_db_path, supplies)
            stats['open_supply_ids'] = [
                str(item.get('id') or '').strip() for item in supplies
                if str(item.get('id') or '').strip() and not bool(item.get('done'))
            ]
            stats['core']['supplies_ok'] = True
            # Persist exact membership. Operational sync must cover *all* open
            # supplies, not merely the first N records: sellers can easily have
            # more than eight FBS warehouses/supplies at once. Keep a bounded
            # tail of recently closed supplies for handover/post-sale continuity.
            if deep:
                sync_limit = max(1, int(config.get('fbs_lifecycle_supply_sync_limit', 60)))
                open_supplies = [s for s in supplies if not bool(s.get('done'))]
                open_ids = {str(s.get('id') or '') for s in open_supplies}
                closed_tail = [
                    s for s in supplies if str(s.get('id') or '') not in open_ids
                ][:sync_limit]
                sync_supplies = open_supplies + closed_tail
            else:
                tail_limit = max(2, min(12, int(config.get('fbs_operational_supply_sync_limit', 8) or 8)))
                open_supplies = [s for s in supplies if not bool(s.get('done'))]
                open_ids = {str(s.get('id') or '') for s in open_supplies}
                closed_tail = [s for s in supplies if str(s.get('id') or '') not in open_ids][:tail_limit]
                sync_supplies = open_supplies + closed_tail
            membership_ok = True
            for supply in sync_supplies:
                sid = str(supply.get('id') or '').strip()
                if not sid:
                    continue
                try:
                    ids = wb.get_supply_order_ids(sid)
                    stats['supply_order_links'] += set_fbs_supply_order_ids(
                        marking_db_path, supply_id=sid, order_ids=ids
                    )
                except Exception as exc:
                    membership_ok = False
                    stats['errors'].append(f'Состав поставки {sid}: {exc}')
            stats['core']['membership_ok'] = membership_ok
        except Exception as exc:
            stats['errors'].append(f'Поставки WB: {exc}')

        try:
            recent_orders = wb.get_orders_last_days(days=(int(config.get('fbs_registry_history_days', 92)) if deep else int(config.get('fbs_operational_history_days', 7))))
            stats['orders'] = upsert_fbs_orders(marking_db_path, recent_orders)
            stats['core']['orders_ok'] = True
        except Exception as exc:
            stats['errors'].append(f'Сборочные задания WB: {exc}')

    if selected_order_ids:
        ids = sorted({int(x) for x in selected_order_ids if int(x) > 0})
        # Explicit lookups are also registered, so the result survives reload.
        upsert_fbs_orders(marking_db_path, [{'id': oid} for oid in ids])
    else:
        sync_date_from = _registry_history_start() if deep else (moscow_now().date() - timedelta(days=int(config.get('fbs_operational_history_days', 7)))).isoformat()
        current_rows = list_fbs_registry_rows(
            marking_db_path,
            limit=max(1, int(config.get('fbs_lifecycle_status_sync_limit', 10000) if deep else config.get('fbs_operational_status_sync_limit', 1200))),
            date_from=sync_date_from,
        )
        ids = [
            int(r.get('order_id') or 0) for r in current_rows
            if int(r.get('order_id') or 0) > 0
            and (deep or str(r.get('wb_status') or '').lower() not in WB_TERMINAL_STATUSES)
        ]

    if ids:
        try:
            wb_rows = wb.get_order_statuses(ids)
            stats['statuses'] = update_fbs_order_statuses(marking_db_path, wb_rows)
            stats['core']['statuses_ok'] = True
            stats['status_updates'] = update_post_sale_wb_statuses(marking_db_path, wb_rows)
        except Exception as exc:
            stats['errors'].append(f'Статусы WB: {exc}')
    else:
        stats['core']['statuses_ok'] = True

    # Recover KIZ values attached manually in WB. For routine sync we only ask
    # metadata for rows that still have no local KIZ; explicit order lookup always
    # asks WB for the requested IDs.
    if selected_order_ids:
        meta_ids = ids
    else:
        current_rows = list_fbs_lifecycle_rows(
            marking_db_path, limit=max(1, int(config.get('fbs_lifecycle_status_sync_limit', 10000))),
            date_from=_registry_history_start(),
        )
        meta_ids = [
            int(r.get('order_id') or 0)
            for r in current_rows
            if int(r.get('order_id') or 0) > 0
            and not r.get('raw_code')
            and (
                str(r.get('wb_status') or '').lower() == 'sold'
                or str(r.get('supplier_status') or '').lower() == 'complete'
            )
        ]
    if meta_ids and (deep or selected_order_ids):
        meta_stats = _recover_wb_sgtin_metadata(meta_ids)
        stats['metadata'] = meta_stats
        stats['errors'].extend(meta_stats.get('errors') or [])

    if include_supplies and not selected_order_ids:
        core = stats.get('core') or {}
        core_ok = bool(
            core.get('supplies_ok')
            and core.get('membership_ok')
            and core.get('orders_ok')
            and core.get('statuses_ok')
        )
        core_errors = [
            str(err) for err in stats.get('errors') or []
            if str(err).startswith((
                'Поставки WB:', 'Состав поставки ',
                'Сборочные задания WB:', 'Статусы WB:',
            ))
        ]
        try:
            record_fbs_sync_state(
                marking_db_path,
                sync_key=('deep' if deep else 'operational'),
                ok=core_ok,
                details={
                    'supplies': int(stats.get('supplies') or 0),
                    'orders': int(stats.get('orders') or 0),
                    'statuses': int(stats.get('statuses') or 0),
                    'supply_order_links': int(stats.get('supply_order_links') or 0),
                    'open_supply_ids': list(stats.get('open_supply_ids') or []),
                    'core': core,
                },
                error_text='; '.join(core_errors[:5]),
            )
        except Exception as exc:
            stats['errors'].append(f'Checkpoint FBS: {exc}')
    return stats


def _post_sale_context(query: str = '', date_from: str = '') -> dict[str, Any]:
    ensure_fbs_registry_from_marking_codes(marking_db_path)
    date_from = str(date_from or '').strip()[:10]
    rows = list_fbs_lifecycle_rows(marking_db_path, limit=10000, query=query, date_from=date_from)
    catalog_now = get_catalog()
    for row in rows:
        product = catalog_now.find_by_article(str(row.get('seller_article') or ''))
        row['display_name'] = str((product or {}).get('display_name') or row.get('seller_article') or 'Товар')
        state = str(row.get('post_sale_status') or '')
        wb_state = str(row.get('wb_status') or '').lower()
        circulation = str(row.get('circulation_status') or '').upper()
        has_kiz = bool(row.get('marking_code_id') and row.get('raw_code'))
        wb_meta_state = str(row.get('wb_sgtin_state') or '').strip()
        wb_meta_state_l = wb_meta_state.lower()
        wb_meta_has_kiz = bool(int(row.get('wb_sgtin_count') or 0) > 0 or wb_meta_state_l in {
            'filled','valid','pending','sgtinmaysell','sgtinintroduced'
        })
        row['wb_meta_has_kiz'] = wb_meta_has_kiz
        row['kiz_source'] = 'WB' if str(row.get('marking_status') or '') == 'recovered_wb' else ('FBE' if has_kiz else '')
        if state in {'retirement_submitted','retired','return_submitted','returned'}:
            effective = state
        elif wb_meta_state_l.startswith('conflict:'):
            effective = 'kiz_conflict'
        elif wb_state == 'sold':
            if not has_kiz and wb_meta_has_kiz:
                effective = 'sold_wb_kiz_unavailable'
            elif not has_kiz:
                effective = 'sold_no_kiz'
            elif circulation != 'INTRODUCED':
                effective = 'sold_not_introduced'
            else:
                effective = 'ready_to_retire'
        elif wb_state == 'canceled':
            effective = 'canceled'
        elif wb_state:
            effective = 'at_wb'
        elif has_kiz:
            effective = 'unknown'
        else:
            effective = 'registry_only'
        row['post_sale_status_effective'] = effective
        row['post_sale_label'] = {
            'at_wb':'Передано / в обработке WB',
            'ready_to_retire':'К выводу',
            'retirement_submitted':'Вывод обрабатывается',
            'retired':'Выведен',
            'canceled':'Отказ / отмена',
            'return_submitted':'Возврат обрабатывается',
            'returned':'Возвращен в оборот',
            'kiz_conflict':'Конфликт КИЗ: WB и FBE различаются',
            'sold_no_kiz':'Продано · WB не вернул КИЗ',
            'sold_wb_kiz_unavailable':'Продано · КИЗ закреплен в WB, но код не возвращен',
            'sold_not_introduced':'Продано · КИЗ найден, проверяем ЧЗ',
            'registry_only':'Задание найдено · КИЗ не назначен',
            'unknown':'Ожидает статуса WB',
        }.get(effective, 'Ожидает статуса WB')
        row['post_sale_class'] = {
            'ready_to_retire':'warn','retired':'ok','canceled':'muted','returned':'ok',
            'retirement_submitted':'info','return_submitted':'info','at_wb':'info',
            'kiz_conflict':'error','sold_no_kiz':'error','sold_wb_kiz_unavailable':'warn','sold_not_introduced':'warn','registry_only':'muted','unknown':'muted',
        }.get(effective,'muted')
        row['can_retire'] = effective == 'ready_to_retire'
        row['can_return'] = effective == 'retired'
    return {
        'rows': rows,
        'at_wb': [r for r in rows if r['post_sale_status_effective'] in {'at_wb','unknown'}],
        'ready': [r for r in rows if r['post_sale_status_effective']=='ready_to_retire'],
        'retired': [r for r in rows if r['post_sale_status_effective'] in {'retired','retirement_submitted'}],
        'returns': [r for r in rows if r['post_sale_status_effective'] in {'return_submitted','returned'}],
        'canceled': [r for r in rows if r['post_sale_status_effective']=='canceled'],
        'issues': [r for r in rows if r['post_sale_status_effective'] in {'kiz_conflict','sold_no_kiz','sold_wb_kiz_unavailable','sold_not_introduced','registry_only'}],
        'query': query,
        'date_from': date_from,
    }


@app.get('/marking/post-sale', response_class=HTMLResponse)
def marking_post_sale_page(
    request: Request, msg: str | None = None, error: str | None = None, q: str = '', date_from: str = ''
):
    ctx = _post_sale_context(query=q, date_from=date_from)
    ctx.update({'request':request,'msg':msg,'error':error,'config':config})
    return templates.TemplateResponse(request, 'post_sale.html', ctx)


def _post_sale_refresh_impl(report: Callable[..., None]) -> dict[str, Any]:
    report(progress=0, total=3, message='Синхронизирую историю WB…')
    sync = _sync_fbs_lifecycle_registry(include_supplies=True, deep=True)
    report(progress=1, total=3, message='Восстанавливаю КИЗы и статусы заданий…')
    true_sync = _reconcile_post_sale_true_statuses()
    report(progress=2, total=3, message='Обновляю локальный реестр…')
    meta = sync.get('metadata') or {}
    recovered = int(meta.get('recovered') or 0)
    meta_only = int(meta.get('meta_only') or 0)
    message = (
        f"WB: {sync['status_updates']} статусов · КИЗ восстановлено: {recovered}"
        f" · КИЗ только отмечено WB: {meta_only}"
        f" · ЧЗ обновлено: {true_sync.get('updated', 0)}"
    )
    report(progress=3, total=3, message=message)
    return {
        'ok': True, 'updated': sync['status_updates'], 'sync': sync,
        'true_sync': true_sync, 'message': message,
    }


@app.post('/api/marking/post-sale/refresh')
def refresh_post_sale_statuses():
    job_id = submit_background_job(
        'Синхронизация продаж и КИЗ',
        _post_sale_refresh_impl,
        operation_key='post-sale-refresh',
    )
    return {'ok': True, 'job_id': job_id, 'message': 'Синхронизация запущена'}


def _check_post_sale_orders_impl(payload: PostSaleLookupPayload) -> dict[str, Any]:
    ids = sorted({int(x) for x in payload.order_ids if int(x) > 0})
    if not ids:
        raise MarkingDbError('Укажите хотя бы одно сборочное задание')
    sync = _sync_fbs_lifecycle_registry(selected_order_ids=ids, include_supplies=False)
    true_sync = _reconcile_post_sale_true_statuses(ids)
    rows = list_fbs_lifecycle_rows(marking_db_path, limit=max(100, len(ids) * 3))
    found = {int(row.get('order_id') or 0) for row in rows}
    meta = sync.get('metadata') or {}
    return {
        'ok': True,
        'requested': len(ids),
        'found': sum(1 for order_id in ids if order_id in found),
        'updated': sync['status_updates'],
        'recovered': int(meta.get('recovered') or 0),
        'message': (
            f"Проверено: {len(ids)} · статусов WB: {sync['status_updates']}"
            f" · КИЗ восстановлено из WB: {int(meta.get('recovered') or 0)}"
            f" · ЧЗ обновлено: {true_sync.get('updated', 0)}"
        ),
    }


@app.post('/api/marking/post-sale/check-orders')
def check_post_sale_orders(payload: PostSaleLookupPayload):
    ids = ','.join(str(value) for value in sorted(set(payload.order_ids))) or 'none'
    job_id = submit_background_job(
        'Проверка заданий WB и КИЗ',
        lambda report: _check_post_sale_orders_impl(payload),
        operation_key=f'post-sale-check:{ids}',
    )
    return {'ok': True, 'job_id': job_id, 'message': 'Проверка запущена'}


def _retire_payload(rows: list[dict[str, Any]], *, participant_inn: str) -> dict[str, Any]:
    """Build the True API LK_RECEIPT body for a WB distance sale."""
    if not rows:
        raise MarkingDbError("Нет КИЗов для документа вывода")
    today = moscow_now().date().isoformat()
    order_ids = sorted({int(row.get("order_id") or 0) for row in rows if int(row.get("order_id") or 0) > 0})
    products = [
        {"cis": suz.identification_code(str(row.get("raw_code") or ""))}
        for row in rows
    ]
    number_tail = f"{order_ids[0]}-{order_ids[-1]}" if order_ids else "KIZ"
    return {
        "inn": participant_inn,
        "action": "DISTANCE",
        "action_date": today,
        "document_type": "OTHER",
        "document_number": f"WB-FBS-{number_tail}",
        "document_date": today,
        "primary_document_custom_name": "Wildberries FBS",
        "products": products,
    }


def _retire_post_sale_impl(payload: PostSaleIdsPayload) -> dict[str, Any]:
    if not payload.confirmed:
        raise MarkingDbError('Подтвердите вывод КИЗ из оборота')
    all_rows = {
        int(row['order_id']): row
        for row in list_post_sale_assignments(marking_db_path, limit=5000)
        if row.get('order_id')
    }
    ids = payload.order_ids or [
        order_id for order_id, row in all_rows.items()
        if str(row.get('post_sale_status') or '') == 'ready_to_retire'
    ]
    selected = [all_rows[order_id] for order_id in ids if order_id in all_rows]
    lifecycle_by_id = {
        int(row.get('order_id') or 0): row
        for row in list_fbs_lifecycle_rows(marking_db_path, limit=10000)
        if row.get('order_id')
    }
    conflicts = [
        row for row in selected
        if str(
            (lifecycle_by_id.get(int(row.get('order_id') or 0)) or {}).get('wb_sgtin_state') or ''
        ).lower().startswith('conflict:')
    ]
    if conflicts:
        raise MarkingDbError('Вывод заблокирован: для части заданий КИЗ в WB отличается от локального КИЗ FBE')
    already_submitted = [
        row for row in selected
        if str(row.get('post_sale_status') or '') in {'retirement_submitted', 'retired'}
    ]
    rows = [
        row for row in selected
        if str(row.get('post_sale_status') or '') == 'ready_to_retire'
    ]
    bad = [row for row in selected if row not in rows and row not in already_submitted]
    if bad:
        raise MarkingDbError('В вывод попали задания без подтвержденного wbStatus=sold')
    if not rows:
        if already_submitted:
            return {'ok': True, 'documents': [], 'already_submitted': len(already_submitted),
                    'message': 'Все выбранные КИЗы уже были отправлены на вывод; повторная отправка не выполнена'}
        raise MarkingDbError('Нет подтвержденных продаж для вывода')
    inn = _validate_inn(
        str(_circulation_defaults().get('participant_inn') or ''), 'ИНН производителя'
    )
    documents: list[str] = []
    for product_group, group_rows in _post_sale_group_rows(rows).items():
        document_payload = _retire_payload(group_rows, participant_inn=inn)
        document = suz.create_true_document(
            product_group=product_group,
            document_type='LK_RECEIPT',
            document_payload=document_payload,
        )
        document_uuid = str(document.get('uuid') or '').strip()
        if not document_uuid:
            raise MarkingDbError(f'True API не вернул UUID документа вывода для {product_group}')
        create_circulation_document(
            marking_db_path, supply_id='POSTSALE', document_uuid=document_uuid,
            document_type='LK_RECEIPT', product_group=product_group,
            payload=document_payload, response=document,
            order_ids=[int(row['order_id']) for row in group_rows],
            submission_kind='retirement',
        )
        documents.append(document_uuid)
    return {
        'ok': True,
        'documents': documents,
        'already_submitted': len(already_submitted),
        'message': f'Отправлено на вывод: {len(rows)} КИЗ',
    }


@app.post('/api/marking/post-sale/retire')
def retire_post_sale(payload: PostSaleIdsPayload):
    if not payload.confirmed:
        return _marking_error_payload(MarkingDbError('Подтвердите вывод КИЗ из оборота'), 400)
    ids = ','.join(str(value) for value in sorted(set(payload.order_ids))) or 'all'
    job_id = submit_background_job(
        'Вывод КИЗ из оборота',
        lambda report: _retire_post_sale_impl(payload),
        operation_key=f'post-sale-retire:{ids}',
    )
    return {'ok': True, 'job_id': job_id, 'message': 'Документы вывода формируются в фоне'}


def _return_payload(rows: list[dict[str, Any]], *, participant_inn: str) -> dict[str, Any]:
    today=moscow_now().date().isoformat(); catalog_now=get_catalog()
    products=[]
    for r in rows:
        p=catalog_now.find_by_article(str(r.get('seller_article') or '')) or {}
        comp=_catalog_compliance(p)
        item={'ki':suz.identification_code(str(r.get('raw_code') or '')),'paid':True,'primary_document_type':'OTHER','primary_document_custom_name':'Wildberries FBS возврат','primary_document_number':str(r.get('order_id') or ''),'primary_document_date':today}
        if comp.get('certificate_type'): item['certificate_type']=comp['certificate_type']
        if comp.get('certificate_number'): item['certificate_number']=comp['certificate_number']
        if comp.get('certificate_date'): item['certificate_date']=comp['certificate_date']
        products.append(item)
    return {'trade_participant_inn':participant_inn,'return_type':'REMOTE_SALE_RETURN','paid':True,'primary_document_type':'OTHER','primary_document_custom_name':'Wildberries FBS возврат','primary_document_number':f'WB-RETURN-{today}','primary_document_date':today,'products_list':products}


def _return_post_sale_impl(payload: PostSaleIdsPayload) -> dict[str, Any]:
    if not payload.confirmed:
        raise MarkingDbError('Подтвердите возврат КИЗ в оборот')
    all_rows = {
        int(row['order_id']): row
        for row in list_post_sale_assignments(marking_db_path, limit=5000)
        if row.get('order_id')
    }
    selected = [all_rows[order_id] for order_id in payload.order_ids if order_id in all_rows]
    already_submitted = [
        row for row in selected
        if str(row.get('post_sale_status') or '') in {'return_submitted', 'returned'}
    ]
    rows = [row for row in selected if str(row.get('post_sale_status') or '') == 'retired']
    bad = [row for row in selected if row not in rows and row not in already_submitted]
    if bad or (not rows and not already_submitted):
        raise MarkingDbError('Возврат доступен только для уже выведенных КИЗ')
    if not rows:
        return {'ok': True, 'documents': [], 'already_submitted': len(already_submitted),
                'message': 'Все выбранные КИЗы уже отправлены на возврат; повторная отправка не выполнена'}
    inn = _validate_inn(
        str(_circulation_defaults().get('participant_inn') or ''), 'ИНН производителя'
    )
    documents: list[str] = []
    for product_group, group_rows in _post_sale_group_rows(rows).items():
        document_payload = _return_payload(group_rows, participant_inn=inn)
        document = suz.create_true_document(
            product_group=product_group,
            document_type='LP_RETURN',
            document_payload=document_payload,
        )
        document_uuid = str(document.get('uuid') or '').strip()
        if not document_uuid:
            raise MarkingDbError(f'True API не вернул UUID документа возврата для {product_group}')
        create_circulation_document(
            marking_db_path, supply_id='POSTSALE', document_uuid=document_uuid,
            document_type='LP_RETURN', product_group=product_group,
            payload=document_payload, response=document,
            order_ids=[int(row['order_id']) for row in group_rows],
            submission_kind='return',
        )
        documents.append(document_uuid)
    return {
        'ok': True,
        'documents': documents,
        'already_submitted': len(already_submitted),
        'message': f'Отправлено на возврат в оборот: {len(rows)} КИЗ',
    }


@app.post('/api/marking/post-sale/return')
def return_post_sale(payload: PostSaleIdsPayload):
    if not payload.confirmed:
        return _marking_error_payload(MarkingDbError('Подтвердите возврат КИЗ в оборот'), 400)
    ids = ','.join(str(value) for value in sorted(set(payload.order_ids))) or 'none'
    job_id = submit_background_job(
        'Возврат КИЗ в оборот',
        lambda report: _return_post_sale_impl(payload),
        operation_key=f'post-sale-return:{ids}',
    )
    return {'ok': True, 'job_id': job_id, 'message': 'Документы возврата формируются в фоне'}
