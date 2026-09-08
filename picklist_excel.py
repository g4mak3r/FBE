from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _norm(value: str) -> str:
    value = str(value or "").strip().lower().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", "", value)


ARTICLE_HEADERS = {
    "артикулпродавца",
    "артикулпоставщика",
    "артикул",
    "sellerarticle",
    "supplierarticle",
    "vendorcode",
    "article",
}

STICKER_HEADERS = {
    "стикер",
    "номерстикера",
    "стикерномер",
    "номеpстикера",
    "номерсборочногозадания",
    "сборочноезадание",
    "номершк",
    "шк",
    "barcode",
    "barcodes",
    "wbкод",
}

NAME_HEADERS = {
    "название",
    "наименование",
    "названиетовара",
    "наименованиетовара",
    "названиесобираемоготовара",
    "товар",
    "product",
    "productname",
}


def _find_header_row(ws, max_rows: int = 40) -> tuple[int, dict[str, int]] | None:
    """Find a WB picklist header row.

    WB portal XLSX files often have service rows above the table:
    date, supply title, blank row, item count, then the real header.
    Some exports also have unreliable sheet dimensions, so do not trust
    max_row/max_column while scanning the top of the worksheet.
    """
    for row_idx in range(1, max_rows + 1):
        values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, 80)]
        headers = [_clean(value) for value in values]
        normalized = {_norm(h): i for i, h in enumerate(headers) if h}
        if any(h in normalized for h in ARTICLE_HEADERS):
            return row_idx, normalized
    return None


def _first_matching_index(headers: dict[str, int], candidates: set[str]) -> int | None:
    for candidate in candidates:
        if candidate in headers:
            return headers[candidate]
    return None


def parse_wb_picklist_xlsx(path: str | Path) -> tuple[list[dict[str, str]], list[str]]:
    """Parse WB подборочный лист and preserve its visible row order.

    Required: seller article column.
    Optional: sticker/assembly number and product name columns.
    The WB portal changes column names from time to time, so the parser accepts
    several Russian/English header variants and returns warnings rather than
    crashing when optional fields are missing.
    """
    warnings: list[str] = []
    rows_out: list[dict[str, str]] = []
    path = Path(path)

    # WB exports picklists with a wrong/too-small worksheet dimension in some files.
    # In openpyxl read_only mode such sheets may report max_row/max_column as 1/1,
    # although the actual table starts lower on the page. Load normally so header
    # detection sees all cells exactly as Excel shows them.
    wb = load_workbook(path, data_only=True, read_only=False)

    selected = None
    for ws in wb.worksheets:
        found = _find_header_row(ws)
        if found:
            selected = (ws, found[0], found[1])
            break

    if not selected:
        return [], ["В листе подбора не найден столбец 'Артикул продавца'."]

    ws, header_row, headers = selected
    article_idx = _first_matching_index(headers, ARTICLE_HEADERS)
    sticker_idx = _first_matching_index(headers, STICKER_HEADERS)
    name_idx = _first_matching_index(headers, NAME_HEADERS)

    if article_idx is None:
        return [], ["В листе подбора не найден столбец 'Артикул продавца'."]

    if sticker_idx is None:
        warnings.append("Не найден столбец с номером стикера/сборочного задания. Порядок всё равно будет взят из строк Excel.")

    for excel_row_num, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        article = _clean(row[article_idx]) if article_idx < len(row) else ""
        if not article:
            continue

        sticker_number = _clean(row[sticker_idx]) if sticker_idx is not None and sticker_idx < len(row) else ""
        product_name = _clean(row[name_idx]) if name_idx is not None and name_idx < len(row) else ""

        rows_out.append(
            {
                "row_number": str(excel_row_num),
                "seller_article": article,
                "sticker_number": sticker_number,
                "product_name": product_name,
            }
        )

    if not rows_out:
        warnings.append("Лист подбора распознан, но строк с артикулами продавца не найдено.")

    return rows_out, warnings
